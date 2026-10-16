"""All maintained rules, independent annotated cases, actual core-task bindings."""
from copy import deepcopy
from dataclasses import replace
from io import BytesIO
import json
from pathlib import Path
import sqlite3
import tempfile
import unittest
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from scripts.check_rules import read_cases,case_values,dataset_for
from src import engine,rule_conversion as conversion,rule_bindings,tax_conditions
from src.kg_retrieval import metrics
from src.kg_retrieval.bundle import split_bundle
from tests.enterprise_support import material_key
from tests.test_core_inputs import book,TAX
from webapp import app as module,policy_graph as pg
from webapp.storage import Store

ROOT=Path(__file__).resolve().parents[1]
RULES={r.id:r for r in engine.load_rules(ROOT/'rules')}


class ConversionTests(unittest.TestCase):
    def setUp(self):
        self.db=sqlite3.connect(':memory:')
        self.db.row_factory=sqlite3.Row
        pg.migrate(self.db)
        self.addCleanup(self.db.close)
        pg.sync_rules(self.db,list(RULES.values()))

    def test_all_96_independent_annotated_cases_match_compiled_conditions(self):
        count=0
        for group in read_cases():
            compiled=conversion.compile_rule(RULES[group['rule_id']])
            for case in group['cases']:
                with self.subTest(rule=group['rule_id'],case=case['name']):
                    values=case_values(group,case)
                    dataset=dataset_for(RULES[group['rule_id']],values)
                    scope={key:getattr(dataset.company,key) for key in ('industry','region','taxpayer_type','business_scope')}
                    facts={m['metric_id']:{'state':'known','value':str(values[name]),'unit':conversion.metric_unit(name),
                           'source_record_ids':['source:'+name], 'metric_period':dataset.metrics[name].period,
                           'reference':dataset.metrics[name].reference,'company_scope':scope,'simulation':True}
                           for name,m in compiled['bases'].items() if name in values}
                    truth={}
                    for calc,condition in zip(compiled['calculations'],compiled['conditions']):
                        result=conversion.calculate(calc,facts,task={
                            'start_date':'2026-01-01','end_date':'2026-12-31','data_scope':'synthetic'})
                        truth[condition['condition_id']]='UNKNOWN' if result['state']!='known' else 'TRUE' if result['value']=='1' else 'FALSE'
                    result=tax_conditions.logic_value(compiled['logic'],truth)
                    if 'UNKNOWN' in truth.values():result='UNKNOWN'
                    self.assertEqual(result,{'hit':'TRUE','pass':'FALSE','skipped':'UNKNOWN'}[case['expected']])
                    count+=1
        self.assertEqual(count,96)

    def test_conversion_is_idempotent_versioned_and_never_approves_policy(self):
        before={t:self.db.execute(f'SELECT count(*) FROM {t}').fetchone()[0] for t in pg.IDS}
        pg.sync_rules(self.db,list(RULES.values()))
        self.assertEqual(before,{t:self.db.execute(f'SELECT count(*) FROM {t}').fetchone()[0] for t in pg.IDS})
        self.assertEqual(self.db.execute('SELECT count(*) FROM TaxRule WHERE conversion_json IS NOT NULL').fetchone()[0],24)
        self.assertEqual(self.db.execute("SELECT count(*) FROM Condition WHERE condition_source='internal_screening'").fetchone()[0],28)
        self.assertEqual(self.db.execute("SELECT count(*) FROM TaxRule WHERE review_status!='candidate'").fetchone()[0],0)
        first=conversion.compile_rule(RULES['R-001'])
        changed=conversion.compile_rule(replace(RULES['R-001'],logic={**RULES['R-001'].logic,'threshold':0.2}))
        self.assertNotEqual(first['rule_id'],changed['rule_id'])

    def test_calculation_registry_rejects_missing_or_forged_dependencies(self):
        calc=conversion.compile_rule(RULES['R-001'])['calculations'][0]
        self.assertEqual(conversion.calculate(calc,{})['state'],'unknown')
        bad={**calc,'calculation_code':'eval(anything)'}
        with self.assertRaises(ValueError):conversion.calculate(bad,{})
        spec=json.loads(calc['source_filter_json']);spec['logic']={'type':'python','code':'evil'}
        with self.assertRaises(engine.RuleError):conversion.validate_calculation({**calc,'source_filter_json':json.dumps(spec)})
        forged={'TaxRule':[{'rule_id':'forged','name':'fake','rule_type':'risk_screen','outcome_text':'fake','conversion_json':{}}]}
        with self.assertRaises(ValueError):pg.import_bundle(self.db,forged)

    def test_unverified_policy_leads_split_without_losing_originals_or_required_context(self):
        data={key:[] for key in ['facts','conditions','exceptions','gaps','relationships','policy_evidence','policy_leads']}
        data.update(task={},check={},bundle_id='test')
        original='必要例外与原文必须完整保留。'*12
        data['policy_leads']=[{'clause_id':'required','required':True,'text':'完整必要依据'}]+[
            {'clause_id':str(i),'required':False,'text':original,'retrieval_matches':['redundant trace']*100} for i in range(4)]
        result=split_bundle(data,2200)
        self.assertGreater(result['part_count'],1)
        self.assertFalse(result['unsplittable'])
        self.assertTrue(all(any(c['clause_id']=='required' for c in p['policy_leads']) for p in result['parts']))
        found={c['clause_id']:c['text'] for p in result['parts'] for c in p['policy_leads']}
        self.assertEqual(found,{c['clause_id']:c['text'] for c in data['policy_leads']})
        self.assertIn('retrieval_matches',data['policy_leads'][1])

    def test_report_keeps_required_and_ai_cited_originals_and_human_source_coordinates(self):
        from webapp.core_reports import report_evidence,report_gaps,source_location
        slot={'bundle':{'policy_evidence':[{'clause_id':'required','required':True}],
                        'policy_leads':[{'clause_id':'cited'},{'clause_id':'extra'}]},
              'ai':{'assessments':[{'clause_ids':['cited']}]}}
        original=deepcopy(slot)
        self.assertEqual([p['clause_id'] for p in report_evidence(slot)],['required','cited'])
        self.assertEqual(slot,original)
        self.assertEqual(source_location(json.dumps({'text_path':'官方原文.txt','start_line':3,'end_line':5})),
                         '官方原文.txt · 第3至5行')
        self.assertEqual(source_location('原表!C6'),'原表!C6')
        gaps=[{'code':'change_pending','clause_id':'one'}, {'code':'change_pending','clause_id':'two'},
              {'code':'policy_review','message':'原文待审核','clause_id':'three'}]
        saved=deepcopy(gaps)
        self.assertEqual(report_gaps(gaps),[{'message':'政策修订关系尚待审核。','count':2},
                                           {'message':'原文待审核','count':1}])
        self.assertEqual(gaps,saved)

    def test_all_cases_are_resolvable_from_existing_complete_material_metrics(self):
        # This independently exercises the Dataset -> retained audit source ->
        # graph metric -> derived condition path for all 24 rules and 96 cases.
        for group in read_cases():
            compiled=conversion.compile_rule(RULES[group['rule_id']])
            definitions={m['metric_id']:m for m in [*compiled['bases'].values(),*compiled['calculations']]}
            for case in group['cases']:
                with self.subTest(rule=group['rule_id'],case=case['name']):
                    dataset=dataset_for(RULES[group['rule_id']],case_values(group,case))
                    dataset.company.taxpayer_id=TAX
                    task={'id':'task','company_id':'client','tenant_id':'org','tax_no':TAX,
                          'start_date':'2026-01-01','end_date':'2026-12-31','data_scope':'synthetic'}
                    entry={'id':'audit','client_id':'client','org_id':'org','dataset':dataset}
                    rows=rule_bindings.audit_records(entry,task)
                    facts=metrics.collect(definitions,{'records':rows},task)
                    truth={c['condition_id']:tax_conditions.compare(c,definitions,facts) for c in compiled['conditions']}
                    actual=tax_conditions.logic_value(compiled['logic'],truth)
                    if any(f['state']!='known' for f in facts.values()):actual='UNKNOWN'
                    self.assertEqual(actual,{'hit':'TRUE','pass':'FALSE','skipped':'UNKNOWN'}[case['expected']])


class ConvertedScopeTests(unittest.TestCase):
    def collect(self,dataset,rule_code,task=None):
        dataset.company.taxpayer_id=TAX
        task=task or {'id':'task','company_id':'client','tenant_id':'org','tax_no':TAX,
                      'start_date':'2026-01-01','end_date':'2026-12-31'}
        entry={'id':'audit','client_id':'client','org_id':'org','dataset':dataset}
        rows=rule_bindings.audit_records(entry,task)
        compiled=conversion.compile_rule(RULES[rule_code])
        definitions={m['metric_id']:m for m in [*compiled['bases'].values(),*compiled['calculations']]}
        return metrics.collect(definitions,{'records':rows},task),compiled,rows

    def test_historical_period_is_preserved_and_wrong_or_missing_period_is_unknown(self):
        group=next(g for g in read_cases() if g['rule_id']=='R-020')
        data=dataset_for(RULES['R-020'],case_values(group,group['cases'][0]))
        good,compiled,rows=self.collect(data,'R-020')
        calculation=compiled['calculations'][0]['metric_id']
        self.assertEqual(good[calculation]['state'],'known')
        history='历史.上年同期收入'
        self.assertEqual(next(r for r in rows if r['metric_name']==history)['metric_period'],'2025-01-01至2025-12-31')
        for actual in ('','2026年度','2025年第1季度'):
            with self.subTest(actual=actual):
                bad=deepcopy(data);bad.metrics[history].period=actual
                facts,_,_=self.collect(bad,'R-020')
                self.assertEqual(facts[calculation]['state'],'unknown')

    def test_benchmark_scope_source_and_period_are_enforced_in_graph_calculation(self):
        from tests.test_business_scope import reference_dataset
        data=reference_dataset()
        good,compiled,_=self.collect(data,'R-003')
        calculation=compiled['calculations'][0]['metric_id']
        self.assertEqual(good[calculation]['state'],'known')
        for key,value in [('region','广东省'),('industry','另一行业'),('end','2026-06-30'),
                          ('source','其他来源'),('definition','另一比率'),('version','')]:
            with self.subTest(key=key):
                bad=deepcopy(data)
                for name in ('参考.税负率下限','参考.税负率上限'):
                    bad.metrics[name].reference[key]=value
                facts,_,_=self.collect(bad,'R-003')
                self.assertEqual(facts[calculation]['state'],'unknown')

    def test_enterprise_task_rejects_demo_benchmarks_even_from_synthetic_archive(self):
        from tests.test_business_scope import reference_dataset
        data=reference_dataset();data.simulation=True
        for name in ('参考.税负率下限','参考.税负率上限'):
            data.metrics[name].reference['nature']='仿真示意'
        facts,compiled,_=self.collect(data,'R-003')
        self.assertEqual(facts[compiled['calculations'][0]['metric_id']]['state'],'unknown')

    def test_reference_revision_changes_signature_and_keeps_frozen_metadata(self):
        from tests.test_business_scope import reference_dataset
        data=reference_dataset()
        before,_,rows=self.collect(data,'R-003');frozen=deepcopy(rows)
        for name in ('参考.税负率下限','参考.税负率上限'):
            data.metrics[name].reference['version']='新的构造版本'
        after,_,_=self.collect(data,'R-003')
        self.assertNotEqual(metrics.signature(before),metrics.signature(after))
        self.assertEqual(rows,frozen)
        self.assertEqual(next(r for r in rows if r['reference'])['reference']['version'],'测试-v1')

    def test_graph_arithmetic_does_not_use_legacy_legal_version_catalog(self):
        from tests.test_business_scope import reference_dataset
        with patch('src.legal_periods.resolve',side_effect=AssertionError('legacy catalog called')):
            facts,compiled,_=self.collect(reference_dataset(),'R-003')
        self.assertEqual(facts[compiled['calculations'][0]['metric_id']]['state'],'known')

    def test_old_calculation_is_unknown_until_rechecked_instead_of_crashing_history(self):
        old=deepcopy(conversion.compile_rule(RULES['R-001'])['calculations'][0])
        old['calculation_code']='yaml-screen-v1'
        spec=json.loads(old['source_filter_json']);spec['registry']='yaml-screen-v1'
        old['source_filter_json']=json.dumps(spec)
        result=conversion.calculate(old,{})
        self.assertEqual(result['state'],'unknown')
        self.assertIn('重新核查',result['reason'])

    def test_evidence_package_and_report_keep_actual_period_and_reference_metadata(self):
        from jinja2 import Environment,FileSystemLoader
        from src.kg_retrieval.bundle import _facts
        from tests.test_business_scope import reference_dataset
        from webapp.core_reports import report_evidence,report_gaps,source_location
        facts,compiled,_=self.collect(reference_dataset(),'R-003')
        definitions={m['metric_id']:m for m in [*compiled['bases'].values(),*compiled['calculations']]}
        task={'period_start':'2026-01-01','period_end':'2026-12-31'}
        catalog={sid for f in facts.values() for sid in f['source_record_ids']}
        normalized,gaps=_facts({'metrics':definitions},facts,task,catalog)
        self.assertFalse(gaps)
        ref=next(f for f in normalized.values() if f.get('reference'))
        self.assertEqual(ref['reference']['source'],'构造来源，不是真实业务验收')
        ref['metric_period']='2025年度'
        package={'task':task,'check':{},'facts':list(normalized.values()),'gaps':[],
                 'conditions':[],'exceptions':[],'policy_evidence':[],'policy_leads':[],
                 'relationships':[],'retrieval_trace':{'excluded':[]}}
        env=Environment(loader=FileSystemLoader(ROOT/'templates'),autoescape=True)
        env.filters.update(report_evidence=report_evidence,report_gaps=report_gaps,source_location=source_location)
        html=env.get_template('core_check_report.html').render(report={
            'task':{},'sources':[],'summary':{},'processing':{'status':'ready_for_analysis','issues':[]},
            'checks':{'test':{'bundle':package}}})
        self.assertIn('来源实际期间：2025年度',html)
        self.assertIn('测试-v1',html)
        self.assertIn('江苏省 / 批发和零售业 / 一般纳税人',html)


@material_key
class BindingHTTPTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.store=Store(Path(self.temp.name)/'test.db')
        p=patch.object(module,'store',self.store);p.start();self.addCleanup(p.stop)
        self.user=self.store.create_user('owner','Rule-conversion-2026!','owner','org_admin','alpha')
        self.store.create_user('other','Rule-conversion-2026!','other','org_admin','beta')
        self.client=TestClient(module.app);self.client.__enter__();self.addCleanup(self.client.__exit__,None,None,None)
        self.login('owner')
        # Retrieval behavior is tested here without depending on a machine's
        # optional embedding model. Real local-model acceptance runs separately.
        p=patch('src.kg_retrieval.semantic_search.search',return_value={'status':'not_ready','hits':[]});p.start();self.addCleanup(p.stop)

    def login(self,name):
        self.client.cookies.clear()
        self.client.cookies.set(module.COOKIE_NAME,self.store.authenticate(name,'Rule-conversion-2026!')[1])

    def upload(self,task,kind,raw):
        r=self.client.post('/api/enterprise/input-tasks/'+task['id']+'/files',data={'material_type':kind,'expected_revision':task['revision']},files={'files':(kind+'.xlsx',raw)})
        self.assertEqual(r.status_code,200,r.text)
        return r.json()

    def task(self):
        task=self.client.post('/api/enterprise/input-tasks',json={'company_name':'纯合成测试规则企业','tax_no':TAX,
              'start_date':'2026-01-01','end_date':'2026-01-31','applicable_forms':[]}).json()
        wb=load_workbook(BytesIO(book('account_balance')));s=wb.active;s['A6']='6001';s['B6']='主营业务收入';s['E6']='980000.00';s.append(['6051','其他业务收入','0','0','0','0'])
        raw=BytesIO();wb.save(raw);wb.close();task=self.upload(task,'account_balance',raw.getvalue())
        wb=load_workbook(BytesIO(book('vat_declaration')));wb.active['C6']='1000000.00';raw=BytesIO();wb.save(raw);wb.close()
        return self.upload(task,'vat_declaration',raw.getvalue())

    def bind(self,task,name,ids,field):
        r=self.client.post('/api/enterprise/input-tasks/'+task['id']+'/metric-bindings',json={'expected_revision':task['revision'],'name':name,'record_ids':ids,'field':field,'aggregation':'sum','basis':'已核实同一主体同期不含税口径，收入剔除结转并包含所需科目。'})
        self.assertEqual(r.status_code,200,r.text)
        return r.json()

    def test_uploaded_sources_compute_real_rule_without_fabricating_policy_approval(self):
        task=self.task()
        records=task['processing']['records']
        task=self.bind(task,'营业收入',[r['id'] for r in records if r['material_type']=='account_balance'],'current_credit')
        task=self.bind(task,'增值税.销售额',[r['id'] for r in records if r['material_type']=='vat_declaration' and r['item_key']=='1'],'declared_value')
        catalog=self.client.get('/api/knowledge/checks').json()
        self.assertEqual(sum(bool(r['computation_converted']) for r in catalog),24)
        rid=next(r['rule_id'] for r in catalog if r['rule_code']=='R-001')
        response=self.client.post('/api/knowledge/retrieve',json={'task_id':task['id'],'expected_revision':task['revision'],'rule_id':rid,'region_code':'CN-31'})
        self.assertEqual(response.status_code,200,response.text)
        result=response.json();check=result['bundle']['check']
        self.assertEqual(check['fact_evaluation'],'FALSE')
        self.assertEqual(check['evaluation'],'UNKNOWN')
        self.assertNotIn('rule_logic_pending',{g['code'] for g in result['bundle']['gaps']})
        self.assertEqual({f['name']:f['value'] for f in result['bundle']['facts'] if not f['is_calculation']},{'营业收入':'980000.00','增值税.销售额':'1000000.00'})
        self.assertIn('rule_unapproved',{g['code'] for g in result['bundle']['gaps']})
        report=self.client.post('/api/enterprise/input-tasks/'+task['id']+'/reports',json={'expected_revision':result['revision']})
        self.assertEqual(report.status_code,200,report.text)
        url='/api/enterprise/input-tasks/'+task['id']+'/reports/'+str(report.json()['version'])
        html=self.client.get(url).text
        self.assertIn('内部规则事实计算',html);self.assertIn('980000.00',html)
        self.login('other');self.assertEqual(self.client.get(url).status_code,404)

    def test_forged_values_foreign_rows_and_stale_revision_are_rejected(self):
        task=self.task();url='/api/enterprise/input-tasks/'+task['id']+'/metric-bindings'
        body={'expected_revision':task['revision'],'name':'营业收入','record_ids':['foreign'],'field':'current_credit','basis':'确认口径'}
        self.assertEqual(self.client.post(url,json=body).status_code,422)
        self.assertEqual(self.client.post(url,json={**body,'value':'999'}).status_code,422)
        self.assertEqual(self.client.post(url,json={**body,'expected_revision':task['revision']-1}).status_code,409)
        self.login('other');self.assertEqual(self.client.post(url,json=body).status_code,404)

    def test_binding_does_not_bypass_identity_missing_values_or_duplicate_versions(self):
        task=self.task();rows=[r for r in task['processing']['records'] if r['material_type']=='account_balance']
        task=self.bind(task,'营业收入',[r['id'] for r in rows],'current_credit')
        definition=conversion.compile_rule(RULES['R-001'])['bases']['营业收入']
        self.assertEqual(rule_bindings.resolve(definition,task['processing'],task['task'])['state'],'known')
        for issue in [{'code':'identity_unverified','source_file_id':rows[0]['source_file_id'],'message':'身份未核实'},
                      {'code':'duplicate_period','source_file_id':rows[0]['source_file_id'],'message':'两个冲突来源'}]:
            processing=deepcopy(task['processing']);processing['issues'].append(issue)
            self.assertEqual(rule_bindings.resolve(definition,processing,task['task'])['state'],'unknown')
        processing=deepcopy(task['processing']);processing['records'][0]['current_credit']='999'
        self.assertEqual(rule_bindings.resolve(definition,processing,task['task'])['state'],'unknown')

    def test_extra_metric_source_is_server_authorized_and_bound_to_company_period(self):
        task=self.task();url='/api/enterprise/input-tasks/'+task['id']+'/audit-metrics'
        self.assertEqual(self.client.post(url,json={'expected_revision':task['revision'],'audit_id':'absent'}).status_code,404)
        group=read_cases()[0];dataset=dataset_for(RULES['R-001'],group['baseline']);dataset.company.taxpayer_id=TAX
        entry={'id':'saved','org_id':'alpha','client_id':task['task']['company_id'],'dataset':dataset}
        with self.assertRaisesRegex(ValueError,'完整检查期间'):rule_bindings.audit_records(entry,task['task'])
        dataset.company.period='2026年1月'
        rows=rule_bindings.audit_records(entry,task['task']);self.assertTrue(rows)
        self.store.save_audit('saved',self.user,task['task']['company_id'],dataset,
                              engine.run([RULES['R-001']],dataset),{},'2026-10-16T00:00:00Z')
        result=self.client.post(url,json={'expected_revision':task['revision'],'audit_id':'saved'})
        self.assertEqual(result.status_code,200,result.text)
        linked=result.json()
        self.assertTrue(any(r['material_type']=='business_metric' for r in linked['processing']['records']))
        definition=conversion.compile_rule(RULES['R-001'])['bases']['营业收入']
        self.assertEqual(rule_bindings.resolve(definition,linked['processing'],linked['task'])['state'],'known')
        changed=deepcopy(linked['processing']);changed['records'][0]['current_credit']='100'
        self.assertEqual(rule_bindings.resolve(definition,changed,linked['task'])['state'],'unknown')

    def test_missing_account_codes_and_conflicting_sources_are_not_silently_summed(self):
        task=self.task();records=[r for r in task['processing']['records'] if r['material_type']=='account_balance']
        response=self.client.post('/api/enterprise/input-tasks/'+task['id']+'/metric-bindings',json={
            'expected_revision':task['revision'],'name':'营业收入','record_ids':[records[0]['id']],
            'field':'current_credit','aggregation':'sum','basis':'确认口径'})
        self.assertEqual(response.status_code,422)
        task=self.bind(task,'营业收入',[r['id'] for r in records],'current_credit')
        definition=conversion.compile_rule(RULES['R-001'])['bases']['营业收入']
        processing=deepcopy(task['processing']);processing['records'].append({'material_type':'business_metric','metric_name':'营业收入','value':'1'})
        self.assertEqual(rule_bindings.resolve(definition,processing,task['task'])['state'],'unknown')

    def test_corpus_import_requires_platform_role_and_current_manifest(self):
        result=self.client.get('/api/knowledge/corpus/status')
        self.assertEqual(result.status_code,200)
        self.assertEqual(self.client.post('/api/knowledge/corpus/import',json={'manifest_sha256':result.json()['manifest_sha256']}).status_code,403)
        user=self.store.create_user('platform','Rule-conversion-2026!','platform','platform_admin','platform')
        self.client.cookies.clear();self.client.cookies.set(module.COOKIE_NAME,self.store.authenticate(user['username'],'Rule-conversion-2026!')[1])
        self.assertEqual(self.client.post('/api/knowledge/corpus/import',json={'manifest_sha256':'0'*64}).status_code,409)


if __name__=='__main__':unittest.main()
