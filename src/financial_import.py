"""Explicit, source-addressed financial export layouts; never rewrite originals.

Only labeled columns are mapped. Amounts retain their source unit until evaluated;
missing cells, formulas without independently verified values, and ambiguous
periods are not replaced with zero or guessed from filenames.
"""
from copy import deepcopy
from calendar import monthrange
from decimal import Decimal
import re

from openpyxl.utils import get_column_letter

from . import config, loader, material_provenance as provenance
from .input_errors import InputError

VERSION = 'financial-export-v2'
FINANCIAL = {config.SHEET_ACCOUNTS, config.SHEET_INCOME,
             config.SHEET_BALANCE, config.SHEET_CASHFLOW, config.SHEET_DECLARATION}
ALIASES = {'科目余额': config.SHEET_ACCOUNTS, '科目发生额及余额表': config.SHEET_ACCOUNTS,
           '损益表': config.SHEET_INCOME}
UNITS = {'元': '1', '人民币元': '1', '千元': '1000', '人民币千元': '1000',
         '万元': '10000', '人民币万元': '10000', '百万元': '1000000',
         '人民币百万元': '1000000', '亿元': '100000000', '人民币亿元': '100000000'}
CURRENCY_LABELS = {'币种', '币别', '本位币', '记账本位币', '货币代码', '货币', 'currency'}
ITEM_ALIASES = {
    config.SHEET_BALANCE: {'资产总计': '资产总额', '负债合计': '负债总额',
        '所有者权益合计': '所有者权益', '股东权益合计': '所有者权益',
        '所有者权益(或股东权益)合计': '所有者权益', '所有者权益（或股东权益）合计': '所有者权益'},
    config.SHEET_CASHFLOW: {'经营活动产生的现金流量净额': '经营净额',
        '投资活动产生的现金流量净额': '投资净额', '筹资活动产生的现金流量净额': '筹资净额',
        '汇率变动对现金及现金等价物的影响': '汇率影响'},
}

# Exact source titles, not a translation of company names or arbitrary cells.
ANNUAL_TITLES = {
    prefix + title: (kind, basis)
    for prefix, basis in [('合併', 'consolidated'), ('母公司', 'parent')]
    for title, kind in [('損益表', config.SHEET_INCOME), ('財務狀況表', config.SHEET_BALANCE),
                        ('現金流量表', config.SHEET_CASHFLOW)]
}
ANNUAL_ITEMS = {
    config.SHEET_INCOME: {'收入': ('利润表.营业收入', '1'),
        '稅前利潤': ('利润表.利润总额', '1'), '所得稅': ('利润表.所得税费用', '-1'),
        '年度盈利': ('利润表.净利润', '1')},
    config.SHEET_BALANCE: {'總資產': ('资产负债表.资产总额', '1'),
        '總負債': ('资产负债表.负债总额', '1'), '總權益': ('资产负债表.所有者权益', '1')},
    config.SHEET_CASHFLOW: {'經營活動所產生的淨現金流入': ('现金流量表.经营净额', '1'),
        '投資活動所支付的淨現金流出': ('现金流量表.投资净额', '1'),
        '融資活動所產生的淨現金流出': ('现金流量表.筹资净额', '1'),
        '外幣匯率變動的影響': ('现金流量表.汇率影响', '1'),
        '現金及現金等價物年初餘額': ('现金.期初现金及等价物', '1'),
        '現金及現金等價物年末餘額': ('现金.期末现金及等价物', '1')},
}


class MappingRequired(InputError):
    """Recoverable choice, distinct from conflicting or invalid source evidence."""


def reanalysis_required(document):
    """Check retained financial candidates without mutating frozen snapshots.

    Mapping reanalysis is the most recent local parser record when present;
    never fall back to an older successful upload record if it is incomplete.
    Version and actual adapter/period source hashes are both necessary.
    """
    model, mapping = document.get('import_model'), document.get('import_mapping')
    if model is None and mapping is None:
        return False
    if (not isinstance(model, dict) or not isinstance(mapping, dict) or
            model.get('version') != VERSION or mapping.get('version') != VERSION):
        return True
    extraction = document.get('extraction', {})
    if not isinstance(extraction, dict):
        return True
    record = extraction.get('mapping_reanalysis', extraction.get('local', {}))
    if not isinstance(record, dict) or record.get('status') != 'succeeded':
        return True
    program = record.get('program')
    saved = program.get('sources') if isinstance(program, dict) else None
    if not isinstance(saved, dict):
        return True
    current_program = provenance.program('local')
    current = current_program['sources']
    dependencies = program.get('dependencies', {})
    if (not isinstance(dependencies, dict) or not current_program['dependencies'].get('openpyxl') or
            dependencies.get('openpyxl') != current_program['dependencies']['openpyxl']):
        return True
    if document.get('kind') == 'xls' and any(not current_program['dependencies'].get(name) or
            dependencies.get(name) != current_program['dependencies'][name] for name in ('xlrd', 'olefile')):
        return True
    return any(not current.get(name) or saved.get(name) != current[name]
               for name in ('financial_import.py', 'periods.py', 'materials.py', 'loader.py',
                            'config.py', 'workbooks.py', 'material_review.py', 'legacy_workbooks.py'))


def options(value, sheets):
    if not isinstance(value, dict) or set(value) - set(sheets):
        raise InputError('财务映射须为原工作表到选项的映射，不能指定未知工作表。')
    for name, choice in value.items():
        if not isinstance(choice, dict) or set(choice) - {'unit', 'period', 'columns', 'excluded'}:
            raise InputError('财务映射含未知选项。')
        if 'excluded' in choice and type(choice['excluded']) is not bool:
            raise InputError('工作表排除选项须为布尔值。')
        if 'unit' in choice and (not isinstance(choice['unit'], str) or choice['unit'] not in UNITS):
            raise InputError('金额单位须为人民币元、千元、万元、百万元或亿元。')
        if 'period' in choice and (not isinstance(choice['period'], str) or len(choice['period']) > 64):
            raise InputError('材料实际期间须为不超过 64 字符的文本。')
        columns = choice.get('columns', {})
        if not isinstance(columns, dict) or len(columns) > 20 or any(
                not isinstance(k, str) or not isinstance(v, str) or not re.fullmatch(r'[A-Z]{1,2}', v)
                for k, v in columns.items()):
            raise InputError('金额栏次须选原表中的列字母。')
    return deepcopy(value)


def token(value):
    return re.sub(r'\s+', '', str(value or '')).replace('：', ':')


def table_kind(rows):
    """Identify a flat export from its source title/header, never its filename."""
    grid = [[token(value) for value in row] for row in rows[:10]]
    header = next((i for i, row in enumerate(grid) if any(v in {
        '项目', '资产', '科目编码', '科目代码', '科目编号'} for v in row)), len(grid))
    titles = {ANNUAL_TITLES[value][0] if value in ANNUAL_TITLES else ALIASES.get(value, value)
              for row in grid[:header] for value in row
              if value in FINANCIAL | ALIASES.keys() | ANNUAL_TITLES.keys()}
    account_headers = [row for row in grid if
        sum(v in {'科目编码', '科目代码', '科目编号'} for v in row) == 1 and
        sum(v in {'科目名称', '会计科目'} for v in row) == 1]
    if len(titles) > 1 or (account_headers and titles and titles != {config.SHEET_ACCOUNTS}):
        raise InputError('CSV/TSV 原文表类冲突，请分别导出，不按文件名选择其中一种。')
    return next(iter(titles), config.SHEET_ACCOUNTS if len(account_headers) == 1 else '')


def number(value, address):
    if isinstance(value, str):
        text = value.strip().replace('，', ',')
        if text in {'', '-', '—', '–'}:
            return None
        if text.startswith('=') or text.startswith('#'):
            raise InputError(f'{address}：公式或错误值不能作为已核实金额，请导出固定值。')
        if text.startswith('(') and text.endswith(')'):
            text = '-' + text[1:-1]
        if ',' in text and not re.fullmatch(r'-?\d{1,3}(?:,\d{3})+(?:\.\d+)?', text):
            raise InputError(f'{address}：金额分隔符不明确，请按小数点和三位千分位核对。')
        value = text.replace(',', '')
    return loader._number(value, address)


def _currency(value):
    value = token(value).replace('幣', '币')
    return 'CNY' if value.upper() in {'人民币', 'CNY', 'RMB', '人民币(CNY)', '人民币（CNY）',
                                    'CNY(人民币)', 'CNY（人民币）'} else value


def _money_metadata(text, title):
    """A scale is not a currency; explicit unsupported facts cannot be mapped away."""
    labels = r'(?:编制单位|单位名称|企业名称|币种|币别|记账本位币|本位币|货币代码|货币单位|金额单位|货币|单位|所属期|报告期|Currency)[:：]'
    fields = re.findall(r'(?<!编制)(金额单位|货币单位|币种|币别|记账本位币|本位币|货币代码|货币|单位|Currency)[:：]\s*(.*?)'
                        r'(?=' + labels + r'|\d{4}\s*(?:年|Q[1-4]|H[12]|[-/.]\d)|[\n,，;；]|$)', text, re.I)
    units, currencies, currency_declared = set(), set(), False
    for label, raw in fields:
        value = token(raw)
        currency_declared |= label.casefold() in CURRENCY_LABELS
        if not value:
            continue
        if label.casefold() in CURRENCY_LABELS or _currency(value) == 'CNY':
            currencies.add(_currency(value))
        else:
            value = re.sub(r'[（(](?:人民币|CNY|RMB)[）)]$', '', value, flags=re.I)
            units.add(value)
            if value.startswith('人民币'):
                currencies.add('CNY')
    if len(units) > 1:
        if len({UNITS.get(u, u) for u in units}) > 1:
            raise InputError(f'{title}：金额单位冲突，不能用用户指定值消除。')
    unit = sorted(units)[0] if units else ''
    return {'unit': unit, 'factor': UNITS.get(unit, ''), 'unit_supported': not unit or unit in UNITS,
            'currency': '、'.join(sorted(currencies)) or ('缺失' if currency_declared else ''),
            'currency_supported': currencies == {'CNY'} or (not currencies and not currency_declared)}


def _period_metadata(text, title):
    """Extract explicit intervals and keep isolated dates as endpoints, not months."""
    from .periods import parse_period, statement_date
    # Do not join adjacent fields/periods into fabricated dates (e.g. 月 + 2026).
    text = str(text).replace('：', ':')
    periods, unresolved = set(), []
    def add(value):
        try:
            p = parse_period(value, f'{title} 原文实际期间')
        except InputError:
            unresolved.append(value)
            return
        label = (f'{p.start.year}-{p.start.month:02d}' if p.group == 'month' else
                 f'{p.start.year}Q{(p.start.month - 1) // 3 + 1}' if p.group == 'quarter' else
                 f'{p.start.year}H{1 if p.start.month == 1 else 2}' if p.group == 'half' else str(p.start.year))
        periods.add(label)
    full_date = r'(?<!\d)\d{4}\s*(?:年|-|/|\.)\s*\d{1,2}\s*(?:月|-|/|\.)\s*\d{1,2}日?(?!\d)'
    def interval(match):
        first, last = statement_date(match[1]), statement_date(match[2])
        if not first or not last:
            raise InputError(f'{title}：材料期间日期无效。')
        add(first + '至' + last)
        return ' '
    text = re.sub('(' + full_date + r')\s*(?:至|~|—)\s*(' + full_date + ')', interval, text)
    dates = set()
    def endpoint(match):
        value = statement_date(match[0])
        if not value:
            raise InputError(f'{title}：报表时点无效。')
        dates.add(value)
        return ' '
    text = re.sub(full_date, endpoint, text)
    if len(dates) > 1:
        raise InputError(f'{title}：报表时点冲突。')
    def month_interval(match):
        year, first, last_year, last = match.groups()
        last_year = last_year or year
        try:
            end_day = monthrange(int(last_year), int(last))[1]
            add(f'{int(year):04d}-{int(first):02d}-01至{int(last_year):04d}-{int(last):02d}-{end_day}')
        except ValueError:
            unresolved.append(match[0])
        return ' '
    text = re.sub(r'(?<!\d)(\d{4})\s*年\s*(\d{1,2})\s*月?\s*(?:至|-|~|—)\s*'
                  r'(?:(\d{4})\s*年\s*)?(\d{1,2})\s*月', month_interval, text)
    patterns = [
        (r'(?<!\d)(\d{4})\s*年?\s*第?\s*([1-4一二三四])\s*季度',
         lambda m: m[1] + 'Q' + str('一二三四'.index(m[2]) + 1 if m[2] in '一二三四' else m[2])),
        (r'(?<!\d)(\d{4})\s*-?\s*Q([1-4])(?!\d)', lambda m: m[1] + 'Q' + m[2]),
        (r'(?<!\d)(\d{4})\s*年?\s*(上|下)半年', lambda m: m[1] + ('H1' if m[2] == '上' else 'H2')),
        (r'(?<!\d)(\d{4})\s*-?\s*H([12])(?!\d)', lambda m: m[1] + 'H' + m[2]),
        (r'(?<!\d)(\d{4})\s*(?:年度|年报|年末)', lambda m: m[1]),
        (r'(?<!\d)(\d{4})\s*年\s*(\d{1,2})\s*月', lambda m: m[1] + '-' + m[2]),
        (r'(?<!\d)(\d{4})\s*[-/.]\s*(\d{1,2})(?![\d./-])', lambda m: m[1] + '-' + m[2]),
    ]
    for pattern, convert in patterns:
        def extract(match):
            add(convert(match))
            return ' '
        text = re.sub(pattern, extract, text, flags=re.I)
    remaining = re.search(r'(?<!\d)\d{4}\s*(?:年\s*(?:第?\s*[一二三四\d].*?(?:季|月)|[上下]半年)|[- ]?\s*[QH]\d)', text, re.I)
    if remaining:
        unresolved.append(remaining[0])
    if len(periods) > 1:
        raise InputError(f'{title}：表头期间不唯一，不能选择一个期间代替。')
    return {'period': next(iter(periods), ''), 'as_of': next(iter(dates), ''),
            'period_unresolved': bool(unresolved), 'period_unresolved_text': '、'.join(unresolved)[:256]}


def _metadata(sheet, kind):
    annual = _annual_metadata(sheet, kind)
    if annual is not None:
        return annual
    grid = _headers(sheet)
    header = next((i for i, row in enumerate(grid) if any(v in {
        '项目', '资产', '科目编码', '科目代码', '科目编号'} for v in row)), len(grid))
    lines = []
    labels = CURRENCY_LABELS | {'货币单位', '金额单位', '单位', '编制单位', '企业名称', '单位名称'}
    for row in sheet.iter_rows(max_row=max(1, header)):
        cells, index = [str(c.value or '') for c in row], 0
        while index < len(cells):
            label = token(cells[index]).rstrip(':')
            if label.casefold() in labels:
                has_value_cell = index + 1 < len(cells)
                lines.append(label + ':' + (cells[index + 1] if has_value_cell else ''))
                index += 2 if has_value_cell else 1
            else:
                lines.append(cells[index])
                index += 1
    text = '\n'.join(lines)
    # Export/print/submission dates are not the represented accounting period.
    text = re.sub(r'(?:打印|导出|上传|填报|编制)日期[:：]?\s*\d{4}(?:\s*年\s*\d{1,2}\s*月'
                  r'(?:\s*\d{1,2}\s*日)?|[-/.]\d{1,2}(?:[-/.]\d{1,2})?)', '', text)
    names = re.findall(r'(?:编制单位|单位名称|企业名称)[:：]([^\s]+)', text)
    if not names:
        title = token(sheet.cell(1, 1).value)
        names = [title[:-len(sheet.title)]] if title.endswith(sheet.title) and title != sheet.title else []
    names = [re.split(r'\d{4}年|单位[:：]', n)[0].strip() for n in names]
    if len(set(names)) > 1:
        raise InputError(f'{sheet.title}：表头企业名称冲突。')
    return {'name': names[0] if len(set(names)) == 1 else '',
            **_period_metadata(text, sheet.title), **_money_metadata(text, sheet.title)}


def _headers(sheet):
    """Expand genuine merged headers only, never forward-fill arbitrary body cells."""
    grid = [[token(c.value) for c in row] for row in sheet.iter_rows(max_row=min(10, sheet.max_row))]
    for merged in sheet.merged_cells.ranges:
        if merged.max_row > len(grid):
            continue
        value = grid[merged.min_row - 1][merged.min_col - 1]
        for r in range(merged.min_row - 1, merged.max_row):
            for c in range(merged.min_col - 1, merged.max_col):
                grid[r][c] = value
    return grid


def _annual_metadata(sheet, kind):
    """Explicit year-column statements with a source end-date and note column."""
    from datetime import date
    grid = _headers(sheet)
    source_grid = [[token(c.value) for c in row] for row in sheet.iter_rows(max_row=min(10, sheet.max_row))]
    # Merged title/date/unit cells are one original assertion, not multiple
    # conflicting assertions manufactured by header expansion.
    titles = [(r, c, v) for r, row in enumerate(source_grid) for c, v in enumerate(row) if v in ANNUAL_TITLES]
    if not titles:
        return None
    if len(titles) != 1 or ANNUAL_TITLES[titles[0][2]][0] != kind:
        raise InputError(f'{sheet.title}：年度报表标题或主体口径冲突。')
    title_row, item_column, title = titles[0]
    digit = dict(zip('零〇一二三四五六七八九', '00123456789'))
    def integer(value):
        if value.isascii() and value.isdigit():
            return int(value)
        if '十' not in value:
            return int(''.join(digit[v] for v in value))
        a, b = value.split('十')
        return (int(digit[a]) if a else 1) * 10 + (int(digit[b]) if b else 0)
    dates = []
    for row in source_grid[title_row + 1:]:
        for value in row:
            match = re.fullmatch(r'([0-9]{4}|[零〇一二三四五六七八九]{4})年'
                                 r'([0-9]{1,2}|[一二三四五六七八九十]{1,3})月'
                                 r'([0-9]{1,2}|[一二三四五六七八九十]{1,3})日(止年度)?', value)
            if match:
                try:
                    dates.append((date(*(integer(match[i]) for i in (1, 2, 3))), bool(match[4])))
                except (ValueError, KeyError):
                    raise InputError(f'{sheet.title}：原文年度报表日期无效。') from None
    if len(dates) != 1:
        raise InputError(f'{sheet.title}：缺少唯一的原文年度报表日期，不以年份列或文件名推测。')
    end, annual = dates[0]
    if (end.month, end.day) != (12, 31) or (kind != config.SHEET_BALANCE and not annual):
        raise InputError(f'{sheet.title}：此年度栏次适配仅支持原文明示的完整自然年度，不把时点当发生额期间。')
    headers = [(r + 1, [(c, v) for c, v in enumerate(row) if re.fullmatch(r'[0-9]{4}年', v)])
               for r, row in enumerate(grid) if any(re.fullmatch(r'[0-9]{4}年', v) for v in row)]
    if len(headers) != 1:
        raise InputError(f'{sheet.title}：年度金额表头不唯一。')
    header, columns = headers[0]
    years = [v for _, v in columns]
    if len(years) != len(set(years)) or years.count(f'{end.year}年') != 1 or any(c <= item_column for c, _ in columns):
        raise InputError(f'{sheet.title}：年份列重复或与原文年度不对应。')
    heading = grid[header - 1]
    if heading[item_column] not in {'', '項目', '项目'} or not any(v in {'附註', '附注'} for v in heading):
        raise InputError(f'{sheet.title}：项目、附注和年份列结构不受支持。')
    money_lines = []
    for row in source_grid[:header - 1]:
        for index, value in enumerate(row):
            simplified = value.translate(str.maketrans('單幣萬圓數種', '单币万元数种'))
            match = re.fullmatch(r'[（(]单位:(.*?)(?:[,，]每股数除外)?[）)]', simplified)
            if match:
                money_lines.append('单位:' + match[1])
            elif simplified.rstrip(':').casefold() in CURRENCY_LABELS | {'金额单位', '单位'}:
                money_lines.append(simplified.rstrip(':') + ':' + (row[index + 1] if index + 1 < len(row) else ''))
            else:
                money_lines.append(simplified)
    money = _money_metadata('\n'.join(money_lines), sheet.title)
    # Unlabeled top-line company is part of this explicit layout, kept verbatim.
    names = [str(sheet.cell(r + 1, item_column + 1).value).strip() for r in range(title_row)
             if token(sheet.cell(r + 1, item_column + 1).value)]
    if len(names) != 1:
        raise InputError(f'{sheet.title}：年度报表须保留唯一的原文企业名称。')
    return {'name': names[0], 'period': str(end.year) if annual else '', 'as_of': end.isoformat(),
            'period_unresolved': False, 'period_unresolved_text': '', **money,
            'statement_basis': ANNUAL_TITLES[title][1],
            'annual_columns': {'header': header, 'item': item_column,
                               'years': {get_column_letter(c + 1): v for c, v in columns},
                               'current': next(c for c, v in columns if v == f'{end.year}年')}}


def statement_basis(document):
    """One scope contract for PDF, native annual exports and unqualified tables."""
    model = document.get('import_model')
    model = model if isinstance(model, dict) else {}
    basis = document.get('pdf_statement_basis') or model.get('statement_basis')
    if basis:
        return basis
    names = {row.get('name', '').split('.', 1)[0] for row in document.get('rows', [])}
    sheets = {s.get('kind') for s in model.get('sheets', []) if not s.get('excluded')}
    return 'unspecified' if (names | sheets) & {config.SHEET_INCOME, config.SHEET_BALANCE, config.SHEET_CASHFLOW} else None


def _address(sheet, row, column):
    return f'{sheet.title}!{get_column_letter(column + 1)}{row}'


def _currency_columns(sheet, header_rows):
    grid, header = _headers(sheet), max(header_rows)
    columns = {c for r in header_rows for c, value in enumerate(grid[r - 1])
               if value.translate(str.maketrans('幣種別貨', '币种别货')).casefold() in
               {'币种', '币别', '货币', '货币代码', 'currency'}}
    return [(c, [area for area in sheet.merged_cells.ranges
                 if area.min_col == area.max_col == c + 1 and area.min_row > header]) for c in sorted(columns)]


def _check_row_currency(sheet, row, columns):
    for column, merged in columns:
        cell = sheet.cell(row, column + 1)
        if cell.value is None:
            area = next((area for area in merged if area.min_row <= row <= area.max_row), None)
            if area:
                cell = sheet.cell(area.min_row, area.min_col)
        if _currency(cell.value) != 'CNY':
            raise MappingRequired(f'{sheet.title}!{cell.coordinate}：原金额行币种「{cell.value or "缺失"}」'
                                  '未证明人民币口径，不按表头单位换算；请排除此表或提供人民币原表。')


def _cell(model, sheet, row, column, label, meta):
    address = _address(sheet, row, column)
    cell = sheet.cell(row, column + 1)
    formula = model.get('formula_values', {}).get(address)
    value = None if formula else number(cell.value, address)
    model['cells'][address] = {'id': address, 'table': sheet.title, 'label': label,
                               'value': str(value) if value is not None else None,
                               'unit': meta['unit']}
    if formula:
        expression = cell.value if isinstance(cell.value, str) else '=' + (formula['formula'] or '（原 XML 未提供表达式文本）')
        cache, status = None, 'missing'
        raw = formula['cached_raw']
        if raw is not None and raw.strip():
            status = 'error' if formula['cached_type'] == 'e' else 'non_numeric'
            if formula['cached_type'] == 'n':
                try:
                    numeric = loader._number(raw, address)
                    if numeric is not None and abs(numeric) <= Decimal('1e18') and numeric.as_tuple().exponent >= -12:
                        cache, status = str(numeric), 'candidate'
                except InputError:
                    pass
        model['cells'][address].update(formula=expression, formula_type=formula['formula_type'],
            formula_attributes=formula['formula_attributes'],
            cached_value=cache, cached_raw=raw, cache_status=status)
    return address


def _accounts(sheet, model, meta):
    grid = _headers(sheet)
    candidates = []
    for index, row in enumerate(grid):
        codes = [c for c, v in enumerate(row) if v in {'科目编码', '科目代码', '科目编号'}]
        names = [c for c, v in enumerate(row) if v in {'科目名称', '会计科目'}]
        if len(codes) != 1 or len(names) != 1:
            continue
        mapping = {}
        consumed = index
        for c, value in enumerate(row):
            split = index + 1 < len(grid) and grid[index + 1][c] in {'借方', '贷方'}
            if value in {'本期借方', '本期贷方'} or (value in {'期初余额', '期末余额'} and not split):
                mapping.setdefault(value, []).append(c)
            if split:
                side = grid[index + 1][c]
                if value in {'本期发生额', '本期发生', '本月发生额'}:
                    mapping.setdefault('本期' + side, []).append(c)
                    consumed = index + 1
                elif value in {'期初余额', '期末余额'}:
                    mapping.setdefault(value + side, []).append(c)
                    consumed = index + 1
        if all(mapping.get(k) for k in ('本期借方', '本期贷方')):
            candidates.append((consumed + 1, codes[0], names[0], mapping, index + 1))
    if len(candidates) != 1:
        raise InputError(f'{sheet.title}：无法唯一识别科目、名称和本期借贷列，不使用累计列替代。')
    header, code_column, name_column, mapping, first_header = candidates[0]
    choice = meta['options'].get('columns', {})
    available = {'account:' + k: [{'value': get_column_letter(c + 1), 'label': k} for c in v]
                 for k, v in mapping.items()}
    model['sheets'][-1]['column_choices'] = available
    if set(choice) - set(available):
        raise InputError('科目映射包含原表不存在的金额字段。')
    missing = []
    for title, columns in mapping.items():
        selected = choice.get('account:' + title)
        if selected:
            valid = [c for c in columns if get_column_letter(c + 1) == selected]
            if not valid:
                raise InputError(f'{sheet.title}：所选 {title} 列不属于原始表头。')
            mapping[title] = valid
        elif len(columns) > 1:
            missing.append(title)
    if missing:
        raise MappingRequired('金额栏次重复，请明确选择：' + '、'.join(missing))
    model['sheets'][-1]['columns'] = {k: get_column_letter(v[0] + 1) for k, v in mapping.items()}
    if not meta['unit'] or not meta['period']:
        raise MappingRequired('金额单位或实际期间缺失，请对照原件明确选择；用户指定不是原文证明。')
    currency_columns = _currency_columns(sheet, range(first_header, header + 1))
    seen = set()
    for r in range(header + 1, sheet.max_row + 1):
        raw_code = sheet.cell(r, code_column + 1).value
        if raw_code is None or token(raw_code) in {'合计', '总计'}:
            continue
        _check_row_currency(sheet, r, currency_columns)
        code = str(int(raw_code)) if isinstance(raw_code, float) and raw_code.is_integer() else str(raw_code).strip()
        if not re.fullmatch(r'\d+(?:[.\-]\d+)*', code) or code in seen:
            raise InputError(f'{_address(sheet, r, code_column)}：科目编码无效或重复，不能静默汇总。')
        seen.add(code)
        name = str(sheet.cell(r, name_column + 1).value or '').strip()
        record = {'code': code, 'name': name, 'factor': meta['factor']}
        for field, title in [('opening', '期初余额'), ('debit', '本期借方'),
                             ('credit', '本期贷方'), ('closing', '期末余额')]:
            # Split balances are not netted when one side is absent: blank is unknown.
            columns = mapping.get(title, []) if title.startswith('本期') else (
                mapping.get(title + '借方', []) + mapping.get(title + '贷方', []))
            if not columns:
                columns = mapping.get(title, [])
            refs = [_cell(model, sheet, r, c, f'{code} {name} · {title}', meta) for c in columns]
            record[field] = refs
        model['accounts'].append(record)
    if not seen:
        raise InputError(f'{sheet.title}：没有可识别的科目数据行。')
    model['warnings'].append('科目编码保持原样；缺少映射科目不视为零。上下级科目不自动合并。')


def _item(value):
    text = token(value)
    return re.sub(r'^(?:[一二三四五六七八九十]+[、.．]|(?:加|减):)', '', text)


def _annual_statement(sheet, model, meta, kind):
    layout = meta['annual_columns']
    header, label_column, value_column = layout['header'], layout['item'], layout['current']
    group_id = f'{header}:{get_column_letter(label_column + 1)}'
    model['sheets'][-1]['column_choices'] = {
        group_id: [{'value': c, 'label': year} for c, year in layout['years'].items()]}
    choices = meta['options'].get('columns', {})
    if set(choices) - {group_id} or choices.get(group_id, get_column_letter(value_column + 1)) != get_column_letter(value_column + 1):
        raise InputError(f'{sheet.title}：所选年份列与原文报表期间不一致，比较年不能代替本期。')
    model['sheets'][-1]['columns'] = [{'header_row': header, 'item': get_column_letter(label_column + 1),
        'amount': get_column_letter(value_column + 1), 'heading': layout['years'][get_column_letter(value_column + 1)]}]
    if not meta['unit']:
        raise MappingRequired('原文年度报表金额单位缺失，请核实后选择，不自动指定人民币元。')
    currency_columns = _currency_columns(sheet, {header})
    seen = {}
    for r in range(header + 1, sheet.max_row + 1):
        label = token(sheet.cell(r, label_column + 1).value)
        mapped = ANNUAL_ITEMS[kind].get(label)
        if not mapped or mapped[0] not in model['keys']:
            if label:
                model['unmapped'].append({'sheet': sheet.title, 'row': r, 'item': label,
                                          'address': _address(sheet, r, value_column)})
            continue
        key, sign = mapped
        _check_row_currency(sheet, r, currency_columns)
        ref = _cell(model, sheet, r, value_column, label + ' · ' + layout['years'][get_column_letter(value_column + 1)], meta)
        original = model['cells'][ref]['value']
        if (label.endswith('淨現金流出') and original is not None and Decimal(original) > 0 or
                label.endswith('淨現金流入') and original is not None and Decimal(original) < 0):
            raise MappingRequired(f'{ref}：净流入/流出文字与原金额符号不一致，不自动翻转或当作已核实净额。')
        if key in seen:
            original_row, record = seen[key]
            intervening = {token(sheet.cell(n, label_column + 1).value) for n in range(original_row + 1, r)}
            # This explicit earnings allocation restates the total. Preserve
            # both addresses and require equality, never sum or choose one.
            if (label != '年度盈利' or len(record['refs']) != 1 or
                    not {'應佔盈利:', '本公司權益持有者', '非控制性權益'} <= intervening):
                raise InputError(f'{sheet.title}：年度项目重复，不静默选择或汇总。')
            record['refs'].append(ref)
            record['operation'] = 'equal'
            record['detail'] += '；分配后重列总额须与首次总额一致，两处均保留来源。'
        else:
            factor = str(Decimal(meta['factor']) * Decimal(sign))
            record = {'name': key, 'refs': [ref], 'factor': factor,
                'detail': f'原表项目 {label}；年度栏次 {layout["years"][get_column_letter(value_column + 1)]}；'
                          f'单位 {meta["unit"]}；原文终点 {meta["as_of"]}；主体口径 {meta["statement_basis"]}' +
                          ('；原表所得税按利润加项列示，映射所得税费用时取相反数。' if sign == '-1' else '')}
            model['rows'].append(record)
            seen[key] = (r, record)
    model['warnings'].append('年度报表只采用原文对应年份；比较年、每股指标及未映射成本明细不参与本期检测，不补造总成本。')


def _statement(sheet, model, meta, kind):
    if meta.get('annual_columns'):
        return _annual_statement(sheet, model, meta, kind)
    grid = _headers(sheet)
    groups = []
    choices = {}
    unresolved = []
    all_headings = {'本月数', '本月金额', '本期金额', '本期数', '金额', '本年累计数', '本年累计金额',
                    '本年金额', '本季数', '本季度金额', '本半年数', '本半年金额',
                    '期末数', '期末余额', '年末余额', '期末金额'}
    for index, row in enumerate(grid):
        # A consecutive repeated header is one group, not a second copy of its
        # data. Keep the final header's source row. Nonconsecutive duplicates
        # remain separate groups and still fail the duplicate-item check.
        if index + 1 < len(grid) and row == grid[index + 1]:
            continue
        labels = [c for c, value in enumerate(row) if value in {
            '项目', '资产', '负债及所有者权益', '负债和所有者权益', '负债及股东权益',
            '负债和所有者权益（或股东权益）', '负债和所有者权益(或股东权益)'}]
        for n, label in enumerate(labels):
            end = labels[n + 1] if n + 1 < len(labels) else len(row)
            if kind == config.SHEET_BALANCE:
                accepted = {'期末数', '期末余额', '年末余额', '期末金额'} | ({'本期金额'} if meta.get('template') else set())
            elif re.fullmatch(r'\d{4}-\d{2}', meta['period']):
                accepted = {'本月数', '本月金额', '本期金额', '本期数', '金额'}
            elif re.fullmatch(r'\d{4}', meta['period']):
                accepted = {'本年累计数', '本年累计金额', '本期金额', '本年金额', '本期数', '金额'}
            elif re.fullmatch(r'\d{4}Q[1-4]', meta['period']):
                accepted = {'本季数', '本季度金额', '本期金额', '本期数', '金额'}
            elif re.fullmatch(r'\d{4}H[12]', meta['period']):
                accepted = {'本半年数', '本半年金额', '本期金额', '本期数', '金额'}
            else:
                accepted = {'本期金额', '本期数', '金额'}
            group_id = f'{index + 1}:{get_column_letter(label + 1)}'
            choices[group_id] = [{'value': get_column_letter(c + 1), 'label': row[c]}
                                 for c in range(label + 1, end) if row[c] in all_headings]
            values = [c for c in range(label + 1, end) if row[c] in accepted]
            selected = meta['options'].get('columns', {}).get(group_id)
            if selected:
                values = [c for c in values if get_column_letter(c + 1) == selected]
                if not values:
                    raise InputError(f'{sheet.title}：所选金额列不属于本期口径，累计不能代替本月。')
            if len(values) != 1:
                unresolved.append(group_id)
            else:
                groups.append((index + 1, label, values[0], row[values[0]]))
    model['sheets'][-1]['column_choices'] = choices
    if set(meta['options'].get('columns', {})) - set(choices):
        raise InputError('财务报表映射包含不存在的表头分组。')
    if unresolved:
        raise MappingRequired('无法唯一确定本期金额列与实际期间，请选择期间和金额栏次。')
    if not groups:
        raise InputError(f'{sheet.title}：没有可核实的项目/金额表头。')
    seen = set()
    model['sheets'][-1]['columns'] = [
        {'header_row': r, 'item': get_column_letter(c + 1), 'amount': get_column_letter(v + 1), 'heading': h}
        for r, c, v, h in groups]
    if not meta['unit'] or (not meta['period'] and (kind != config.SHEET_BALANCE or not meta['as_of'])):
        raise MappingRequired('金额单位或实际期间缺失，请对照原件明确选择；用户指定不是原文证明。')
    currency_columns = _currency_columns(sheet, {group[0] for group in groups})
    for header, label_column, value_column, heading in groups:
        for r in range(header + 1, sheet.max_row + 1):
            label = _item(sheet.cell(r, label_column + 1).value)
            if not label:
                continue
            mapped_label = ITEM_ALIASES.get(kind, {}).get(label, label)
            key = f'{kind}.{mapped_label}'
            supported = label in config.DECLARATION_ITEMS.values() if kind == config.SHEET_DECLARATION else key in model['keys']
            # Unknown rows remain visible in the mapping, never masquerade as rule inputs.
            if not supported:
                model['unmapped'].append({'sheet': sheet.title, 'row': r, 'item': label,
                                          'address': _address(sheet, r, value_column)})
                continue
            _check_row_currency(sheet, r, currency_columns)
            if key in seen:
                raise InputError(f'{sheet.title}：项目重复「{label}」，不静默选择首行。')
            seen.add(key)
            ref = _cell(model, sheet, r, value_column, label + ' · ' + heading, meta)
            target = model['declarations'] if kind == config.SHEET_DECLARATION else model['rows']
            target.append({'name': label if kind == config.SHEET_DECLARATION else key,
                                  'refs': [ref], 'factor': meta['factor'],
                                  'detail': f'原表项目 {label}；栏次 {heading}；单位 {meta["unit"]}；期间 {meta["period"] or "未提供"}'})


def evaluate(document, edits=None):
    """Recalculate only source-addressed candidates, without reopening or altering originals."""
    candidate = deepcopy(document)
    model = document['import_model']
    if model['pending']:
        # Neither attribution nor numeric edits can promote unresolved source
        # layouts or mixed scopes into usable financial inputs.
        candidate.update(deepcopy(model['base']))
        return candidate
    values = {key: number((edits or {}).get(key, item['value']), key)
              for key, item in model['cells'].items()}

    def amount(refs, factor, operation='net'):
        if not refs or any(values[ref] is None for ref in refs):
            return None
        if operation == 'equal' and len({values[ref] for ref in refs}) != 1:
            raise InputError('原表重复列示的年度总额不一致，不能选择其中一处或分别修正为不同金额。')
        result = values[refs[0]] if operation == 'equal' else sum((values[ref] for ref in refs), Decimal(0)) if operation == 'sum' else (
            values[refs[0]] if len(refs) == 1 else values[refs[0]] - values[refs[1]])
        return str(result * Decimal(factor))

    def source(refs, factor):
        notes = []
        for ref in refs:
            cell = model['cells'][ref]
            if cell.get('formula'):
                notes.append(f'{ref}（原公式 {cell["formula"]}；文件缓存 {cell.get("cached_raw") or "缺失"}，未验证新鲜度；'
                             f'用户核对值 {(edits or {}).get(ref, "留空")}；不执行公式、未改原件）')
            else:
                notes.append(f'{ref}（原文 {cell["value"]}；用户修正为 {edits[ref]}，未改原件）'
                             if ref in (edits or {}) else ref)
        return '、'.join(notes) + f'；原单位换算 ×{factor}'

    base = model['base']
    candidate['accounts'] = deepcopy(base['accounts']) + [{**{k: r[k] for k in ('code', 'name')},
        **{k: amount(r[k], r['factor']) for k in ('opening', 'debit', 'credit', 'closing')}}
        for r in model['accounts']]
    candidate['account_cell_sources'] = {**base.get('account_cell_sources', {}), **{r['code']: {k: source(r[k], r['factor'])
        for k in ('opening', 'debit', 'credit', 'closing')} for r in model['accounts']}
    }
    candidate['declarations'] = deepcopy(base['declarations'])
    candidate['declaration_cell_sources'] = deepcopy(base.get('declaration_cell_sources', {}))
    for row in model['declarations']:
        value = amount(row['refs'], row['factor'])
        if value is not None:
            if row['name'] in candidate['declarations']:
                raise InputError('标准与导出申报表项目重复，不能静默覆盖。')
            candidate['declarations'][row['name']] = value
            candidate['declaration_cell_sources'][row['name']] = source(row['refs'], row['factor'])
    candidate['rows'] = deepcopy(base['rows'])
    for row in model['rows']:
        value = amount(row['refs'], row['factor'], row.get('operation', 'net'))
        if value is not None:
            candidate['rows'].append({'name': row['name'], 'value': value,
                                      'source': source(row['refs'], row['factor']), 'detail': row['detail']})
    return candidate


def read(workbook, company, keys, import_options=None, formula_values=None):
    """Return None for unchanged standard contracts; exported grids use explicit layouts."""
    known = {value for name, value in vars(config).items() if name.startswith('SHEET_') and isinstance(value, str)}
    kinds = {}
    for sheet in workbook:
        title = token(sheet.title)
        if title in FINANCIAL | ALIASES.keys():
            kinds[sheet.title] = ALIASES.get(title, title)
        elif title not in known:
            # Software-specific worksheet names are not business evidence;
            # require an explicit source title or unique account header instead.
            kind = table_kind([[c.value for c in row] for row in sheet.iter_rows(max_row=min(10, sheet.max_row))])
            if kind:
                kinds[sheet.title] = kind
    sheets = [s for s in workbook if s.title in kinds]
    if not sheets:
        return None
    standard = {config.SHEET_ACCOUNTS: config.COL_ACCOUNTS,
                config.SHEET_DECLARATION: ['项目', '金额'],
                **{name: config.COL_STATEMENT for name in FINANCIAL - {config.SHEET_ACCOUNTS, config.SHEET_DECLARATION}}}
    templates = {s.title for s in sheets if s.title in standard and
                 [c.value for c in s[1]][:len(standard[s.title])] == standard[s.title]}
    exports = [s for s in sheets if s.title not in templates or
               any(address.rsplit('!', 1)[0] == s.title for address in (formula_values or {}))]
    if not exports:
        if import_options:
            raise InputError('标准模板不能接受财务导出映射覆盖。')
        return None
    selected = options(import_options or {}, [s.title for s in exports])
    model = {'version': VERSION, 'cells': {}, 'accounts': [], 'rows': [], 'declarations': [],
             'keys': sorted(keys or []), 'warnings': [], 'sheets': [], 'unmapped': [], 'pending': [],
             'formula_values': formula_values or {}}
    companies = {company['name']} if company.get('name') else set()
    source_periods = {company['period']} if company.get('period') else set()
    for sheet in exports:
        kind = kinds[sheet.title]
        meta = _metadata(sheet, kind)
        meta['template'] = sheet.title in templates
        if meta['template']:
            meta.update(unit='元', factor='1')  # Explicit standard-template contract.
        origin = deepcopy(meta)
        choice = selected.get(sheet.title, {})
        excluded = choice.get('excluded', False)
        if not meta['period'] and not meta['as_of'] and not meta['period_unresolved'] and company.get('period'):
            meta['period'] = company['period']
        if choice.get('unit'):
            if meta['unit'] and UNITS.get(meta['unit']) != UNITS[choice['unit']]:
                raise InputError(f'{sheet.title}：原文金额单位不能被映射覆盖。')
            meta['unit'], meta['factor'] = choice['unit'], UNITS[choice['unit']]
        if choice.get('period'):
            from .periods import parse_period
            entered = parse_period(choice['period'], '用户指定材料实际期间')
            if meta['period']:
                original = parse_period(meta['period'], '原文实际期间')
                if (entered.start, entered.end) != (original.start, original.end):
                    raise InputError(f'{sheet.title}：原文实际期间不能被映射覆盖。')
            if meta['as_of'] and entered.end.isoformat() != meta['as_of']:
                raise InputError(f'{sheet.title}：原文报表时点不能被映射覆盖。')
            meta['period'] = entered.label
        meta['options'] = choice
        model['sheets'].append({'sheet': sheet.title, 'kind': kind, **{k: v for k, v in meta.items() if k != 'options'},
            'source_unit': origin['unit'], 'source_period': origin['period'], 'excluded': excluded,
            'unit_origin': 'template' if meta['template'] else 'source' if origin['unit'] else 'user' if choice.get('unit') else 'missing',
            'period_origin': 'source' if origin['period'] else 'user' if choice.get('period') else
                'source_company' if company.get('period') and not meta['as_of'] else 'missing'})
        if excluded:
            continue
        if meta['period_unresolved']:
            model['pending'].append({'sheet': sheet.title, 'message':
                f'原文期间暂不支持或不完整：{meta["period_unresolved_text"]}。不能手填其他期间覆盖，须排除此表或提供明确期间原表。'})
            continue
        if not meta['currency_supported'] or not meta['unit_supported']:
            model['pending'].append({'sheet': sheet.title, 'message':
                f'原文币种/金额单位不支持：{origin["currency"] or "未声明币种"}、{origin["unit"] or "未声明单位"}。'
                '财务导出仅按人民币口径检测；不能手填单位覆盖，须排除此表或提供人民币原表。'})
            continue
        if meta['name']:
            companies.add(meta['name'])
        if meta['period']:
            source_periods.add(meta['period'])
        try:
            if kind == config.SHEET_ACCOUNTS:
                _accounts(sheet, model, meta)
            else:
                _statement(sheet, model, meta, kind)
        except MappingRequired as exc:
            model['pending'].append({'sheet': sheet.title, 'message': str(exc)})
    active = [s for s in model['sheets'] if not s['excluded']]
    def pending(message):
        model['pending'].append({'sheet': '工作簿', 'message': message})

    if len(companies) > 1:
        pending('企业名称冲突：' + '、'.join(sorted(companies)) + '。请逐工作表排除其他企业材料，不能手填覆盖。')
    bases = {s.get('statement_basis', 'unspecified') for s in active
             if s['kind'] in {config.SHEET_INCOME, config.SHEET_BALANCE, config.SHEET_CASHFLOW}}
    if any(s.title in templates and s.title not in {s.title for s in exports} and
           kinds[s.title] in {config.SHEET_INCOME, config.SHEET_BALANCE, config.SHEET_CASHFLOW} for s in sheets):
        bases.add('unspecified')  # Standard financial sheets still have their own unresolved scope.
    if len(bases) > 1:
        pending('母公司/合并或未明确主体口径的财务报表不能混用，请排除不同口径工作表后重检。')
    model['statement_basis'] = next(iter(bases), '') if len(bases) <= 1 else ''
    from .periods import parse_period
    periods = {(p.start, p.end): p.label for p in
               (parse_period(value, '材料所属期') for value in sorted(source_periods))}
    if len(periods) > 1:
        pending('实际期间冲突：' + '、'.join(sorted(periods.values())) + '。请逐工作表排除其他期间，不混入本期检测。')
    period = next(iter(periods.values()), '') if len(periods) <= 1 else company.get('period', '')
    if period:
        end = parse_period(period, '材料所属期').end.isoformat()
        if any(s['as_of'] and s['as_of'] != end for s in active):
            pending('报表时点与发生额期间终点不一致，请排除不同期工作表后重检。')
    for kind in sorted(FINANCIAL):
        matches = [s['sheet'] for s in active if s['kind'] == kind]
        if len(matches) > 1:
            pending('相同表类存在多个工作表：' + '、'.join(matches) + '。请核对并保留本次使用的一份，不静默覆盖或汇总。')
    model['company'] = {**company, 'name': next(iter(companies), '') if len(companies) <= 1 else company.get('name', ''),
                        'period': period}
    model['ignored_sheets'] = [s.title for s in workbook if token(s.title) not in known and s.title not in kinds]
    # The old chart is explicit and name-checked, not a renaming of original codes.
    accounts = {r['code']: r for r in model['accounts']}
    for name, side, codes, names in [
        ('营业收入', 'credit', ('5101', '5102'), ('主营业务收入', '其他业务收入')),
        ('营业成本', 'debit', ('5401', '5402'), ('主营业务成本', '其他业务支出'))]:
        if all(code in accounts and token(accounts[code]['name']) == expected for code, expected in zip(codes, names)):
            selected = [accounts[c] for c in codes]
            if len({a['factor'] for a in selected}) != 1:
                raise InputError('映射科目的金额单位不一致。')
            model['rows'].append({'name': name, 'refs': [a[side][0] for a in selected],
                'operation': 'sum', 'factor': selected[0]['factor'],
                'detail': '原始科目显式映射：' + ' + '.join(codes) + ' 本期' + ('贷方' if side == 'credit' else '借方')})
    model.pop('formula_values')
    return model
