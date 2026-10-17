"""Business acceptance before tax rules: core inputs, evidence, gaps and tenant ACL."""

import base64
from copy import deepcopy
from io import BytesIO
import json
import os
from pathlib import Path
import sqlite3
import struct
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import Workbook

from src import core_materials, core_mapping
from src.input_errors import InputError
from webapp import app as module, material_batches
from webapp.storage import Store

TAX = "91310000123456789A"
OTHER = "91310000987654321B"


def book(
    kind,
    *,
    start="2026-01-01",
    end="2026-01-31",
    tax=TAX,
    direction="销项",
    blank=False,
    form="增值税纳税申报表主表",
    alias=None,
):
    wb = Workbook()
    sheet = wb.active
    sheet.title = (
        "科目余额表" if kind == "account_balance" else form if kind == "vat_declaration" else direction + "发票明细"
    )
    sheet.append([sheet.title])
    sheet.append(["纳税人识别号：" + tax])
    sheet.append([("税款所属期" if kind == "vat_declaration" else "统计期间") + f"：{start} 至 {end}"])
    sheet.append(["金额单位：元"])
    if kind == "account_balance":
        sheet.append(["科目编码", "科目名称", "期初余额", "本期借方", "本期贷方", "期末余额"])
        sheet.append(["1001", "库存现金", "0", "100.0123", None if blank else "20", "80.0123"])
    elif kind == "vat_declaration":
        sheet.append(["项目", "栏次", alias or "本期数"])
        sheet.append(["应税销售额", "1", "100.0123"])
        sheet.append(["销项税额", "11", "13.0016"])
    else:
        sheet.append(["票号", "开具日期", "销方税号", "购方税号", "金额", "税额", "红字状态", "作废状态", "进销项"])
        sheet.append(
            [
                "12345678901234567890",
                start,
                TAX if direction == "销项" else OTHER,
                OTHER if direction == "销项" else TAX,
                "100.0123",
                "13.0016",
                "蓝字",
                "正常",
                direction,
            ]
        )
    raw = BytesIO()
    wb.save(raw)
    wb.close()
    return raw.getvalue()


def mapper():
    db = sqlite3.connect(":memory:")
    db.row_factory = sqlite3.Row
    core_mapping.migrate(db)
    return db, core_mapping.Mapper(db, "org", "task", "sha")


def pdf_table(kind="vat_declaration"):
    """Constructed PDF fixture with a ruled table and fixed ASCII field names."""
    contents = ["0.5 w"]
    columns = {"vat_declaration":[40, 200, 280, 430], "account_balance":[10, 140, 230, 320, 410, 500],
               "invoice_detail":[10,180,280,400,490]}[kind]
    for x in columns:
        contents.append(f"{x} 600 m {x} 690 l S")
    for y in [600, 630, 660, 690]:
        contents.append(f"{columns[0]} {y} m {columns[-1]} {y} l S")
    rows = [
        (675, ["item_label", "item_key", "declared_value"]),
        (645, ["Sales", "1", "100.0123"]),
        (615, ["Tax", "11", "13.0016"]),
    ] if kind == "vat_declaration" else [
        (675, ["account_name", "opening_balance", "current_debit", "current_credit", "closing_balance"]),
        (645, ["Cash", "10", "100", "20", "90"]),
        (615, ["Bank", "5", "20", "10", "15"]),
    ] if kind == "account_balance" else [
        (675, ["invoice_no", "issue_date", "amount_excluding_tax", "tax_amount"]),
        (645, ["12345678901234567890", "2026-01-15", "100", "13"]),
        (615, ["22345678901234567890", "2026-01-16", "200", "26"]),
    ]
    for y, cols in rows:
        for x, value in zip([n+5 for n in columns], cols):
            contents.append(f"BT /F1 8 Tf {x} {y} Td ({value}) Tj ET")
    content = "\n".join(contents).encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 500 750] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        f"<< /Length {len(content)} >>\nstream\n".encode() + content + b"\nendstream",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for i, obj in enumerate(objects, 1):
        offsets.append(len(output))
        output.extend(f"{i} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = len(output)
    output.extend(b"xref\n0 6\n0000000000 65535 f \n")
    for n in offsets[1:]:
        output.extend(f"{n:010} 00000 n \n".encode())
    output.extend(f"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode())
    return bytes(output)


def invoice_pdf(*, number="12345678901234567890", total="113.00", other_number=None):
    """Chinese ticket with side-by-side parties and item amounts unlike totals."""
    lines = [(40, 710, "电子发票（普通发票）"), (40, 680, "发票号码："+number),
             (330, 680, "开票日期：2026年01月15日"),
             (40, 640, "购买方信息"), (330, 640, "销售方信息"),
             (40, 620, "名称：购方企业"), (330, 620, "名称：销方企业"),
             (40, 600, "统一社会信用代码："+OTHER), (330, 600, "统一社会信用代码："+TAX),
             (40, 550, "货物名称 金额 税额"), (40, 530, "商品甲 60.00 7.80"), (40, 510, "商品乙 40.00 5.20"),
             (40, 470, "合 计 ￥100.00 ￥13.00"), (40, 440, "价税合计（小写）￥"+total)]
    if other_number:
        lines.append((40, 400, "发票号码："+other_number))
    return chinese_pdf(lines)


def chinese_pdf(lines, drawing=""):
    content = (drawing + "\n" + "\n".join(f"BT /F1 10 Tf {x} {y} Td <{value.encode('utf-16-be').hex()}> Tj ET" for x,y,value in lines)).encode()
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
               b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 640 750] /Resources << /Font << /F1 4 0 R >> >> /Contents 6 0 R >>",
               b"<< /Type /Font /Subtype /Type0 /BaseFont /STSong-Light /Encoding /UniGB-UCS2-H /DescendantFonts [5 0 R] >>",
               b"<< /Type /Font /Subtype /CIDFontType0 /BaseFont /STSong-Light /FontDescriptor 7 0 R /CIDSystemInfo << /Registry (Adobe) /Ordering (GB1) /Supplement 4 >> >>",
               f"<< /Length {len(content)} >>\nstream\n".encode()+content+b"\nendstream",
               b"<< /Type /FontDescriptor /FontName /STSong-Light /Flags 6 /FontBBox [-250 -250 1200 1000] /ItalicAngle 0 /Ascent 880 /Descent -120 /CapHeight 880 /StemV 80 >>"]
    output=bytearray(b"%PDF-1.4\n");offsets=[0]
    for i,obj in enumerate(objects,1):
        offsets.append(len(output));output.extend(f"{i} 0 obj\n".encode()+obj+b"\nendobj\n")
    xref=len(output);output.extend(f"xref\n0 {len(objects)+1}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        output.extend(f"{offset:010} 00000 n \n".encode())
    output.extend(f"trailer\n<< /Size {len(objects)+1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode())
    return bytes(output)


def grouped_business_pdf(kind, *, unknown=False):
    """Actual PDF cells with wrapped subheaders, as on native printed forms."""
    title = "科目余额表" if kind == "account_balance" else "增值税纳税申报表主表"
    lines = [(40,710,title), (40,690,"纳税人识别号："+TAX), (40,670,"金额单位：元"),
             (40,650,("会计期间" if kind=="account_balance" else "税款所属期")+"：2026-01-01 至 2026-01-31")]
    columns = [40,150,220,290,360,430,500,570] if kind=="account_balance" else [40,180,290,400]
    drawing = "0.5 w\n"+"\n".join(f"{x} 480 m {x} 630 l S" for x in columns)
    drawing += "\n"+"\n".join(f"40 {y} m {columns[-1]} {y} l S" for y in [480,530,580,630])
    if kind=="account_balance":
        parents=["科目名称","期初余额","","本期发生额","","期末余额",""]
        children=["","借\n方","贷\n方","借\n方","贷\n方","借\n方","贷\n方"]
        values=["现金","10","0","100","20","90","0"]
    else:
        parents=["项目","一般项目",""]
        children=["","本期\n数","本年\n累计数"]
        values=["应税销售额","100","900"]
    if unknown:
        if kind=="account_balance":
            parents[3]="发生额来源"
        else:
            children=["","申报\n值","累计\n值"]
    for x,parent,child,value in zip(columns,parents,children,values):
        if parent:lines.append((x+5,610,parent))
        for offset,part in enumerate(child.split("\n")):
            if part:lines.append((x+5,565-offset*13,part))
        lines.append((x+5,510,value))
    return chinese_pdf(lines,drawing)


def balance_xls():
    from tests.test_legacy_workbooks import record, compound

    rows = [
        ["科目余额表"],
        ["纳税人识别号：" + TAX],
        ["统计期间：2026-01-01 至 2026-01-31"],
        ["金额单位：元"],
        ["科目名称", "期初余额", "本期借方", "本期贷方", "期末余额"],
        ["库存现金", 0, 100.0123, 20, 80.0123],
    ]
    global_records = (
        record(0x809, struct.pack("<HHHH", 0x500, 5, 0, 0))
        + record(0x42, struct.pack("<H", 936))
        + record(0xE0, b"\0" * 16)
        + record(0x4D, b"\0" * 4200)
    )
    name = b"Sheet1"
    start = len(global_records) + 4 + 7 + len(name) + 4
    bound = record(0x85, struct.pack("<IBBB", start, 0, 0, len(name)) + name)
    cells = b""
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            if isinstance(value, str):
                encoded = value.encode("gbk")
                cells += record(0x204, struct.pack("<HHHH", r, c, 0, len(encoded)) + encoded)
            else:
                cells += record(0x203, struct.pack("<HHHd", r, c, 0, value))
    stream = (
        global_records
        + bound
        + record(0xA)
        + record(0x809, struct.pack("<HHHH", 0x500, 0x10, 0, 0))
        + record(0x200, struct.pack("<HHHHH", 0, 6, 0, 5, 0))
        + cells
        + record(0xA)
    )
    return compound(stream)


class CoreParserTests(unittest.TestCase):
    def setUp(self):
        self.db, self.mapper = mapper()
        self.addCleanup(self.db.close)
        self.task = {
            "id": "task",
            "company_id": "company",
            "tenant_id": "org",
            "tax_no": TAX,
            "start_date": "2026-01-01",
            "end_date": "2026-01-31",
            "applicable_forms": [],
        }

    def parse(self, kind, **kwargs):
        return core_materials.parse(book(kind, **kwargs), "资料.xlsx", kind, kind + str(kwargs), self.mapper)

    def complete(self):
        return [
            self.parse("account_balance"),
            self.parse("vat_declaration"),
            self.parse("invoice_detail"),
            self.parse("invoice_detail", direction="进项"),
        ]

    def test_three_materials_ready_and_precision_sources(self):
        result = core_materials.prepare(self.task, self.complete())
        self.assertEqual(result["issues"], [])
        self.assertEqual(result["status"], "ready_for_analysis")
        row = result["records"][0]
        self.assertEqual(row["current_debit"], "100.0123")
        self.assertEqual(row["opening_balance"], "0")
        self.assertEqual(row["field_sources"]["current_debit"]["position"], "科目余额表!D6")
        self.assertEqual(row["raw_record"]["3"], "100.0123")
        self.assertEqual(result["invoice_summaries"][0]["amount_excluding_tax"], "100.0123")
        self.assertEqual(result["evidence_links"]["comparison_status"], "requires_confirmed_rules")

    def test_missing_category_and_main_form_are_blocked(self):
        result = core_materials.prepare(
            self.task, [self.parse("account_balance"), self.parse("vat_declaration", form="附表一")]
        )
        self.assertEqual(result["status"], "needs_completion")
        self.assertTrue({"missing_material", "main_form_gap"} <= {i["code"] for i in result["issues"]})

    def test_blank_is_not_zero(self):
        result = core_materials.prepare(self.task, [self.parse("account_balance", blank=True)])
        row = result["records"][0]
        self.assertIsNone(row["current_credit"])
        self.assertEqual(row["value_states"]["current_credit"], "missing")
        self.assertIn("record_fields", {i["code"] for i in result["issues"]})

    def test_subject_conflict_invoice_roles_and_outside_dates(self):
        docs = self.complete()
        docs[0] = self.parse("account_balance", tax=OTHER)
        docs.append(self.parse("invoice_detail", start="2025-12-01", end="2025-12-31"))
        result = core_materials.prepare(self.task, docs)
        self.assertIn("company_conflict", {i["code"] for i in result["issues"]})
        self.assertNotIn("invoice_company_conflict", {i["code"] for i in result["issues"]})
        self.assertFalse(result["records"][-1]["in_scope"])

    def test_range_modes_require_complete_calendar_intervals(self):
        base = {"company_name": "企业", "tax_no": TAX}
        self.assertEqual(core_materials.scope_errors(
            {**base, "start_date": "2026-01-01", "end_date": "2026-01-31", "range_kind": "month"}), {})
        self.assertIn("range_kind", core_materials.scope_errors(
            {**base, "start_date": "2026-01-01", "end_date": "2026-02-28", "range_kind": "month"}))
        self.assertEqual(core_materials.scope_errors(
            {**base, "start_date": "2026-01-01", "end_date": "2026-03-31", "range_kind": "quarter"}), {})
        self.assertIn("range_kind", core_materials.scope_errors(
            {**base, "start_date": "2026-02-01", "end_date": "2026-04-30", "range_kind": "quarter"}))
        self.assertEqual(core_materials.scope_errors(
            {**base, "start_date": "2026-01-01", "end_date": "2026-12-31", "range_kind": "year"}), {})
        self.assertIn("range_kind", core_materials.scope_errors(
            {**base, "start_date": "2026-01-01", "end_date": "2026-01-31", "range_kind": "year"}))
        self.assertEqual(core_materials.scope_errors(
            {**base, "start_date": "2026-01-15", "end_date": "2026-02-20", "range_kind": "custom"}), {})
        self.assertIn("range_kind", core_materials.scope_errors(
            {**base, "start_date": "2026-01-01", "end_date": "2026-01-31", "range_kind": "week"}))

    def test_mixed_invoice_file_blocks_and_keeps_out_of_scope_record(self):
        docs = self.complete()
        invoice = docs[2]
        original_id = invoice["records"][0]["id"]
        outside = deepcopy(invoice["records"][0])
        outside.update(id="outside-record", invoice_no="OUTSIDE-2026-02",
                       issue_date="2026-02-15", source_position="Sheet1!A9")
        invoice["records"].append(outside)
        result = core_materials.prepare(self.task, docs)
        codes = {i["code"] for i in result["issues"]}
        self.assertIn("record_outside_scope", codes)
        self.assertIn("mixed_scope_file", codes)
        statuses = {r["id"]: r["scope_status"] for r in result["records"]
                    if r["material_type"] == "invoice_detail"}
        self.assertEqual(statuses[original_id], "in_scope")
        self.assertEqual(statuses["outside-record"], "outside_scope")
        kept = next(r for r in result["records"] if r["id"] == "outside-record")
        self.assertEqual(kept["source_position"], "Sheet1!A9")
        self.assertEqual(kept["date_basis"], "issue_date")
        self.assertFalse(kept["in_scope"])
        self.assertNotIn("outside-record", {rid for group in result["invoice_summaries"]
                                            for rid in group["source_record_ids"]})

    def test_unknown_invoice_date_blocks_and_is_not_in_scope(self):
        docs = self.complete()
        record = docs[2]["records"][0]
        record["issue_date"] = None
        result = core_materials.prepare(self.task, docs)
        self.assertIn("record_date_unknown", {i["code"] for i in result["issues"]})
        kept = next(r for r in result["records"] if r["id"] == record["id"])
        self.assertEqual(kept["scope_status"], "unknown_date")
        self.assertEqual(kept["date_basis"], "issue_date")
        self.assertFalse(kept["in_scope"])

    def test_user_confirmed_period_is_labelled_separately(self):
        docs = self.complete()
        expected = {}
        for record in docs[0]["records"]:
            record["confirmed_metadata"] = {"period_start": "2026-01-01", "period_end": "2026-01-31"}
            expected[record["id"]] = "user_confirmed"
        result = core_materials.prepare(self.task, docs)
        labels = {r["id"]: r["date_source"] for r in result["records"] if r["id"] in expected}
        self.assertEqual(labels, expected)

    def test_months_merge_quarter_and_gap_remains(self):
        self.task["end_date"] = "2026-03-31"
        docs = []
        for start, end in [("2026-01-01", "2026-01-31"), ("2026-02-01", "2026-02-28"), ("2026-03-01", "2026-03-31")]:
            docs.extend(self.parse(kind, start=start, end=end) for kind in ["account_balance", "vat_declaration"])
            for direction in ["进项", "销项"]:
                doc = self.parse("invoice_detail", start=start, end=end, direction=direction)
                doc["records"][0]["invoice_no"] += start
                docs.append(doc)
        self.assertEqual(core_materials.prepare(self.task, docs)["issues"], [])
        docs = [
            d
            for d in docs
            if not (d["material_type"] == "vat_declaration" and d["parts"][0]["period"][0] == "2026-02-01")
        ]
        result = core_materials.prepare(self.task, docs)
        self.assertTrue(any("2026-02-01 至 2026-02-28" in i["message"] for i in result["issues"]))

    def test_duplicates_keep_all_sources_but_count_once_conflict_blocks(self):
        docs = self.complete()
        duplicate = deepcopy(docs[2])
        duplicate["id"] = "duplicate"
        duplicate["sha256"] = "different-source"
        for r in duplicate["records"]:
            r["id"] = "duplicate-record"
            r["source_file_id"] = "duplicate"
        docs.append(duplicate)
        result = core_materials.prepare(self.task, docs)
        self.assertEqual(sum(g["count"] for g in result["invoice_summaries"]), 2)
        self.assertEqual(result["records"][-1]["duplicate_of"], docs[2]["records"][0]["id"])
        duplicate["records"][0]["tax_amount"] = "99"
        self.assertIn("invoice_conflict", {i["code"] for i in core_materials.prepare(self.task, docs)["issues"]})

    def test_applicable_forms_not_invented_and_quarter_not_repeated(self):
        self.task["applicable_forms"] = None
        self.assertIn(
            "forms_unconfirmed", {i["code"] for i in core_materials.prepare(self.task, self.complete())["issues"]}
        )
        self.task["end_date"] = "2026-03-31"
        self.task["applicable_forms"] = ["附表一"]
        doc = self.parse("vat_declaration", end="2026-03-31")
        result = core_materials.prepare(self.task, [doc])
        self.assertEqual(len(result["records"]), 2)
        self.assertIn("missing_form", {i["code"] for i in result["issues"]})

    def test_structured_xml_and_unsafe_entities(self):
        xml = f"<Invoice><InvoiceNumber>0012</InvoiceNumber><IssueDate>2026-01-05</IssueDate><SellerTaxID>{TAX}</SellerTaxID><BuyerTaxID>{OTHER}</BuyerTaxID><TotalAmount>10.01</TotalAmount><TotalTax>1.30</TotalTax><RedFlag>蓝字</RedFlag><VoidFlag>正常</VoidFlag><Items><Item><Name>服务</Name></Item></Items></Invoice>".encode()
        doc = core_materials.parse(xml, "票.xml", "invoice_detail", "xml", self.mapper)
        self.assertEqual(doc["records"][0]["invoice_no"], "0012")
        self.assertEqual(doc["records"][0]["field_sources"]["invoice_no"]["position"], "/Invoice/InvoiceNumber[1]")
        self.assertTrue(any(item["tag"] == "Name" for item in doc["records"][0]["raw_record"]))
        with self.assertRaises(InputError):
            core_materials.parse(
                b'<!DOCTYPE x [<!ENTITY e "boom">]><Invoice/>', "票.xml", "invoice_detail", "bad", self.mapper
            )

    def test_legacy_xls_text_pdf_and_multicolumn_vat_sources(self):
        doc = core_materials.parse(balance_xls(), "余额.xls", "account_balance", "xls", self.mapper)
        self.assertEqual(doc["records"][0]["current_debit"], "100.0123")
        doc = core_materials.parse(pdf_table(), "申报.pdf", "vat_declaration", "pdf", self.mapper)
        self.assertEqual(doc["records"][0]["declared_value"], "100.0123")
        self.assertIn("PDF第1页/表1", doc["records"][0]["source_position"])
        from openpyxl import load_workbook

        wb = load_workbook(BytesIO(book("vat_declaration")))
        s = wb.active
        s.cell(5, 4, "本年累计数")
        s.cell(6, 4, "200.0246")
        s.cell(7, 4, "26.0032")
        out = BytesIO()
        wb.save(out)
        wb.close()
        doc = core_materials.parse(out.getvalue(), "申报.xlsx", "vat_declaration", "wide", self.mapper)
        self.assertEqual(len(doc["records"]), 4)
        self.assertEqual({r["column_scope"] for r in doc["records"]}, {"本期数", "本年累计数"})
        self.assertEqual({r["declared_value"] for r in doc["records"]}, {"100.0123", "13.0016", "200.0246", "26.0032"})

    def test_entirely_unknown_headers_can_use_bounded_ai_mapping(self):
        from openpyxl import load_workbook

        wb = load_workbook(BytesIO(book("account_balance")))
        s = wb.active
        for c, value in enumerate(["编号别名", "账户别名", "开端数值", "借入数值", "贷出数值", "终端数值"], 1):
            s.cell(5, c, value)
        out = BytesIO()
        wb.save(out)
        wb.close()
        targets = [
            "account_code",
            "account_name",
            "opening_balance",
            "current_debit",
            "current_credit",
            "closing_balance",
        ]
        response = {
            "mappings": [
                {"source_column": i, "target_field": target, "status": "mapped", "explanation": "context"}
                for i, target in enumerate(targets)
            ]
        }
        self.mapper.ai = True
        with (
            patch.dict(os.environ, {"TAXPEARLS_AI_ENABLED": "1", "TAXPEARLS_AI_API_KEY": "synthetic"}),
            patch("src.core_mapping.chat_content", return_value=json.dumps(response)) as call,
        ):
            doc = core_materials.parse(out.getvalue(), "全别名.xlsx", "account_balance", "unknown", self.mapper)
        self.assertEqual(doc["issues"], [])
        self.assertEqual(doc["records"][0]["current_debit"], "100.0123")
        self.assertEqual(call.call_count, 1)
        self.assertNotIn(TAX, json.dumps(call.call_args.args[1]))
        self.assertEqual(doc["records"][0]["field_sources"]["current_debit"]["origin"], "ai_candidate")

    def test_native_grouped_headers_keep_balance_sides_and_vat_scopes(self):
        wb = Workbook()
        s = wb.active
        s.title = "科目余额表"
        for row in [
            ["科目余额表"],
            ["纳税人识别号：" + TAX],
            ["统计期间：2026-01-01 至 2026-01-31"],
            ["金额单位：元"],
            ["科目名称", "期初余额", None, "本期发生额", None, "期末余额", None],
            [None, "借方", "贷方", "借方", "贷方", "借方", "贷方"],
            ["库存现金", "10", "0", "20", "5", "25", "0"],
        ]:
            s.append(row)
        out = BytesIO()
        wb.save(out)
        wb.close()
        doc = core_materials.parse(out.getvalue(), "双栏.xlsx", "account_balance", "sides", self.mapper)
        self.assertEqual(doc["issues"], [])
        r = doc["records"][0]
        self.assertEqual(r["opening_balance"], {"debit": "10", "credit": "0"})
        self.assertEqual(r["field_sources"]["opening_balance"]["positions"], ["科目余额表!B7", "科目余额表!C7"])
        wb = Workbook()
        s = wb.active
        s.title = "增值税纳税申报表主表"
        for row in [
            [s.title],
            ["税款所属期：2026-01-01 至 2026-01-31"],
            ["金额单位：元"],
            ["项目", "栏次", "一般项目", None, "即征即退项目", None],
            [None, None, "本期数", "本年累计数", "本期数", "本年累计数"],
            ["销售额", "1", "10", "20", "30", "40"],
        ]:
            s.append(row)
        out = BytesIO()
        wb.save(out)
        wb.close()
        doc = core_materials.parse(out.getvalue(), "分项目.xlsx", "vat_declaration", "scopes", self.mapper)
        self.assertEqual(doc["issues"], [])
        self.assertEqual(
            {r["column_scope"] for r in doc["records"]},
            {"一般项目本期数", "一般项目本年累计数", "即征即退项目本期数", "即征即退项目本年累计数"},
        )

    def test_ocr_untrusted_candidates_require_review_and_no_default_calls(self):
        raw = (Path(__file__).parent / "fixtures" / "materials-scanned.pdf").read_bytes()
        with patch("src.core_materials.chat_content") as call:
            doc = core_materials.parse(raw, "扫描.pdf", "vat_declaration", "scan", self.mapper)
            call.assert_not_called()
            self.assertFalse(doc["records"])
        response = {
            "form_name": "增值税纳税申报表主表",
            "form_type": "main",
            "tax_no": TAX,
            "period_start": "2026-01-01",
            "period_end": "2026-01-31",
            "unit": "元",
            "quote": "税款所属期 2026-01",
            "rows": [
                {
                    "item_label": "销售额",
                    "item_key": "1",
                    "declared_value": "10.01",
                    "column_scope": "本期数",
                    "quote": "销售额 10.01",
                }
            ],
        }
        with (
            patch.dict(os.environ, {"TAXPEARLS_AI_ENABLED": "1", "TAXPEARLS_AI_API_KEY": "synthetic"}),
            patch("src.core_materials.chat_content", return_value=json.dumps(response)),
        ):
            doc = core_materials.parse(raw, "扫描.pdf", "vat_declaration", "scan", self.mapper, ocr=True)
        self.assertTrue(doc["ocr_candidate"])
        self.assertEqual(doc["records"][0]["declared_value"], "10.01")
        self.assertIn("ocr_review", {i["code"] for i in doc["issues"]})
        self.assertTrue(doc["records"][0]["requires_source_review"])
        response["rows"][0]["declared_value"] = "99"
        with (
            patch.dict(os.environ, {"TAXPEARLS_AI_ENABLED": "1", "TAXPEARLS_AI_API_KEY": "synthetic"}),
            patch("src.core_materials.chat_content", return_value=json.dumps(response)),
        ):
            doc = core_materials.parse(raw, "扫描.pdf", "vat_declaration", "scan", self.mapper, ocr=True)
        self.assertFalse(doc["records"])
        self.assertIn("ocr_failed", {i["code"] for i in doc["issues"]})

    def test_pdf_invoice_ticket_totals_parties_and_unknown_statuses(self):
        with patch("src.core_materials.chat_content") as model:
            doc = core_materials.parse(invoice_pdf(), "电子发票.pdf", "invoice_detail", "ticket", self.mapper, ocr=True)
            model.assert_not_called()
        self.assertEqual(len(doc["records"]), 1, doc)
        record = doc["records"][0]
        self.assertEqual(record["invoice_no"], "12345678901234567890")
        self.assertEqual(record["issue_date"], "2026-01-15")
        self.assertEqual((record["seller_tax_no"], record["buyer_tax_no"]), (TAX, OTHER))
        self.assertEqual((record["amount_excluding_tax"], record["tax_amount"], record["amount_including_tax"]), ("100.00", "13.00", "113.00"))
        self.assertIsNone(record["red_letter_status"])
        self.assertIsNone(record["void_status"])
        self.assertTrue(record["requires_source_review"])
        self.assertIn("PDF第1页", record["field_sources"]["invoice_no"]["position"])
        self.assertIn("合计", record["field_sources"]["tax_amount"]["quote"])
        self.assertEqual(doc["parts"][0]["tax_nos"], [])
        self.assertIsNone(doc["parts"][0]["period"])
        self.assertIn("invoice_status_missing", {i["code"] for i in doc["issues"]})
        checked = core_materials.prepare(self.task, [doc])
        self.assertEqual(checked["records"][0]["direction"], "output")
        self.assertIn("invoice_coverage_gap", {i["code"] for i in checked["issues"]})

    def test_pdf_invoice_conflicting_totals_and_multiple_tickets_block(self):
        doc = core_materials.parse(invoice_pdf(total="114.00"), "票.pdf", "invoice_detail", "bad", self.mapper)
        self.assertIn("invoice_total_conflict", {i["code"] for i in doc["issues"]})
        doc = core_materials.parse(invoice_pdf(other_number="22345678901234567890"), "票.pdf", "invoice_detail", "multi", self.mapper)
        self.assertFalse(doc["records"])
        self.assertIn("pdf_ambiguous", {i["code"] for i in doc["issues"]})

    def test_multipage_invoice_sources_duplicate_once_and_pdf_listing_status_gap(self):
        import pypdfium2
        output=BytesIO()
        with pypdfium2.PdfDocument.new() as merged, pypdfium2.PdfDocument(invoice_pdf()) as ticket:
            merged.import_pages(ticket)
            merged.import_pages(ticket)
            merged.save(output)
        doc=core_materials.parse(output.getvalue(),"两页.pdf","invoice_detail","multi",self.mapper)
        self.assertEqual(len(doc["records"]),2)
        self.assertIn("PDF第2页",doc["records"][1]["source_position"])
        result=core_materials.prepare(self.task,[doc])
        self.assertIn("duplicate_invoice",{i["code"] for i in result["notices"]})
        self.assertEqual(result["invoice_summaries"][0]["count"],1)
        doc=core_materials.parse(pdf_table("invoice_detail"),"发票清单.pdf","invoice_detail","grid",self.mapper)
        self.assertEqual(len(doc["records"]),2,doc)
        self.assertNotIn("record_fields",{i["code"] for i in doc["issues"]})
        self.assertNotIn("missing_columns",{i["code"] for i in doc["issues"]})
        self.assertIn("invoice_status_missing",{i["code"] for i in doc["issues"]})

    def test_balance_pdf_table_and_scanned_business_candidates(self):
        doc = core_materials.parse(pdf_table("account_balance"), "余额.pdf", "account_balance", "balance", self.mapper)
        self.assertEqual(len(doc["records"]), 2, doc)
        self.assertEqual(doc["records"][0]["closing_balance"], "90")
        raw = (Path(__file__).parent / "fixtures" / "materials-scanned.pdf").read_bytes()
        for kind, row in [("account_balance", {"account_name":"现金", "opening_balance":"10", "current_debit":"100", "current_credit":"20", "closing_balance":"90"}),
                          ("invoice_detail", {"invoice_no":"12345678901234567890", "issue_date":"2026-01-15", "seller_tax_no":TAX,"buyer_tax_no":OTHER,"amount_excluding_tax":"100", "tax_amount":"13"})]:
            answer={"unit":"元", "rows":[{"fields":row,"quotes":{key:value for key,value in row.items()}}]}
            with self.subTest(kind=kind), patch.dict(os.environ, {"TAXPEARLS_AI_ENABLED":"1", "TAXPEARLS_AI_API_KEY":"synthetic"}), patch("src.core_materials.chat_content",return_value=json.dumps(answer)):
                doc=core_materials.parse(raw,"扫描.pdf",kind,"scan-"+kind,self.mapper,ocr=True)
            self.assertEqual(len(doc["records"]),1,doc)
            self.assertTrue(doc["records"][0]["requires_source_review"])
            self.assertIn("ocr_review",{i["code"] for i in doc["issues"]})
        answer["rows"][0]["quotes"]["tax_amount"]="其他金额 99"
        with patch.dict(os.environ, {"TAXPEARLS_AI_ENABLED":"1", "TAXPEARLS_AI_API_KEY":"synthetic"}), patch("src.core_materials.chat_content",return_value=json.dumps(answer)):
            doc=core_materials.parse(raw,"扫描.pdf","invoice_detail","bad",self.mapper,ocr=True)
        self.assertFalse(doc["records"])
        self.assertIn("ocr_failed",{i["code"] for i in doc["issues"]})

    def test_other_two_pdf_uploads_preserve_wrapped_group_headers(self):
        for kind in ["account_balance", "vat_declaration"]:
            with self.subTest(kind=kind), patch("src.core_materials.chat_content") as model:
                doc=core_materials.parse(grouped_business_pdf(kind),kind+".pdf",kind,kind,self.mapper,ocr=True)
                model.assert_not_called()
                self.assertEqual(doc["issues"],[],doc)
                if kind=="account_balance":
                    self.assertEqual(len(doc["records"]),1)
                    self.assertEqual(doc["records"][0]["opening_balance"],{"debit":"10","credit":"0"})
                    self.assertEqual(doc["records"][0]["closing_balance"],{"debit":"90","credit":"0"})
                else:
                    self.assertEqual({r["column_scope"]:r["declared_value"] for r in doc["records"]},
                                     {"一般项目本期数":"100","一般项目本年累计数":"900"})

    def test_other_two_partial_pdf_tables_use_ocr_and_keep_gaps_when_unavailable(self):
        balance={"account_name":"现金","opening_balance":"10","current_debit":"100","current_credit":"20","closing_balance":"90"}
        for kind in ["account_balance", "vat_declaration"]:
            raw=grouped_business_pdf(kind,unknown=True)
            answer={"tax_no":TAX,"period_start":"2026-01-01","period_end":"2026-01-31","unit":"元","quote":"来源表头"}
            if kind=="account_balance":
                answer["rows"]=[{"fields":balance,"quotes":balance}]
                field="current_debit"
            else:
                answer.update(form_name="增值税纳税申报表主表",form_type="main",rows=[{"item_label":"应税销售额","declared_value":"100","column_scope":"一般项目本期数","quote":"应税销售额 100"}])
                field="declared_value"
            with self.subTest(kind=kind), patch.dict(os.environ,{"TAXPEARLS_AI_ENABLED":"1","TAXPEARLS_AI_API_KEY":"synthetic"}),patch("src.core_materials.chat_content",return_value=json.dumps(answer)) as model:
                doc=core_materials.parse(raw,"非标准.pdf",kind,"partial-"+kind,self.mapper,ocr=True)
                self.assertEqual(model.call_count,1)
                self.assertTrue(doc["ocr_candidate"],doc)
                self.assertEqual(doc["records"][0][field],"100")
                self.assertEqual({i["code"] for i in doc["issues"]},{"ocr_review"})
                self.assertTrue(doc["records"][0]["requires_source_review"])
            with patch.dict(os.environ,{"TAXPEARLS_AI_ENABLED":"0"}),patch("src.core_materials.chat_content") as model:
                doc=core_materials.parse(raw,"非标准.pdf",kind,"no-model-"+kind,self.mapper,ocr=True)
                model.assert_not_called()
                self.assertNotIn("ocr_candidate",doc)
                self.assertTrue(doc["records"])
                self.assertIsNone(doc["records"][0].get(field))
                self.assertTrue({"missing_columns","ocr_unavailable"}<={i["code"] for i in doc["issues"]})


class MappingTests(unittest.TestCase):
    def setUp(self):
        self.db, _ = mapper()
        self.addCleanup(self.db.close)

    def resolve(self, sha, target="declared_value"):
        mapper = core_mapping.Mapper(self.db, "org", "task", sha, ai=True)
        response = json.dumps(
            {
                "mappings": [
                    {"source_column": 1, "target_field": target, "status": "mapped", "explanation": "column context"}
                ]
            }
        )
        with (
            patch.dict(os.environ, {"TAXPEARLS_AI_ENABLED": "1", "TAXPEARLS_AI_API_KEY": "synthetic"}),
            patch("src.core_mapping.chat_content", return_value=response) as call,
        ):
            result = mapper.resolve("vat_declaration", ["项目", "本期申报值"], "main", [["收入", "12.345"]])
            return result, call

    def test_five_distinct_hashes_promote_repeated_hash_does_not(self):
        for _ in range(5):
            self.resolve("same")
        self.assertEqual(self.db.execute("SELECT COUNT(*) FROM core_mapping_observations").fetchone()[0], 1)
        for i in range(4):
            self.resolve("independent" + str(i))
        result, call = self.resolve("sixth")
        self.assertEqual(result[1]["origin"], "promoted")
        call.assert_not_called()

    def test_same_source_candidate_survives_local_reparse_without_cross_source_reuse(self):
        self.resolve("original")
        with patch("src.core_mapping.chat_content") as call:
            for ai in [False, True]:
                m = core_mapping.Mapper(self.db, "org", "task", "original", ai=ai)
                result = m.resolve("vat_declaration", ["项目", "本期申报值"], "main", [["收入", "12.345"]])
                self.assertEqual(result[1]["origin"], "ai_candidate")
                self.assertEqual(result[1]["independent_sources"], 1)
            for org, sha, context in [("org", "different", "main"), ("foreign", "original", "main"),
                                      ("org", "original", "supplement")]:
                m = core_mapping.Mapper(self.db, org, "task", sha)
                self.assertNotIn(1, m.resolve("vat_declaration", ["项目", "本期申报值"], context, [["收入", "12.345"]]))
            call.assert_not_called()

    def test_mapping_conflict_never_promotes_and_values_not_sent(self):
        _, call = self.resolve("one")
        self.assertNotIn("12.345", json.dumps(call.call_args.args[1]))
        self.resolve("two", "item_key")
        for i in range(6):
            result, _ = self.resolve(str(i))
        self.assertNotIn(1, result)

    def test_known_alias_never_calls_ai_ambiguous_duplicate_unmatched(self):
        m = core_mapping.Mapper(self.db, "org", "task", "sha", ai=True)
        with patch("src.core_mapping.chat_content") as call:
            result = m.resolve("account_balance", ["科目名称", "期初余额", "期末余额"], "account_balance", [])
            self.assertEqual(len(result), 3)
            call.assert_not_called()
            result = m.resolve("vat_declaration", ["本期数", "本月数"], "main", [])
            self.assertEqual(result, {})


class CoreAPITests(unittest.TestCase):
    def setUp(self):
        self.tmp = TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        env = patch.dict(
            os.environ,
            {
                "TAXPEARLS_FIELD_KEY": base64.b64encode(b"f" * 32).decode(),
                "TAXPEARLS_MATERIAL_KEY": base64.b64encode(b"c" * 32).decode(),
                "TAXPEARLS_AI_ENABLED": "0",
                "TAXPEARLS_MATERIAL_QUEUE_ENABLED": "0",
                "TAXPEARLS_NOTIFICATION_EMAIL_ENABLED": "0",
            },
        )
        env.start()
        self.addCleanup(env.stop)
        self.store = Store(Path(self.tmp.name) / "test.db")
        p = patch.object(module, "store", self.store)
        p.start()
        self.addCleanup(p.stop)
        self.admin = self.store.create_user("core-admin", "Core-test-2026!", "管理员", "org_admin", "core-org")
        self.foreign = self.store.create_user("core-other", "Core-test-2026!", "其他", "org_admin", "other-org")
        self.client = TestClient(module.app, raise_server_exceptions=False)
        self.client.__enter__()
        self.addCleanup(self.client.__exit__, None, None, None)
        self.login(self.admin)

    def login(self, user):
        self.client.cookies.clear()
        token = self.store.authenticate(user["username"], "Core-test-2026!")[1]
        self.client.cookies.set(module.COOKIE_NAME, token)

    def create(self):
        r = self.client.post(
            "/api/enterprise/input-tasks",
            json={
                "company_name": "核心资料测试企业",
                "tax_no": TAX,
                "start_date": "2026-01-01",
                "end_date": "2026-01-31",
                "applicable_forms": [],
            },
        )
        self.assertEqual(r.status_code, 200, r.text)
        return r.json()

    def upload(self, task, kind, **kwargs):
        r = self.client.post(
            "/api/enterprise/input-tasks/" + task["id"] + "/files",
            data={"expected_revision": str(task["revision"]), "material_type": kind},
            files=[("files", ("同名.xlsx", book(kind, **kwargs)))],
        )
        self.assertEqual(r.status_code, 200, r.text)
        return r.json()

    def scope(self, task, **changes):
        fields = {key: task["task"][key] for key in core_materials.SCOPE_FIELDS}
        return self.client.post(f"/api/enterprise/input-tasks/{task['id']}/scope",
                                json={**fields, "expected_revision": task["revision"], **changes})

    def test_material_export_contains_normalized_facts_and_enforces_acl_and_revision(self):
        task = self.upload(self.create(), "invoice_detail")
        url = f"/api/enterprise/input-tasks/{task['id']}"
        export_url = url + f"/data?expected_revision={task['revision']}"
        before = self.client.get(url).json()
        response = self.client.get(export_url)
        self.assertEqual(response.status_code, 200, response.text)
        result = response.json()
        self.assertEqual(result["schema_version"], "taxpearls.material-data.v1")
        self.assertEqual(result["records"][0]["fields"]["invoice_no"], "12345678901234567890")
        self.assertEqual(result["records"][0]["fields"]["amount_excluding_tax"], "100.0123")
        self.assertNotIn("documents", result)
        self.assertNotIn("events", result)
        self.assertNotIn("versions", result)
        self.assertEqual(set(result), {"schema_version", "task", "records", "sources"})
        self.assertNotIn("quality", result)
        self.assertTrue(before["processing"]["issues"])
        after = self.client.get(url).json()
        # Authorized reads append access audit events; the material snapshot stays unchanged.
        self.assertEqual({k: v for k, v in after.items() if k != "events"},
                         {k: v for k, v in before.items() if k != "events"})
        self.assertEqual(after["processing"]["issues"], before["processing"]["issues"])
        self.assertEqual(self.client.get(url + "/data?expected_revision=0").status_code, 409)
        self.login(self.foreign)
        self.assertEqual(self.client.get(export_url).status_code, 404)
        self.client.cookies.clear()
        self.assertEqual(self.client.get(export_url).status_code, 401)

    def test_blank_draft_upload_autofill_and_explicit_confirmation(self):
        task = self.client.post("/api/enterprise/input-tasks", json={}).json()
        self.assertFalse(task["task"]["scope_confirmed"])
        self.assertIsNone(task["task"]["company_id"])
        self.assertEqual(self.store.list_clients(self.admin), [])
        task = self.upload(task, "account_balance")
        self.assertEqual(task["task"]["tax_no"], TAX)
        # The source period is a candidate only; the program never selects the
        # inspection range from material dates.
        self.assertEqual(task["task"]["start_date"], "")
        self.assertEqual(task["task"]["end_date"], "")
        self.assertEqual(task["task"]["scope_candidates"]["start_date"][0]["value"], "2026-01-01")
        self.assertEqual(task["task"]["scope_candidates"]["end_date"][0]["value"], "2026-01-31")
        self.assertNotIn("start_date", task["task"]["scope_origins"])
        self.assertEqual(task["task"]["scope_origins"]["tax_no"]["sources"][0]["file_id"], task["documents"][0]["id"])
        from webapp.core_inputs import require_scope
        with self.assertRaises(ValueError):
            require_scope(task)
        result = self.scope(task, company_name="文件读取测试企业",
                            start_date="2026-01-01", end_date="2026-01-31", confirm=True)
        self.assertEqual(result.status_code, 200, result.text)
        confirmed = result.json()
        require_scope(confirmed)
        self.assertTrue(confirmed["task"]["scope_confirmed"])
        self.assertEqual(len(self.store.list_clients(self.admin)), 1)
        self.assertTrue(confirmed["processing"]["records"])
        with self.store.connect() as db:
            batch = material_batches._authorize(db, self.admin, task["id"])
            self.assertEqual(batch["client_id"], confirmed["task"]["company_id"])
        self.assertEqual(self.scope(task, confirm=True).status_code, 409)

    def test_source_period_is_candidate_and_never_overrides_user_range(self):
        task = {key: "" for key in core_materials.SCOPE_FIELDS}
        doc = {"id": "source", "material_type": "account_balance", "parts": [{"name": "Sheet1", "tax_nos": [TAX],
            "period": ["2026-02-01", "2026-02-28"], "parsed_metadata": {"period": ["2026-01-01", "2026-01-31"]}}]}
        core_materials.autofill_scope(task, [doc])
        self.assertEqual(task["start_date"], "")
        self.assertEqual(task["end_date"], "")
        self.assertEqual(task["scope_candidates"]["start_date"][0]["value"], "2026-01-01")
        self.assertEqual(task["scope_candidates"]["end_date"][0]["value"], "2026-01-31")
        self.assertNotIn("start_date", task["scope_origins"])
        # A user-selected range survives later parses even when the source differs.
        task.update(start_date="2026-01-15", end_date="2026-01-20")
        core_materials.autofill_scope(task, [doc])
        self.assertEqual((task["start_date"], task["end_date"]), ("2026-01-15", "2026-01-20"))

    def test_legacy_file_selected_range_is_cleared_for_reconfirmation(self):
        task = {
            "start_date": "2026-01-01",
            "end_date": "2026-01-31",
            "scope_confirmed": True,
            "scope_origins": {"start_date": {"kind": "file"}, "end_date": {"kind": "file"}},
        }
        doc = {"id": "source", "material_type": "account_balance", "parts": [{"name": "Sheet1", "tax_nos": [TAX],
            "period": ["2026-01-01", "2026-01-31"], "parsed_metadata": {"period": ["2026-01-01", "2026-01-31"]}}]}
        core_materials.autofill_scope(task, [doc])
        self.assertEqual(task["start_date"], "")
        self.assertEqual(task["end_date"], "")
        self.assertFalse(task["scope_confirmed"])
        self.assertEqual(task["scope_candidates"]["start_date"][0]["value"], "2026-01-01")

    def test_draft_name_is_read_from_explicit_file_metadata(self):
        from openpyxl import load_workbook
        task = self.client.post("/api/enterprise/input-tasks", json={}).json()
        wb = load_workbook(BytesIO(book("account_balance")))
        wb.active.insert_rows(2)
        wb.active.cell(2, 1, "企业名称：自动读取有限公司")
        raw = BytesIO()
        wb.save(raw)
        wb.close()
        r = self.client.post(f"/api/enterprise/input-tasks/{task['id']}/files",
            data={"material_type": "account_balance", "expected_revision": str(task["revision"])},
            files={"files": ("余额.xlsx", raw.getvalue())})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.json()["task"]["company_name"], "自动读取有限公司")
        self.assertFalse(r.json()["task"]["scope_confirmed"])

    def test_manual_values_preserved_and_conflict_blocks_confirmation(self):
        task = self.create()
        task = self.upload(task, "account_balance", tax=OTHER)
        self.assertEqual(task["task"]["tax_no"], TAX)
        self.assertIn("tax_no", task["scope_errors"])
        r = self.scope(task, confirm=True)
        self.assertEqual(r.status_code, 200, r.text)
        self.assertFalse(r.json()["task"]["scope_confirmed"])
        from webapp.core_inputs import require_scope
        with self.assertRaises(ValueError):
            require_scope(r.json())

    def test_multiple_source_periods_are_not_combined_and_exclusion_rederives_candidates(self):
        task = self.client.post("/api/enterprise/input-tasks", json={}).json()
        task = self.upload(task, "account_balance")
        task = self.upload(task, "vat_declaration", start="2026-02-01", end="2026-02-28")
        self.assertEqual(task["task"]["start_date"], "")
        self.assertEqual(len(task["task"]["scope_candidates"]["start_date"]), 2)
        file_id = task["documents"][-1]["id"]
        r = self.client.post(f"/api/enterprise/input-tasks/{task['id']}/prepare",
                             json={"expected_revision": task["revision"], "source_options": {file_id: {"excluded": True}}})
        self.assertEqual(r.status_code, 200, r.text)
        # Excluding a source updates the candidates but still does not choose
        # the inspection range for the user.
        self.assertEqual(r.json()["task"]["start_date"], "")
        self.assertEqual(len(r.json()["task"]["scope_candidates"]["start_date"]), 1)

    def test_invalid_range_mode_blocks_confirmation(self):
        task = self.create()
        r = self.scope(task, start_date="2026-01-01", end_date="2026-02-28",
                       range_kind="quarter", confirm=True)
        self.assertEqual(r.status_code, 200, r.text)
        value = r.json()
        self.assertIn("range_kind", value["scope_errors"])
        self.assertFalse(value["task"]["scope_confirmed"])

    def test_range_change_reruns_checks_and_retains_history(self):
        task = self.create()
        task = self.upload(task, "account_balance")
        confirmed = self.scope(task, range_kind="month", confirm=True)
        self.assertEqual(confirmed.status_code, 200, confirmed.text)
        confirmed = confirmed.json()
        self.assertTrue(confirmed["task"]["scope_confirmed"])
        self.assertEqual(confirmed["task"]["range_kind"], "month")

        changed = self.scope(confirmed, start_date="2026-02-01", end_date="2026-02-28",
                             range_kind="month")
        self.assertEqual(changed.status_code, 200, changed.text)
        changed = changed.json()
        self.assertFalse(changed["task"]["scope_confirmed"])
        self.assertEqual(changed["task"]["scope_revision"], 1)
        self.assertEqual(changed["task"]["scope_history"][-1]["start_date"], "2026-01-01")
        self.assertEqual(changed["task"]["scope_history"][-1]["end_date"], "2026-01-31")
        self.assertGreater(changed["processing"]["scope"]["outside_records"], 0)
        codes = {i["code"] for i in changed["processing"]["issues"]}
        self.assertTrue({"period_alignment", "record_outside_scope"} <= codes)
        self.assertEqual(changed["versions"][-1]["revision"], changed["revision"])

    def test_existing_company_selection_and_cross_tenant_rejection(self):
        existing = self.store.upsert_client(self.admin, "已有企业", TAX)
        value = {"client_id": existing["id"], "company_name": existing["name"], "tax_no": TAX}
        r = self.client.post("/api/enterprise/input-tasks", json=value)
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.json()["task"]["company_id"], existing["id"])
        self.login(self.foreign)
        self.assertEqual(self.client.post("/api/enterprise/input-tasks", json=value).status_code, 404)
        self.assertEqual(self.scope(r.json(), confirm=True).status_code, 404)
        self.login(self.admin)
        task = r.json()
        bad = self.scope(task, client_id=existing["id"], tax_no=OTHER, confirm=True)
        self.assertEqual(bad.status_code, 422, bad.text)
        self.assertEqual(self.client.get(f"/api/enterprise/input-tasks/{task['id']}").json()["revision"], task["revision"])

    def test_invalid_scope_keeps_existing_client_authorization(self):
        task = self.create()
        r = self.scope(task, tax_no="bad", confirm=True)
        self.assertEqual(r.status_code, 200, r.text)
        self.assertFalse(r.json()["task"]["scope_confirmed"])
        self.assertIn("tax_no", r.json()["scope_errors"])
        self.assertEqual(r.json()["task"]["company_id"], task["task"]["company_id"])
        with self.store.connect() as db:
            batch = material_batches._authorize(db, self.admin, task["id"])
            self.assertEqual(batch["client_id"], task["task"]["company_id"])

    def test_direct_ai_request_cannot_bypass_missing_scope(self):
        task = self.client.post("/api/enterprise/input-tasks", json={}).json()
        with patch("webapp.knowledge_evidence.analyze") as model:
            r = self.client.post("/api/knowledge/retrieve", json={"task_id": task["id"],
                "expected_revision": task["revision"], "rule_id": "VAT-R01", "use_ai": True})
            self.assertEqual(r.status_code, 422, r.text)
            self.assertIn("企业与检查区间", r.text)
            model.assert_not_called()

    def test_mapping_and_ocr_use_admin_extraction_keys_instead_of_other_purposes(self):
        root = self.store.create_user("core-root", "Core-test-2026!", "平台", "platform_admin", "platform")
        self.login(root)
        for usage in ["extraction", "knowledge"]:
            response = self.client.post("/api/admin/ai-keys", json={
                "name": usage, "api_key": "synthetic-" + usage,
                "base_url": "https://" + usage + ".example/v1", "model": usage + "-model",
                "usage": usage, "disable_thinking": False, "vision": True,
            })
            self.assertEqual(response.status_code, 200, response.text)
        self.login(self.admin)
        task = self.create()
        url = "/api/enterprise/input-tasks/" + task["id"] + "/files"
        with patch.dict(os.environ, {"TAXPEARLS_AI_ENABLED": "1", "TAXPEARLS_AI_API_KEY": "synthetic-environment"}):
            answer = json.dumps({"mappings": [{"source_column": 2, "target_field": "declared_value",
                                               "status": "mapped", "explanation": "source heading"}]})
            with patch("src.core_mapping.chat_content", return_value=answer) as mapping:
                response = self.client.post(url, data={"expected_revision": str(task["revision"]),
                    "material_type": "vat_declaration", "mapping_mode": "ai"},
                    files={"files": ("非标准.xlsx", book("vat_declaration", alias="本期申报数额"))})
                self.assertEqual(response.status_code, 200, response.text)
                task = response.json()
                self.assertEqual(task["documents"][0]["records"][0]["declared_value"], "100.0123")
                self.assertEqual(mapping.call_args.args[0].api_keys, ("synthetic-extraction",))
                self.assertEqual(mapping.call_args.args[0].model, "extraction-model")
            answer = json.dumps({"form_name": "增值税纳税申报表主表", "form_type": "main", "tax_no": TAX,
                "period_start": "2026-01-01", "period_end": "2026-01-31", "unit": "元", "quote": "申报主表",
                "rows": [{"item_label": "销售额", "item_key": "1", "declared_value": "100.0123",
                          "column_scope": "本期数", "quote": "销售额 100.0123"}]})
            with patch("src.core_materials.chat_content", return_value=answer) as ocr:
                response = self.client.post(url, data={"expected_revision": str(task["revision"]),
                    "material_type": "vat_declaration", "ocr": "1"}, files={"files": ("扫描.pdf",
                    (Path(__file__).parent / "fixtures" / "materials-scanned.pdf").read_bytes())})
                self.assertEqual(response.status_code, 200, response.text)
                self.assertTrue(any(d.get("ocr_candidate") for d in response.json()["documents"]))
                self.assertEqual(ocr.call_args.args[0].api_keys, ("synthetic-extraction",))
                self.assertEqual(ocr.call_args.args[0].base_url, "https://extraction.example/v1")

    def test_task_upload_reopen_supplement_and_retained_original_no_risk(self):
        task = self.create()
        self.assertEqual(task["processing"]["status"], "needs_completion")
        for kind in ["account_balance", "vat_declaration", "invoice_detail"]:
            task = self.upload(task, kind)
        incomplete = deepcopy(task)
        task = self.upload(task, "invoice_detail", direction="进项")
        self.assertEqual(task["processing"]["status"], "ready_for_analysis", task["processing"]["issues"])
        url = "/api/enterprise/input-tasks/" + task["id"]
        fresh = self.client.get(url).json()
        self.assertEqual(fresh["processing"], task["processing"])
        file = fresh["files"][0]
        r = self.client.get(url + "/originals/" + file["id"])
        self.assertEqual(r.status_code, 200)
        self.assertEqual(__import__("hashlib").sha256(r.content).hexdigest(), file["sha256"])
        with self.store.connect() as db:
            self.assertEqual(db.execute("SELECT COUNT(*) FROM audits").fetchone()[0], 0)
            self.assertEqual(db.execute("SELECT COUNT(*) FROM material_executions").fetchone()[0], 0)
            cipher = db.execute("SELECT bytes_cipher FROM material_originals LIMIT 1").fetchone()[0]
            self.assertNotEqual(cipher[:2], b"PK")
        old = material_batches.read(self.store, self.admin, task["id"])
        self.assertTrue(
            any(v["payload"].get("processing", {}).get("status") == "needs_completion" for v in old["versions"])
        )
        r = self.client.post(
            "/api/enterprise/materials/" + task["id"] + "/confirm", json={"expected_revision": task["revision"]}
        )
        self.assertEqual(r.status_code, 409, r.text)
        r = self.client.post(url + "/prepare", json={"expected_revision": incomplete["revision"]})
        self.assertEqual(r.status_code, 409)

    def test_conflicting_manual_metadata_is_not_ready_and_malformed_json_is_rejected(self):
        task = self.upload(self.create(), "account_balance")
        fid = task["documents"][0]["id"]
        url = "/api/enterprise/input-tasks/" + task["id"] + "/prepare"
        r = self.client.post(
            url,
            json={
                "expected_revision": task["revision"],
                "source_options": {fid: {"confirmed_metadata": {"unit": "万元"}}},
            },
        )
        self.assertEqual(r.status_code, 200, r.text)
        task = r.json()
        self.assertIn("metadata_conflict", {i["code"] for i in task["processing"]["issues"]})
        r = self.client.post(
            url,
            json={"expected_revision": task["revision"], "source_options": {fid: {"confirmed_metadata": {"unit": []}}}},
        )
        self.assertEqual(r.status_code, 422, r.text)
        r = self.client.post(url, content=b"x" * 70000, headers={"Content-Type": "application/json"})
        self.assertEqual(r.status_code, 422, r.text)
        self.assertEqual(
            self.client.get("/api/enterprise/input-tasks/" + task["id"]).json()["revision"], task["revision"]
        )

    def test_tenant_scope_and_invalid_inputs(self):
        task = self.create()
        self.login(self.foreign)
        self.assertEqual(self.client.get("/api/enterprise/input-tasks/" + task["id"]).status_code, 404)
        self.assertEqual(self.client.get("/api/enterprise/input-tasks").json(), [])
        self.login(self.admin)
        r = self.client.post(
            "/api/enterprise/input-tasks",
            json={"company_name": "x", "tax_no": TAX, "start_date": "2026-02-01", "end_date": "2026-01-01"},
        )
        self.assertEqual(r.status_code, 200)
        self.assertIn("end_date", r.json()["scope_errors"])
        r = self.client.post(
            "/api/enterprise/input-tasks/" + task["id"] + "/files",
            data={"expected_revision": str(task["revision"]), "material_type": "account_balance"},
            files={"files": ("假的.xlsx", b"bad")},
        )
        self.assertEqual(r.status_code, 422, r.text)

    def ready_task(self):
        task = self.create()
        for kind in ["account_balance", "vat_declaration", "invoice_detail"]:
            task = self.upload(task, kind)
        return self.upload(task, "invoice_detail", direction="进项")

    @patch("webapp.material_batches.members.now", return_value="2026-01-01T00:00:00+00:00")
    def test_repeated_upload_exclusion_and_restoration_preserve_versions_and_totals(self, _now):
        # All uploads share a timestamp; IDs sort opposite to arrival order.
        # A later duplicate must not replace the earlier source or reorder totals.
        file_ids = iter([letter * 32 for letter in ("f", "e", "d", "c", "0")])
        retain_files = material_batches._files

        def retain_with_ids(*args, **kwargs):
            with patch("webapp.material_batches.secrets.token_hex", side_effect=lambda _: next(file_ids)):
                return retain_files(*args, **kwargs)

        retain_patch = patch("webapp.material_batches._files", side_effect=retain_with_ids)
        retain_patch.start()
        self.addCleanup(retain_patch.stop)
        task = self.ready_task()
        original_summaries = deepcopy(task["processing"]["invoice_summaries"])
        earlier = deepcopy(task)
        url = "/api/enterprise/input-tasks/" + task["id"]
        output_doc = next(
            d
            for d in task["documents"]
            if d["material_type"] == "invoice_detail" and d["records"][0]["direction"] == "销项"
        )
        raw_output = self.client.get(url + "/originals/" + output_doc["id"]).content
        response = self.client.post(
            url + "/files",
            data={"expected_revision": str(task["revision"]), "material_type": "invoice_detail"},
            files={"files": ("重复原件.xlsx", raw_output)},
        )
        self.assertEqual(response.status_code, 200, response.text)
        task = response.json()
        self.assertEqual(task["processing"]["status"], "ready_for_analysis")
        self.assertEqual(task["processing"]["invoice_summaries"], original_summaries)
        self.assertEqual(len(task["files"]), 5)
        self.assertEqual([f["id"] for f in task["files"]], [letter * 32 for letter in ("f", "e", "d", "c", "0")])
        self.assertTrue(any(i["code"] == "duplicate_file" for i in task["processing"]["notices"]))
        input_doc = next(
            d
            for d in task["documents"]
            if d["material_type"] == "invoice_detail" and d["records"][0]["direction"] == "进项"
        )
        url = "/api/enterprise/input-tasks/" + task["id"]
        raw_input = self.client.get(url + "/originals/" + input_doc["id"]).content

        def exclude(value):
            response = self.client.post(
                url + "/prepare",
                json={"expected_revision": task["revision"], "source_options": {input_doc["id"]: {"excluded": value}}},
            )
            self.assertEqual(response.status_code, 200, response.text)
            return response.json()

        task = exclude(True)
        self.assertEqual(task["processing"]["status"], "needs_completion")
        self.assertEqual(sum(g["count"] for g in task["processing"]["invoice_summaries"]), 1)
        task = exclude(False)
        self.assertEqual(task["processing"]["status"], "ready_for_analysis")
        self.assertEqual(task["processing"]["invoice_summaries"], original_summaries)
        versions = material_batches.read(self.store, self.admin, task["id"])["versions"]
        prior = next(v["payload"] for v in versions if v["revision"] == earlier["revision"])
        self.assertEqual(prior["processing"], earlier["processing"])
        self.assertEqual(self.client.get(url + "/originals/" + input_doc["id"]).content, raw_input)

    def test_quarter_task_requires_each_applicable_form_to_cover_each_month(self):
        response = self.client.post(
            "/api/enterprise/input-tasks",
            json={
                "company_name": "季度合成验收企业",
                "tax_no": TAX,
                "start_date": "2026-01-01",
                "end_date": "2026-03-31",
                "applicable_forms": ["附表一"],
            },
        )
        self.assertEqual(response.status_code, 200, response.text)
        task = response.json()
        for start, end in [("2026-01-01", "2026-01-31"), ("2026-02-01", "2026-02-28"), ("2026-03-01", "2026-03-31")]:
            task = self.upload(task, "account_balance", start=start, end=end)
            task = self.upload(task, "vat_declaration", start=start, end=end)
            if start != "2026-02-01":
                task = self.upload(task, "vat_declaration", start=start, end=end, form="附表一")
        self.assertTrue(
            any(
                i["code"] == "form_coverage_gap" and "2026-02-01 至 2026-02-28" in i["message"]
                for i in task["processing"]["issues"]
            )
        )
        self.assertFalse(
            any(i["code"] == "coverage_gap" and "申报表" in i["message"] for i in task["processing"]["issues"])
        )
        task = self.upload(task, "vat_declaration", start="2026-02-01", end="2026-02-28", form="附表一")
        self.assertFalse(
            any(
                i["code"] in {"main_form_gap", "form_coverage_gap", "missing_form"}
                for i in task["processing"]["issues"]
            )
        )
        self.assertEqual(len([r for r in task["processing"]["records"] if r["material_type"] == "vat_declaration"]), 12)

    def test_client_reassignment_revokes_task_download_and_upload_immediately(self):
        first = self.store.create_user(
            "core-accountant-one", "Core-test-2026!", "会计一", "accountant", self.admin["org_id"]
        )
        second = self.store.create_user(
            "core-accountant-two", "Core-test-2026!", "会计二", "accountant", self.admin["org_id"]
        )
        task = self.upload(self.create(), "account_balance")
        self.store.upsert_client(self.admin, task["task"]["company_name"], TAX, first["id"])
        url = "/api/enterprise/input-tasks/" + task["id"]
        fid = task["files"][0]["id"]
        self.login(first)
        self.assertEqual(self.client.get(url).status_code, 200)
        self.assertEqual(self.client.get(url + "/originals/" + fid).status_code, 200)
        self.store.upsert_client(self.admin, task["task"]["company_name"], TAX, second["id"])
        self.assertEqual(self.client.get(url).status_code, 404)
        self.assertEqual(self.client.get(url + "/originals/" + fid).status_code, 404)
        self.assertEqual(self.client.get("/api/enterprise/input-tasks").json(), [])
        with patch("src.core_materials.parse") as parser:
            response = self.client.post(
                url + "/files",
                data={"expected_revision": str(task["revision"]), "material_type": "invoice_detail"},
                files={"files": ("票.xlsx", book("invoice_detail"))},
            )
            self.assertEqual(response.status_code, 404)
            parser.assert_not_called()
        self.login(second)
        self.assertEqual(self.client.get(url).status_code, 200)

    def test_wrong_slot_and_unsafe_xml_leave_no_partial_files_or_revisions(self):
        task = self.create()
        url = "/api/enterprise/input-tasks/" + task["id"]
        for kind, name, raw in [
            ("account_balance", "票.xml", b"<Invoice/>"),
            ("invoice_detail", "恶意.xml", b'<!DOCTYPE x [<!ENTITY e "boom">]><Invoice/>'),
        ]:
            response = self.client.post(
                url + "/files",
                data={"expected_revision": str(task["revision"]), "material_type": kind},
                files={"files": (name, raw)},
            )
            self.assertEqual(response.status_code, 422, response.text)
        fresh = self.client.get(url).json()
        self.assertEqual(fresh["revision"], task["revision"])
        self.assertEqual(fresh["files"], [])

    def test_all_pdf_slots_preserve_original_and_invoice_review_roundtrip(self):
        task = self.create()
        url = "/api/enterprise/input-tasks/" + task["id"]
        for kind, raw in [("account_balance", pdf_table("account_balance")), ("vat_declaration", pdf_table()), ("invoice_detail", invoice_pdf())]:
            response = self.client.post(url+"/files", data={"expected_revision":str(task["revision"]),"material_type":kind,"ocr":"1"}, files={"files":(kind+".pdf",raw)})
            self.assertEqual(response.status_code, 200, response.text)
            task=response.json()
            doc=next(d for d in task["documents"] if d["material_type"]==kind)
            self.assertTrue(doc["records"],doc)
            original=self.client.get(url+"/originals/"+doc["id"])
            self.assertEqual(original.status_code,200)
            self.assertEqual(original.content,raw)
        doc=next(d for d in task["documents"] if d["material_type"]=="invoice_detail")
        file_id,record_id=doc["id"],doc["records"][0]["id"]
        status={record_id:{"red_letter_status":"蓝字","void_status":"正常"}}
        def save(change):
            return self.client.post(url+"/prepare",json={"expected_revision":task["revision"],"source_options":{file_id:change}})
        response=save({"invoice_statuses":status,"source_reviewed":True})
        self.assertEqual(response.status_code,200,response.text)
        task=response.json()
        doc=next(d for d in task["documents"] if d["id"]==file_id)
        self.assertFalse({"pdf_source_review","ocr_review","invoice_status_missing"}&{i["code"] for i in doc["issues"]})
        self.assertEqual(doc["records"][0]["red_letter_status"],"蓝字")
        self.assertIsNone(doc["records"][0]["parsed_invoice_statuses"]["red_letter_status"])
        self.assertFalse(doc["records"][0]["requires_source_review"])
        self.assertIn("invoice_coverage_gap",{i["code"] for i in task["processing"]["issues"]})
        self.assertEqual(self.client.get(url).json()["source_options"][file_id]["invoice_statuses"],status)
        # Clearing a prior confirmation restores the missing-state blocker.
        response=save({"invoice_statuses":{record_id:{"red_letter_status":"","void_status":""}}})
        self.assertEqual(response.status_code,200,response.text)
        task=response.json()
        doc=next(d for d in task["documents"] if d["id"]==file_id)
        self.assertIn("invoice_status_missing",{i["code"] for i in doc["issues"]})
        self.assertTrue(doc["records"][0]["requires_source_review"])
        response=save({"invoice_statuses":{"foreign-record":{"void_status":"正常"}}})
        self.assertEqual(response.status_code,422)

    def test_other_two_scanned_pdf_upload_review_reopen_and_revoke(self):
        raw=(Path(__file__).parent/"fixtures"/"materials-scanned.pdf").read_bytes()
        balance={"account_name":"现金","opening_balance":"10","current_debit":"100","current_credit":"20","closing_balance":"90"}
        for kind in ["account_balance","vat_declaration"]:
            task=self.create()
            url="/api/enterprise/input-tasks/"+task["id"]
            answer={"tax_no":TAX,"period_start":"2026-01-01","period_end":"2026-01-31","unit":"元","quote":"来源表头"}
            if kind=="account_balance":
                answer["rows"]=[{"fields":balance,"quotes":balance}]
            else:
                answer.update(form_name="增值税纳税申报表主表",form_type="main",rows=[{"item_label":"应税销售额","declared_value":"100","column_scope":"本期数","quote":"应税销售额 100"}])
            with self.subTest(kind=kind),patch.dict(os.environ,{"TAXPEARLS_AI_ENABLED":"1","TAXPEARLS_AI_API_KEY":"synthetic"}),patch("src.core_materials.chat_content",return_value=json.dumps(answer)) as model:
                response=self.client.post(url+"/files",data={"expected_revision":str(task["revision"]),"material_type":kind,"ocr":"1"},files={"files":(kind+".pdf",raw)})
                self.assertEqual(response.status_code,200,response.text)
                task=response.json();doc=task["documents"][0];fid=doc["id"]
                self.assertEqual(self.client.get(url+"/originals/"+fid).content,raw)
                self.assertTrue(doc["records"][0]["requires_source_review"])
                for reviewed in [True,False]:
                    response=self.client.post(url+"/prepare",json={"expected_revision":task["revision"],"source_options":{fid:{"source_reviewed":reviewed}}})
                    self.assertEqual(response.status_code,200,response.text)
                    task=response.json();doc=task["documents"][0]
                    self.assertEqual(doc["records"][0]["requires_source_review"],not reviewed)
                    self.assertEqual("ocr_review" in {i["code"] for i in doc["issues"]},not reviewed)
                    self.assertEqual(self.client.get(url).json()["documents"][0]["records"],doc["records"])
                self.assertEqual(model.call_count,1)

    def test_scanned_invoice_review_uses_cache_and_revokes_without_model(self):
        task=self.create()
        url="/api/enterprise/input-tasks/"+task["id"]
        fields={"invoice_no":"12345678901234567890","issue_date":"2026-01-15","buyer_tax_no":OTHER,"seller_tax_no":TAX,"amount_excluding_tax":"100","tax_amount":"13"}
        answer={"unit":"元","rows":[{"fields":fields,"quotes":fields}]}
        raw=(Path(__file__).parent/"fixtures"/"materials-scanned.pdf").read_bytes()
        with patch.dict(os.environ,{"TAXPEARLS_AI_ENABLED":"1","TAXPEARLS_AI_API_KEY":"synthetic"}),patch("src.core_materials.chat_content",return_value=json.dumps(answer)) as model:
            response=self.client.post(url+"/files",data={"expected_revision":str(task["revision"]),"material_type":"invoice_detail","ocr":"1"},files={"files":("扫描票.pdf",raw)})
            self.assertEqual(response.status_code,200,response.text)
            task=response.json();doc=task["documents"][0]
            self.assertTrue(doc["ocr_candidate"])
            rid=doc["records"][0]["id"]
            for status,reviewed in [({"red_letter_status":"蓝字","void_status":"正常"},True),({"red_letter_status":"","void_status":""},False)]:
                response=self.client.post(url+"/prepare",json={"expected_revision":task["revision"],"source_options":{doc["id"]:{"invoice_statuses":{rid:status},"source_reviewed":reviewed}}})
                self.assertEqual(response.status_code,200,response.text)
                task=response.json();doc=task["documents"][0]
                self.assertEqual(doc["records"][0]["requires_source_review"],not reviewed)
                self.assertEqual("invoice_status_missing" in {i["code"] for i in doc["issues"]},not reviewed)
            self.assertEqual(model.call_count,1)

    def test_ocr_review_cache_revoke_and_metadata_restore_without_more_calls(self):
        task = self.ready_task()
        vat = next(d for d in task["documents"] if d["material_type"] == "vat_declaration")
        url = "/api/enterprise/input-tasks/" + task["id"]
        response = self.client.post(
            url + "/prepare",
            json={"expected_revision": task["revision"], "source_options": {vat["id"]: {"excluded": True}}},
        )
        self.assertEqual(response.status_code, 200, response.text)
        task = response.json()
        answer = {
            "form_name": "增值税纳税申报表主表",
            "form_type": "main",
            "tax_no": TAX,
            "period_start": "2026-01-01",
            "period_end": "2026-01-31",
            "unit": "元",
            "quote": "申报主表",
            "rows": [
                {
                    "item_label": "销售额",
                    "item_key": "1",
                    "declared_value": "100.0123",
                    "column_scope": "本期数",
                    "quote": "销售额 100.0123",
                }
            ],
        }
        with (
            patch.dict(os.environ, {"TAXPEARLS_AI_ENABLED": "1", "TAXPEARLS_AI_API_KEY": "synthetic"}),
            patch("src.core_materials.chat_content", return_value=json.dumps(answer)) as call,
        ):
            response = self.client.post(
                url + "/files",
                data={"expected_revision": str(task["revision"]), "material_type": "vat_declaration", "ocr": "1"},
                files={
                    "files": ("扫描.pdf", (Path(__file__).parent / "fixtures" / "materials-scanned.pdf").read_bytes())
                },
            )
            self.assertEqual(response.status_code, 200, response.text)
            task = response.json()
            scan = next(d for d in task["documents"] if d["name"] == "扫描.pdf")
            fid = scan["id"]
            self.assertEqual(call.call_count, 1)
            self.assertEqual(task["processing"]["status"], "needs_completion")
            response = self.client.post(
                url + "/prepare",
                json={"expected_revision": task["revision"], "source_options": {fid: {"source_reviewed": True}}},
            )
            self.assertEqual(response.status_code, 200, response.text)
            task = response.json()
            self.assertEqual(task["processing"]["status"], "ready_for_analysis")
            self.assertEqual(call.call_count, 1)
            response = self.client.post(
                url + "/prepare",
                json={
                    "expected_revision": task["revision"],
                    "source_options": {fid: {"confirmed_metadata": {"unit": "万元"}}},
                },
            )
            self.assertEqual(response.status_code, 200, response.text)
            task = response.json()
            self.assertEqual(task["processing"]["status"], "needs_completion")
            self.assertTrue(any(i["code"] == "ocr_review" for i in task["processing"]["issues"]))
            response = self.client.post(
                url + "/prepare",
                json={
                    "expected_revision": task["revision"],
                    "source_options": {fid: {"confirmed_metadata": {}, "source_reviewed": True}},
                },
            )
            self.assertEqual(response.status_code, 200, response.text)
            task = response.json()
            scan = next(d for d in task["documents"] if d["id"] == fid)
            self.assertEqual(scan["records"][0]["unit"], "元")
            self.assertEqual(call.call_count, 1)

    def test_changed_invoice_and_red_void_values_never_auto_adjust_or_overwrite(self):
        from openpyxl import load_workbook

        task = self.ready_task()
        url = "/api/enterprise/input-tasks/" + task["id"]
        wb = load_workbook(BytesIO(book("invoice_detail")))
        s = wb.active
        s.cell(6, 6, "99")
        out = BytesIO()
        wb.save(out)
        wb.close()
        response = self.client.post(
            url + "/files",
            data={"expected_revision": str(task["revision"]), "material_type": "invoice_detail"},
            files={"files": ("更正票.xlsx", out.getvalue())},
        )
        self.assertEqual(response.status_code, 200, response.text)
        task = response.json()
        self.assertTrue(any(i["code"] == "invoice_conflict" for i in task["processing"]["issues"]))
        original = next(
            d
            for d in task["documents"]
            if d["material_type"] == "invoice_detail" and d["records"][0]["tax_amount"] == "13.0016"
        )
        self.assertEqual(original["records"][0]["raw_record"]["5"], "13.0016")
        wb = load_workbook(BytesIO(book("invoice_detail")))
        s = wb.active
        s.cell(6, 1, "RED-NEW-ID")
        s.cell(6, 5, "-50.01")
        s.cell(6, 6, "-6.50")
        s.cell(6, 7, "红字")
        s.cell(6, 8, "作废")
        out = BytesIO()
        wb.save(out)
        wb.close()
        response = self.client.post(
            url + "/files",
            data={"expected_revision": str(task["revision"]), "material_type": "invoice_detail"},
            files={"files": ("红字作废.xlsx", out.getvalue())},
        )
        self.assertEqual(response.status_code, 200, response.text)
        task = response.json()
        group = next(g for g in task["processing"]["invoice_summaries"] if g["red_letter_status"] == "红字")
        self.assertEqual(group["amount_excluding_tax"], "-50.01")
        self.assertEqual(group["tax_amount"], "-6.50")
        self.assertEqual(group["void_status"], "作废")
        with self.store.connect() as db:
            self.assertEqual(db.execute("SELECT COUNT(*) FROM audits").fetchone()[0], 0)


if __name__ == "__main__":
    unittest.main()
