"""Actual upload/retention/review/check/report paths for VAT-R01, no policy approval shortcut."""
from copy import deepcopy
from io import BytesIO
import unittest
from unittest.mock import patch
from openpyxl import load_workbook

from tests import test_kg_retrieval as kg
from tests.test_core_inputs import book, TAX, OTHER
from src import vat_reconciliation as vat
from src.kg_retrieval import outbound
from webapp import knowledge_evidence as service


def invoices(rows):
    workbook = load_workbook(BytesIO(book("invoice_detail")))
    sheet = workbook.active
    sheet.delete_rows(5, sheet.max_row)
    sheet.append(["票号", "发票明细行号", "开具日期", "销方税号", "购方税号", "金额", "税额", "红字状态", "作废状态", "进销项"])
    for number, line, amount, red, status in rows:
        sheet.append([number, line, "2026-01-01", TAX, OTHER, amount, "0", red, status, "销项"])
    content = BytesIO()
    workbook.save(content)
    workbook.close()
    return content.getvalue()


@kg.material_key
class VATReconciliationTests(unittest.TestCase):
    setUp = kg.RetrievalHTTPTests.setUp
    login = kg.RetrievalHTTPTests.login
    task = kg.RetrievalHTTPTests.task
    def upload(self, task, content, name="发票.xlsx", kind="invoice_detail"):
        response = self.client.post("/api/enterprise/input-tasks/"+task["id"]+"/files",
                    data={"expected_revision":task["revision"],"material_type":kind}, files={"files":(name,content)})
        self.assertEqual(response.status_code,200,response.text)
        return response.json()

    def review(self, task, **overrides):
        row = next(r for r in task["processing"]["records"] if r["material_type"]=="vat_declaration" and r["item_key"]=="1")
        values = {"expected_revision":task["revision"],"mapping_record_id":row["id"],"form_revision":"DEMO-2026-A",
                  "taxpayer_category":"合成一般纳税人","comparable":True,"basis":"合成同主体同期间不含税口径，已核对调整事项。",**overrides}
        response = self.client.post("/api/enterprise/input-tasks/"+task["id"]+"/vat-review",json=values)
        self.assertEqual(response.status_code,200,response.text)
        return self.client.get("/api/enterprise/input-tasks/"+task["id"]).json()

    def run_check(self, task):
        response = self.client.post("/api/knowledge/retrieve",json={"rule_id":vat.RULE,"task_id":task["id"],
                         "expected_revision":task["revision"],"region_code":"CN-31"})
        self.assertEqual(response.status_code,200,response.text)
        return response.json()

    def test_vat_r01_retained_inputs_to_dual_state_and_frozen_report(self):
        task = self.upload(self.task(), invoices([("I1","L1","200.0123","蓝字","正常")]))
        before = self.run_check(task)
        self.assertEqual(before["bundle"]["fact_assessment"]["status"],"FACT_INCOMPLETE")
        task = self.client.get("/api/enterprise/input-tasks/"+task["id"]).json()
        task = self.review(task)
        self.assertTrue(task['vat_review_current'])
        result = self.run_check(task)
        facts = result["bundle"]["fact_assessment"]
        self.assertEqual(facts["status"],"FACT_CONFIRMED")
        self.assertEqual(facts["amount"],"100.0000")
        self.assertEqual(result["bundle"]["risk_assessment"]["status"],"NEEDS_REVIEW")
        self.assertEqual(result["bundle"]["check"]["review_status"],"candidate")
        url="/api/enterprise/input-tasks/"+task["id"]+"/reports"
        report=self.client.post(url,json={"expected_revision":result["revision"]})
        self.assertEqual(report.status_code,200,report.text)
        html=self.client.get(url+"/1?format=html")
        self.assertIn("已核实差额",html.text)
        self.assertIn("100.0000",html.text)
        self.assertIn("风险仍待核查",html.text)
        self.login("kg-other")
        self.assertEqual(self.client.get(url+"/1?format=html").status_code,404)

    def test_multiline_red_void_duplicate_and_unknown_status(self):
        workbook=load_workbook(BytesIO(book("vat_declaration")))
        workbook.active['C6']='1300';content=BytesIO();workbook.save(content);workbook.close()
        # This case tests net amounts; cross-task declaration conflicts have their own coverage.
        created=self.client.post('/api/enterprise/input-tasks',json={'company_name':'纯合成测试净额','tax_no':TAX,'start_date':'2026-01-01','end_date':'2026-01-31','applicable_forms':[]}).json()
        task=self.upload(created,content.getvalue(),'主表.xlsx','vat_declaration')
        task=self.upload(task,invoices([('I5','L1','1000','蓝字','正常'),('I5','L2','500','蓝字','正常'),('I6','L1','-200','红字','正常'),('I7','L1','999','蓝字','作废'),('I5','L2','500','蓝字','正常')]))
        self.assertFalse(any(i['code']=='invoice_conflict' for i in task['processing']['issues']))
        result=self.run_check(self.review(task))
        self.assertEqual(result['bundle']['fact_assessment']['amount'],'0')
        self.assertEqual(result['bundle']['risk_assessment']['status'],'NO_RISK_SIGNAL')
        task=self.client.get('/api/enterprise/input-tasks/'+task['id']).json()
        task=self.upload(task,invoices([('I8','L1','1','蓝字','unknown')]),'未知状态.xlsx')
        result=self.run_check(self.review(task))
        self.assertEqual(result['bundle']['fact_assessment']['status'],'FACT_UNVERIFIED')
        self.assertIsNone(next(f for f in result['bundle']['facts'] if f['metric_id']==vat.INVOICE)['value'])

    def test_not_comparable_material_update_invalidates_review(self):
        task=self.upload(self.task(),invoices([('I1','L1','100.0123','蓝字','正常')]))
        task=self.review(task,comparable=False)
        result=self.run_check(task)
        self.assertEqual(result['bundle']['fact_assessment']['status'],'NOT_COMPARABLE')
        task=self.client.get('/api/enterprise/input-tasks/'+task['id']).json()
        task=self.review(task,comparable=True)
        result=self.run_check(task)
        self.assertEqual(result['bundle']['fact_assessment']['amount'],'0.0000')
        task=self.client.get('/api/enterprise/input-tasks/'+task['id']).json()
        task=self.upload(task,invoices([('I2','L1','1','蓝字','正常')]),'补传.xlsx')
        self.assertFalse(task['vat_review_current'])
        self.assertTrue(task['tax_checks'][vat.RULE]['stale'])
        result=self.run_check(task)
        self.assertEqual(result['bundle']['fact_assessment']['status'],'FACT_INCOMPLETE')

    def test_review_rejects_foreign_source_and_version(self):
        task=self.task()
        values={'expected_revision':task['revision'],'mapping_record_id':'not-accessible','form_revision':'DEMO','taxpayer_category':'DEMO','comparable':True,'basis':'DEMO'}
        url='/api/enterprise/input-tasks/'+task['id']+'/vat-review'
        self.assertEqual(self.client.post(url,json=values).status_code,422)
        values['expected_revision']=0
        self.assertEqual(self.client.post(url,json=values).status_code,409)

    def test_negative_disclaimer_does_not_mask_positive_legal_claim(self):
        for text in ['不确认违法','不能确认违法','无法据此确认违法','不代表全面合规']:
            self.assertFalse(service.legal_claim(text),text)
        for text in ['确定违法','不确认违法，但确定违法','不代表全面合规；应补缴税款']:
            self.assertTrue(service.legal_claim(text),text)

    def test_outbound_readiness_and_authorization_are_independent(self):
        task=self.task();result=self.client.post('/api/knowledge/retrieve',json={'rule_id':'R08','task_id':task['id'],'expected_revision':task['revision'],'region_code':'CN-33'}).json()
        package=result['bundle']
        delivery=outbound.prepare(package,authorized=True,destination='https://api.deepseek.com')
        self.assertEqual(delivery['outbound_gate'],'allowed')
        self.assertEqual(delivery['analysis_mode'],'review_only')
        self.assertEqual(outbound.prepare(package,authorized=False,destination='https://api.deepseek.com')['outbound_gate'],'blocked')
        poisoned=deepcopy(package)
        poisoned['batching']['parts'][0]['policy_leads'].append({'text':'91310000123456789A'})
        self.assertEqual(outbound.prepare(poisoned,authorized=True,destination='https://api.deepseek.com')['outbound_gate'],'blocked')

    def test_invalid_ai_preserves_confirmed_difference(self):
        task=self.review(self.upload(self.task(),invoices([('I1','L1','200.0123','蓝字','正常')])))
        model,_=kg.semantic_search.available()
        gate={'passed':True,'model_fingerprints':[model.fingerprint if model else None]}
        with patch.object(service,'acceptance_gate',return_value=gate),patch('webapp.knowledge._chat_json',return_value=({'assessments':[]},'unit-provider')):
            response=self.client.post('/api/knowledge/retrieve',json={'rule_id':vat.RULE,'task_id':task['id'],
                'expected_revision':task['revision'],'region_code':'CN-31','use_ai':True})
        self.assertEqual(response.status_code,200,response.text)
        result=response.json();self.assertFalse(result['ai']['validated'])
        self.assertEqual(result['bundle']['fact_assessment']['status'],'FACT_CONFIRMED')
        self.assertEqual(result['bundle']['fact_assessment']['amount'],'100.0000')
        self.assertEqual(result['ai']['risk_assessment']['status'],'NEEDS_REVIEW')

    def test_zero_proof_cannot_mask_another_unparsed_original(self):
        from tests import vat_v12 as fixture
        scenario=next(s for s in fixture.data('demo_business_scenarios.json') if s['scenario_id']=='S05')
        task=self.task()
        # The original January table is retained and independently confirmed.
        raw=fixture.workbook(scenario,'invoice_detail',period=('2026-01-01','2026-01-31'))
        task=self.upload(task,raw,'空明细.xlsx')
        proof=task['documents'][-1]['id'];task=self.review(task,zero_proof_file_id=proof)
        payload=deepcopy(task)
        payload['documents'].append({'id':'unparsed','material_type':'invoice_detail','parts':[],
                                     'records':[],'issues':[{'code':'parse_failed'}]})
        payload['check_reviews'][vat.RULE]['source_signature']=vat.source_signature(payload)
        with self.store.connect() as db:
            definitions=kg.relation_search.rule_context(db,vat.RULE)['metrics']
        self.assertEqual(vat.resolve(definitions,task)[1][vat.INVOICE]['value'],'0.00')
        self.assertEqual(vat.resolve(definitions,payload)[1][vat.INVOICE]['state'],'unknown')
        workbook=load_workbook(BytesIO(raw))
        workbook.create_sheet('未识别明细').append(['无法识别的业务字段','仍有金额 999'])
        content=BytesIO();workbook.save(content);workbook.close()
        task=self.upload(task,content.getvalue(),'空表加未识别页.xlsx')
        self.assertTrue(any(i['code']=='unreadable_sheet' for i in task['documents'][-1]['issues']))
        task=self.review(task,zero_proof_file_id=proof)
        result=self.run_check(task)
        self.assertEqual(next(f for f in result['bundle']['facts'] if f['metric_id']==vat.INVOICE)['state'],'unknown')
