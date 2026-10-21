"""Cross-task reuse, conservative identity, versions and immutable evidence."""
from copy import deepcopy
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
import json
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from src import core_dedup as dedup
from tests import test_core_inputs as core
from webapp import material_batches


TAX = "91310000123456789A"
TASK = {"id": "task", "tenant_id": "org", "company_id": "client", "tax_no": TAX}


def invoice(identifier="a", amount="100", **changes):
    return {"id": identifier, "material_type": "invoice_detail", "source_file_id": identifier,
            "source_position": "Sheet1!A6", "invoice_no": "12345678901234567890",
            "seller_tax_no": TAX, "buyer_tax_no": "91310000987654321B", "issue_date": "2026-01-01",
            "amount_excluding_tax": amount, "tax_amount": "13", "unit": "元", "direction": "output",
            "red_letter_status": "蓝字", "void_status": "正常", "in_scope": True, **changes}


class IdentityTests(unittest.TestCase):
    def test_same_business_different_file_row_and_decimal_format(self):
        rows = [invoice(), invoice("b", "100.0000", source_position="other!A15")]
        result, issues = dedup.classify(rows, TASK)
        self.assertEqual(result["counts"]["task_duplicate"], 1)
        self.assertEqual(rows[1]["duplicate_of"], "a")
        self.assertEqual(issues, [])
        processing = {"records": rows, "issues": []}
        dedup.refresh(processing)
        self.assertEqual(processing["invoice_summaries"][0]["amount_excluding_tax"], "100")

    def test_amount_is_a_version_not_a_new_business_key(self):
        rows = [invoice(), invoice("b", "120")]
        result, issues = dedup.classify(rows, TASK)
        self.assertEqual(result["counts"]["conflict"], 2)
        self.assertEqual(len({r["record_key"] for r in result["records"]}), 1)
        self.assertEqual(len({r["version_key"] for r in result["records"]}), 2)
        self.assertEqual(issues[0]["code"], "record_version_conflict")

    def test_blank_zero_large_decimal_line_and_red_invoice_are_distinct(self):
        self.assertNotEqual(dedup.describe(invoice(tax_amount=None), TASK), dedup.describe(invoice(tax_amount="0"), TASK))
        self.assertEqual(dedup.number("1234567890123456789012345678901234567890.1200"),
                         "1234567890123456789012345678901234567890.12")
        rows = [invoice(invoice_line_no="001"), invoice("b", invoice_line_no="002"),
                invoice("red", "-100", invoice_no="22345678901234567890", red_letter_status="红字")]
        result, _ = dedup.classify(rows, TASK)
        self.assertEqual(result["counts"]["new"], 3)

    def test_missing_identity_does_not_guess_from_date_or_amount(self):
        for changes in ({"invoice_no": "12345678"}, {"seller_tax_no": None},
                        {"invoice_no": None}, {"requires_source_review": True}, {"identity_incomplete": True}):
            self.assertIsNone(dedup.describe(invoice(**changes), TASK))
        self.assertIsNotNone(dedup.describe(invoice(invoice_no="00000001", invoice_code="012345678901"), TASK))

    def test_account_period_ledger_currency_and_auxiliary_dimensions_are_identity(self):
        row = {"material_type": "account_balance", "account_code": "01001", "account_name": "现金",
               "period_start": "2026-01-01", "period_end": "2026-01-31", "closing_balance": "100", "unit": "元"}
        original = dedup.describe(row, TASK)
        for changes in ({"period_end": "2026-03-31"}, {"ledger_id": "账套2"}, {"currency": "USD"},
                        {"auxiliary_code": "客户2"}, {"account_code": "1001"}):
            self.assertNotEqual(original[0], dedup.describe({**row, **changes}, TASK)[0])
        self.assertEqual(original[0], dedup.describe({**row, "closing_balance": "200"}, TASK)[0])

    def test_declaration_form_column_period_and_line_must_match(self):
        row = {"material_type": "vat_declaration", "period_start": "2026-01-01", "period_end": "2026-01-31",
               "form_type": "main", "form_name": "主表", "item_key": "1", "column_scope": "本期数", "declared_value": "100"}
        original = dedup.describe(row, TASK)
        for changes in ({"form_name": "附表一"}, {"column_scope": "本年累计"}, {"item_key": "2"}):
            self.assertNotEqual(original[0], dedup.describe({**row, **changes}, TASK)[0])
        self.assertIsNone(dedup.describe({**row, "column_scope": None}, TASK))

    def test_historical_reuse_still_included_in_current_task(self):
        past, _ = dedup.classify([invoice()], TASK)
        entry = past["records"][0]
        history = {entry["record_key"]: [{"task_id": "completed", "version_key": entry["version_key"]}]}
        rows = [invoice("new-upload")]
        result, issues = dedup.classify(rows, TASK, history=history)
        self.assertEqual(result["counts"]["historical_duplicate"], 1)
        self.assertEqual(result["counts"]["effective"], 1)
        self.assertFalse(rows[0].get("duplicate_of"))
        self.assertEqual(issues, [])

    def test_explicit_choice_covers_only_observed_versions(self):
        rows = [invoice(), invoice("b", "120")]
        result, _ = dedup.classify(rows, TASK)
        a, b = result["records"]
        choices = {a["record_key"]: {"selected": b["version_key"], "versions": [a["version_key"], b["version_key"]]}}
        decided, issues = dedup.classify(rows, TASK, decisions=choices)
        self.assertEqual(decided["counts"]["superseded"], 1)
        self.assertEqual(issues, [])
        rows.append(invoice("c", "130"))
        reopened, _ = dedup.classify(rows, TASK, decisions=choices)
        self.assertEqual(reopened["counts"]["conflict"], 3)

    def test_apply_is_idempotent_and_does_not_edit_source_documents(self):
        documents = [invoice(), invoice("b")]
        processing = {"records": deepcopy(documents), "issues": [], "status": "ready_for_analysis"}
        dedup.apply(processing, TASK)
        before = deepcopy(processing)
        dedup.apply(processing, TASK)
        self.assertEqual(processing, before)
        self.assertNotIn("dedup", documents[0])

    def test_display_only_history_labels_do_not_invalidate_linked_material(self):
        from src.rule_bindings import material_signature
        processing = {"records": [invoice()]}
        original = material_signature(processing)
        processing["records"][0]["dedup"] = {"record_key": "opaque", "version_key": "version", "status": "new"}
        self.assertEqual(material_signature(processing), original)
        processing["records"][0]["dedup"]["status"] = "historical_duplicate"
        self.assertEqual(material_signature(processing), original)
        processing["records"][0]["dedup_conflict"] = True
        self.assertNotEqual(material_signature(processing), original)


class RecordIndexAPITests(unittest.TestCase):
    setUp = core.CoreAPITests.setUp
    login = core.CoreAPITests.login
    create = core.CoreAPITests.create
    upload = core.CoreAPITests.upload

    def test_empty_draft_and_unadopted_period_never_appear_ready(self):
        response = self.client.post("/api/enterprise/input-tasks", json={})
        self.assertEqual(response.status_code, 200, response.text)
        task = response.json()
        self.assertEqual(task["processing"]["status"], "needs_completion")
        task = self.upload(task, "account_balance")
        self.assertEqual(task["task"]["start_date"], "")
        self.assertIn("start_date", task["scope_errors"])
        self.assertEqual(task["processing"]["status"], "needs_completion")

    def read(self, task):
        response = self.client.get("/api/enterprise/input-tasks/" + task["id"])
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    def changed_invoice(self, task, amount="120", number=None):
        book = load_workbook(BytesIO(core.book("invoice_detail")))
        book.active.cell(6, 5, amount)
        if number:
            book.active.cell(6, 1, number)
        stream = BytesIO()
        book.save(stream)
        book.close()
        response = self.client.post("/api/enterprise/input-tasks/" + task["id"] + "/files",
            data={"expected_revision": str(task["revision"]), "material_type": "invoice_detail"},
            files={"files": ("后续上传.xlsx", stream.getvalue())})
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    def test_completed_history_new_task_reuses_and_old_revision_stays_exact(self):
        first = self.upload(self.create(), "invoice_detail")
        view = material_batches.read(self.store, self.admin, first["id"])
        payload = deepcopy(view["versions"][-1]["payload"])
        payload["completion"] = {"report_version": 1, "content_hash": "frozen", "completed_at": "2026-10-20"}
        with self.store.connect() as db:
            db.execute("BEGIN IMMEDIATE")
            batch = material_batches._authorize(db, self.admin, first["id"])
            material_batches._append(db, batch, self.admin, "confirmation", payload, lifecycle="complete")
        old = self.read(first)
        second = self.upload(self.create(), "invoice_detail")
        self.assertEqual(second["processing"]["deduplication"]["counts"]["historical_duplicate"], 1)
        self.assertEqual(second["processing"]["invoice_summaries"][0]["count"], 1)
        self.assertFalse(second["processing"]["records"][0].get("duplicate_of"))
        self.assertTrue(second["processing"]["deduplication"]["records"][0]["historical_sources"][0]["completed"])
        self.changed_invoice(self.create())
        self.assertEqual(self.read(first), old)
        with self.store.connect() as db:
            self.assertEqual(db.execute("SELECT COUNT(*) FROM core_business_records").fetchone()[0], 1)
            self.assertEqual(db.execute("SELECT COUNT(*) FROM core_business_versions").fetchone()[0], 2)

    def test_all_three_materials_are_deduplicated_across_files_and_tasks(self):
        for kind, count in (("invoice_detail", 1), ("account_balance", 1), ("vat_declaration", 2)):
            with self.subTest(kind=kind):
                first = self.upload(self.create(), kind)
                duplicate = self.upload(first, kind)
                counts = duplicate["processing"]["deduplication"]["counts"]
                self.assertEqual(counts["task_duplicate"], count)
                self.assertEqual(len(duplicate["processing"]["records"]), count * 2)
                self.assertFalse(any(i["code"] == "duplicate_period" for i in duplicate["processing"]["issues"]))
                later = self.upload(self.create(), kind)
                self.assertEqual(later["processing"]["deduplication"]["counts"]["historical_duplicate"], count)

    def test_conflict_choice_is_revision_and_observation_bound_and_audited(self):
        first = self.upload(self.create(), "invoice_detail")
        second = self.changed_invoice(self.create())
        self.assertEqual(second["processing"]["deduplication"]["counts"]["conflict"], 1)
        self.assertEqual(second["processing"]["invoice_summaries"], [])
        self.assertEqual(self.read(first)["processing"]["deduplication"]["counts"]["conflict"], 1)
        entry = second["processing"]["deduplication"]["records"][0]
        data = {"expected_revision": second["revision"], "source_record_id": entry["source_record_id"],
                "expected_versions": entry["conflict_versions"], "confirm": True}
        url = "/api/enterprise/input-tasks/" + second["id"] + "/deduplication/resolve"
        self.assertEqual(self.client.post(url, json={**data, "expected_revision": second["revision"] - 1}).status_code, 409)
        self.assertEqual(self.client.post(url, json={**data, "confirm": False}).status_code, 422)
        self.assertEqual(self.client.post(url, json={**data, "expected_versions": ["0" * 64]}).status_code, 409)
        response = self.client.post(url, json=data)
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["processing"]["deduplication"]["counts"]["conflict"], 0)
        self.assertEqual(response.json()["processing"]["invoice_summaries"][0]["amount_excluding_tax"], "120")
        with self.store.connect() as db:
            self.assertEqual(db.execute("SELECT COUNT(*) FROM material_events WHERE batch_id=? AND action='dedup_version_choice'", (second["id"],)).fetchone()[0], 1)
        third = self.changed_invoice(self.create(), "130")
        self.assertEqual(third["processing"]["deduplication"]["counts"]["conflict"], 1)
        self.assertEqual(self.read(second)["processing"]["deduplication"]["counts"]["conflict"], 1)

    def test_exclusion_removes_only_its_source_membership(self):
        first = self.upload(self.create(), "invoice_detail")
        second = self.upload(self.create(), "invoice_detail")
        file_id = first["documents"][0]["id"]
        response = self.client.post("/api/enterprise/input-tasks/" + first["id"] + "/prepare",
            json={"expected_revision": first["revision"], "source_options": {file_id: {"excluded": True}}})
        self.assertEqual(response.status_code, 200, response.text)
        third = self.upload(self.create(), "invoice_detail")
        histories = third["processing"]["deduplication"]["records"][0]["historical_sources"]
        self.assertEqual({r["task_id"] for r in histories}, {second["id"]})

    def test_source_delete_and_task_delete_do_not_remove_other_task_evidence(self):
        first = self.upload(self.create(), "invoice_detail")
        second = self.upload(self.create(), "invoice_detail")
        file_id = first["documents"][0]["id"]
        material_batches.delete_original(self.store, self.admin, first["id"], file_id, first["revision"])
        self.assertEqual(self.client.get("/api/enterprise/input-tasks/" + second["id"] + "/originals/" + second["documents"][0]["id"]).status_code, 200)
        third = self.upload(self.create(), "invoice_detail")
        history = third["processing"]["deduplication"]["records"][0]["historical_sources"]
        self.assertEqual({h["task_id"] for h in history}, {second["id"]})
        response = self.client.request("DELETE", "/api/enterprise/input-tasks/" + second["id"],
            json={"expected_revision": second["revision"], "confirm": True})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(self.read(third)["processing"]["deduplication"]["counts"]["historical_duplicate"], 0)

    def test_tenant_client_assignment_and_index_are_isolated(self):
        first = self.upload(self.create(), "invoice_detail")
        self.login(self.foreign)
        second = self.upload(self.create(), "invoice_detail")
        self.assertEqual(second["processing"]["deduplication"]["counts"]["historical_duplicate"], 0)
        self.assertNotEqual(first["processing"]["records"][0]["dedup"]["record_key"],
                            second["processing"]["records"][0]["dedup"]["record_key"])
        self.assertIn(self.client.get("/api/enterprise/input-tasks/" + first["id"]).status_code, (403, 404))
        self.login(self.admin)
        accountant = self.store.create_user("dedup-accountant", "Core-test-2026!", "会计", "accountant", self.admin["org_id"])
        self.login(accountant)
        self.assertIn(self.client.get("/api/enterprise/input-tasks/" + first["id"]).status_code, (403, 404))
        self.store.upsert_client(self.admin, first["task"]["company_name"], TAX, accountant["id"])
        self.assertEqual(self.read(first)["id"], first["id"])

    def test_union_uses_records_once_and_does_not_inherit_unselected_tasks(self):
        a = self.upload(self.create(), "invoice_detail")
        b = self.upload(self.create(), "invoice_detail")
        self.changed_invoice(self.create(), number="22345678901234567890")
        response = self.client.post("/api/enterprise/input-tasks/record-union", json={"task_ids": [a["id"], b["id"]]})
        self.assertEqual(response.status_code, 200, response.text)
        value = response.json()["processing"]
        self.assertEqual(value["deduplication"]["counts"]["task_duplicate"], 1)
        self.assertEqual(value["invoice_summaries"][0]["count"], 1)
        self.assertEqual(value["invoice_summaries"][0]["amount_excluding_tax"], "100.0123")
        self.assertEqual(len(value["records"]), 2)

    def test_legacy_index_backfill_preserves_completed_payload_and_is_not_limited_to_history_page(self):
        first = self.upload(self.create(), "invoice_detail")
        with self.store.connect() as db:
            db.execute("DELETE FROM core_business_sources")
            db.execute("DELETE FROM core_business_index_tasks")
        second = self.upload(self.create(), "invoice_detail")
        self.assertEqual(second["processing"]["deduplication"]["counts"]["historical_duplicate"], 1)
        with self.store.connect() as db:
            self.assertEqual(db.execute("SELECT revision FROM material_batches WHERE id=?", (first["id"],)).fetchone()[0], first["revision"])

    def test_index_and_payload_roll_back_together_on_audit_failure(self):
        first = self.upload(self.create(), "invoice_detail")
        event = material_batches._event
        def failed(db, batch, actor, action, detail, **kw):
            if action == "edit":
                raise RuntimeError("synthetic audit failure")
            return event(db, batch, actor, action, detail, **kw)
        view = material_batches.read(self.store, self.admin, first["id"])
        payload = deepcopy(view["versions"][-1]["payload"])
        payload["processing"]["records"][0]["amount_excluding_tax"] = "999"
        with patch("webapp.material_batches._event", side_effect=failed), self.assertRaises(RuntimeError):
            material_batches.append_revision(self.store, self.admin, first["id"], first["revision"], "edit", payload)
        with self.store.connect() as db:
            self.assertEqual(db.execute("SELECT COUNT(*) FROM core_business_versions").fetchone()[0], 1)
        self.assertEqual(self.read(first)["revision"], first["revision"])

    def test_index_has_no_plaintext_tax_ids_amounts_or_names(self):
        self.upload(self.create(), "invoice_detail")
        with self.store.connect() as db:
            rows = [list(r) for r in db.execute("SELECT record_key,scope_key FROM core_business_records")]
            self.assertNotIn(TAX, json.dumps(rows))
            self.assertTrue(all(len(v) == 64 for r in rows for v in r))

    def test_history_change_is_revalidated_before_analysis_even_without_task_revision_change(self):
        from webapp import core_record_index
        from webapp.access import AccessDenied
        first = self.upload(self.create(), "invoice_detail")
        payload = deepcopy(material_batches.read(self.store, self.admin, first["id"])["versions"][-1]["payload"])
        self.changed_invoice(self.create(), "120")
        with self.store.connect() as db:
            db.execute("BEGIN IMMEDIATE")
            batch = material_batches._authorize(db, self.admin, first["id"])
            self.assertEqual(batch["revision"], first["revision"])
            with self.assertRaises(AccessDenied) as rejected:
                core_record_index.require_current(db, batch, self.admin, payload)
            self.assertEqual(rejected.exception.status, 409)

    def test_more_than_one_hundred_tasks_backfill_oldest_source(self):
        first = self.upload(self.create(), "invoice_detail")
        with self.store.connect() as db:
            for n in range(101):
                identifier = "empty-history-" + str(n)
                db.execute("INSERT INTO material_batches VALUES (?,?,?,?,0,'now')",
                           (identifier, self.admin["org_id"], self.admin["id"], first["task"]["company_id"]))
                db.execute("INSERT INTO core_input_tasks VALUES (?)", (identifier,))
            db.execute("DELETE FROM core_business_sources")
            db.execute("DELETE FROM core_business_index_tasks")
        second = self.upload(self.create(), "invoice_detail")
        self.assertEqual(second["processing"]["deduplication"]["counts"]["historical_duplicate"], 1)
        self.assertEqual(second["processing"]["deduplication"]["records"][0]["historical_sources"][0]["task_id"], first["id"])

    def test_concurrent_tasks_share_one_business_identity_and_revalidate_reads(self):
        first = self.upload(self.create(), "invoice_detail")
        second = self.upload(self.create(), "invoice_detail")
        def write(task, amount):
            view = material_batches.read(self.store, self.admin, task["id"])
            payload = deepcopy(view["versions"][-1]["payload"])
            payload["processing"]["records"][0]["amount_excluding_tax"] = amount
            return material_batches.append_revision(self.store, self.admin, task["id"], task["revision"], "edit", payload)
        with ThreadPoolExecutor(max_workers=2) as pool:
            results = list(pool.map(lambda pair: write(*pair), [(first, "100"), (second, "120")]))
        self.assertEqual(results, [first["revision"] + 1, second["revision"] + 1])
        with self.store.connect() as db:
            self.assertEqual(db.execute("SELECT COUNT(*) FROM core_business_records").fetchone()[0], 1)
        self.assertEqual(self.read(first)["processing"]["deduplication"]["counts"]["conflict"], 1)
        self.assertEqual(self.read(second)["processing"]["deduplication"]["counts"]["conflict"], 1)

    def test_other_company_same_amount_is_not_duplicate(self):
        self.upload(self.create(), "account_balance")
        response = self.client.post("/api/enterprise/input-tasks", json={"company_name": "另一合成企业", "tax_no": core.OTHER,
            "start_date": "2026-01-01", "end_date": "2026-01-31", "applicable_forms": []})
        self.assertEqual(response.status_code, 200, response.text)
        other = self.upload(response.json(), "account_balance", tax=core.OTHER)
        self.assertEqual(other["processing"]["deduplication"]["counts"]["historical_duplicate"], 0)

    def test_worksheet_rename_does_not_change_declaration_identity(self):
        first = self.upload(self.create(), "vat_declaration")
        book = load_workbook(BytesIO(core.book("vat_declaration")))
        book.active.title = "导出工作表"
        stream = BytesIO(); book.save(stream); book.close()
        response = self.client.post("/api/enterprise/input-tasks/" + first["id"] + "/files",
            data={"expected_revision": first["revision"], "material_type": "vat_declaration"},
            files={"files": ("另一个文件名.xlsx", stream.getvalue())})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["processing"]["deduplication"]["counts"]["task_duplicate"], 2)

    def test_report_freeze_survives_later_conflicting_upload_and_new_check_is_unknown(self):
        from tests.test_kg_retrieval import CoreReportTests
        helper = CoreReportTests()
        helper.setUp()
        try:
            tid, checked = helper.checked()
            self.assertEqual(next(f for f in checked["bundle"]["facts"] if f["metric_id"] == "M01")["state"], "known")
            base = "/api/enterprise/input-tasks/" + tid
            report = helper.client.post(base + "/reports", json={"expected_revision": checked["revision"]})
            self.assertEqual(report.status_code, 200, report.text)
            frozen = helper.client.get(base + "/reports/1").content
            completed = helper.client.post(base + "/complete", json={"expected_revision": checked["revision"], "report_version": 1})
            self.assertEqual(completed.status_code, 200, completed.text)
            created = helper.client.post("/api/enterprise/input-tasks", json={"company_name": "纯合成测试企业", "tax_no": TAX,
                "start_date": "2026-01-01", "end_date": "2026-01-31", "applicable_forms": []}).json()
            workbook = load_workbook(BytesIO(core.book("vat_declaration")))
            workbook.active["C6"] = "200.0123"
            stream = BytesIO(); workbook.save(stream); workbook.close()
            uploaded = helper.client.post("/api/enterprise/input-tasks/" + created["id"] + "/files",
                data={"expected_revision": created["revision"], "material_type": "vat_declaration"},
                files={"files": ("更正申报.xlsx", stream.getvalue())})
            self.assertEqual(uploaded.status_code, 200, uploaded.text)
            task = uploaded.json()
            self.assertEqual(task["processing"]["deduplication"]["counts"]["conflict"], 1)
            checked_again = helper.client.post("/api/knowledge/retrieve", json={"rule_id": "R01", "task_id": task["id"],
                "expected_revision": task["revision"], "region_code": "CN-33"})
            self.assertEqual(checked_again.status_code, 200, checked_again.text)
            self.assertEqual(next(f for f in checked_again.json()["bundle"]["facts"] if f["metric_id"] == "M01")["state"], "unknown")
            self.assertEqual(helper.client.get(base + "/reports/1").content, frozen)
            self.assertEqual(helper.client.get(base).json()["completion"], completed.json()["completion"])
            report_json = helper.client.get(base + "/reports/1?format=json").json()
            self.assertTrue(all("business_record" in source for source in report_json["sources"]))
        finally:
            helper.doCleanups()
