"""Bounded, in-memory material import. PDF candidates always require review."""
from __future__ import annotations

from dataclasses import asdict
from decimal import Decimal
from hashlib import sha256
from io import BytesIO
from pathlib import PurePosixPath
import re
from zipfile import ZipFile, BadZipFile

import pdfplumber
from openpyxl import load_workbook

from . import config, loader
from .models import Account, Company, Dataset, Metric
from .ai_extraction import ExtractionError

InputError = loader.InputError
MAX_FILE = 10 * 1024 * 1024
MAX_TOTAL = 50 * 1024 * 1024
MAX_FILES = 20
COMPANY_KEYS = ("name", "taxpayer_id", "industry", "period")


def _text(value):
    return "" if value is None else str(value).strip()


def _serial(value):
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, dict):
        return {k: _serial(v) for k, v in value.items()}
    if isinstance(value, list):
        return [_serial(v) for v in value]
    return value


def expand_uploads(files):
    """Never extract archives to disk; validate all entries before decompressing."""
    if not files or len(files) > MAX_FILES:
        raise InputError("每次请选择 1–20 个文件。")
    if sum(len(data) for _, data in files) > MAX_TOTAL:
        raise InputError("上传总大小不能超过 50MB。")
    output, total = [], 0
    for name, data in files:
        if not data or len(data) > MAX_FILE:
            raise InputError(f"{name}：文件为空或超过 10MB。")
        name = name.replace("\\", "/").split("/")[-1]
        if not name.lower().endswith(".zip"):
            output.append((name, data))
            total += len(data)
        else:
            try:
                with ZipFile(BytesIO(data)) as archive:
                    entries = archive.infolist()
                    if len(entries) > 100:
                        raise InputError(f"{name}：ZIP 条目过多。")
                    candidates = []
                    for item in entries:
                        path = PurePosixPath(item.filename.replace("\\", "/"))
                        if path.is_absolute() or ".." in path.parts or ":" in item.filename:
                            raise InputError(f"{name}：ZIP 含不安全路径。")
                        if item.is_dir():
                            continue
                        if item.flag_bits & 1 or (item.external_attr >> 16) & 0o170000 == 0o120000:
                            raise InputError(f"{name}：不支持加密文件或符号链接。")
                        if item.filename.lower().endswith(".zip"):
                            raise InputError(f"{name}：不支持嵌套 ZIP。")
                        if item.file_size > MAX_FILE or item.file_size > max(1, item.compress_size) * 250:
                            raise InputError(f"{name}：解压大小或压缩比例超限。")
                        candidates.append(item)
                    if not candidates:
                        raise InputError(f"{name}：ZIP 中没有文件。")
                    total += sum(i.file_size for i in candidates)
                    if total > MAX_TOTAL or len(output) + len(candidates) > MAX_FILES:
                        raise InputError("解压后最多 20 个文件、合计 50MB。")
                    for item in candidates:
                        output.append((f"{name}/{item.filename}", archive.read(item)))
            except (BadZipFile, RuntimeError, NotImplementedError) as exc:
                raise InputError(f"{name}：ZIP 无法读取或已损坏。") from exc
        if total > MAX_TOTAL or len(output) > MAX_FILES:
            raise InputError("解压后最多 20 个文件、合计 50MB。")
    return output


def _excel(data, doc):
    # XLSX itself is a ZIP. Bound expansion before openpyxl allocates XML trees.
    with ZipFile(BytesIO(data)) as archive:
        if len(archive.infolist()) > 1000 or sum(i.file_size for i in archive.infolist()) > MAX_TOTAL:
            raise InputError("Excel 解压内容超过限制。")
    wb = load_workbook(BytesIO(data), data_only=False)
    try:
        if any(ws.max_row > 30000 or ws.max_column > 100 for ws in wb):
            raise InputError("每张 Excel 工作表最多 30000 行、100 列。")
        company = Company("", "", "", "")
        if config.SHEET_COMPANY in wb.sheetnames:
            company = loader._read_company(wb)
        elif config.SHEET_SUPPLEMENT in wb.sheetnames:
            ws = loader._sheet(wb, config.SHEET_SUPPLEMENT, config.COL_SUPPLEMENT)
            periods = {_text(row[3]) for row in ws.iter_rows(min_row=2, values_only=True) if row[0] is not None}
            periods.discard("")
            if len(periods) != 1:
                raise InputError("补充指标必须具有唯一的核对所属期。")
            company.period = periods.pop()
        accounts = loader._read_accounts(wb) if config.SHEET_ACCOUNTS in wb.sheetnames else []
        declarations = loader._read_declarations(wb) if config.SHEET_DECLARATION in wb.sheetnames else {}
        metrics = {}
        for sheet in (config.SHEET_INCOME, config.SHEET_BALANCE, config.SHEET_CASHFLOW):
            loader._read_statement(wb, sheet, sheet, metrics)
        loader._read_history(wb, metrics)
        loader._read_invoices(wb, metrics)
        loader._read_bank(wb, metrics)
        loader._read_supplement(wb, company, metrics)
        if not accounts and not declarations and not metrics:
            raise InputError("没有找到支持的账表；请保留标准工作表名称和列名。")
        doc.update(company=asdict(company), accounts=_serial([asdict(a) for a in accounts]),
                   declarations=_serial(declarations), rows=[_serial(asdict(m)) for m in metrics.values()])
        doc["summary"] = f"{len(accounts)} 行科目、{len(declarations)} 项申报、{len(metrics)} 项补充/报表指标"
    finally:
        wb.close()


def _pdf(data, doc, keys):
    aliases = dict(zip(config.COMPANY_FIELDS, COMPANY_KEYS))
    aliases.update({"纳税人名称": "name", "编制单位": "name", "统一社会信用代码": "taxpayer_id", "税号": "taxpayer_id"})
    seen = set()
    with pdfplumber.open(BytesIO(data)) as pdf:
        if not 1 <= len(pdf.pages) <= 50:
            raise InputError("PDF 须为 1–50 页。")
        doc["page_count"] = len(pdf.pages)
        for page_no, page in enumerate(pdf.pages, 1):
            text = (page.extract_text() or "")
            if len(text) > 30000:
                raise InputError(f"PDF 第 {page_no} 页文字过多。")
            doc["pages"].append({"page": page_no, "text": text})
            if not text.strip():
                doc["warnings"].append(f"第 {page_no} 页无可提取文字，可能为扫描页；请对照原件手工录入，当前未接入 OCR。")
                continue
            for line in text.splitlines():
                for label, key in aliases.items():
                    match = re.fullmatch(r"\s*" + re.escape(label) + r"\s*[:：]\s*(.+?)\s*", line)
                    if match:
                        value = match[1]
                        if doc["company"][key] and doc["company"][key] != value:
                            raise InputError(f"PDF 中出现不一致的{label}，请拆分为同一企业、同一期间的材料。")
                        doc["company"][key] = value
            # Only extract an explicit single value, never choose between amount columns.
            prefix = next((p for p in ("利润表", "资产负债表", "现金流量表") if p in text[:300]), "")
            is_vat = "增值税" in text[:300]
            candidates = []
            for line in text.splitlines():
                match = re.fullmatch(r"\s*([^:：]+?)\s*[:：]\s*([-+]?\d[\d,]*(?:\.\d+)?)\s*", line)
                if match:
                    candidates.append((match[1].strip(), match[2], line))
            for table in page.extract_tables():
                if not table or len(table[0]) != 2:
                    continue
                header = [_text(v).replace("\n", "") for v in table[0]]
                if header[0] not in {"项目", "指标"} or header[1] not in {"本期金额", "金额", "数值"}:
                    continue
                candidates.extend((_text(row[0]).replace("\n", ""), _text(row[1]), " | ".join(_text(v) for v in row)) for row in table[1:] if len(row) == 2)
            # Units must be reviewed; do not infer or scale amounts automatically.
            if "万元" in text or "千元" in text:
                doc["warnings"].append(f"第 {page_no} 页含非元单位；候选数值保留原数，提交前必须换算为元。")
            for label, value, excerpt in candidates:
                key = f"{prefix}.{label}" if prefix and f"{prefix}.{label}" in keys else f"增值税.{label}" if is_vat and f"增值税.{label}" in keys else label if label in keys else ""
                if key not in keys or (page_no, key, value) in seen:
                    continue
                try:
                    number = loader._number(value.replace(",", ""), f"PDF 第 {page_no} 页")
                except InputError:
                    continue
                if number is None:
                    continue
                seen.add((page_no, key, value))
                doc["rows"].append({"name": key, "value": str(number), "page": page_no, "detail": excerpt})
        if len(doc["rows"]) > 500:
            raise InputError("PDF 候选指标超过 500 项，请拆分材料。")
    doc["warnings"].insert(0, "PDF 结果是待核对候选值：请核实企业、期间、单位、金额列及指标口径；未识别项目可手工添加。金额统一为元，空白保持缺失。")
    doc["summary"] = f"{doc['page_count']} 页 PDF，识别到 {len(doc['rows'])} 项候选指标"


def _unmapped_excel(data, doc):
    """Expose cell addresses for AI without bypassing standard-table validation."""
    known = {config.SHEET_ACCOUNTS, config.SHEET_COMPANY, config.SHEET_DECLARATION,
             config.SHEET_INCOME, config.SHEET_BALANCE, config.SHEET_CASHFLOW,
             config.SHEET_HISTORY, config.SHEET_INVOICES, config.SHEET_BANK, config.SHEET_SUPPLEMENT}
    with ZipFile(BytesIO(data)) as archive:
        if sum(i.file_size for i in archive.infolist()) > MAX_TOTAL or len(archive.infolist()) > 1000:
            raise InputError("Excel 解压内容超过限制。")
    wb = load_workbook(BytesIO(data), data_only=False, read_only=True)
    try:
        if known.intersection(wb.sheetnames):
            raise InputError("标准工作表校验失败，请先修正格式/重复项；AI 不覆盖原始输入错误。")
        for page_no, ws in enumerate(wb, 1):
            if ws.max_row > 2000 or ws.max_column > 50:
                raise InputError("AI Excel 单表最多 2000 行、50 列，请拆分。")
            lines = []
            for row in ws:
                cells = []
                for cell in row:
                    if cell.data_type == "f":
                        raise InputError("AI Excel 含未计算公式，请先导出固定数值快照。")
                    if cell.value is not None:
                        cells.append(f"{cell.coordinate}={cell.value}")
                if cells:
                    lines.append(" | ".join(cells))
            text = "\n".join(lines)
            if len(text) > 60000:
                raise InputError("AI Excel 单表内容过长，请拆分。")
            doc["pages"].append({"page": page_no, "text": text, "label": f"工作表：{ws.title}"})
        doc["page_count"] = len(doc["pages"])
    finally:
        wb.close()


def preview(files, keys, extractor=None):
    docs = []
    for index, (name, data) in enumerate(expand_uploads(files)):
        doc = {"id": str(index), "name": name, "fingerprint": sha256(data).hexdigest()[:16],
               "kind": "pdf" if name.lower().endswith(".pdf") else "xlsx", "company": dict.fromkeys(COMPANY_KEYS, ""),
               "accounts": [], "declarations": {}, "rows": [], "pages": [], "warnings": [], "error": "",
               "extraction": {"method": "local"}}
        try:
            if name.lower().endswith(".xlsx"):
                try:
                    _excel(data, doc)
                except InputError:
                    if extractor is None:
                        raise
                    _unmapped_excel(data, doc)
                    extractor.enrich(doc, data)
            elif name.lower().endswith(".pdf"):
                _pdf(data, doc, keys)
                if extractor is not None:
                    try:
                        extractor.enrich(doc, data)
                    except ExtractionError as exc:
                        doc["extraction"] = {"method": "ai_failed"}
                        doc["warnings"].insert(0, f"AI 未完成：{exc} 当前仅显示本地解析候选，需人工核对或重新上传。")
            else:
                raise InputError("不支持此文件类型；请选择 .xlsx 或 .pdf（也可放入 ZIP）。")
        except InputError as exc:
            doc["error"] = str(exc)
        except ExtractionError as exc:
            doc["error"] = f"AI 提取失败：{exc}"
            doc["extraction"] = {"method": "ai_failed"}
        except Exception:
            doc["error"] = "文件无法解析；请检查是否损坏、加密或格式不符。"
        docs.append(doc)
    return docs


def _merge_value(target, key, value, location):
    if key in target and target[key] != value:
        raise InputError(f"{location}：{key} 存在冲突值；请核对后选择一份正确材料，不能自动相加或覆盖。")
    target[key] = value


def build_dataset(documents, selections, company_override, keys):
    """Validate edits server-side and merge same-scope evidence without summation."""
    company_data, accounts, declarations, metrics = {}, {}, {}, {}
    account_sources, declaration_sources = {}, {}
    for doc in documents:
        selection = selections[doc["id"]]
        if doc["error"]:
            raise InputError(f"{doc['name']}：{doc['error']}")
        company = dict(doc["company"])
        edits = selection.get("company", {})
        if not isinstance(edits, dict):
            raise InputError("企业信息格式错误。")
        for key in COMPANY_KEYS:
            val = _text(edits.get(key, company[key]))
            if len(val) > 200:
                raise InputError("企业信息过长。")
            editable = doc["kind"] == "pdf" or doc.get("review_required", False)
            if not editable and company[key] and val != company[key]:
                raise InputError("Excel 中已有企业信息不能在核对页改写。")
            company[key] = val or _text(company_override.get(key, ""))
            if company[key]:
                _merge_value(company_data, key, company[key], "企业或期间不一致")
        source = f"{doc['name']} [SHA256:{doc['fingerprint']}]"
        for row in doc["accounts"]:
            account = Account(row["code"], row["name"], *(loader._number(row[k], source) for k in ("opening", "debit", "credit", "closing")))
            _merge_value(accounts, account.code, account, source)
            account_sources.setdefault(account.code, []).append(source)
        for key, val in doc["declarations"].items():
            _merge_value(declarations, key, loader._number(val, source), source)
            declaration_sources.setdefault(key, []).append(source)
        rows = doc["rows"]
        if editable:
            if selection.get("reviewed") is not True:
                raise InputError(f"{doc['name']}：请先核对提取结果的期间、单位和金额口径。")
            rows = selection.get("rows", rows)
            if not isinstance(rows, list) or len(rows) > 500:
                raise InputError("PDF 指标列表格式错误或超过 500 项。")
        for row in rows:
            if not isinstance(row, dict):
                raise InputError("指标行格式错误。")
            key = _text(row.get("name"))
            value = loader._number(row.get("value"), source)
            if value is None:
                continue
            if editable and key not in keys:
                raise InputError(f"{key}：请选择规则支持的标准指标。")
            if key.startswith("人力.") and key.endswith("人数") and (value < 0 or value != value.to_integral_value()):
                raise InputError("人数必须为非负整数。")
            detail = _text(row.get("detail"))
            if len(detail) > 2000:
                raise InputError("口径说明超过 2000 字。")
            if editable:
                page = row.get("page")
                if type(page) is not int or not 1 <= page <= doc["page_count"] or not detail:
                    raise InputError("PDF 指标须填写有效页码和口径说明。")
                origin = f"PDF 第 {page} 页" if doc["kind"] == "pdf" else doc["pages"][page - 1]["label"]
                row_source = f"{source} / {origin}（人工核对/录入）"
                original = [r for r in doc["rows"] if r["name"] == key and r["page"] == page]
                detail += "；识别候选原值：" + ("、".join(r["value"] for r in original) or "无，人工补录")
                if doc.get("extraction", {}).get("method") == "ai":
                    row_source += f"；AI 提取模型 {doc['extraction']['model']}"
                    detail += "；AI 原始证据：" + "；".join(f"{r.get('ai_raw_value')} {r.get('ai_unit')}；{r.get('ai_quote')}" for r in original)
            else:
                row_source = f"{source} / {row['source']}"
            if key in metrics:
                if doc["kind"] != "pdf" and row.get("source", "").startswith(("发票明细", "银行流水")):
                    raise InputError(f"指标「{key}」来自多份明细汇总；请先合并原始明细并去重，再上传一份完整清单。")
                if metrics[key].value != value:
                    raise InputError(f"指标「{key}」在材料中存在冲突值，请核对后重选材料。")
                metrics[key].source += "；" + row_source
            else:
                metrics[key] = Metric(key, value, row_source, detail)
    missing = [label for key, label in zip(COMPANY_KEYS, config.COMPANY_FIELDS) if not company_data.get(key)]
    if missing:
        raise InputError("请补充企业信息：" + "、".join(missing))
    derived = loader._build_metrics(list(accounts.values()), declarations)
    for key, metric in derived.items():
        if key in config.ACCOUNT_MAP:
            sources = [s for code in config.ACCOUNT_MAP[key]["accounts"] for s in account_sources[code]]
        else:
            sources = declaration_sources[config.DECLARATION_ITEMS[key]]
        metric.source = "；".join(dict.fromkeys(sources)) + " / " + metric.source
        if key in metrics:
            if metric.value != metrics[key].value:
                raise InputError(f"指标「{key}」与原始账表计算值冲突，请核对。")
            metric.source += "；" + metrics[key].source
        metrics[key] = metric
    if not metrics:
        raise InputError("没有可执行核对的指标；请补充至少一个有效指标，缺失值不会当成零。")
    return Dataset(Company(*(company_data[k] for k in COMPANY_KEYS)), list(accounts.values()), declarations, metrics)
