"""Repair the collected Guangdong source package; never open a business database.

Run with --package PATH --fetch once to retain official attachments, then use
--offline to reproduce text, dates, source checks and conservative review queues.
Collection and machine screening never approve policy applicability.
"""
from __future__ import annotations

import argparse
from collections import Counter
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date, datetime, timedelta, timezone
import hashlib
from html.parser import HTMLParser
import io
import json
from pathlib import Path
import re
import struct
import urllib.parse
import urllib.request
import xml.etree.ElementTree as ET
import zipfile


VERSION = 3
CST = timezone(timedelta(hours=8))
ACTIVE = {"有效", "全文有效", "常年有效"}
EXPIRED = {"无效", "失效", "全文失效", "已废止失效"}
PARTIAL = {"部分有效", "条款失效"}
TAXES = ["土地增值税", "增值税", "消费税", "企业所得税", "个人所得税", "房产税",
         "城镇土地使用税", "耕地占用税", "契税", "车船税", "水资源税", "资源税",
         "环境保护税", "印花税", "城市维护建设税", "车辆购置税", "烟叶税", "船舶吨税", "关税"]
ALIASES = {"环保税": "环境保护税", "城建税": "城市维护建设税", "个税": "个人所得税"}
FEES = ["地方教育附加", "教育费附加", "社会保险费", "水土保持补偿费", "残疾人就业保障金"]
FILE = re.compile(r"\.(pdf|docx?|xlsx?|zip|rar)(?:$|[?#])", re.I)
VOID = {"area", "base", "br", "col", "embed", "hr", "img", "input", "link", "meta", "param", "source", "wbr"}
DAY = r"(\d{4})年(\d{1,2})月(\d{1,2})日"
DAY_RE = re.compile(DAY)
SELF = r"本(?:通知|公告|办法|决定|细则|规定|规程|意见)"
NUMBER = {str(i): i for i in range(10)} | dict(zip("零一二三四五六七八九", range(10))) | {"〇": 0}


def sha(data):
    return hashlib.sha256(data).hexdigest()


def read_json(path):
    return json.loads(path.read_text(encoding="utf-8-sig"))


def write_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def checked_path(package, relative):
    path = (package / relative).resolve()
    if Path(relative).is_absolute() or not path.is_relative_to(package.resolve()) or path == package.resolve():
        raise ValueError(f"来源路径越界：{relative}")
    return path


def read_verified(package, relative, checksum):
    data = checked_path(package, relative).read_bytes()
    if not checksum or sha(data) != checksum:
        raise ValueError(f"留存来源哈希不一致：{relative}")
    return data


class Node:
    def __init__(self, tag="root", attrs=()):
        self.tag, self.attrs, self.children = tag, dict(attrs), []

    def nodes(self):
        yield self
        for child in self.children:
            if isinstance(child, Node):
                yield from child.nodes()

    def text(self):
        if self.tag in {"script", "style", "noscript"}:
            return ""
        if self.tag == "br":
            return "\n"
        if self.tag == "tr":
            cells = [c for c in self.children if isinstance(c, Node) and c.tag in {"td", "th"}]
            return "\t".join(re.sub(r"\s+", " ", c.text()).strip() for c in cells) + "\n"
        sep = "\t" if self.tag == "tr" else ""
        result = sep.join(c.text() if isinstance(c, Node) else c for c in self.children)
        return result + ("\n" if self.tag in {"p", "div", "tr", "li", "table", "h1", "h2", "h3"} else "")


class DOM(HTMLParser):
    def __init__(self, html):
        super().__init__(convert_charrefs=True)
        self.root = Node()
        self.stack = [self.root]
        self.feed(html)

    def handle_starttag(self, tag, attrs):
        node = Node(tag, attrs)
        self.stack[-1].children.append(node)
        if tag not in VOID:
            self.stack.append(node)

    def handle_startendtag(self, tag, attrs):
        self.handle_starttag(tag, attrs)
        if tag not in VOID:
            self.handle_endtag(tag)

    def handle_endtag(self, tag):
        for i in range(len(self.stack) - 1, 0, -1):
            if self.stack[i].tag == tag:
                del self.stack[i:]
                return

    def handle_data(self, data):
        self.stack[-1].children.append(data)


def clean(text):
    lines = [re.sub(r"[ \r\f\v\u00a0\u3000]+", " ", x).strip() for x in text.splitlines()]
    return "\n".join(x for x in lines if x)


def html_parts(html):
    root = DOM(html).root
    body = next((n for n in root.nodes() if n.attrs.get("id") == "zoomcon"), None)
    text = clean(body.text()) if body else ""
    attachments = []
    for node in root.nodes():
        href = node.attrs.get("href", "")
        if node.tag == "a" and FILE.search(href):
            attachments.append({"link": href, "label": clean(node.text()), "kind": "attachment"})
    images = []
    if body:
        for node in body.nodes():
            src = node.attrs.get("src", "")
            if node.tag == "img" and src and not re.search(r"(?:icon\d*|sysimage|spacer|logo)", src, re.I):
                images.append({"link": src, "label": node.attrs.get("alt", ""), "kind": "body_image"})
    return text, attachments, images


def tax_tags(title):
    probe = title.replace("有关税收", "").replace("有关税务", "")
    result = []
    for name in TAXES:
        if name in probe:
            result.append(name)
            probe = probe.replace(name, "")
    for alias, name in ALIASES.items():
        if alias in probe and name not in result:
            result.append(name)
    return result


def document_kind(title):
    if any(x in title for x in ["清理结果", "失效废止", "废止和修改", "全文废止", "废止三份"]) or ("规范性文件" in title and any(x in title for x in ["清理", "废止", "失效"])):
        return "cleanup"
    if any(x in title for x in ["指引", "指南", "汇编", "名单", "征集", "春风行动", "若干措施"]):
        return "reference"
    return "policy_candidate"


def normal_dates(text):
    # Digit-by-digit Chinese dates, not an unanchored conversion of all numerals.
    def replace(match):
        try:
            year = "".join(str(NUMBER[c]) for c in match[1])
            def num(s):
                if "十" in s:
                    a, b = s.split("十")
                    return (NUMBER[a] if a else 1) * 10 + (NUMBER[b] if b else 0)
                return int("".join(str(NUMBER[c]) for c in s))
            return f"{year}年{num(match[2])}月{num(match[3])}日"
        except (KeyError, ValueError):
            return match[0]
    return re.sub(r"([〇零一二三四五六七八九]{4})年([一二三四五六七八九十]{1,3})月([一二三四五六七八九十]{1,3})日", replace, text)


def iso(groups):
    try:
        return date(*map(int, groups)).isoformat()
    except ValueError:
        return None


def add_years_end(start, years):
    d = date.fromisoformat(start)
    try:
        end = d.replace(year=d.year + years)
    except ValueError:
        end = d.replace(year=d.year + years, day=28)
    return (end - timedelta(days=1)).isoformat()


def extract_dates(text, publish_date, as_of, page_validity=None, *, kind="policy_candidate"):
    result = {"effective_from": None, "effective_from_source": None, "effective_from_evidence": None,
              "effective_to": None, "effective_to_evidence": None, "effective_to_source": None,
              "applicable_periods": [], "repeals": [], "validity_observed": page_validity,
              "status_as_of": "unknown", "as_of": as_of, "gaps": []}
    sentences = [x for x in re.split(r"(?<=[。；])", re.sub(r"\s+", "", normal_dates(text))) if x]
    if kind in {"policy_candidate", "cleanup"}:
        for sentence in sentences:
            # Only document-self assertions. Never borrow a date from a quoted
            # law in a compilation or from a superseded document's title.
            own = re.search(SELF, sentence)
            if not own:
                continue
            assertion = sentence[own.start():]
            start = re.search(SELF + r"(?:自|从)" + DAY + r"(?:[（(]税款所属(?:时期|期)[）)])?起(?:开始)?(?:施行|执行|实施|试行)", assertion)
            period = re.search(SELF + r"(?:执行期限|有效期|适用期)(?:为|自)" + DAY + r"至" + DAY, assertion)
            if period:
                a, b = iso(period.groups()[:3]), iso(period.groups()[3:])
                if a and b:
                    result.update(effective_from=a, effective_to=b, effective_from_source="text_period",
                                  effective_to_source="text_period", effective_from_evidence=sentence,
                                  effective_to_evidence=sentence)
            elif start and not result["effective_from"]:
                result.update(effective_from=iso(start.groups()), effective_from_source="text_assertion",
                              effective_from_evidence=sentence)
                tail = assertion[start.end():]
                end = re.match(r"至" + DAY, tail)
                if end:
                    result.update(effective_to=iso(end.groups()), effective_to_source="text_period",
                                  effective_to_evidence=sentence)
            duration = re.search(r"(?:有效期(?:为)?|试行)([一二三四五六七八九十\d]+)年", assertion)
            if duration and result["effective_from"] and not result["effective_to"]:
                n = duration[1]
                years = int(n) if n.isdigit() else NUMBER.get(n)
                if years:
                    result.update(effective_to=add_years_end(result["effective_from"], years),
                                  effective_to_source="duration_calculation", effective_to_evidence=sentence)
                    result["gaps"].append("duration_end_requires_review")
            # Relative enactment is supported only when the same sentence
            # explicitly ties this document to its issue date.
            if re.search(SELF + r"自(?:制发|发布|公布)之日起施行", assertion) and publish_date:
                result.update(effective_from=publish_date, effective_from_source="relative_issue_date",
                              effective_from_evidence=sentence)
                if duration and not result["effective_to"]:
                    n = duration[1]
                    years = int(n) if n.isdigit() else NUMBER.get(n)
                    if years:
                        result.update(effective_to=add_years_end(publish_date, years),
                                      effective_to_source="duration_calculation", effective_to_evidence=sentence)
                        result["gaps"].append("duration_end_requires_review")
        if "延续实施至" in text:
            match = re.search(r"延续实施至" + DAY, normal_dates(re.sub(r"\s+", "", text)))
            if match:
                result.update(effective_to=iso(match.groups()), effective_to_source="extension_assertion",
                              effective_to_evidence=next(s for s in sentences if "延续实施至" in s))
    for s in sentences:
        if any(x in s for x in ["废止", "停止执行", "失效"]):
            result["repeals"].append({"evidence": s, "review_status": "needs_review", "target_scope": "unresolved"})
        if any(x in s for x in ["执行期限", "所属期", "试行", "参照", "起执行至", "有效期", "申报期"]):
            for m in re.finditer(DAY + r"(?:起执行)?至" + DAY, s):
                a, b = iso(m.groups()[:3]), iso(m.groups()[3:])
                if a and b:
                    result["applicable_periods"].append({"from": a, "to": b, "evidence": s})
    if page_validity in EXPIRED:
        result["status_as_of"] = "expired"
    elif page_validity in PARTIAL:
        result["status_as_of"] = "partial"
    elif result["effective_from"] and result["effective_from"] > as_of:
        result["status_as_of"] = "pending"
    elif result["effective_to"] and result["effective_to"] < as_of:
        result["status_as_of"] = "expired"
    # A page's active label is an observation, never current legal approval.
    for key in ["effective_from", "effective_to"]:
        if not result[key]:
            result["gaps"].append(key + "_unknown")
    return result


def official_url(url):
    parts = urllib.parse.urlsplit(url)
    host = (parts.hostname or "").lower()
    return parts.scheme in {"http", "https"} and (host.endswith(".gov.cn") or host.endswith(".chinatax.gov.cn"))


def fetch(url):
    if not official_url(url):
        raise ValueError("附件不是官方政府域名")
    encoded = urllib.parse.quote(url, safe=":/?=&%#;+@")
    request = urllib.request.Request(encoded, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(request, timeout=18) as response:
        if not official_url(response.geturl()):
            raise ValueError("附件重定向至非官方域名")
        data = response.read(32 * 1024 * 1024 + 1)
        if len(data) > 32 * 1024 * 1024:
            raise ValueError("附件超过32MiB，需单独核对")
        if b"<html" in data[:512].lower() and FILE.search(url):
            raise ValueError("附件返回HTML而不是文件")
        return data


def doc_piece_text(word, table):
    """Read Word 97+ text via its FIB/CLX piece table, never byte-string recovery.

    MS-DOC 2.4.1: this retains logical text order. Table layout, fields, text
    boxes and rendered pages still require visual review; no macros execute.
    """
    def u16(data, at):
        return struct.unpack_from('<H', data, at)[0]
    if len(word) < 34 or u16(word, 0) != 0xA5EC or u16(word, 2) < 0xC1:
        raise ValueError('不是支持的 Word97+ 主文档流。')
    if u16(word, 10) & 0x0100:
        raise ValueError('Word document is encrypted')
    at = 34 + 2 * u16(word, 32)
    long_count = u16(word, at)
    # FibRgLw97.ccpText counts the main story. The remaining pieces may contain
    # headers, footers, field instructions and text boxes, not policy body text.
    body_count = struct.unpack_from('<I', word, at + 2 + 3 * 4)[0] if long_count >= 4 else None
    at += 2 + 4 * long_count
    if u16(word, at) <= 33:
        raise ValueError('Word 文本索引缺失。')
    fc, size = struct.unpack_from('<II', word, at + 2 + 33 * 8)
    if not size or fc + size > len(table):
        raise ValueError('Word 文本索引越界。')
    clx, pos = table[fc:fc + size], 0
    while pos < len(clx) and clx[pos] == 1:
        pos += 3 + u16(clx, pos + 1)
    if pos + 5 > len(clx) or clx[pos] != 2:
        raise ValueError('Word 文本分片表缺失。')
    length = struct.unpack_from('<I', clx, pos + 1)[0]
    plc = clx[pos + 5:pos + 5 + length]
    if len(plc) != length or length < 16 or (length - 4) % 12:
        raise ValueError('Word 文本分片表长度无效。')
    n = (length - 4) // 12
    cps = struct.unpack_from('<' + 'I' * (n + 1), plc)
    if cps[0] != 0 or cps[-1] > 16_000_000 or any(a >= b for a, b in zip(cps, cps[1:])):
        raise ValueError('Word 字符区间无效。')
    chunks = []
    for i in range(n):
        pointer = struct.unpack_from('<I', plc, 4 * (n + 1) + 8 * i + 2)[0]
        compressed = bool(pointer & 0x40000000)
        offset = (pointer & 0x3FFFFFFF) // (2 if compressed else 1)
        count = (cps[i + 1] - cps[i]) * (1 if compressed else 2)
        if pointer & 0x80000000 or offset + count > len(word):
            raise ValueError('Word 文本分片越界。')
        chunks.append(word[offset:offset + count].decode('cp1252' if compressed else 'utf-16-le'))
    text = ''.join(chunks)
    if body_count is not None:
        if body_count > cps[-1]:
            raise ValueError('Word 正文字符数量越界。')
        # CP offsets count UTF-16 code units, including surrogate pairs.
        text = text.encode('utf-16-le')[:body_count * 2].decode('utf-16-le')
    return text.replace('\r', '\n').replace('\x07', '\t').replace('\x0b', '\n').replace('\x0c', '\n')


def extract_doc(data):
    import olefile
    with olefile.OleFileIO(io.BytesIO(data)) as document:
        word = document.openstream('WordDocument').read()
        name = '1Table' if struct.unpack_from('<H', word, 10)[0] & 0x0200 else '0Table'
        return doc_piece_text(word, document.openstream(name).read())


def source_scope(text):
    """Keep explicit scope quotations as review evidence, never infer eligibility."""
    scope = {'included_regions': ['CN-44'], 'excluded_regions': [], 'conditions': [], 'evidence': []}
    for m in re.finditer(r'[^。\n]*(?:不含深圳(?:市)?|深圳(?:市)?除外)[^。\n]*[。]?', text):
        if '广东' in m[0] or '全省' in m[0] or '我省' in m[0]:
            if 'CN-44-4403' not in scope['excluded_regions']:
                scope['excluded_regions'].append('CN-44-4403')
            scope['evidence'].append({'kind': 'excluded_region', 'quote': m[0],
                                      'start_char': m.start(), 'end_char_exclusive': m.end()})
    for m in re.finditer(r'[^。\n]*(?:试点纳税人|试点企业|试点范围|具体范围由|汇总[缴申]|核定扣除|小规模纳税人|一般纳税人)[^。\n]*[。]?', text):
        if len(scope['conditions']) >= 12:
            break
        scope['conditions'].append({'quote': m[0], 'start_char': m.start(),
                                     'end_char_exclusive': m.end(), 'review_status': 'needs_review'})
    return scope


def extract_asset(data, suffix):
    if suffix == '.doc' and data.startswith(bytes.fromhex('d0cf11e0a1b11ae1')):
        return extract_doc(data), 'extracted_needs_layout_review'
    if data.startswith(b"%PDF"):
        import pdfplumber
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            text = "\n\n".join(f"【PDF第{i+1}页】\n" + (p.extract_text() or "") for i, p in enumerate(pdf.pages))
        return text, "extracted_needs_layout_review" if len(re.sub(r"【PDF第\d+页】|\s", "", text)) >= 50 else "needs_ocr"
    if suffix in {".xls", ".xlsx"}:
        if data.startswith(b"PK"):
            from openpyxl import load_workbook
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
            text = "\n".join(f"【工作表:{s.title}】\n" + "\n".join("\t".join(str(v) if v is not None else "" for v in row) for row in s.values) for s in book)
            book.close()
        else:
            import xlrd
            book = xlrd.open_workbook(file_contents=data)
            text = "\n".join(f"【工作表:{s.name}】\n" + "\n".join("\t".join(str(s.cell_value(i, j)) for j in range(s.ncols)) for i in range(s.nrows)) for s in book.sheets())
        return text, "extracted_needs_table_review"
    if suffix == ".docx" and data.startswith(b"PK"):
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            root = ET.fromstring(z.read("word/document.xml"))
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        text = "\n".join("".join(n.text or "" for n in p.findall(".//w:t", ns)) for p in root.findall(".//w:p", ns))
        return text, "extracted_needs_layout_review"
    return "", "needs_ocr" if suffix in {".jpg", ".png", ".gif", ".bmp", ".jpeg"} else "needs_extraction"


def asset_record(package, url, kind, label, old, online):
    suffix = Path(urllib.parse.urlsplit(url).path).suffix.lower()
    relative = f"法规原文/附件/{sha(url.encode())[:24]}{suffix or '.bin'}"
    path = checked_path(package, relative)
    item = {"url": url, "kind": kind, "label": label, "raw_path": relative}
    stage = "retention"
    try:
        if path.exists():
            data = path.read_bytes()
            if old and old.get("raw_sha256") and sha(data) != old["raw_sha256"]:
                raise ValueError("附件哈希不一致，拒绝静默重写")
        elif online:
            data = fetch(url)
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(data)
        else:
            return {**item, "status": "missing", "error": "未留存；离线不下载"}
        item.update(raw_sha256=sha(data), raw_bytes=len(data))
        item["retained_at"] = (old or {}).get("retained_at") or datetime.fromtimestamp(path.stat().st_mtime, CST).isoformat()
        stage = "extraction"
        if old and old.get("text_path") and old.get("status", "").startswith("extracted"):
            text = read_verified(package, old["text_path"], old["text_sha256"]).decode("utf-8")
            status = old["status"]
        else:
            text, status = extract_asset(data, suffix)
        item["status"] = status
        if text:
            tp = relative + ".txt"
            checked_path(package, tp).write_bytes(text.encode("utf-8"))
            item.update(text_path=tp, text_sha256=sha(text.encode()), text_chars=len(text))
        if old and old.get('quality_review'):
            review = old['quality_review']
            if review.get('raw_sha256') == item.get('raw_sha256') and review.get('text_sha256') == item.get('text_sha256'):
                item['quality_review'] = review
    except Exception as exc:
        item.update(status="needs_password" if "encrypted" in str(exc).lower() else "error",
                    error_stage=stage, error=f"{type(exc).__name__}: {exc}")
    return item


def repair_deed(data):
    import pdfplumber
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        notice, page = pdf.pages[47], pdf.pages[48]
        matches = page.search("根据《中华人民共和国契税法》")
        if not matches:
            raise ValueError("契税公报版式与已核对的第49页不一致")
        y = matches[0]["top"] - 1
        w, h = page.width, page.height
        header = page.crop((0, 0, w, y)).extract_text() or ""
        left = page.crop((0, y, w / 2, h - 40)).extract_text() or ""
        right = page.crop((w / 2, y, w, h - 40)).extract_text() or ""
        return f"【PDF第48页】\n{notice.extract_text()}\n【PDF第49页，正文左栏后右栏】\n{header}\n{left}\n{right}"


def norm_title(value):
    return re.sub(r"[^\w]", "", value)


def norm_no(value):
    value = re.sub(r"[\s〔〕［］\[\]（）()]", "", value or "")
    return re.sub(r"(?<=字)第|(?<=年)第", "", value)


def sheet_tables(data):
    if data.startswith(b"PK"):
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        result = [(s.title, [[str(v) if v is not None else "" for v in row] for row in s.values]) for s in book]
        book.close()
        return result
    import xlrd
    book = xlrd.open_workbook(file_contents=data)
    return [(s.name, [[str(s.cell_value(i, j)).strip() for j in range(s.ncols)] for i in range(s.nrows)]) for s in book.sheets()]


def cleanup_catalogs(package, pages, assets):
    catalogs = []
    for page in pages:
        for url in page["attachment_urls"]:
            asset = assets[url]
            name = asset["label"] or urllib.parse.unquote(url)
            kind = ("full" if re.search(r"全文.*(?:废止|失效)", name) else
                    "partial" if re.search(r"部分条款.*(?:废止|失效)", name) else
                    "modified" if "修改" in name else "keep" if "继续执行" in name or "现行有效" in name else "unknown")
            item = {"page_title": page["title"], "page_url": page["source_url"],
                    "catalog_date": page.get("publish_date"), "kind": kind, "attachment_url": url,
                    "raw_path": asset["raw_path"], "raw_sha256": asset.get("raw_sha256"), "records": []}
            try:
                if Path(asset["raw_path"]).suffix.lower() not in {".xls", ".xlsx"}:
                    raise ValueError("目录非Excel，保留原件待复核")
                tables = sheet_tables(read_verified(package, asset["raw_path"], asset.get("raw_sha256")))
                for sheet, rows in tables:
                    header = next((i for i, row in enumerate(rows) if "标题" in row or "文件名称" in row), None)
                    if header is None:
                        raise ValueError(f"未识别清理目录表头：{sheet}")
                    names = rows[header]
                    for index, row in enumerate(rows[header + 1:], header + 2):
                        values = dict(zip(names, row))
                        title = values.get("标题") or values.get("文件名称")
                        if not title:
                            continue
                        number = next((values.get(k) for k in ["文件字号", "发文字号", "文号", "文件文号"] if values.get(k)), "")
                        item["records"].append({"title": title, "document_no": number, "row": index,
                                                "sheet": sheet, "values": row, "headers": names})
                item["status"] = "parsed"
            except Exception as exc:
                item.update(status="needs_review", error=str(exc))
            catalogs.append(item)
    return catalogs


def match_cleanup(document, catalogs):
    hits = []
    for catalog in catalogs:
        dt = catalog.get("catalog_date")
        if not dt or (document.get("publish_date") and document["publish_date"] > dt):
            continue
        for row in catalog["records"]:
            # A title with a conflicting document number may be a later edition.
            title_match = norm_title(row["title"]) == norm_title(document["title"])
            a, b = norm_no(row["document_no"]), norm_no(document.get("document_no"))
            short_a = bool(re.fullmatch(r"\d{4}年\d+号", a))
            short_b = bool(re.fullmatch(r"\d{4}年\d+号", b))
            compatible = a == b or (title_match and ((short_a and b.endswith(a)) or (short_b and a.endswith(b))))
            number_match = bool(a and b and a == b and not short_a)
            if number_match or (title_match and not (a and b and not compatible)):
                hits.append({"kind": catalog["kind"], "catalog_date": dt,
                             "page_url": catalog["page_url"], "attachment_url": catalog["attachment_url"],
                             "raw_path": catalog["raw_path"], "raw_sha256": catalog["raw_sha256"], **row})
    kinds = {h["kind"] for h in hits}
    return {"cleanup_verdict": "expired" if "full" in kinds else "partial" if "partial" in kinds else
            "modified" if "modified" in kinds else "historically_kept" if "keep" in kinds else None,
            "matches": hits, "note": "目录只证明该次清理结果；不自动确认当前或历史适用期间。"}


def grade(item, cross):
    dc = item["date_coverage"]
    if cross["cleanup_verdict"] == "expired" or item.get("validity_observed") in EXPIRED or dc["status_as_of"] == "expired":
        return "历史资料", ["页面/清理目录失效或正文期限已届满；保留原文供历史期间复核"]
    if item["text_status"] != "parsed" or item["content_gaps"]:
        return "待补齐", [*item["content_gaps"], "正文或附件仍需补齐/复核"]
    if item["document_kind"] in {"reference", "cleanup"}:
        return "背景资料", ["指引、名单、清理目录等保留为背景或效力证据，不能替代具体政策条款"]
    if cross["cleanup_verdict"] in {"partial", "modified"} or item.get("validity_observed") in PARTIAL:
        return "待复核", ["存在部分失效或修改，需定位到条款及适用期间"]
    if not dc["effective_from"] or dc["status_as_of"] == "pending" or "duration_end_requires_review" in dc["gaps"]:
        return "待复核", ["施行/终止日期、期限折算或尚未生效状态需复核"]
    return "优先复核", ["正文和来源已留存，有日期证据；仍未批准效力、条款或企业适用性"]


def apply_research(package, document, dates):
    """Replay retained, manually selected official evidence without approving it."""
    research = document.get('legal_research', {})
    for event in research.get('events', []):
        read_verified(package, event['raw_path'], event['raw_sha256'])
        text = read_verified(package, event['text_path'], event['text_sha256']).decode('utf-8')
        if text[event['start_char']:event['end_char_exclusive']] != event['quote']:
            raise ValueError('效力检索证据与留存原文不一致。')
        if event.get('effective_on'):
            date.fromisoformat(event['effective_on'])
    if research.get('events'):
        dates['research_events'] = research['events']
        kinds = {e['kind'] for e in research['events']}
        if 'full_repeal' in kinds:
            dates['status_as_of'] = 'expired'
        elif kinds & {'partial_repeal', 'amendment'} and dates['status_as_of'] != 'expired':
            dates['status_as_of'] = 'partial'
    override = research.get('own_enactment')
    if override:
        if override['source_url'] != document['source_url'] or override['raw_sha256'] != document['raw_sha256']:
            raise ValueError('施行日起期必须来自本文件自身的证据。')
        text = read_verified(package, override['text_path'], override['text_sha256']).decode('utf-8')
        if text[override['start_char']:override['end_char_exclusive']] != override['quote']:
            raise ValueError('施行日期引文与原文不一致。')
        value = date.fromisoformat(override['value']).isoformat()
        dates.update(effective_from=value, effective_from_source='retained_own_enactment_review',
                     effective_from_evidence=override['quote'])
        dates['gaps'] = [g for g in dates['gaps'] if g != 'effective_from_unknown']
    termination = research.get('reviewed_termination')
    if termination:
        # This selection is recorded by the source review, never inferred from
        # a directory's publication date or a partial repeal.
        index = termination.get('event_index')
        events = research.get('events', [])
        if not isinstance(index, int) or isinstance(index, bool) or not 0 <= index < len(events):
            raise ValueError('终止日期未绑定已有全文废止证据。')
        event = events[index]
        if event.get('kind') != 'full_repeal' or event.get('target_document_id') != document.get('document_id'):
            raise ValueError('终止日期只能采用本文件的全文废止证据。')
        on = date.fromisoformat(event['effective_on'])
        literal = f'{on.year}年{on.month}月{on.day}日'
        if literal not in re.sub(r'\s+', '', normal_dates(event['quote'])):
            raise ValueError('废止引文没有明确的变更日期，仍须保留待核实。')
        value = (on - timedelta(days=1)).isoformat()
        if termination.get('effective_to') != value or dates.get('effective_from') and value < dates['effective_from']:
            raise ValueError('终止日期与明确废止日期或起期不一致。')
        dates.update(effective_to=value, effective_to_source='reviewed_explicit_whole_repeal',
                     effective_to_evidence=event['quote'])
        dates['gaps'] = [g for g in dates['gaps'] if g != 'effective_to_unknown']
    return dates


def backup(package):
    path = package / "修正前快照.zip"
    if path.exists():
        return
    files = [p for p in package.iterdir() if p.is_file() and p.suffix in {".json", ".csv", ".md", ".py"}]
    files.extend((package / "法规原文").rglob("*.txt"))
    with zipfile.ZipFile(path, "x", zipfile.ZIP_DEFLATED) as z:
        for p in files:
            z.write(p, p.relative_to(package))


def repair(package, *, online=False, index_path=None, as_of=None):
    package = package.resolve()
    manifest = read_json(package / "来源清单.json")
    as_of = as_of or manifest["as_of"]
    date.fromisoformat(as_of)
    docs = manifest["documents"]
    # Verify all inputs before changing any metadata or text.
    for d in docs:
        for kind in ["raw", "text"]:
            read_verified(package, d[kind + "_path"], d[kind + "_sha256"])
    previous_path = package / "附件清单.json"
    previous_assets = read_json(previous_path).get("assets", []) if previous_path.exists() else []
    old_assets = {x["url"]: x for x in previous_assets}
    for asset in previous_assets:
        if asset.get("raw_sha256"):
            read_verified(package, asset["raw_path"], asset["raw_sha256"])
    backup(package)
    saved_index = package / "采集索引.json"
    if not saved_index.exists() and index_path:
        write_json(saved_index, read_json(index_path))
    index = read_json(saved_index) if saved_index.exists() else {"documents": []}
    pages_path = package / "清理依据页面.json"
    pages = read_json(pages_path).get("pages", []) if pages_path.exists() else []
    by_url = {x["source_url"]: x for x in pages}
    fetch_errors = []
    targets = [x for x in index["documents"] if x.get("local_direct") and
               any(k in x["title"] for k in ["清理", "废止", "失效", "公布现行有效"]) and
               any(k in x["title"] for k in ["规范性文件", "税费规范"])]
    for target in targets:
        url = target["url"]
        if url in by_url:
            continue
        try:
            existing = next((d for d in docs if d["source_url"] == url), None)
            if existing:
                data = read_verified(package, existing["raw_path"], existing["raw_sha256"])
                relative = existing["raw_path"]
            elif online:
                data = fetch(url)
                relative = f"法规原文/清理依据/{sha(url.encode())[:24]}.html"
                path = checked_path(package, relative)
                path.parent.mkdir(parents=True, exist_ok=True)
                path.write_bytes(data)
            else:
                raise ValueError("离线未留存清理页面")
            by_url[url] = {"source_url": url, "title": target["title"], "publish_date": target.get("publish_date"),
                           "raw_path": relative, "raw_sha256": sha(data)}
        except Exception as exc:
            fetch_errors.append({"url": url, "error": str(exc)})
    pages = list(by_url.values())
    requests = {}
    for d in [*docs, *pages]:
        raw = read_verified(package, d["raw_path"], d["raw_sha256"])
        if not d["raw_path"].endswith(".html"):
            continue
        body_text, links, images = html_parts(raw.decode("utf-8", "replace"))
        if d in pages:
            d["text_path"] = d["raw_path"] + ".txt"
            d["text_sha256"] = sha(body_text.encode())
            checked_path(package, d["text_path"]).write_bytes(body_text.encode())
            d["review_status"] = "needs_review"
        d["attachment_urls"] = []
        d["body_image_urls"] = []
        for link in [*links, *images]:
            url = urllib.parse.urljoin(d["source_url"], link["link"])
            requests.setdefault(url, {"url": url, "kind": link["kind"], "label": link["label"]})
            key = "attachment_urls" if link["kind"] == "attachment" else "body_image_urls"
            if url not in d[key]:
                d[key].append(url)
    assets = {}
    with ThreadPoolExecutor(max_workers=4) as pool:
        futures = {pool.submit(asset_record, package, **req, old=old_assets.get(url), online=online): url
                   for url, req in requests.items()}
        for i, future in enumerate(as_completed(futures), 1):
            url = futures[future]
            assets[url] = future.result()
            if i % 25 == 0:
                print(f"附件检查 {i}/{len(futures)}", flush=True)
    assets = dict(sorted(assets.items()))
    write_json(previous_path, {"as_of": as_of, "assets": list(assets.values())})
    write_json(pages_path, {"as_of": as_of, "pages": pages, "missing_pages": fetch_errors})
    catalogs = cleanup_catalogs(package, pages, assets)
    write_json(package / "清理目录记录.json", {"catalogs": catalogs})
    cross_results = []
    for d in docs:
        raw = read_verified(package, d["raw_path"], d["raw_sha256"])
        original_text = read_verified(package, d["text_path"], d["text_sha256"]).decode("utf-8")
        if "pre_repair_text_sha256" not in d:
            d["pre_repair_text_sha256"] = d["text_sha256"]
        text = original_text
        if d["raw_path"].endswith(".html"):
            text, _, _ = html_parts(raw.decode("utf-8", "replace"))
        elif d.get("source_kind") == "provincial_people_congress" and "契税" in d["title"]:
            text = repair_deed(raw)
            d["text_extraction"] = "PDF第48页及第49页标题后按左栏、右栏重排；未人工批准条款"
        tags = tax_tags(d["title"])
        d["tax_types"] = tags
        d["fee_types"] = [f for f in FEES if f in d["title"]]
        d["topics"] = ["税收优惠"] if any(x in d["title"] for x in ["优惠", "减免", "退役", "创业就业"]) else []
        d["tax_type"] = tags[0] if len(tags) == 1 else "跨税种" if tags else "征管通用"
        d["document_kind"] = document_kind(d["title"])
        d["region_code"] = "CN-GD"
        d["region_review_status"] = "needs_review"
        gaps = []
        meaningful = len(re.sub(r"附件下载|\s", "", text)) >= 50
        urls = d.get("attachment_urls", [])
        if not meaningful and urls:
            pieces = []
            for url in urls:
                asset = assets[url]
                if asset.get("text_path"):
                    pieces.append(f"【附件来源:{url}】\n" + read_verified(package, asset["text_path"], asset["text_sha256"]).decode("utf-8"))
            if pieces:
                text = "\n\n".join(pieces)
                meaningful = len(text) >= 50
                gaps.append("attachment_body_requires_layout_review")
        for url in urls:
            if assets[url]["status"] in {"missing", "error", "needs_extraction", "needs_ocr", "needs_password"}:
                gaps.append("attachment_incomplete:" + url)
            else:
                gaps.append("attachment_requires_review:" + url)
        for url in d.get("body_image_urls", []):
            gaps.append("body_image_requires_review:" + url)
        if not meaningful:
            gaps.append("body_text_missing")
        if "水资源税改革试点" in d["title"] and "地下水" not in text.split("附件：")[-1].split("备注")[0]:
            gaps.append("source_table_underground_water_rates_missing")
        d["content_gaps"] = sorted(set(gaps))
        d["text_status"] = "parsed" if meaningful else "needs_ocr" if d.get("body_image_urls") else "attachment_only" if urls else "missing_text"
        d["text_chars"] = len(text)
        d["text_sha256"] = sha(text.encode())
        checked_path(package, d["text_path"]).write_bytes(text.encode("utf-8"))
        dc = extract_dates(text, d.get("publish_date"), as_of, d.get("validity_observed"), kind=d["document_kind"])
        dc = apply_research(package, d, dc)
        if "房产税施行细则" in d["title"] and d.get("source_kind") == "provincial_government_regulation":
            dc["original_effective_from"] = dc["effective_from"]
            dc.update(effective_from=None, effective_from_source=None, effective_from_evidence=None,
                      version_amended_on="2021-09-29")
            dc["gaps"] = sorted(set([*dc["gaps"], "effective_from_unknown", "version_effective_from_requires_decree_review"]))
        d["date_coverage"] = dc
        if '增值税' in d['tax_types']:
            d['scope'] = source_scope(text)
        d["validity_status"] = "unverified"
        d["extraction_status"] = "needs_review"
        cross = match_cleanup(d, catalogs)
        d["review_queue"], d["review_reasons"] = grade(d, cross)
        cross_results.append({"document_id": d["document_id"], **cross})
    manifest.update(as_of=as_of, repair_version=VERSION, documents=docs,
                    repair_boundary="只修正候选资料；未导入业务库、未批准效力或扩大核查范围")
    write_json(package / "来源清单.json", manifest)
    write_json(package / "效力核对结果.json", {"as_of": as_of, "catalogs": [
        {k: v for k, v in c.items() if k != "records"} | {"count": len(c["records"])} for c in catalogs],
        "missing_pages": fetch_errors, "documents": cross_results})
    # Keep the provincial supplement consistent with the main manifest.
    supplement = package / "补充清单.json"
    if supplement.exists():
        extra = read_json(supplement)
        extra.update(as_of=as_of, documents=[d for d in docs if d.get("source_kind", "").startswith("provincial_")])
        write_json(supplement, extra)
    return reports(package, manifest, assets, catalogs, fetch_errors, pages)


def reports(package, manifest, assets, catalogs, errors, pages):
    import csv
    docs = manifest["documents"]
    fields = ["编号", "税种", "涉及税种", "费用", "标题", "文号", "发文日期", "施行日期", "终止日期", "页面状态", "审查队列", "理由", "来源URL"]
    rows = [[d["document_id"], d["tax_type"], "/".join(d["tax_types"]), "/".join(d["fee_types"]), d["title"],
             d.get("document_no"), d.get("publish_date"), d["date_coverage"]["effective_from"],
             d["date_coverage"]["effective_to"], d.get("validity_observed"), d["review_queue"],
             "；".join(d["review_reasons"]), d["source_url"]] for d in docs]
    for name in ["政策索引.csv", "入图谱建议.csv"]:
        with (package / name).open("w", encoding="utf-8-sig", newline="") as f:
            writer = csv.writer(f)
            writer.writerow(fields)
            writer.writerows(rows)
    dates = [[d["document_id"], d["date_coverage"]["effective_from"], d["date_coverage"]["effective_to"],
              d["date_coverage"]["effective_from_source"], d["date_coverage"]["effective_to_source"],
              d["date_coverage"]["effective_from_evidence"], d["date_coverage"]["effective_to_evidence"],
              json.dumps(d["date_coverage"]["applicable_periods"], ensure_ascii=False),
              ";".join(d["date_coverage"]["gaps"])] for d in docs]
    header = ["编号", "施行日期", "终止日期", "起期出处", "止期出处", "起期证据", "止期证据", "适用期间", "日期缺口"]
    with (package / "日期覆盖矩阵.csv").open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(header)
        writer.writerows(dates)
    (package / "日期覆盖矩阵.md").write_text("# 广东候选资料日期证据\n\n期限折算为待复核线索；未知值不补造。\n\n| " + " | ".join(header) + " |\n|" + "---|" * len(header) + "\n" + "\n".join("| " + " | ".join(str(v or "").replace("|", "/").replace("\n", " ") for v in row) + " |" for row in dates) + "\n", encoding="utf-8")
    totals = dict(Counter(d["review_queue"] for d in docs))
    report = {"as_of": manifest["as_of"], "repair_version": VERSION, "documents": len(docs),
              "tax_type_counts": dict(Counter(t for d in docs for t in d["tax_types"])),
              "text_status_counts": dict(Counter(d["text_status"] for d in docs)),
              "review_queues": totals, "assets": len(assets),
              "assets_retained": sum(bool(a.get("raw_sha256")) for a in assets.values()),
              "asset_status_counts": dict(Counter(a["status"] for a in assets.values())),
              "cleanup_pages": len(pages), "cleanup_attachments": len(catalogs),
              "cleanup_rows": sum(len(c["records"]) for c in catalogs), "missing_cleanup_pages": errors,
              "has_effective_from": sum(bool(d["date_coverage"]["effective_from"]) for d in docs),
              "has_effective_to": sum(bool(d["date_coverage"]["effective_to"]) for d in docs),
              "formal_approvals": 0, "business_database_writes": 0,
              "checks": {"all_candidates_unverified": all(d["validity_status"] == "unverified" for d in docs),
                         "no_direct_use_queue": all(d["review_queue"] != "可直接用" for d in docs),
                         "no_vat_substring_misclassification": all("土地增值税" not in d["title"] or "土地增值税" in d["tax_types"] for d in docs),
                         "dates_have_evidence": all(all(not d["date_coverage"][k] or d["date_coverage"][k + "_evidence"] for k in ["effective_from", "effective_to"]) for d in docs),
                         "all_document_hashes_match": all(sha(checked_path(package, d[k + "_path"]).read_bytes()) == d[k + "_sha256"] for d in docs for k in ["raw", "text"]),
                         "all_retained_asset_hashes_match": all(sha(checked_path(package, a[k + "_path"]).read_bytes()) == a[k + "_sha256"] for a in assets.values() for k in ["raw", "text"] if a.get(k + "_sha256")),
                         "all_cleanup_source_hashes_match": all(sha(checked_path(package, p[k + "_path"]).read_bytes()) == p[k + "_sha256"] for p in pages for k in ["raw", "text"])}}
    write_json(package / "验收报告.json", report)
    write_json(package / "入图谱建议汇总.json", {"totals": totals, "formal_approvals": 0,
                                                    "note": "撤销可直接用；分级仅安排补齐、复核和历史留存。"})
    boundary = ["仅修正现有331份候选及其附件、清理依据；未导入业务库。",
                "税收优惠为主题，跨税种记录保留多个tax_types；费用单独记录。",
                "页面有效仅作观察值；清理目录只证明当次处理，当前效力仍未审核。",
                "缺失、扫描件、旧DOC/压缩包、表格及PDF版式均保留待复核状态。",
                "省/市/深圳及特殊区域的适用边界未自动确认。",
                "转发文、地市政策和后续文件仍需补采；资料包不保证广东政策全覆盖。"]
    lines = ["# 广东省税务政策资料包（修正后，候选待复核）", "", f"基准日：{manifest['as_of']}；文件：{len(docs)}份。", "",
             "## 修正结果", "", "| 项目 | 数量 |", "|---|---:|",
             *[f"| {k} | {v} |" for k, v in totals.items()],
             f"| 附件/正文图片链接 | {len(assets)} |", f"| 有起期证据 | {report['has_effective_from']} |",
             f"| 有止期证据 | {report['has_effective_to']} |", f"| 清理目录记录 | {report['cleanup_rows']} |", "",
             "## 使用边界", "", *["- " + x for x in boundary], "",
             "## 离线复验", "", "```powershell", "python 构建脚本.py --offline", "```", "",
             "清理页面、目录附件和采集索引已随包留存；无需tmp目录。联网补齐使用 --fetch，失败逐项留痕。",
             "运行依赖：Python 3.10+、pdfplumber、xlrd、openpyxl。parsed仅表示有可提取文本，不能作为材料完整或政策适用性审核结果。",
             "修正前快照.zip保留旧元数据、脚本和提取文本；原始HTML/PDF不改写。",
             "旧原文路径保持稳定；税种归属以来源清单的tax_types为准。", ""]
    (package / "README.md").write_text("\n".join(lines), encoding="utf-8")
    return report


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--package", type=Path, default=Path(__file__).resolve().parent)
    parser.add_argument("--index", type=Path)
    parser.add_argument("--fetch", action="store_true", help="仅下载未留存的官方附件/清理页面")
    parser.add_argument("--offline", action="store_true", help="从留存原件重建；不联网")
    parser.add_argument("--as-of")
    args = parser.parse_args(argv)
    if args.fetch and args.offline:
        parser.error("--fetch 和 --offline 不能并用")
    report = repair(args.package, online=args.fetch, index_path=args.index, as_of=args.as_of)
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
