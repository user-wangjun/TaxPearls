"""Financial export layouts through the actual material parser and review gates."""
from copy import deepcopy
from decimal import Decimal
from hashlib import sha256
from io import BytesIO
import unittest
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook

from src import config, financial_import, material_formats, material_review, material_security, materials
from src.input_errors import InputError
from tests.test_materials import COMPANY, KEYS
from webapp import enterprise_analysis


def export_book(*, period='2026年1月', unit='万元', legacy=False, balance=False, standard=False,
                value='10', code_sheet='科目余额表', formula=False, currency='', current_heading='本月数'):
    # Internal regression fixture only, not a transformed public source or deliverable.
    book = Workbook()
    book.remove(book.active)
    meta = f'编制单位：{COMPANY["name"]} {period} 单位：{unit}' + (f' 币种：{currency}' if currency else '')
    if standard:
        company = book.create_sheet('企业信息')
        company.append(['项目', '内容'])
        for label, key in zip(config.COMPANY_FIELDS, materials.COMPANY_KEYS):
            company.append([label, COMPANY[key]])
        decl = book.create_sheet('增值税申报')
        decl.append(['项目', '金额'])
        decl.append(['销售额', 100000])
    account = book.create_sheet(code_sheet)
    account.append([meta])
    account.append(['科目代码', '会计科目', '期初余额', None, '本期发生额', None, '期末余额'])
    account.append([None, None, '借方', '贷方', '借方', '贷方', '借方', '贷方'])
    for area in ('C2:D2', 'E2:F2', 'G2:H2'):
        account.merge_cells(area)
    for code, name, debit, credit in (
        ('5101' if legacy else '6001', '主营业务收入', 0, value),
        ('5102' if legacy else '6051', '其他业务收入', 0, 0),
        ('5401' if legacy else '6401', '主营业务成本', 6, 0),
        ('5402' if legacy else '6402', '其他业务支出' if legacy else '其他业务成本', 0, 0)):
        account.append([code, name, None, None, debit, credit, None, None])
    profit = book.create_sheet('利润表')
    profit.append([meta])
    profit.append(['项 目', '行次', current_heading, '本年累计数'])
    profit.append(['一、营业收入', 1, '=1+1' if formula else value, 900])
    profit.append(['减：营业成本', 2, 6, 800])
    profit.append(['不属于规则输入的小计', 3, 99, 99])
    if balance:
        sheet = book.create_sheet('资产负债表')
        sheet.append([f'编制单位：{COMPANY["name"]} 2026年1月31日 单位：{unit}'])
        sheet.append(['资产', '期末余额', '年初余额', '负债及所有者权益', '期末余额', '年初余额'])
        sheet.append(['资产总计', 20, 100, '负债合计', 8, 90])
        sheet.append([None, None, None, '所有者权益合计', 12, 10])
    stream = BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def mixed_export(change, *, standard=False):
    from openpyxl import load_workbook
    book = load_workbook(BytesIO(export_book(balance=True, standard=standard)))
    if change == 'company':
        book['利润表']['A1'] = '编制单位：其他企业 2026年1月 单位：万元'
    elif change == 'period':
        book['利润表']['A1'] = f'编制单位：{COMPANY["name"]} 2026年2月 单位：万元'
    elif change == 'date':
        book['资产负债表']['A1'] = f'编制单位：{COMPANY["name"]} 2026年2月28日 单位：万元'
    elif change == 'currency':
        book['利润表']['A1'] = f'编制单位：{COMPANY["name"]} 2026年1月 单位：万元 币种：USD'
    elif change == 'row_currency':
        sheet = book['科目余额表']
        sheet['I2'] = '币种'
        for row in range(4, 8):
            sheet.cell(row, 9, 'USD' if row == 4 else 'CNY')
    else:
        raise ValueError(change)
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def parse(raw):
    return materials.preview([('原始导出.xlsx', raw)], KEYS, allow_incomplete_company=True, capture_standard=True)[0]


def formula_book(*, cache='10', cache_type='n', shared=False, standard=False):
    """In-memory OOXML test fixture, not a real/certified financial source.

    Deliberately stale cache (formula computes 2, stored candidate is 10) proves
    the parser does not silently execute or endorse either value.
    """
    if standard:
        from tests.test_materials import workbook
        raw = workbook('利润表', [['项目', '本期金额'], ['营业收入', '=1+1'], ['营业成本', 6]], COMPANY)
        path, cells = 'xl/worksheets/sheet2.xml', ('B2',)
    else:
        raw = export_book(formula=True)
        path, cells = 'xl/worksheets/sheet2.xml', ('C3', 'C4') if shared else ('C3',)
    output = BytesIO()
    ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with ZipFile(BytesIO(raw)) as original, ZipFile(output, 'w', ZIP_DEFLATED) as updated:
        for item in original.infolist():
            content = original.read(item)
            if item.filename == path:
                root = ET.fromstring(content)
                for index, address in enumerate(cells):
                    cell = root.find(f'.//s:c[@r="{address}"]', ns)
                    formula = cell.find('s:f', ns)
                    if formula is None:
                        formula = ET.SubElement(cell, '{' + ns['s'] + '}f')
                    formula.text = 'SUM(A3:B3)' if shared and index == 0 else None if shared else '1+1'
                    if shared:
                        formula.attrib.update(t='shared', si='0')
                        if index == 0:
                            formula.attrib['ref'] = 'C3:C4'
                    cell.attrib['t'] = cache_type
                    value = cell.find('s:v', ns)
                    if value is None:
                        value = ET.SubElement(cell, '{' + ns['s'] + '}v')
                    value.text = cache
                content = ET.tostring(root, encoding='utf-8')
            updated.writestr(item, content)
    return output.getvalue()


def scope():
    return {'name': COMPANY['name'], 'taxpayer_id': COMPANY['taxpayer_id'], 'industry': COMPANY['industry'],
            'period_start': '2026-01-01', 'period_end': '2026-01-31'}


def annual_book(kind='income', *, basis='合併', currency='', unit='人民幣百萬元', date='二零二六年十二月三十一日',
                years=('2026年', '2025年'), duplicate='3', note=True, revenue=12):
    """Constructed year-column regression layout, never a rewritten real source."""
    title = {'income': '損益表', 'balance': '財務狀況表', 'cashflow': '現金流量表'}[kind]
    book = Workbook()
    sheet = book.active
    sheet.title = '原表'
    sheet.append([COMPANY['name']])
    sheet.append([])
    sheet.append([basis + title])
    sheet.append([date + ('止年度' if kind != 'balance' else '')])
    sheet.append(['（單位：' + unit + ('，每股數除外）' if kind == 'income' else '）')])
    sheet.append(['幣種：' + currency] if currency else [])
    sheet.append([None, None, '截至12月31日止年度'])
    sheet.append([None, '附註' if note else None, *years])
    sheet.append([])
    items = {'income': [('收入', revenue, 90), ('稅前利潤', 4, 80), ('所得稅', -1, 70),
              ('年度盈利', 3, 60), ('應佔盈利：', None, None), ('本公司權益持有者', 2, 50),
              ('非控制性權益', 1, 10), ('年度盈利', duplicate, 60),
              ('每股盈利－基本（人民幣元）', '0.1', '0.2')],
        'balance': [('總資產', 20, 90), ('總負債', 8, 70), ('總權益', 12, 20)],
        'cashflow': [('經營活動所產生的淨現金流入', 6, 50), ('投資活動所支付的淨現金流出', -3, 40),
                     ('融資活動所產生的淨現金流出', -2, 30), ('外幣匯率變動的影響', 0, 10),
                     ('現金及現金等價物年初餘額', 4, 80), ('現金及現金等價物年末餘額', 5, 90)]}[kind]
    for label, value, previous in items:
        sheet.append([label, None, value, previous])
    out = BytesIO()
    book.save(out)
    book.close()
    return out.getvalue()


class FinancialExports(unittest.TestCase):
    def good(self, raw):
        document = parse(raw)
        self.assertEqual(document['error'], '', document['error'])
        self.assertEqual(document['extraction']['local']['status'], 'succeeded')
        self.assertIn('finished_at', document['extraction']['local'])
        return document

    def test_annual_native_units_years_and_exact_source_meaning(self):
        doc = self.good(annual_book())
        self.assertEqual(doc['company']['name'], COMPANY['name'])
        self.assertEqual(doc['company']['period'], '2026')
        self.assertEqual(doc['import_model']['statement_basis'], 'consolidated')
        rows = {r['name']: r for r in doc['rows']}
        self.assertEqual(rows['利润表.营业收入']['value'], '12000000')
        self.assertEqual(rows['利润表.所得税费用']['value'], '1000000')
        self.assertEqual(rows['利润表.净利润']['value'], '3000000')
        self.assertIn('原表!C13、原表!C17', rows['利润表.净利润']['source'])
        self.assertFalse(any('D' in r['source'] for r in rows.values()))
        self.assertEqual(set(rows), {'利润表.营业收入', '利润表.利润总额', '利润表.所得税费用', '利润表.净利润'})
        self.assertTrue(any('每股盈利' in r['item'] for r in doc['import_mapping']['unmapped']))

    def test_annual_restatement_mismatch_and_corrections_cannot_choose_one(self):
        self.assertIn('不一致', parse(annual_book(duplicate='7'))['error'])
        doc = self.good(annual_book())
        with self.assertRaisesRegex(InputError, '不一致'):
            material_review.apply(doc, {'原表!C17': '7'})
        changed, _ = material_review.apply(doc, {'原表!C13': '7', '原表!C17': '7'})
        self.assertEqual(next(r['value'] for r in changed['rows'] if r['name'] == '利润表.净利润'), '7000000')
        # Clearing either amount leaves the total unavailable, not zero or a first-row fallback.
        changed, _ = material_review.apply(doc, {'原表!C17': None})
        self.assertFalse(any(r['name'] == '利润表.净利润' for r in changed['rows']))

    def test_annual_date_year_and_structure_are_not_guessed(self):
        for changes in [{'date': '二零二六年六月三十日'}, {'date': '二零二六年十三月三十一日'},
                        {'date': '未知日期'}, {'years': ('2025年', '2024年')},
                        {'years': ('2026年', '2026年')}, {'note': False}]:
            with self.subTest(changes=changes):
                self.assertTrue(parse(annual_book(**changes))['error'])
        doc = self.good(annual_book())
        with self.assertRaisesRegex(InputError, '比较年'):
            materials._excel(annual_book(), doc, keys=KEYS, allow_incomplete_company=True,
                             import_options={'原表': {'columns': {'8:A': 'D'}}})
        with self.assertRaisesRegex(InputError, '实际期间不能'):
            materials._excel(annual_book(), doc, keys=KEYS, allow_incomplete_company=True,
                             import_options={'原表': {'period': '2025'}})

    def test_annual_currency_units_remain_blocking_and_cannot_be_mapped_away(self):
        for changes in [{'currency': 'USD'}, {'unit': '港幣百萬元'}, {'unit': '未知单位'}]:
            with self.subTest(changes=changes):
                doc = self.good(annual_book(**changes))
                self.assertTrue(doc['import_mapping']['pending'])
                self.assertFalse(doc['rows'])
        doc = self.good(annual_book(currency='USD'))
        materials._excel(annual_book(currency='USD'), doc, keys=KEYS, allow_incomplete_company=True,
                         import_options={'原表': {'unit': '人民币百万元'}})
        self.assertTrue(doc['import_mapping']['pending'])
        self.assertFalse(doc['rows'])

    def test_annual_balance_is_only_a_source_endpoint_and_cashflow_keeps_signs(self):
        doc = self.good(annual_book('balance'))
        self.assertEqual(doc['company']['period'], '')
        self.assertEqual(doc['import_mapping']['sheets'][0]['as_of'], '2026-12-31')
        self.assertEqual({r['name']: r['value'] for r in doc['rows']},
                         {'资产负债表.资产总额': '20000000', '资产负债表.负债总额': '8000000', '资产负债表.所有者权益': '12000000'})
        cash = self.good(annual_book('cashflow'))
        self.assertEqual(next(r['value'] for r in cash['rows'] if r['name'] == '现金流量表.投资净额'), '-3000000')
        self.assertEqual(next(r['value'] for r in cash['rows'] if r['name'] == '现金流量表.汇率影响'), '0')

    def test_annual_native_basis_cross_file_and_endpoint_guards(self):
        docs = materials.preview([('合併.xlsx', annual_book()), ('母公司.xlsx', annual_book(basis='母公司'))], KEYS,
                                 allow_incomplete_company=True)
        target = {**scope(), 'period_start': '2026-01-01', 'period_end': '2026-12-31'}
        result = enterprise_analysis.analyze(docs, target, {}, [])
        self.assertIn('statement_basis_conflict', {n['code'] for n in result['feedback']['blocking']})
        excluded = enterprise_analysis.analyze(docs, target, {docs[1]['id']: {'purpose': 'excluded'}}, [])
        self.assertNotIn('statement_basis_conflict', {n['code'] for n in excluded['feedback']['blocking']})
        balance = self.good(annual_book('balance'))
        choice = {balance['id']: {'company': {**COMPANY, 'period': '2025'}}}
        with self.assertRaisesRegex(InputError, '时点'):
            materials.build_dataset([balance], choice, {**COMPANY, 'period': '2025'}, KEYS)

    def test_annual_explicit_basis_cannot_mix_with_unqualified_native_or_pdf(self):
        from tests.test_materials import workbook
        native = parse(workbook('利润表', [config.COL_STATEMENT, ['营业收入', 1]], {**COMPANY, 'period': '2026'}))
        native['id'] = 'other'
        annual = self.good(annual_book('balance'))
        target = {**scope(), 'period_start': '2026-01-01', 'period_end': '2026-12-31'}
        for other in [native, {**deepcopy(native), 'kind': 'pdf', 'pdf_statement_basis': 'parent'}]:
            result = enterprise_analysis.analyze([annual, other], target, {}, [])
            self.assertIn('statement_basis_conflict', {n['code'] for n in result['feedback']['blocking']})

    def test_annual_formula_amount_stays_missing_until_explicit_cell_review(self):
        raw = annual_book(revenue='=1+1')
        doc = self.good(raw)
        field = next(f for f in material_review.fields(doc) if f['id'] == '原表!C10')
        self.assertEqual(field['formula'], '=1+1')
        self.assertIsNone(field['value'])
        self.assertFalse(any(r['name'] == '利润表.营业收入' for r in doc['rows']))
        changed, _ = material_review.apply(doc, {'原表!C10': '2'})
        self.assertEqual(next(r['value'] for r in changed['rows'] if r['name'] == '利润表.营业收入'), '2000000')
        target = {**scope(), 'period_start': '2026-01-01', 'period_end': '2026-12-31'}
        choice = {'standard_edits': {'原表!C10': '2'}}
        requirements = material_security.review_requirements(doc, choice, target)
        choice['evidence_reviews'] = [r['id'] for r in requirements if not r['key'].startswith('formula:')]
        self.assertTrue(material_security.blockers(doc, choice, target))
        self.assertEqual(doc['sha256'], sha256(raw).hexdigest())

    def test_annual_row_currency_and_flow_direction_cannot_be_ignored(self):
        from openpyxl import load_workbook
        for change in ('currency', 'positive_outflow'):
            book = load_workbook(BytesIO(annual_book('cashflow')))
            sheet = book.active
            if change == 'currency':
                sheet['E8'] = '幣種'
                for r in range(10, 16):
                    sheet.cell(r, 5, 'USD' if r == 10 else '人民幣')
            else:
                sheet['C11'] = 3
            raw = BytesIO()
            book.save(raw)
            book.close()
            doc = self.good(raw.getvalue())
            self.assertTrue(doc['import_mapping']['pending'])
            self.assertFalse(doc['rows'])

    def test_annual_genuine_merged_metadata_is_one_source_assertion(self):
        from openpyxl import load_workbook
        book = load_workbook(BytesIO(annual_book()))
        for row in (1, 3, 4, 5):
            book.active.merge_cells(f'A{row}:D{row}')
        output = BytesIO()
        book.save(output)
        book.close()
        doc = self.good(output.getvalue())
        self.assertFalse(doc['import_mapping']['pending'])
        self.assertEqual(doc['company']['period'], '2026')
        self.assertEqual(next(r['value'] for r in doc['rows'] if r['name'] == '利润表.营业收入'), '12000000')

    def test_annual_cannot_hide_unqualified_standard_financial_sheet_in_same_workbook(self):
        from openpyxl import load_workbook
        book = load_workbook(BytesIO(annual_book()))
        sheet = book.create_sheet('资产负债表')
        sheet.append(config.COL_STATEMENT)
        sheet.append(['资产总额', 20])
        output = BytesIO()
        book.save(output)
        book.close()
        doc = self.good(output.getvalue())
        self.assertTrue(any('主体口径' in p['message'] for p in doc['import_mapping']['pending']))
        target = {**scope(), 'period_start': '2026-01-01', 'period_end': '2026-12-31'}
        choice = {'evidence_reviews': [r['id'] for r in material_security.review_requirements(doc, {}, target)]}
        self.assertFalse(enterprise_analysis.analyze([doc], target, {doc['id']: choice}, [])['can_confirm'])

    def test_split_headers_source_units_and_month_not_cumulative(self):
        raw = export_book()
        doc = self.good(raw)
        self.assertEqual(doc['sha256'], sha256(raw).hexdigest())
        self.assertEqual(doc['company']['period'], '2026-01')
        self.assertEqual(doc['accounts'][0]['credit'], '100000')
        self.assertIsNone(doc['accounts'][0]['opening'])
        row = next(r for r in doc['rows'] if r['name'] == '利润表.营业收入')
        self.assertEqual(row['value'], '100000')
        self.assertIn('利润表!C3', row['source'])
        self.assertFalse(any('小计' in r['name'] for r in doc['rows']))
        fields = {f['id']: f for f in material_review.fields(doc)}
        self.assertEqual(fields['利润表!C3']['value'], '10')
        self.assertIn('万元', fields['利润表!C3']['label'])
        self.assertEqual(doc['import_mapping']['sheets'][0]['columns']['本期贷方'], 'F')

    def test_currency_is_not_a_scale_and_cannot_be_overridden_or_acknowledged(self):
        for currency in ('USD', '美元', 'HKD', '港币', 'EUR', 'CHF', 'BTC', 'CNY、USD'):
            with self.subTest(currency=currency):
                raw = export_book(currency=currency)
                doc = self.good(raw)
                self.assertIn('币种', str(doc['import_mapping']['pending']))
                self.assertFalse(doc['rows'])
                self.assertFalse(doc['accounts'])
                selection = {'evidence_reviews': [r['id'] for r in material_security.review_requirements(doc, {}, scope())]}
                self.assertFalse(enterprise_analysis.analyze([doc], scope(), {doc['id']: selection}, [])['can_confirm'])
                materials._excel(raw, doc, keys=KEYS, allow_incomplete_company=True,
                    import_options={'科目余额表': {'unit': '万元'}, '利润表': {'unit': '万元'}})
                self.assertTrue(doc['import_mapping']['pending'])
                self.assertFalse(doc['rows'])
                self.assertEqual(doc['sha256'], sha256(raw).hexdigest())

    def test_foreign_sheet_can_be_excluded_but_restoring_it_blocks_again(self):
        raw = mixed_export('currency')
        doc = self.good(raw)
        self.assertTrue(doc['import_mapping']['pending'])
        materials._excel(raw, doc, keys=KEYS, allow_incomplete_company=True,
            import_options={'利润表': {'excluded': True}})
        self.assertFalse(doc['import_mapping']['pending'])
        self.assertEqual(doc['accounts'][0]['credit'], '100000')
        self.assertFalse(any(r['name'].startswith('利润表.') for r in doc['rows']))
        materials._excel(raw, doc, keys=KEYS, allow_incomplete_company=True)
        self.assertTrue(doc['import_mapping']['pending'])
        self.assertFalse(doc['accounts'])

    def test_rmb_scales_and_aliases_are_exact_not_prefix_matches(self):
        for unit, expected in [('人民币万元', '100000'), ('百万元', '10000000'), ('亿元', '1000000000')]:
            for currency in ('人民币', 'CNY', 'RMB'):
                with self.subTest(unit=unit, currency=currency):
                    doc = self.good(export_book(unit=unit, currency=currency))
                    self.assertEqual(doc['accounts'][0]['credit'], expected)
                    self.assertEqual(doc['import_mapping']['sheets'][0]['currency'], 'CNY')
        for unit in ('美元', '万元美元', '元（美元）', '未知单位'):
            with self.subTest(unit=unit):
                raw = export_book(unit=unit)
                doc = self.good(raw)
                self.assertTrue(doc['import_mapping']['pending'])
                self.assertFalse(doc['rows'])
                with self.assertRaisesRegex(InputError, '原文金额单位不能'):
                    materials._excel(raw, doc, keys=KEYS, allow_incomplete_company=True,
                        import_options={'科目余额表': {'unit': '元'}, '利润表': {'unit': '元'}})

    def test_isolated_date_needs_explicit_flow_period_and_endpoint_must_match(self):
        raw = export_book(period='2026年1月31日', current_heading='本期金额')
        doc = self.good(raw)
        self.assertEqual(doc['company']['period'], '')
        self.assertEqual(doc['import_mapping']['sheets'][1]['as_of'], '2026-01-31')
        self.assertFalse(doc['rows'])
        self.assertTrue(doc['import_mapping']['pending'])
        choices = {s: {'period': '2026-01'} for s in ('科目余额表', '利润表')}
        materials._excel(raw, doc, keys=KEYS, allow_incomplete_company=True, import_options=choices)
        self.assertFalse(doc['import_mapping']['pending'])
        self.assertEqual(doc['company']['period'], '2026-01')
        self.assertEqual(doc['import_mapping']['sheets'][1]['source_period'], '')
        self.assertEqual(doc['import_mapping']['sheets'][1]['period_origin'], 'user')
        self.assertEqual(doc['rows'][0]['value'], '100000')
        with self.assertRaisesRegex(InputError, '时点不能'):
            materials._excel(raw, doc, keys=KEYS, allow_incomplete_company=True,
                import_options={'利润表': {'period': '2026-02'}})

    def test_explicit_quarter_half_and_intervals_do_not_use_month_or_cumulative(self):
        for period, normalized, heading in [('2026年第一季度', '2026Q1', '本季数'),
                ('2026年第二季度', '2026Q2', '本季度金额'), ('2026年上半年', '2026H1', '本半年数'),
                ('2026H2', '2026H2', '本半年金额'),
                ('2026年1月1日至2026年3月31日', '2026Q1', '本期金额'),
                ('2026年1至3月', '2026Q1', '本季数'),
                ('2026年1月至2026年6月', '2026H1', '本半年数'),
                ('2026 - Q2', '2026Q2', '本季度金额')]:
            with self.subTest(period=period):
                raw = export_book(period=period, current_heading=heading)
                doc = self.good(raw)
                self.assertFalse(doc['import_mapping']['pending'])
                self.assertEqual(doc['company']['period'], normalized)
                self.assertEqual(doc['rows'][0]['value'], '100000')
                with self.assertRaisesRegex(InputError, '期间不能'):
                    materials._excel(raw, doc, keys=KEYS, allow_incomplete_company=True,
                        import_options={'利润表': {'period': '2026-01'}})
        for period in ('2026Q2', '2026H2'):
            doc = self.good(export_book(period=period))
            self.assertTrue(doc['import_mapping']['pending'])
            self.assertFalse(doc['rows'])

    def test_export_dates_are_not_source_endpoints_and_conflicting_periods_block(self):
        for stamp in ('2026-02-02', '2026/02/02', '2026年2月2日'):
            doc = self.good(export_book(period='2026年1月 导出日期：' + stamp))
            self.assertEqual(doc['company']['period'], '2026-01')
            self.assertEqual(doc['import_mapping']['sheets'][0]['as_of'], '')
        for period in ('2026Q1 2026Q2', '2026年1月 2026年上半年'):
            self.assertIn('期间不唯一', parse(export_book(period=period))['error'])

    def test_currency_field_order_and_aliases_do_not_consume_period_fields(self):
        for label in ('币别', '本位币', '记账本位币', 'Currency'):
            meta = financial_import._money_metadata(f'{label}：CNY 2026Q1 单位：人民币万元', '表')
            self.assertEqual((meta['currency'], meta['unit']), ('CNY', '人民币万元'))
            meta = financial_import._money_metadata(f'{label}：USD 2026Q1 单位：万元', '表')
            self.assertFalse(meta['currency_supported'])
        meta = financial_import._money_metadata('货币单位：人民币 2026Q1', '表')
        self.assertTrue(meta['currency_supported'])
        self.assertEqual(meta['unit'], '')

    def test_row_currency_is_verified_and_only_genuine_merges_supply_missing_values(self):
        from openpyxl import load_workbook
        for currency in ('USD', None, 'CNY'):
            for merged in (False, True):
                book = load_workbook(BytesIO(export_book(currency='CNY')))
                sheet = book['科目余额表']
                sheet['I2'] = '币种'
                sheet['I4'] = currency
                if merged:
                    sheet.merge_cells('I4:I7')
                else:
                    for row in range(5, 8):
                        sheet.cell(row, 9, 'CNY')
                raw = BytesIO()
                book.save(raw)
                book.close()
                doc = self.good(raw.getvalue())
                if currency == 'CNY':
                    self.assertFalse(doc['import_mapping']['pending'])
                    self.assertEqual(doc['accounts'][0]['credit'], '100000')
                else:
                    self.assertIn('币种', str(doc['import_mapping']['pending']))
                    self.assertFalse(doc['accounts'])
                    self.assertFalse(doc['rows'])

    def test_adjacent_metadata_cells_are_not_mistaken_for_currency_data_columns(self):
        from openpyxl import load_workbook
        for currency in ('USD', 'CNY', None):
            book = load_workbook(BytesIO(export_book()))
            sheet = book['科目余额表']
            sheet['I1'], sheet['J1'] = '币种', currency
            output = BytesIO()
            book.save(output)
            book.close()
            doc = self.good(output.getvalue())
            if currency == 'CNY':
                self.assertFalse(doc['import_mapping']['pending'])
                self.assertEqual(doc['accounts'][0]['credit'], '100000')
            else:
                self.assertTrue(doc['import_mapping']['pending'])
                self.assertFalse(doc['accounts'])

    def test_unsupported_original_interval_cannot_inherit_or_accept_other_period(self):
        for period in ('2026年1至9月', '2026年第一季', '2026Q5'):
            raw = export_book(period=period, standard=True, current_heading='本期金额')
            doc = self.good(raw)
            self.assertTrue(doc['import_mapping']['pending'])
            self.assertTrue(doc['import_mapping']['sheets'][0]['period_unresolved'])
            self.assertEqual(doc['import_mapping']['sheets'][0]['source_period'], '')
            self.assertFalse(doc['rows'])
            materials._excel(raw, doc, keys=KEYS, allow_incomplete_company=True,
                import_options={s: {'period': '2026-01'} for s in ('科目余额表', '利润表')})
            self.assertTrue(doc['import_mapping']['pending'])
            self.assertFalse(doc['rows'])

    def test_legacy_codes_and_name_checked_mapping_are_preserved(self):
        doc = self.good(export_book(legacy=True, code_sheet='科目余额'))
        self.assertEqual(doc['accounts'][0]['code'], '5101')
        rows = {r['name']: r for r in doc['rows']}
        self.assertEqual(rows['营业收入']['value'], '100000')
        self.assertIn('科目余额!F4', rows['营业收入']['source'])
        self.assertEqual(rows['营业成本']['value'], '60000')

    def test_stock_date_end_columns_and_aliases(self):
        doc = self.good(export_book(balance=True))
        rows = {r['name']: r for r in doc['rows']}
        self.assertEqual(rows['资产负债表.资产总额']['value'], '200000')
        self.assertEqual(rows['资产负债表.负债总额']['value'], '80000')
        self.assertEqual(doc['import_mapping']['sheets'][-1]['as_of'], '2026-01-31')
        self.assertEqual(doc['import_mapping']['sheets'][-1]['period'], '')

    def test_mixed_standard_and_export_edits_keep_other_tables(self):
        doc = self.good(export_book(standard=True))
        saved = deepcopy(doc)
        candidate, changes = material_review.apply(doc, {'利润表!C3': '12', '增值税申报!B2': '110000'})
        self.assertEqual(doc, saved)
        self.assertEqual(candidate['declarations']['销售额'], '110000')
        self.assertEqual(next(r for r in candidate['rows'] if r['name'] == '利润表.营业收入')['value'], '120000')
        self.assertEqual(len(changes), 2)
        self.assertIn('原文 10', next(r for r in candidate['rows'] if r['name'] == '利润表.营业收入')['source'])

    def test_clearing_and_zero_have_different_meanings(self):
        doc = self.good(export_book())
        cleared, _ = material_review.apply(doc, {'利润表!C3': None})
        zeroed, _ = material_review.apply(doc, {'利润表!C3': '0'})
        self.assertFalse(any(r['name'] == '利润表.营业收入' for r in cleared['rows']))
        self.assertEqual(next(r for r in zeroed['rows'] if r['name'] == '利润表.营业收入')['value'], '0')
        with self.assertRaises(InputError):
            material_review.apply(doc, {'利润表!D3': '1'})

    def test_numbers_missing_metadata_and_formulas(self):
        doc = self.good(export_book(value='1,234.50'))
        self.assertEqual(Decimal(doc['accounts'][0]['credit']), Decimal('12345000'))
        negative = self.good(export_book(value='(12.3)'))
        self.assertEqual(Decimal(negative['accounts'][0]['credit']), Decimal('-123000'))
        missing = self.good(export_book(unit=''))
        self.assertTrue(missing['import_mapping']['pending'])
        self.assertTrue(material_security.blockers(missing, {}, scope()))
        formula = self.good(export_book(formula=True))
        field = next(f for f in material_review.fields(formula) if f['id'] == '利润表!C3')
        self.assertEqual(field['formula'], '=1+1')
        self.assertIsNone(field['value'])
        self.assertEqual(field['cache_status'], 'missing')
        self.assertFalse(any(r['name'] == '利润表.营业收入' for r in formula['rows']))
        self.assertIn('期间', parse(export_book(period='2025年12月 2026年1月'))['error'])
        self.assertIn('分隔符', parse(export_book(value='1.234,50'))['error'])
        self.assertIn('分隔符', parse(export_book(value='1,23'))['error'])

    def test_mapping_review_is_bound_to_values_scope_and_purpose(self):
        doc = self.good(export_book(legacy=True))
        selection = {'purpose': 'current', 'standard_edits': {}}
        requirements = material_security.review_requirements(doc, selection, scope())
        selection['evidence_reviews'] = [r['id'] for r in requirements]
        self.assertFalse(material_security.blockers(doc, selection, scope()))
        selection['standard_edits'] = {'科目余额表!F4': '11'}
        self.assertTrue(material_security.blockers(doc, selection, scope()))
        selection['standard_edits'] = {}
        selection['purpose'] = 'history'
        self.assertTrue(material_security.blockers(doc, selection, scope()))

    def test_reviewed_mapping_can_build_actual_rule_dataset_after_correction(self):
        from tests.test_materials import ROOT
        from src import engine
        rules = engine.load_rules(ROOT / 'rules')
        doc = self.good(export_book(legacy=True))
        selection = {'standard_edits': {'利润表!C3': '12'}}
        selection['evidence_reviews'] = [r['id'] for r in material_security.review_requirements(doc, selection, scope())]
        result = enterprise_analysis.analyze([doc], scope(), {doc['id']: selection}, rules)
        self.assertTrue(result['can_confirm'], result['feedback']['blocking'])
        self.assertEqual(Decimal(result['dataset']['metrics']['营业收入']['value']), Decimal('100000'))
        self.assertEqual(Decimal(result['dataset']['metrics']['利润表.营业收入']['value']), Decimal('120000'))

    def test_annual_values_are_not_monthly_values(self):
        doc = self.good(export_book(period='2026年度'))
        row = next(r for r in doc['rows'] if r['name'] == '利润表.营业收入')
        self.assertEqual(doc['company']['period'], '2026')
        self.assertEqual(row['value'], '9000000')
        self.assertIn('利润表!D3', row['source'])

    def test_source_company_conflict_and_duplicate_accounts_are_not_hidden(self):
        from openpyxl import load_workbook
        book = load_workbook(BytesIO(export_book(balance=True)))
        book['科目余额表'].append(['6001', '主营业务收入', 0, 0, 0, 1, 0, 0])
        output = BytesIO()
        book.save(output)
        book.close()
        self.assertIn('重复', parse(output.getvalue())['error'])
        for change, expected, excluded in [('company', '企业名称', '利润表'),
                ('period', '实际期间', '利润表'), ('date', '时点', '资产负债表')]:
            with self.subTest(change=change):
                raw = mixed_export(change)
                doc = self.good(raw)
                self.assertIn(expected, str(doc['import_mapping']['pending']))
                self.assertFalse(doc['accounts'])
                self.assertFalse(doc['rows'])
                # Hand-entered scope and acknowledging every item do not clear
                # contradictory facts in the original workbook.
                selection = {'evidence_reviews': [r['id'] for r in material_security.review_requirements(doc, {}, scope())]}
                self.assertFalse(enterprise_analysis.analyze([doc], scope(), {doc['id']: selection}, [])['can_confirm'])
                self.assertTrue(material_security.blockers(doc, selection, scope(), require_review=False))
                materials._excel(raw, doc, allow_incomplete_company=True, capture_standard=True,
                    keys=KEYS, import_options={excluded: {'excluded': True}})
                self.assertFalse(doc['import_mapping']['pending'])
                self.assertEqual(doc['company']['name'], COMPANY['name'])
                self.assertEqual(doc['company']['period'], COMPANY['period'])
                self.assertEqual(doc['sha256'], sha256(raw).hexdigest())
                self.assertFalse(any(r['name'].startswith(excluded + '.') for r in doc['rows']))
                self.assertIn('已排除', next(r for r in material_security.review_requirements(doc, {}, scope())
                    if r['label'].endswith(excluded))['reason'])

    def test_original_title_handles_software_sheet_names_and_duplicate_kinds(self):
        from openpyxl import load_workbook
        book = load_workbook(BytesIO(export_book()))
        book['科目余额表'].title = '余额导出_202601'
        profit = book['利润表']
        profit.insert_rows(1)
        profit['A1'] = '利润表'
        profit.title = 'sheet1'
        output = BytesIO()
        book.save(output)
        doc = self.good(output.getvalue())
        self.assertEqual({s['sheet']: s['kind'] for s in doc['import_mapping']['sheets']},
                         {'余额导出_202601': '科目余额表', 'sheet1': '利润表'})
        self.assertFalse(doc['import_mapping']['ignored_sheets'])
        self.assertIn('sheet1!C4', next(r for r in doc['rows'] if r['name'] == '利润表.营业收入')['source'])
        copy = book.copy_worksheet(profit)
        copy.title = 'sheet2'
        output = BytesIO()
        book.save(output)
        book.close()
        raw = output.getvalue()
        doc = self.good(raw)
        self.assertIn('相同表类', str(doc['import_mapping']['pending']))
        self.assertFalse(doc['rows'])
        materials._excel(raw, doc, allow_incomplete_company=True, keys=KEYS,
                         import_options={'sheet2': {'excluded': True}})
        self.assertFalse(doc['import_mapping']['pending'])
        self.assertEqual(sum(r['name'] == '利润表.营业收入' for r in doc['rows']), 1)

    def test_generic_worksheet_is_not_typed_from_a_familiar_amount_label(self):
        from openpyxl import load_workbook
        book = load_workbook(BytesIO(export_book()))
        book['利润表'].title = '无法确认表类'
        output = BytesIO()
        book.save(output)
        book.close()
        doc = self.good(output.getvalue())
        self.assertIn('无法确认表类', doc['import_mapping']['ignored_sheets'])
        self.assertFalse(any(r['name'].startswith('利润表.') for r in doc['rows']))

    def test_company_sheet_is_an_anchor_not_removed_with_conflicting_export(self):
        raw = mixed_export('company', standard=True)
        doc = self.good(raw)
        materials._excel(raw, doc, allow_incomplete_company=True, keys=KEYS,
                         import_options={'科目余额表': {'excluded': True}, '资产负债表': {'excluded': True}})
        self.assertIn('企业名称', str(doc['import_mapping']['pending']))
        self.assertEqual(doc['company']['name'], COMPANY['name'])
        self.assertFalse(any(r['name'].startswith('利润表.') for r in doc['rows']))

    def test_missing_units_can_be_explicitly_mapped_but_source_cannot_be_overridden(self):
        raw = export_book(unit='')
        doc = self.good(raw)
        self.assertFalse(doc['accounts'])
        original = deepcopy(doc)
        choices = {'科目余额表': {'unit': '万元'}, '利润表': {'unit': '万元'}}
        materials._excel(raw, doc, allow_incomplete_company=True, capture_standard=True, keys=KEYS, import_options=choices)
        self.assertFalse(doc['import_mapping']['pending'])
        self.assertEqual(doc['accounts'][0]['credit'], '100000')
        self.assertEqual(doc['import_mapping']['sheets'][0]['unit_origin'], 'user')
        self.assertEqual(doc['import_mapping']['sheets'][0]['source_unit'], '')
        self.assertEqual(doc['sha256'], original['sha256'])
        with self.assertRaisesRegex(InputError, '单位不能'):
            materials._excel(export_book(), deepcopy(doc), allow_incomplete_company=True, keys=KEYS,
                             import_options={'利润表': {'unit': '元'}})
        with self.assertRaisesRegex(InputError, '期间不能'):
            materials._excel(export_book(), deepcopy(doc), allow_incomplete_company=True, keys=KEYS,
                             import_options={'利润表': {'period': '2025-12'}})

    def test_missing_period_and_ambiguous_columns_require_original_column_choices(self):
        from openpyxl import load_workbook
        book = load_workbook(BytesIO(export_book(period='')))
        book['利润表']['C2'] = '本期金额'
        book['利润表']['D2'] = '本期金额'
        output = BytesIO()
        book.save(output)
        book.close()
        raw = output.getvalue()
        doc = self.good(raw)
        self.assertTrue(doc['import_mapping']['pending'])
        choices = {'科目余额表': {'period': '2026-01'},
                   '利润表': {'period': '2026-01', 'columns': {'2:A': 'D'}}}
        materials._excel(raw, doc, allow_incomplete_company=True, capture_standard=True, keys=KEYS, import_options=choices)
        self.assertFalse(doc['import_mapping']['pending'])
        self.assertEqual(next(r for r in doc['rows'] if r['name'] == '利润表.营业收入')['value'], '9000000')
        self.assertEqual(doc['import_mapping']['sheets'][1]['period_origin'], 'user')
        choices['利润表']['columns']['2:A'] = 'B'
        with self.assertRaisesRegex(InputError, '所选金额列'):
            materials._excel(raw, doc, allow_incomplete_company=True, keys=KEYS, import_options=choices)

    def test_cumulative_column_and_forged_mapping_keys_cannot_override_month(self):
        raw = export_book()
        doc = self.good(raw)
        for choice in [{'利润表': {'columns': {'2:A': 'D'}}}, {'利润表': {'columns': {'9:Z': 'C'}}},
                       {'不存在': {'unit': '万元'}}, {'利润表': {'source': '伪造来源'}}]:
            with self.assertRaises(InputError):
                materials._excel(raw, deepcopy(doc), allow_incomplete_company=True, keys=KEYS, import_options=choice)

    def test_csv_tsv_source_title_and_blank_records_preserve_coordinates(self):
        rows = ['', '利润表', f'编制单位：{COMPANY["name"]} 2026年1月 单位：万元',
                '项目,行次,本月数,本年累计数', ',,,', '营业收入,1,10,900']
        for kind, encoding in [('csv', 'utf-8-sig'), ('csv', 'gb18030'), ('tsv', 'utf-16')]:
            text = '\r\n'.join(rows)
            if kind == 'tsv':
                text = text.replace(',', '\t')
            raw = text.encode(encoding)
            # The misleading filename must not choose a balance-sheet interpretation.
            doc = materials.preview([('资产负债表.' + kind, raw)], KEYS,
                                    allow_incomplete_company=True, capture_standard=True)[0]
            self.assertEqual(doc['error'], '', doc['error'])
            self.assertEqual(doc['sha256'], sha256(raw).hexdigest())
            self.assertEqual(doc['import_mapping']['sheets'][0]['kind'], '利润表')
            income = next(r for r in doc['rows'] if r['name'] == '利润表.营业收入')
            self.assertEqual(income['value'], '100000')
            self.assertIn('利润表!C6', income['source'])
            self.assertEqual(material_review.fields(doc)[0]['id'], '利润表!C6')

    def test_untitled_account_csv_requires_explicit_missing_metadata(self):
        raw = '科目代码,会计科目,本期借方,本期贷方\n6001,主营业务收入,0,10'.encode()
        doc = materials.preview([('随机文件.csv', raw)], KEYS,
                                allow_incomplete_company=True, capture_standard=True)[0]
        self.assertEqual(doc['error'], '', doc['error'])
        self.assertTrue(doc['import_mapping']['pending'])
        choices = {'科目余额表': {'unit': '万元', 'period': '2026-01'}}
        materials._excel(materials._table_workbook('随机文件.csv', raw, 'csv'), doc,
                         allow_incomplete_company=True, capture_standard=True, keys=KEYS, import_options=choices)
        self.assertEqual(doc['accounts'][0]['credit'], '100000')
        self.assertEqual(doc['import_mapping']['sheets'][0]['period_origin'], 'user')
        self.assertEqual(material_review.fields(doc)[1]['id'], '科目余额表!D2')

    def test_csv_ambiguous_titles_and_untyped_units_are_rejected(self):
        with self.assertRaisesRegex(InputError, '表类冲突'):
            financial_import.table_kind([['利润表', ''], ['资产负债表', ''], ['项目', '本期金额']])
        # Generic statement columns do not establish which financial statement this is.
        self.assertEqual(financial_import.table_kind([['项目', '本期金额'], ['营业收入', '10']]), '')
        for value in [[], {}, None, True, 10000, '美元']:
            with self.subTest(unit=value), self.assertRaises(InputError):
                financial_import.options({'利润表': {'unit': value}}, ['利润表'])
        with self.assertRaisesRegex(InputError, '表头及数据'):
            material_formats.table_rows(b'\n\na,b\n\n', 'csv')
        with self.assertRaisesRegex(InputError, '超过'):
            material_formats.table_rows(b'\n' * 30000 + b'a,b\n1,2', 'csv')
        for raw in ['利润表\n项目,本月数,本年累计数\n收入,10',
                    '利润表\n项目,本月数\n收入,10,900', '利润表\n项目,本月数\n收入,=1+1']:
            with self.subTest(raw=raw), self.assertRaises(InputError):
                material_formats.table_rows(raw.encode(), 'csv')

    def test_formula_cache_is_candidate_only_and_every_amount_requires_review(self):
        raw = formula_book()
        doc = self.good(raw)
        field = next(f for f in material_review.fields(doc) if f['id'] == '利润表!C3')
        self.assertEqual((field['formula'], field['cached_value'], field['value']), ('=1+1', '10', None))
        self.assertFalse(any(r['name'] == '利润表.营业收入' for r in doc['rows']))
        selection = {'standard_edits': {'利润表!C3': '10'}}
        requirements = material_security.review_requirements(doc, selection, scope())
        selection['evidence_reviews'] = [r['id'] for r in requirements if not r['key'].startswith('formula:')]
        self.assertTrue(material_security.blockers(doc, selection, scope(), require_review=False))
        selection['evidence_reviews'] = [r['id'] for r in requirements]
        self.assertFalse(material_security.blockers(doc, selection, scope()))
        candidate, changes = material_review.apply(doc, selection['standard_edits'])
        row = next(r for r in candidate['rows'] if r['name'] == '利润表.营业收入')
        self.assertEqual(row['value'], '100000')
        self.assertIn('原公式 =1+1', row['source'])
        self.assertIn('用户核对值 10', row['source'])
        self.assertEqual(changes[0]['old'], None)
        self.assertEqual(doc['sha256'], sha256(raw).hexdigest())
        selection['standard_edits']['利润表!C3'] = '12'
        self.assertTrue(material_security.blockers(doc, selection, scope()))

    def test_formula_missing_error_text_and_zero_are_distinct(self):
        for raw, kind, candidate in [(None, 'n', None), ('#VALUE!', 'e', None),
                ('TRUE', 'b', None), ('10', 'str', None), ('NaN', 'n', None),
                ('1,000', 'n', None), ('1e30', 'n', None), ('0', 'n', '0'), ('-2.5', 'n', '-2.5')]:
            with self.subTest(cache=raw, kind=kind):
                doc = self.good(formula_book(cache=raw, cache_type=kind))
                field = next(f for f in material_review.fields(doc) if f['id'] == '利润表!C3')
                self.assertIsNone(field['value'])
                self.assertEqual(field['cached_value'], candidate)
                if candidate is not None:
                    parsed, _ = material_review.apply(doc, {'利润表!C3': candidate})
                    self.assertEqual(Decimal(next(r for r in parsed['rows'] if r['name'] == '利润表.营业收入')['value']),
                                     Decimal(candidate) * 10000)
        doc = self.good(formula_book(cache='0'))
        zero = material_review.normalize(doc, {'利润表!C3': '0'})
        self.assertEqual(zero, {'利润表!C3': '0'})
        cleared, _ = material_review.apply(doc, {'利润表!C3': None})
        self.assertFalse(any(r['name'] == '利润表.营业收入' for r in cleared['rows']))

    def test_shared_formulas_keep_translated_source_addresses_without_execution(self):
        doc = self.good(formula_book(shared=True))
        fields = {f['id']: f for f in material_review.fields(doc)}
        self.assertEqual(fields['利润表!C3']['formula'], '=SUM(A3:B3)')
        self.assertEqual(fields['利润表!C4']['formula'], '=SUM(A4:B4)')
        self.assertEqual(fields['利润表!C3']['formula_attributes'], {'t':'shared','si':'0','ref':'C3:C4'})
        self.assertEqual(fields['利润表!C4']['formula_attributes'], {'t':'shared','si':'0'})
        self.assertEqual({f['id'] for f in doc['import_mapping']['formula_cells']}, {'利润表!C3', '利润表!C4'})
        selection = {'standard_edits': {'利润表!C3': '10', '利润表!C4': '6'}}
        self.assertEqual(sum(r['key'].startswith('formula:') for r in
                            material_security.review_requirements(doc, selection, scope())), 2)

    def test_standard_financial_formula_keeps_template_unit_and_can_be_reviewed(self):
        doc = self.good(formula_book(standard=True))
        sheet = doc['import_mapping']['sheets'][0]
        self.assertEqual((sheet['unit'], sheet['unit_origin']), ('元', 'template'))
        fields = {f['id']: f for f in material_review.fields(doc)}
        self.assertEqual(fields['利润表!B2']['cached_value'], '10')
        candidate, _ = material_review.apply(doc, {'利润表!B2': '12'})
        self.assertEqual(next(r['value'] for r in candidate['rows'] if r['name'] == '利润表.营业收入'), '12')
