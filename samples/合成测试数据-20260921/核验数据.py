r"""只读验证16套合成Excel；运行当前项目的导入、规则和HTML报告链路。

从项目根目录运行：.venv\Scripts\python.exe samples/合成测试数据-20260921/核验数据.py
只写同目录核验结果.json，不修改输入工作簿、业务代码或现有报告。
"""
from __future__ import annotations

from collections import defaultdict
from datetime import datetime
from decimal import Decimal
import hashlib
from io import BytesIO
import json
from pathlib import Path
import re
import sys

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(ROOT))
from openpyxl import load_workbook
from src import loader, engine, render

D = lambda x: Decimal(str(x or 0))
manifest = json.loads((HERE / '预期结果.json').read_text(encoding='utf-8'))
rules = engine.load_rules(ROOT / 'rules')
checks, failures, results = 0, [], []


def check(condition, label):
    global checks
    checks += 1
    if not condition:
        failures.append(label)


def hashes():
    files = [ROOT / p for p in ('src/config.py', 'src/loader.py', 'src/engine.py', 'src/models.py', 'src/render.py', 'templates/report.html')]
    files += sorted((ROOT / 'rules').glob('*.yaml'))
    return {str(p.relative_to(ROOT)): hashlib.sha256(p.read_bytes()).hexdigest() for p in files}


def main():
    start_hashes = hashes()
    check(len(rules)==24, '规则数量应为24；若规则更新，请复核本包预期')
    hit_coverage, pass_coverage, skipped_coverage = set(), set(), set()
    for case in manifest['cases']:
        label = case['id']
        file = HERE / case['file']
        raw = file.read_bytes()
        wb = load_workbook(BytesIO(raw), data_only=False)
        check(len(raw) < 10*1024*1024, label+'上传文件小于10MB')
        check(wb.sheetnames == ['企业信息','科目余额表','增值税申报','补充指标','模拟业务底稿','模拟凭证'], label+'工作表完整')
        for sheet in wb:
            for row in sheet:
                for cell in row:
                    check(cell.data_type!='f', label+f'输入快照不得含Excel公式：{sheet.title}!{cell.coordinate}')
        company = dict(wb['企业信息'].iter_rows(min_row=2,values_only=True))
        check('纯合成测试' in company['企业名称'] and company['纳税人识别号'].startswith('SIM-'), label+'合成标记')
        accts = list(wb['科目余额表'].iter_rows(min_row=2,values_only=True))
        vouchers, totals = defaultdict(lambda:[Decimal(0),Decimal(0)]), defaultdict(lambda:[Decimal(0),Decimal(0)])
        for row in wb['模拟凭证'].iter_rows(min_row=2,values_only=True):
            no,when,summary,code,name,debit,credit = row
            check(when.year==2026 and when.month==8, label+'凭证所属月')
            check(D(debit)>=0 and D(credit)>=0, label+'凭证金额非负')
            vouchers[no][0] += D(debit); vouchers[no][1] += D(credit)
            totals[str(code)][0] += D(debit); totals[str(code)][1] += D(credit)
        check(len(vouchers)==case['voucher_count'], label+'凭证数量')
        for no,(debit,credit) in vouchers.items():
            check(debit==credit,label+no+'凭证借贷平衡')
        # Duplicate-code case intentionally breaks the whole-ledger aggregate.
        for code,name,opening,debit,credit,closing in accts:
            check(D(opening)+D(debit)-D(credit)==D(closing),label+code+'余额滚动')
            check([D(debit),D(credit)]==totals[str(code)],label+code+'余额表发生额与凭证一致')
        if label!='15':
            check(sum(D(r[2]) for r in accts)==0,label+'期初借贷平衡')
            check(sum(D(r[3])-D(r[4]) for r in accts)==0,label+'本期借贷平衡')
            check(sum(D(r[5]) for r in accts)==0,label+'期末借贷平衡')
        for key,value,source,period,detail in wb['补充指标'].iter_rows(min_row=2,values_only=True):
            match=re.fullmatch(r'(.+)!([A-Z]+[0-9]+)（合成）',source)
            check(match is not None,label+key+'明确底稿定位')
            if match:
                check(wb[match[1]][match[2]].value==value,label+key+'底稿与导入值一致')
            if label!='16':check(period==company['所属期'],label+key+'所属期一致')
            check(bool(detail),label+key+'有口径说明')
        wb.close()
        item={'id':label,'file':case['file'],'sha256':hashlib.sha256(raw).hexdigest(),'bytes':len(raw)}
        error = None
        try:
            dataset = loader.load(file)
        except loader.InputError as exc:
            error = str(exc)
        if case['expected_error']:
            check(error is not None and case['expected_error'] in error,label+'按预期拒绝文件导入')
            try:
                loader.load_bytes(raw)
                check(False,label+'内存导入应拒绝')
            except loader.InputError as exc:
                check(case['expected_error'] in str(exc),label+'内存导入拒绝原因')
            item.update(import_status='rejected_as_expected',error=error)
        else:
            check(error is None,label+'导入成功')
            if error:continue
            memory_data=loader.load_bytes(raw)
            check(dataset.values()==memory_data.values(),label+'文件和内存导入结果一致')
            findings=engine.run(rules,dataset)
            actual={f.rule.id:f.status for f in findings}
            for rid,status in case['expected_status'].items():
                check(actual.get(rid)==status,label+rid+'预期'+status+'实际'+str(actual.get(rid)))
            for f in findings:
                if f.status=='hit':hit_coverage.add(f.rule.id)
                if f.status=='pass':pass_coverage.add(f.rule.id)
                if f.status=='skipped':skipped_coverage.add(f.rule.id)
                if f.executed:
                    check(bool(f.calculation) and all(e.source for e in f.evidence),label+f.rule.id+'计算与证据来源')
                else:check(bool(f.skip_reason),label+f.rule.id+'未执行理由')
                if f.rule.id in case['expected_measured']:
                    expected=case['expected_measured'][f.rule.id]
                    check(f.measured is None if expected is None else f.measured is not None and abs(f.measured-expected)<1e-9,label+f.rule.id+'独立算值')
            vm=render.build_view_model(dataset,findings)
            check(sum(vm['summary'][x] for x in ['hit','pass','skipped'])==24,label+'报告状态计数')
            html,_=render.render_html(dataset,findings,write=False)
            check(company['企业名称'] in html and all(r.id in html for r in rules),label+'HTML报告含主体和全部规则')
            item.update(import_status='accepted',summary=vm['summary'],findings=[{'id':f.rule.id,'status':f.status,'measured':f.measured,'calculation':f.calculation,'skip_reason':f.skip_reason} for f in findings])
        results.append(item)
    check(hit_coverage=={r.id for r in rules},'全部24条规则均有命中样例')
    check(pass_coverage=={r.id for r in rules},'全部24条规则均有通过样例')
    check(start_hashes==hashes(),'核验期间业务代码及规则未变化')
    report={'data_kind':'synthetic','verified_at':datetime.now().isoformat(timespec='seconds'),
            'checks':checks,'failed':len(failures),'failures':failures,'cases':len(results),
            'accepted':sum(x['import_status']=='accepted' for x in results),'rejected_as_expected':sum(x['import_status']=='rejected_as_expected' for x in results),
            'rule_evaluations':sum(len(x.get('findings',[])) for x in results),
            'hit_coverage':sorted(hit_coverage),'pass_coverage':sorted(pass_coverage),'skipped_coverage':sorted(skipped_coverage),
            'verification_scope':'本地文件导入、内存导入、规则引擎和HTML报告；未验证Web HTTP接口、浏览器界面或PDF导出。',
            'code_sha256':start_hashes,'results':results}
    (HERE/'核验结果.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({k:report[k] for k in ['checks','failed','failures','cases','accepted','rejected_as_expected','rule_evaluations','hit_coverage']},ensure_ascii=False))
    return bool(failures)


if __name__=='__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    raise SystemExit(main())
