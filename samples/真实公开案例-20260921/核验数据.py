"""本次公开数据交付的离线核验。使用项目 .venv Python 运行，不修改业务代码。"""
from pathlib import Path
from decimal import Decimal
import hashlib
import json
import sys

HERE = Path(__file__).resolve().parent
REPO = HERE.parent.parent
sys.path.insert(0, str(REPO))
from openpyxl import load_workbook

D = lambda x: Decimal(str(x))
cases = [json.loads(p.read_text(encoding='utf-8')) for p in sorted(HERE.glob('*/公开数据.json'))]
manifest = json.loads((HERE / '来源清单.json').read_text(encoding='utf-8'))
checks = []

def check(name, ok, details=''):
    checks.append(dict(name=name, passed=bool(ok), details=details))
    if not ok:
        raise AssertionError(f'{name}: {details}')

for src in manifest:
    if 'path' not in src:
        continue
    p = HERE / src['path']
    check('来源文件SHA256：' + src['path'], hashlib.sha256(p.read_bytes()).hexdigest() == src['sha256'])

wbv = load_workbook(HERE / '真实案例数据总表.xlsx', data_only=True)
wbf = load_workbook(HERE / '真实案例数据总表.xlsx', data_only=False)
count = 0
for c in cases:
    ws, wf = wbv[c['folder']], wbf[c['folder']]
    check(c['case_id'] + '不伪装为完整账套', c['can_import_current_loader'] is False)
    for i, f in enumerate(c['facts'], 9):
        count += 1
        check(f['fact_id'] + '原值一致', ws.cell(i, 3).value == f['value'])
        if f['value'] is None:
            check(f['fact_id'] + '缺项保持空白', ws.cell(i, 5).value is None and f['normalized_value'] is None)
        else:
            expected = D(f['value']) * D(f['multiplier'])
            check(f['fact_id'] + 'JSON单位换算', abs(D(f['normalized_value']) - expected) < D('0.0000001'))
            check(f['fact_id'] + 'Excel公式及缓存', wf.cell(i, 5).data_type == 'f' and isinstance(ws.cell(i, 5).value, (float, int)) and abs(D(ws.cell(i, 5).value) - expected) < D('0.00001'))
        check(f['fact_id'] + '定位完整', bool(f['source_locator']) and c['source']['url'].startswith('https://'))

def val(c, name, period=None):
    fs = [f for f in c['facts'] if f['name'] == name and (period is None or f['period'] == period)]
    assert len(fs) == 1
    return D(fs[0]['value'])

j = cases[0]
for year in ('2016', '2017', '2018'):
    check('嘉元' + year + '收入差额', val(j, '申报材料营业收入', year) - val(j, '申报表计入营业收入合计', year) == val(j, '收入与申报差额（披露）', year))
    # 2016/2017原文上期留抵为横杠，核验不把其设为真实0。
check('嘉元2018增值税含转出及上期留抵', val(j, '销项税额', '2018') - val(j, '进项税额', '2018') + val(j, '进项税额转出', '2018') - val(j, '上期留抵税金', '2018') == val(j, '本期应交增值税', '2018'))
s = cases[3]
check('星印查实收入减申报收入', val(s, '查实销售收入') - val(s, '已申报销售额') == val(s, '未按规定申报收入'))
check('春风人数差（实际用工口径）', val(cases[2], '个税代扣代缴申报人数') - val(cases[2], '实际用工人数') == 25)
check('禁止实际人数代替社保', not any(f['metric_mapping'] == '人力.社保参保人数' for c in cases for f in c['facts']))
lx = cases[-1]
check('隆旭发票批次份数', sum(b['count'] for b in lx['invoice_batches']) == 263)
for b in lx['invoice_batches']:
    check('隆旭价税合计：' + b['seller'], D(b['net']) + D(b['tax']) == D(b['gross']))
for row, b in zip((33, 34, 35), lx['invoice_batches']):
    ws = wbv[lx['folder']]
    check('发票代码文本保真：' + b['seller'], ws.cell(row, 5).value == b['invoice_code'] and ws.cell(row, 5).data_type == 's')
    check('发票号码前导零保真：' + b['seller'], ws.cell(row, 7).value == b['invoice_numbers'])
check('隆旭批次进项税额合计', sum(D(b['tax']) for b in lx['invoice_batches']) == val(lx, '应补缴增值税'))
check('隆旭税费合计', sum(val(lx, n) for n in ['应补缴增值税', '城市维护建设税', '教育费附加', '地方教育附加', '应补缴企业所得税']) == val(lx, '少缴税费合计'))
check('隆旭冲突保留', val(lx, '富顺康税额（冲突概述）') != D(lx['invoice_batches'][1]['tax']))

arithmetic = []
for year in ('2016', '2017', '2018'):
    revenue = val(j, '申报材料营业收入', year)
    declared = val(j, '申报表计入营业收入合计', year)
    ratio = abs(revenue - declared) / abs(declared)
    arithmetic.append(dict(case_id='01', period=year, formula='abs(披露营业收入-申报销售额)/abs(申报销售额)', result=float(ratio), threshold=0.1, above_threshold=ratio>D('0.1'), status='独立算术复核，非调用项目规则引擎'))
    check('嘉元' + year + '独立偏离度复算', ratio < D('0.1'))

# 业务代码同时在其他任务中修改，仅记录当前快照，不修复或代替实际运行。
engine_path = REPO / 'src/engine.py'
try:
    compile(engine_path.read_text(encoding='utf-8'), str(engine_path), 'exec')
    engine_syntax = dict(passed=True, details='仅语法编译；未执行端到端或规则引擎集成测试')
except SyntaxError as e:
    engine_syntax = dict(passed=False, details=f'{type(e).__name__}: {e.msg}; line={e.lineno}')
check('总表未伪装MVP三工作表', not {'企业信息', '科目余额表', '增值税申报'}.issubset(set(wbv.sheetnames)))

report = dict(verified_date='2026-09-21', case_count=len(cases), facts=count,
              complete_importable_ledger_sets=0, checks_passed=len(checks),
              checks=checks, engine_syntax_at_verification=engine_syntax,
              engine_scope='未完成实际引擎及端到端验证；数据不满足现有完整账套入口。',
              independent_arithmetic=arithmetic,
              code_sha256={str(p.relative_to(REPO)): hashlib.sha256(p.read_bytes()).hexdigest() for p in [REPO/'src/engine.py', REPO/'src/loader.py', REPO/'src/config.py', *sorted((REPO/'rules').glob('*.yaml'))]},
              caveats=['历史披露不验证2026年适用法规。', '所有案例缺原始账套，不是六套端到端输入。', 'R-003通过仅说明当前演示区间内，不能作为制造业真实负例。'])
(HERE / '核验结果.json').write_text(json.dumps(report, ensure_ascii=True, indent=2, allow_nan=False), encoding='utf-8')
print(json.dumps({k: report[k] for k in ['case_count', 'facts', 'complete_importable_ledger_sets', 'checks_passed']}, ensure_ascii=False))
print('引擎快照：', json.dumps(engine_syntax, ensure_ascii=True))
