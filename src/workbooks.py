"""Shared Excel resource limits, with a controlled fixed-value legacy reader."""
from io import BytesIO
from pathlib import PurePosixPath
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

MAX_FILE = 10 * 1024 * 1024
MAX_EXPANDED = 50 * 1024 * 1024


def open_workbook(data, *, read_only=False):
    """Check the container before openpyxl allocates workbook objects."""
    if not data or len(data) > MAX_FILE:
        raise ValueError("Excel 文件为空或超过 10MB。")
    if data.startswith(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'):
        from .legacy_workbooks import open_legacy_workbook
        return open_legacy_workbook(data)
    try:
        with ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 1000 or sum(item.file_size for item in entries) > MAX_EXPANDED:
                raise ValueError("Excel 解压内容超过限制。")
        workbook = load_workbook(BytesIO(data), data_only=False, read_only=read_only)
    except (BadZipFile, KeyError, OSError) as exc:
        raise ValueError("Excel 工作簿损坏或格式无效。") from exc
    if any(sheet.max_row > 30000 or sheet.max_column > 100 for sheet in workbook):
        workbook.close()
        raise ValueError("每张 Excel 工作表最多 30000 行、100 列。")
    return workbook


def formula_values(data):
    """Read saved formula caches, never calculate or save the source workbook.

    Call after the normal container/security checks. Cache existence and numeric
    shape do not establish freshness, correctness, or financial attribution.
    """
    from .material_formats import archive_entries, xml_root
    namespace = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    output = {}
    with ZipFile(BytesIO(data)) as archive:
        archive_entries(archive, workbook=True)
        relationships = xml_root(archive.read('xl/_rels/workbook.xml.rels'))
        targets = {}
        for relationship in relationships:
            if relationship.attrib.get('TargetMode', '').casefold() == 'external':
                raise ValueError('Excel 工作表关系不能指向外部来源。')
            targets[relationship.attrib.get('Id')] = relationship.attrib.get('Target', '')
        book = xml_root(archive.read('xl/workbook.xml'))
        for sheet in book.findall('s:sheets/s:sheet', namespace):
            relation = sheet.attrib.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
            target = targets.get(relation)
            if not target:
                raise ValueError('Excel 工作表缺少可核对的内部关系。')
            path = PurePosixPath(target.lstrip('/') if target.startswith('/') else 'xl/' + target)
            if '..' in path.parts or not str(path).startswith('xl/'):
                raise ValueError('Excel 工作表关系路径无效。')
            root = xml_root(archive.read(str(path)))
            for cell in root.findall('s:sheetData/s:row/s:c', namespace):
                formula = cell.find('s:f', namespace)
                if formula is None:
                    continue
                value = cell.find('s:v', namespace)
                address = sheet.attrib['name'] + '!' + cell.attrib['r']
                if address in output:
                    raise ValueError('Excel 公式单元格位置重复。')
                output[address] = {'formula': formula.text or '', 'formula_type': formula.attrib.get('t', 'normal'),
                    'formula_attributes': dict(formula.attrib),
                    'cached_raw': value.text if value is not None else None, 'cached_type': cell.attrib.get('t', 'n')}
    return output
