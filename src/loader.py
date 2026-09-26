"""Excel 材料标准化；空值保持缺失，来源与期间随指标保留。"""
from __future__ import annotations

from io import BytesIO
from decimal import Decimal, InvalidOperation
from pathlib import Path
from openpyxl import load_workbook
from . import config
from .models import Account, Company, Dataset, Metric


class InputError(Exception):
    """输入材料不符合模板要求。"""


def _sheet(wb, name, header):
    if name not in wb.sheetnames:
        raise InputError(f"缺少工作表「{name}」。当前工作表：{wb.sheetnames}")
    ws = wb[name]
    actual = [c.value for c in ws[1]]
    if actual[:len(header)] != header:
        raise InputError(f"「{name}」表头不符。期望：{header}；实际：{actual}")
    return ws


def _number(value, location):
    if value is None or isinstance(value, str) and not value.strip():
        return None
    if isinstance(value, bool):
        raise InputError(f"{location}：布尔值不能作为金额或人数")
    try:
        number = Decimal(str(value))
    except (InvalidOperation, ValueError):
        raise InputError(f"{location}：不是有效数值：{value!r}") from None
    if not number.is_finite():
        raise InputError(f"{location}：数值必须有限")
    return number


def _read_company(wb):
    ws = _sheet(wb, config.SHEET_COMPANY, ["项目", "内容"])
    data = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        key = str(row[0]).strip()
        if key in data:
            raise InputError(f"企业信息项目重复：{key}")
        data[key] = "" if row[1] is None else str(row[1]).strip()
    missing = [key for key in config.COMPANY_FIELDS if not data.get(key)]
    if missing:
        raise InputError(f"企业信息缺少字段：{missing}")
    return Company(*(data[k] for k in config.COMPANY_FIELDS))


def _read_accounts(wb):
    ws = _sheet(wb, config.SHEET_ACCOUNTS, config.COL_ACCOUNTS)
    out, seen = [], set()
    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        if row[0] is None:
            continue
        code = str(row[0]).strip()
        if code in seen:
            raise InputError(f"科目余额表 A{idx}：科目编码重复 {code}；请先汇总到唯一口径")
        seen.add(code)
        numbers = [_number(row[col], f"科目余额表 {chr(65+col)}{idx}") for col in range(2, 6)]
        out.append(Account(code, str(row[1] or "").strip(), *numbers))
    if not out:
        raise InputError("科目余额表没有数据行")
    return out


def _read_declarations(wb):
    ws = _sheet(wb, config.SHEET_DECLARATION, ["项目", "金额"])
    data, seen = {}, set()
    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        if row[0] is None:
            continue
        key = str(row[0]).strip()
        if key in seen:
            raise InputError(f"增值税申报 A{idx}：项目重复 {key}")
        seen.add(key)
        value = _number(row[1], f"增值税申报 B{idx}")
        if value is not None:
            data[key] = value
    return data


def _build_metrics(accounts, declarations):
    metrics = {}
    by_code = {a.code: a for a in accounts}
    for name, spec in config.ACCOUNT_MAP.items():
        selected = [by_code.get(code) for code in spec["accounts"]]
        if any(a is None or getattr(a, spec["side"]) is None for a in selected):
            continue
        amounts = [getattr(a, spec["side"]) for a in selected]
        metrics[name] = Metric(
            name, sum(amounts, Decimal(0)), "科目余额表—" + spec["label"],
            " + ".join(f"{a.code} {a.name} {v:,.2f}" for a, v in zip(selected, amounts)),
        )
    for name, item in config.DECLARATION_ITEMS.items():
        if item in declarations:
            metrics[name] = Metric(name, declarations[item], f"增值税申报—{item}", f"{item} {declarations[item]:,.2f}")
    return metrics


def _read_supplement(wb, company, metrics):
    if config.SHEET_SUPPLEMENT not in wb.sheetnames:
        return
    ws = _sheet(wb, config.SHEET_SUPPLEMENT, config.COL_SUPPLEMENT)
    reserved = set(config.ACCOUNT_MAP) | set(config.DECLARATION_ITEMS) | set(metrics)
    seen = set()
    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        if row[0] is None:
            continue
        key = str(row[0]).strip()
        if key in seen or key in reserved:
            raise InputError(f"补充指标 A{idx}：指标重复或试图覆盖原始表指标「{key}」")
        seen.add(key)
        value = _number(row[1], f"补充指标 B{idx}")
        if value is None:
            continue
        source, period, detail = [str(v).strip() if v is not None else "" for v in row[2:5]]
        if not source or not detail:
            raise InputError(f"补充指标第 {idx} 行「{key}」必须填写来源和口径说明")
        if period != company.period:
            raise InputError(f"补充指标第 {idx} 行「{key}」所属期必须与企业信息一致；历史实际期间写入口径说明")
        if key.startswith("人力.") and key.endswith("人数") and (value < 0 or value != value.to_integral_value()):
            raise InputError(f"补充指标第 {idx} 行「{key}」人数必须为非负整数")
        metrics[key] = Metric(key, value, f"补充指标!B{idx} ← {source}；核对期 {period}", detail)


def _read_statement(wb, sheet_name, prefix, metrics):
    """Read a normalized financial statement without guessing vendor layouts."""
    if sheet_name not in wb.sheetnames:
        return
    ws = _sheet(wb, sheet_name, config.COL_STATEMENT)
    seen = set()
    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        if row[0] is None:
            continue
        item = str(row[0]).strip()
        key = f"{prefix}.{item}"
        if not item or key in seen or key in metrics:
            raise InputError(f"{sheet_name} A{idx}：项目为空、重复或覆盖已有指标「{key}」")
        seen.add(key)
        value = _number(row[1], f"{sheet_name} B{idx}")
        if value is not None:
            metrics[key] = Metric(key, value, f"{sheet_name}!B{idx}—{item}", f"{item} {value:,.2f}")


def _read_history(wb, metrics):
    if config.SHEET_HISTORY not in wb.sheetnames:
        return
    ws = _sheet(wb, config.SHEET_HISTORY, config.COL_HISTORY)
    seen = set()
    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        if row[0] is None:
            continue
        key = str(row[0]).strip()
        if not (key.startswith("历史.") or key.startswith("年度.")):
            raise InputError(f"历史指标 A{idx}：指标须以「历史.」或「年度.」开头")
        if key in seen or key in metrics:
            raise InputError(f"历史指标 A{idx}：指标重复「{key}」")
        seen.add(key)
        value = _number(row[1], f"历史指标 B{idx}")
        source, period, detail = [str(v).strip() if v is not None else "" for v in row[2:5]]
        if value is None:
            continue
        if not source or not period or not detail:
            raise InputError(f"历史指标第 {idx} 行「{key}」必须填写来源、实际所属期和口径说明")
        metrics[key] = Metric(key, value, f"历史指标!B{idx} ← {source}；实际期间 {period}", detail)


def _read_invoices(wb, metrics):
    if config.SHEET_INVOICES not in wb.sheetnames:
        return
    ws = _sheet(wb, config.SHEET_INVOICES, config.COL_INVOICES)
    seen, totals, tax_total = set(), {"销项": Decimal(0), "采购": Decimal(0)}, Decimal(0)
    counts = {"销项": 0, "采购": 0}
    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        if row[0] is None:
            continue
        number = str(row[0]).strip()
        if not number or number in seen:
            raise InputError(f"发票明细 A{idx}：发票号码为空或重复")
        seen.add(number)
        kind, status = str(row[2] or "").strip(), str(row[5] or "").strip()
        if kind not in totals or status not in {"正常", "红字", "作废"}:
            raise InputError(f"发票明细第 {idx} 行：类型仅支持销项/采购，状态仅支持正常/红字/作废")
        amount = _number(row[3], f"发票明细 D{idx}")
        tax = _number(row[4], f"发票明细 E{idx}")
        if amount is None or tax is None:
            raise InputError(f"发票明细第 {idx} 行：不含税金额和税额不能为空")
        if status == "作废":
            continue
        sign = Decimal(-1) if status == "红字" else Decimal(1)
        totals[kind] += sign * abs(amount)
        counts[kind] += 1
        if kind == "采购":
            tax_total += sign * abs(tax)
    for kind, key in (("销项", "发票.销项净额"), ("采购", "发票.采购不含税净额")):
        if counts[kind]:
            metrics[key] = Metric(key, totals[kind], f"发票明细—{kind}有效发票（含红字、剔除作废）", f"{counts[kind]} 张净额合计 {totals[kind]:,.2f}")
    if counts["采购"]:
        metrics["凭证.本期确认抵扣税额"] = Metric(
            "凭证.本期确认抵扣税额", tax_total,
            "发票明细—采购发票税额（仅作为本期确认抵扣清单）", f"采购有效发票税额净计 {tax_total:,.2f}",
        )


def _read_bank(wb, metrics):
    if config.SHEET_BANK not in wb.sheetnames:
        return
    ws = _sheet(wb, config.SHEET_BANK, config.COL_BANK)
    receipts, payments, count = Decimal(0), Decimal(0), 0
    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        if all(value is None for value in row):
            continue
        income = _number(row[2], f"银行流水 C{idx}") or Decimal(0)
        expense = _number(row[3], f"银行流水 D{idx}") or Decimal(0)
        category = str(row[4] or "").strip()
        if not category:
            raise InputError(f"银行流水 E{idx}：必须填写分类，避免把借款或内部转账误当收入")
        if income < 0 or expense < 0 or (income > 0 and expense > 0):
            raise InputError(f"银行流水第 {idx} 行：收支须为非负数，且不能同时大于零")
        receipts += income
        payments += expense
        count += 1
    if count:
        metrics["银行.收入流水合计"] = Metric("银行.收入流水合计", receipts, "银行流水—收入金额列", f"{count} 行流水；收入合计 {receipts:,.2f}")
        metrics["银行.支出流水合计"] = Metric("银行.支出流水合计", payments, "银行流水—支出金额列", f"{count} 行流水；支出合计 {payments:,.2f}")


def _from_workbook(wb) -> Dataset:
    try:
        # data_only=False 使未重算公式显式报错，而不是把空缓存误认为零。
        company = _read_company(wb)
        accounts = _read_accounts(wb)
        declarations = _read_declarations(wb)
        metrics = _build_metrics(accounts, declarations)
        _read_statement(wb, config.SHEET_INCOME, "利润表", metrics)
        _read_statement(wb, config.SHEET_BALANCE, "资产负债表", metrics)
        _read_statement(wb, config.SHEET_CASHFLOW, "现金流量表", metrics)
        _read_history(wb, metrics)
        _read_invoices(wb, metrics)
        _read_bank(wb, metrics)
        _read_supplement(wb, company, metrics)
        return Dataset(company, accounts, declarations, metrics)
    finally:
        wb.close()


def load(path: str | Path) -> Dataset:
    path = Path(path)
    if not path.is_file():
        raise InputError(f"文件不存在：{path}")
    try:
        wb = load_workbook(path, data_only=False)
    except Exception as exc:
        raise InputError(f"无法读取 Excel 工作簿：{exc}") from exc
    return _from_workbook(wb)


def load_bytes(data: bytes) -> Dataset:
    """从内存解析审计材料（Web 上传路径：原始文件不落盘，导入即解析）。"""
    if not data:
        raise InputError("上传文件为空")
    try:
        wb = load_workbook(BytesIO(data), data_only=False)
    except Exception as exc:
        raise InputError(f"无法读取 Excel 工作簿：请确认上传的是 .xlsx 格式的审计材料。原始错误：{exc}") from exc
    return _from_workbook(wb)
