from __future__ import annotations

import tempfile
import unittest
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook

from src import config, engine, loader, materials


ROOT = Path(__file__).resolve().parent.parent
SAMPLE = ROOT / "samples" / "样例企业-审计材料.xlsx"


def human_sheet(workbook, rows):
    if config.SHEET_HUMAN in workbook.sheetnames:
        del workbook[config.SHEET_HUMAN]
    ws = workbook.create_sheet(config.SHEET_HUMAN)
    ws.append(config.COL_HUMAN)
    for row in rows:
        ws.append(row)


def as_bytes(workbook) -> bytes:
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


class HumanRecordParsingTests(unittest.TestCase):
    def test_monthly_dedup_status_amount_and_rule(self):
        workbook = load_workbook(SAMPLE)
        if config.SHEET_SUPPLEMENT in workbook.sheetnames:
            sheet = workbook[config.SHEET_SUPPLEMENT]
            for row in range(sheet.max_row, 1, -1):
                if sheet.cell(row, 1).value in {"人力.个税申报人数", "人力.社保参保人数"}:
                    sheet.delete_rows(row)
        human_sheet(workbook, [
            ["个税申报", "张三", "ID-001", "2026-05", "已申报", 10000, "自然人电子税务局"],
            ["个税申报", "李四", "ID-002", "2026-05", "已申报", 9000, "自然人电子税务局"],
            ["社保参保", "张三", "ID-001", "2026-05", "正常参保", 1800, "社保平台"],
            ["社保参保", "王五", "ID-003", "2026-05", "停保", 0, "社保平台"],
            ["公积金缴存", "张三", "ID-001", "2026-05", "正常缴存", 1200, "公积金中心"],
        ])
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "human.xlsx"
            workbook.save(path)
            workbook.close()
            dataset = loader.load(path)
        self.assertEqual(dataset.get("人力.个税申报人数"), 2)
        self.assertEqual(dataset.get("人力.社保参保人数"), 1)
        self.assertEqual(dataset.get("人力.公积金缴存人数"), 1)
        self.assertEqual(dataset.get("个税.工资薪金申报收入"), 19000)
        self.assertNotIn("ID-001", dataset.metrics["人力.个税申报人数"].detail)
        rule = next(rule for rule in engine.load_rules(ROOT / "rules") if rule.id == "R-004")
        self.assertEqual(engine.evaluate(rule, dataset).status, "hit")

    def test_latest_common_month_and_no_common_month_rejected(self):
        workbook = Workbook()
        human_sheet(workbook, [
            ["个税申报", "甲", "A-001", "2026-01", "已申报", 1, "个税"],
            ["社保参保", "甲", "A-001", "2026-01", "在保", 1, "社保"],
            ["个税申报", "甲", "A-001", "2026-02", "已申报", 1, "个税"],
            ["社保参保", "甲", "A-001", "2026-03", "在保", 1, "社保"],
        ])
        records = loader._read_human_records(workbook)
        company = loader.Company("测试", "T", "测试", "2026H1")
        metrics = loader._human_metrics(company, records)
        self.assertIn("核对月 2026-01", metrics["人力.个税申报人数"].detail)
        records = [record for record in records if record.month != "2026-01"]
        with self.assertRaisesRegex(loader.InputError, "没有相同所属月"):
            loader._human_metrics(company, records)
        workbook.close()

    def test_duplicate_negative_and_out_of_period_are_rejected(self):
        for rows, message in (
            ([
                ["社保参保", "甲", "A-001", "2026-05", "在保", 1, "社保"],
                ["社保参保", "甲", "A-001", "2026-05", "在保", 1, "社保"],
            ], "重复"),
            ([["社保参保", "甲", "A-001", "2026-05", "在保", -1, "社保"]], "不能为负数"),
        ):
            workbook = Workbook()
            human_sheet(workbook, rows)
            with self.assertRaisesRegex(loader.InputError, message):
                loader._read_human_records(workbook)
            workbook.close()
        record = loader._HumanRecord("社保", "a" * 64, "2025-12", True, None, "测试")
        with self.assertRaisesRegex(loader.InputError, "不在核对期间"):
            loader._human_metrics(loader.Company("测试", "T", "测试", "2026H1"), [record])

    def test_multifile_preview_keeps_only_hashed_identity_and_merges(self):
        source = Workbook()
        human_sheet(source, [
            ["个税申报", "张三", "SECRET-ID-001", "2026-05", "已申报", 10000, "个税系统"],
            ["社保参保", "张三", "SECRET-ID-001", "2026-05", "在保", 1800, "社保平台"],
        ])
        human_bytes = as_bytes(source)
        sample_bytes = SAMPLE.read_bytes()
        docs = materials.preview([("sample.xlsx", sample_bytes), ("human.xlsx", human_bytes)], set())
        self.assertFalse(any(doc["error"] for doc in docs), docs)
        serialized = str(docs[1]["human_records"])
        self.assertNotIn("SECRET-ID-001", serialized)
        self.assertNotIn("张三", serialized)
        selections = {doc["id"]: {"company": {}, "rows": doc["rows"]} for doc in docs}
        dataset = materials.build_dataset(docs, selections, {}, set())
        self.assertEqual(dataset.get("人力.个税申报人数"), 1)
        self.assertEqual(dataset.get("人力.社保参保人数"), 1)


if __name__ == "__main__":
    unittest.main()
