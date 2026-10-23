"""Local retrieval contracts and real ONNX/FAISS acceptance, separate from legal validation."""
from copy import deepcopy
import hashlib
from io import BytesIO
import json
from pathlib import Path
import sqlite3
import tempfile
import unittest
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from src.kg_retrieval import applicability, bundle, evaluate, fixtures, keyword_search, metrics, relation_search, semantic_search
from src.kg_retrieval.common import clause_record, implementation_fingerprint
from webapp import app as app_module, knowledge_evidence as service
from webapp.storage import Store
from tests.enterprise_support import material_key

ROOT = Path(__file__).resolve().parents[1]
GOLD = ROOT / "tests/fixtures/kg_retrieval/gold_cases.json"


class RetrievalTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.db = sqlite3.connect(":memory:")
        self.db.row_factory = sqlite3.Row
        self.db.execute("PRAGMA foreign_keys=ON")
        self.addCleanup(self.db.close)
        fixtures.seed(self.db, Path(self.tmp.name) / "originals")
        self.cases = {c["id"]: c for c in json.loads(GOLD.read_text(encoding="utf-8"))["cases"]}

    def package(self, case_id="K01", **kwargs):
        case = self.cases[case_id]
        task, facts, catalog = evaluate.case_input(self.db, case)
        return bundle.build(self.db, task, case["rule_id"], facts, topic=case["topic"], source_catalog=catalog,
                            model_dir=Path(self.tmp.name) / "no-model", token_budget=case.get("token_budget", 6000), **kwargs)

    def test_twenty_synthetic_policies_are_explicitly_not_human_approved(self):
        metadata = json.loads(self.db.execute("SELECT metadata FROM kg_fixture_identity").fetchone()[0])
        self.assertEqual(metadata["policy_count"], 20)
        self.assertFalse(metadata["legal_review"])
        self.assertTrue(all(r[0].startswith("local-fixture://") for r in self.db.execute("SELECT source_url FROM Policy")))

    def context_package(self):
        package = self.package()
        package['policy_evidence'] = [deepcopy(package['policy_evidence'][0])]
        package['policy_evidence'][0]['text'] = '完整合成条款，不得遗漏。' * 900
        package['policy_leads'] = []
        package['gaps'] = []
        package['risk_assessment']['missing_evidence'] = []
        package['batching'] = bundle.split_bundle(package, 128000)
        return package

    @staticmethod
    def context_answer(_system, part, _failure):
        return {'assessments': [{'type': part['response_constraints']['required_type'],
            'text': '本批证据须联合核查。', 'rule_id': part['check']['rule_id'],
            'clause_ids': [p['clause_id'] for p in part['policy_evidence']],
            'source_ids': [], 'amounts': [], 'dates': [], 'required_materials': []}]}, 'synthetic-provider'

    def test_final_outbound_does_not_restore_full_missing_evidence_into_each_batch(self):
        from src.kg_retrieval import batching, outbound
        package = self.package()
        package['gaps'] = [{'code': 'pending', 'message': '待核实事项' + str(i)} for i in range(300)]
        package['risk_assessment']['missing_evidence'] = [g['message'] for g in package['gaps']]
        package['batching'] = bundle.split_bundle(package, 6000)
        self.assertFalse(package['batching']['unsplittable'])
        delivery = outbound.prepare(package, authorized=True, destination='https://provider.example')
        self.assertGreater(len(delivery['parts']), 1)
        self.assertTrue(all(batching.encoded_size(p) <= 6000 for p in delivery['parts']))
        gaps = [g for p in delivery['parts'] for g in p['gaps']]
        self.assertEqual(gaps, package['gaps'])
        self.assertTrue(all('missing_evidence' not in p['risk_assessment'] for p in delivery['parts']))
        self.assertEqual(len(package['risk_assessment']['missing_evidence']), 300)

    def test_explicit_capacity_rejection_refines_losslessly_and_resumes_completed_children(self):
        from src.kg_retrieval import batching
        from webapp.knowledge import ContextLengthExceeded
        package = self.context_package()
        seen, successful = [], []
        failures = [True]

        def provider(system, part, failure):
            seen.append(deepcopy(part))
            if batching.encoded_size(part) > 14000:
                raise ContextLengthExceeded()
            if len(successful) == 1 and failures[0]:
                failures[0] = False
                raise service.HTTPException(502, '模拟网络中断')
            successful.append(deepcopy(part))
            return self.context_answer(system, part, failure)

        gate = {'passed': True, 'model_fingerprints': [None]}
        with patch.object(service, 'acceptance_gate', return_value=gate), \
                patch('webapp.knowledge._chat_json', side_effect=provider):
            partial = service.analyze(package)
        self.assertFalse(partial['validated'])
        self.assertEqual(len(partial['batch_results']), 1)
        self.assertTrue(partial['context_splits'])
        rejected = {json.dumps(p, sort_keys=True) for p in seen if batching.encoded_size(p) > 14000}
        seen.clear()
        with patch.object(service, 'acceptance_gate', return_value=gate), \
                patch('webapp.knowledge._chat_json', side_effect=provider):
            result = service.analyze(package, prior=partial)
        self.assertTrue(result['validated'])
        self.assertEqual(result['reused_batches'], 1)
        self.assertFalse(rejected & {json.dumps(p, sort_keys=True) for p in seen})
        pieces = [p for payload in successful for p in payload['policy_evidence']]
        pieces.sort(key=lambda p: p['source_char_start'])
        self.assertEqual(''.join(p['text'] for p in pieces), package['policy_evidence'][0]['text'])
        self.assertEqual(pieces[0]['source_char_start'], 0)
        self.assertTrue(all(a['source_char_end'] == b['source_char_start'] for a, b in zip(pieces, pieces[1:])))
        expected_hash = hashlib.sha256(package['policy_evidence'][0]['text'].encode()).hexdigest()
        self.assertTrue(all(p['original_text_sha256'] == expected_hash for p in pieces))
        self.assertEqual(result['status'], 'needs_review')
        self.assertTrue(all(p['response_constraints']['required_type'] == 'insufficient_evidence' for p in successful))

    def test_preflight_uses_full_wire_payload_and_recovers_before_sending(self):
        from src.ai_transport import chat_payload
        from src.settings import AISettings
        from webapp.knowledge import SELECTED_SETTINGS
        package = self.context_package()
        settings = AISettings(model='synthetic-model', enabled=True, api_key='synthetic-key')
        package['batching']['model_context'] = {'model': settings.effective_model,
            'planning_context_window': 18000, 'output_reserve': 1024, 'safety_reserve': 1000}
        observed = []

        def provider(system, payload, failure):
            wire = chat_payload(SELECTED_SETTINGS.get(), [
                {'role': 'system', 'content': system},
                {'role': 'user', 'content': json.dumps(payload, ensure_ascii=False)},
            ], temperature=0.2)
            self.assertLessEqual(len(wire) + 1024 + 1000, 18000)
            observed.append(payload)
            return self.context_answer(system, payload, failure)

        with patch.object(service, 'acceptance_gate', return_value={'passed': True, 'model_fingerprints': [None]}), \
                patch('webapp.ai_keys.apply', return_value=settings), \
                patch('webapp.knowledge._chat_json', side_effect=provider):
            result = service.analyze(package)
        self.assertTrue(result['validated'])
        self.assertGreater(len(observed), 1)
        self.assertTrue(result['context_splits'])

    def test_capacity_recovery_is_bounded_and_reauthorizes_each_child(self):
        from webapp.knowledge import ContextLengthExceeded
        package = self.context_package()
        gate = {'passed': True, 'model_fingerprints': [None]}
        with patch.object(service, 'acceptance_gate', return_value=gate), \
                patch('webapp.knowledge._chat_json', side_effect=ContextLengthExceeded()) as provider:
            result = service.analyze(package)
        self.assertFalse(result['validated'])
        self.assertLessEqual(provider.call_count, 5)
        self.assertTrue(any(b['status'] == 'failed' for b in result['batches']))
        self.assertEqual(result['assessments'], [])
        authorizations = []

        def authorize():
            authorizations.append(True)
            if len(authorizations) > 1:
                raise service.HTTPException(409, '资料已变化')

        token = service._AUTHORIZE_BATCH.set(authorize)
        try:
            with patch.object(service, 'acceptance_gate', return_value=gate), \
                    patch('webapp.knowledge._chat_json', side_effect=ContextLengthExceeded()) as provider:
                with self.assertRaises(service.HTTPException) as caught:
                    service.analyze(package)
                self.assertEqual(caught.exception.status_code, 409)
                self.assertEqual(provider.call_count, 1)
        finally:
            service._AUTHORIZE_BATCH.reset(token)

    def test_execution_deadline_stops_before_another_provider_request(self):
        package = self.context_package()
        with patch.object(service, 'acceptance_gate', return_value={'passed': True, 'model_fingerprints': [None]}), \
                patch.object(service, 'monotonic', side_effect=[0, 10000]), \
                patch('webapp.knowledge._chat_json') as provider:
            result = service.analyze(package)
        self.assertFalse(result['validated'])
        self.assertIn('执行时限', result['reason'])
        provider.assert_not_called()

    def test_saved_split_plans_do_not_exhaust_new_recovery_allowance(self):
        from src.kg_retrieval import batching
        from webapp.knowledge import ContextLengthExceeded
        package = self.context_package()
        original = package['batching']['parts'][0]
        package['batching'].update(part_count=9, parts=[
            {**deepcopy(original), 'part_number': index, 'part_count': 9} for index in range(1, 10)])

        def provider(system, part, failure):
            if batching.encoded_size(part) > 14000:
                raise ContextLengthExceeded()
            return self.context_answer(system, part, failure)

        prior = None
        with patch.object(service, 'acceptance_gate', return_value={'passed': True, 'model_fingerprints': [None]}), \
                patch('webapp.knowledge._chat_json', side_effect=provider):
            for _ in range(9):
                prior = service.analyze(package, prior=prior)
                if prior['validated']:
                    break
        self.assertTrue(prior['validated'])
        self.assertGreater(prior['reused_batches'], 0)

    def test_acceptance_fingerprint_ignores_checkout_newlines_but_detects_source_changes(self):
        original = Path.read_bytes
        expected = implementation_fingerprint()
        with patch.object(Path, 'read_bytes', lambda path: original(path).replace(b'\r\n', b'\n').replace(b'\n', b'\r\n')):
            self.assertEqual(implementation_fingerprint(), expected)
        with patch.object(Path, 'read_bytes', lambda path: original(path) + b'\n# changed source\n'):
            self.assertNotEqual(implementation_fingerprint(), expected)

    def test_dependency_order_and_legacy_review_gap_order_do_not_invalidate_content(self):
        self.db.execute("UPDATE Relationship SET review_status='candidate' WHERE relation_type='AMENDS'")
        first = relation_search.dependency_closure(self.db, ['CL05A', 'CL20'])
        self.assertEqual(first, relation_search.dependency_closure(self.db, ['CL20', 'CL05A']))
        package = self.package('K05')
        previous = deepcopy(package)
        previous['gaps'].reverse()
        previous['relationships'].reverse()
        previous['content_hash'] = 'legacy-order-dependent-hash'
        self.assertEqual(bundle.content_hash(package), bundle.content_hash(previous))
        self.assertNotEqual(package['gaps'], previous['gaps'])
        previous['gaps'][0]['message'] = 'changed review requirement'
        self.assertNotEqual(bundle.content_hash(package), bundle.content_hash(previous))
        previous = deepcopy(package)
        previous['facts'][0]['value'] = '999999'
        self.assertNotEqual(bundle.content_hash(package), bundle.content_hash(previous))

    def test_trigram_and_two_character_fallback_and_safe_query(self):
        results = keyword_search.search(self.db, "红字发票")
        self.assertIn("CL02", [h["clause_id"] for h in results["hits"]])
        short = keyword_search.search(self.db, "尾差")
        self.assertTrue(any(h["clause_id"] == "CL15" and h["reason"]["mode"] == "short_term_like" for h in short["hits"]))
        keyword_search.search(self.db, '\" OR * --', keywords=['%_', '"'])
        self.assertEqual(self.db.execute("SELECT count(*) FROM Policy").fetchone()[0], 20)

    def test_fts_insert_update_delete_are_synchronized(self):
        self.db.execute("INSERT INTO Clause(clause_id,policy_id,clause_label,sort_index,content) VALUES ('extra','P01','额外条',99,'独特测试检索甲')")
        self.assertTrue(any(h["clause_id"] == "extra" for h in keyword_search.search(self.db, "独特测试检索甲")["hits"]))
        self.db.execute("UPDATE Clause SET content='替换后的检索乙' WHERE clause_id='extra'")
        self.assertFalse(any(h["clause_id"] == "extra" for h in keyword_search.search(self.db, "独特测试检索甲")["hits"]))
        self.db.execute("DELETE FROM Clause WHERE clause_id='extra'")
        self.assertFalse(any(h["clause_id"] == "extra" for h in keyword_search.search(self.db, "替换后的检索乙")["hits"]))

    def test_required_exception_bypasses_route_and_review_limits(self):
        result = self.package("K04", candidate_limit=1, review_limit=1)
        self.assertTrue({"CL04", "CL04E"} <= {r["clause_id"] for r in result["policy_evidence"]})
        self.assertEqual(result["exceptions"][0]["exception_id"], "X04")

    def test_candidate_relationship_is_a_lead_not_verified_support(self):
        self.db.execute("INSERT INTO Relationship(relationship_id,from_type,from_id,relation_type,to_type,to_id) VALUES ('pending','Clause','CL03','PROVIDES_CONTEXT','TaxRule','R01')")
        result = self.package()
        self.assertIn("CL03", {r["clause_id"] for r in result["policy_leads"]})
        self.assertNotIn("CL03", {r["clause_id"] for r in result["policy_evidence"]})
        self.assertEqual(result["bundle_status"], "needs_review")

    def test_change_cycle_stops_and_keeps_evidence(self):
        self.db.execute("INSERT INTO Relationship(relationship_id,from_type,from_id,relation_type,to_type,to_id,evidence_clause_id,effective_from,review_status) VALUES ('back','Clause','CL05A','INTERPRETS','Clause','CL20','CL20','2026-07-01','approved')")
        closure = relation_search.dependency_closure(self.db, ["CL05A"])
        self.assertEqual(set(closure["clause_ids"]), {"CL05A", "CL20"})
        self.assertEqual(len(closure["relationships"]), 2)

    def test_more_than_eight_categorized_batches_are_analyzed_with_part_local_citations(self):
        package=self.package()
        original=deepcopy(package['batching']['parts'][0])
        package['batching'].update(unsplittable=False,part_count=9,requires_joint_review=True,
            parts=[{**deepcopy(original),'part_number':i,'part_count':9,'batch_category':'policy_basis',
                    'batch_category_name':'政策依据','batch_group':f'policy-{i}'} for i in range(1,10)])
        def answer(system, part, failure):
            self.assertEqual(part['response_constraints']['required_type'],'insufficient_evidence')
            return {'assessments':[{'type':'insufficient_evidence','text':'本批依据仍须与其他批次联合核查。',
                'rule_id':package['check']['rule_id'],'clause_ids':[part['policy_evidence'][0]['clause_id']],
                'source_ids':[],'amounts':[],'dates':[],'required_materials':[]}]},'test-provider'
        with patch.object(service,'acceptance_gate',return_value={'passed':True,'model_fingerprints':[None]}), \
             patch('webapp.knowledge._chat_json',side_effect=answer) as ai:
            result=service.analyze(package)
        self.assertEqual(ai.call_count,9)
        self.assertEqual(len(result['batches']),9)
        self.assertTrue(all(b['status']=='completed' for b in result['batches']))
        self.assertTrue(result['validated'])
        self.assertEqual(result['status'],'needs_review')
        successful_answer=answer
        attempt=[0]
        def intermittent(system, part, failure):
            attempt[0]+=1
            if attempt[0]==3:
                raise service.HTTPException(502,'模拟上游连接失败')
            return successful_answer(system,part,failure)
        with patch.object(service,'acceptance_gate',return_value={'passed':True,'model_fingerprints':[None]}), \
             patch('webapp.knowledge._chat_json',side_effect=intermittent):
            partial=service.analyze(package)
        self.assertFalse(partial['validated'])
        self.assertEqual(len(partial['batch_results']),2)
        self.assertEqual(partial['batches'][2]['status'],'failed')
        with patch.object(service,'acceptance_gate',return_value={'passed':True,'model_fingerprints':[None]}), \
             patch('webapp.knowledge._chat_json',side_effect=successful_answer) as resumed:
            completed=service.analyze(package,prior=partial)
        self.assertTrue(completed['validated'])
        self.assertEqual(completed['reused_batches'],2)
        self.assertEqual(resumed.call_count,7)
        self.assertTrue(all(b['status']=='completed' for b in completed['batches']))
        changed_profile=deepcopy(partial)
        changed_profile['batch_profile']='previous-model-or-output-configuration'
        with patch.object(service,'acceptance_gate',return_value={'passed':True,'model_fingerprints':[None]}), \
             patch('webapp.knowledge._chat_json',side_effect=successful_answer) as changed:
            service.analyze(package,prior=changed_profile)
        self.assertEqual(changed.call_count,9)
        permission_checks=[0]
        def authorize():
            permission_checks[0]+=1
            if permission_checks[0]>1:
                raise service.HTTPException(409,'资料或权限已变化')
        token=service._AUTHORIZE_BATCH.set(authorize)
        try:
            with patch.object(service,'acceptance_gate',return_value={'passed':True,'model_fingerprints':[None]}), \
                 patch('webapp.knowledge._chat_json',side_effect=successful_answer) as stopped:
                with self.assertRaises(service.HTTPException):
                    service.analyze(package)
            self.assertEqual(stopped.call_count,1)
        finally:
            service._AUTHORIZE_BATCH.reset(token)

    def test_unreviewed_or_unsourced_changes_remain_explained_review_gaps(self):
        for review, source in (("candidate", "CL01"), ("approved", None), ("candidate", None)):
            with self.subTest(review=review, source=source):
                self.db.execute("DELETE FROM Relationship WHERE relationship_id='pending-change'")
                self.db.execute("INSERT INTO Relationship(relationship_id,from_type,from_id,relation_type,to_type,to_id,evidence_clause_id,review_status) "
                                "VALUES ('pending-change','Clause','CL01','SUPPLEMENTS','Clause','CL03',?,?)", (source, review))
                result = self.package()
                codes = {g["code"] for g in result["gaps"]}
                if review == "candidate":
                    self.assertIn("change_pending", codes)
                if source is None:
                    self.assertIn("change_source_missing", codes)
                self.assertTrue(all(g.get("message") for g in result["gaps"]))
                self.assertEqual(result["check"]["evaluation"], "UNKNOWN")
                self.assertNotEqual(result["bundle_status"], "ready")
                self.assertEqual(result["risk_assessment"]["missing_evidence"], [g["message"] for g in result["gaps"]])

    def test_historical_repeal_region_unknown_and_null_end(self):
        old, current = self.package("K06"), self.package("K07")
        self.assertIn("CL06", {p["clause_id"] for p in old["policy_evidence"]})
        self.assertNotIn("CL06", {p["clause_id"] for p in current["policy_evidence"]})
        self.assertEqual(applicability.region_matches("CN-33", "CN-33-3301"), True)
        self.assertIsNone(applicability.region_matches("CN-33", "CN"))
        self.db.execute("UPDATE Policy SET effective_to=NULL WHERE policy_id='P01'")
        result = self.package()
        self.assertNotEqual(result["bundle_status"], "ready")
        self.assertIn("CL01", {p["clause_id"] for p in result["policy_leads"]})

    def test_national_default_uses_national_evidence_and_leaves_local_policy_pending(self):
        def national_package(rule_id):
            task, facts, catalog = evaluate.case_input(
                self.db, {"id": "national-default", "rule_id": rule_id, "region_code": "CN"})
            return bundle.build(self.db, task, rule_id, facts, source_catalog=catalog,
                                model_dir=Path(self.tmp.name) / "no-model", offline=True)

        national, local = national_package("R01"), national_package("R07")
        self.assertIn("CL01", {p["clause_id"] for p in national["policy_evidence"]})
        self.assertNotIn("CL07", {p["clause_id"] for p in local["policy_evidence"]})
        self.assertIn("CL07", {p["clause_id"] for p in local["policy_leads"]})
        self.assertEqual(local["check"]["evaluation"], "UNKNOWN")

    def test_crossing_change_date_requires_segments_and_old_rule_review(self):
        task, _, _ = evaluate.case_input(self.db, self.cases["K05"])
        task.update(period_start="2026-06-15", period_end="2026-07-15")
        clause, policy = clause_record(self.db, "CL05A")
        closure = relation_search.dependency_closure(self.db, ["CL05A"])
        decision = applicability.check(clause, policy, task, closure["relationships"])
        self.assertEqual(decision["status"], "needs_review")
        self.assertTrue(decision["segments"])
        self.db.execute("UPDATE Condition SET source_clause_id='CL05A' WHERE condition_id='C05'")
        result = self.package("K05")
        self.assertEqual(result["check"]["evaluation"], "UNKNOWN")
        self.assertIn("rule_source_changed", {gap["code"] for gap in result["gaps"]})

    def test_source_tampering_and_foreign_fact_ids_stay_unknown(self):
        policy = self.db.execute("SELECT raw_file_path FROM Policy WHERE policy_id='P01'").fetchone()
        Path(policy[0]).write_text("changed", encoding="utf-8")
        self.assertNotEqual(self.package()["bundle_status"], "ready")
        task, facts, _ = evaluate.case_input(self.db, self.cases["K01"])
        result = bundle.build(self.db, task, "R01", facts, source_catalog=[], model_dir=Path(self.tmp.name) / "missing")
        self.assertEqual(result["facts"][0]["state"], "unknown")
        self.assertEqual(result["facts"][0]["source_ids"], [])

    def test_deduplication_and_complete_budget_parts(self):
        result = self.package("K12")
        self.assertTrue(result["batching"]["unsplittable"])
        reconstructed = {p["clause_id"]: p["text"] for part in result["batching"]["parts"] for p in part["policy_evidence"]}
        for item in result["policy_evidence"]:
            self.assertEqual(reconstructed[item["clause_id"]], item["text"])
        self.assertTrue(all(part["exceptions"] == result["exceptions"] for part in result["batching"]["parts"]))
        again = self.package()
        self.assertEqual(len(again["policy_evidence"]), len({p["clause_id"] for p in again["policy_evidence"]}))

    def test_missing_model_cannot_pass_semantic_acceptance(self):
        report, _ = evaluate.evaluate(self.db, list(self.cases.values()), model_dir=Path(self.tmp.name) / "missing")
        self.assertFalse(report["gate_passed"])
        self.assertFalse(report["metrics"]["semantic_ready"])
        self.assertEqual(report["metrics"]["external_api_calls"], 0)

    def require_model(self):
        model, reason = semantic_search.available()
        if model is None:
            self.skipTest(reason + "; full acceptance CLI remains not_ready")
        return model

    def test_real_local_model_recalls_paraphrase_offline(self):
        self.require_model()
        case = self.cases["K03"]
        with patch("socket.create_connection", side_effect=AssertionError("offline network call")), patch("urllib.request.urlopen", side_effect=AssertionError("offline download")):
            hits = semantic_search.search(self.db, case["topic"])
        self.assertEqual(hits["status"], "ready")
        self.assertEqual(hits["hits"][0]["clause_id"], "CL02")
        self.assertNotIn("CL02", {h["clause_id"] for h in keyword_search.search(self.db, case["topic"])["hits"]})

    def test_incremental_vectors_versions_and_faiss_export(self):
        model = self.require_model()
        semantic_search.synchronize(self.db, model)
        before = {r["vector_id"]: bytes(r["vector"]) for r in self.db.execute("SELECT * FROM kg_clause_vectors WHERE clause_id!='CL15'")}
        self.db.execute("UPDATE Clause SET content=content||'精确核对尾数。',version='2' WHERE clause_id='CL15'")
        updated = semantic_search.synchronize(self.db, model)
        self.assertEqual(updated["updated_clause_ids"], ["CL15"])
        after = {r["vector_id"]: bytes(r["vector"]) for r in self.db.execute("SELECT * FROM kg_clause_vectors WHERE clause_id!='CL15'")}
        self.assertEqual(before, after)
        self.assertEqual(semantic_search.synchronize(self.db, model)["updated_clause_ids"], [])
        exported = semantic_search.export_index(self.db, Path(self.tmp.name) / "vectors")
        metadata = json.loads(Path(exported["index_path"]).with_suffix(".json").read_text(encoding="utf-8"))
        self.assertEqual(len(metadata["mapping"]), exported["vector_count"])
        self.assertEqual(metadata["index_sha256"], hashlib.sha256(Path(exported["index_path"]).read_bytes()).hexdigest())
        self.assertGreater(exported["vector_count"], self.db.execute("SELECT count(*) FROM Clause").fetchone()[0])

    def test_full_gold_set_uses_real_model_and_zero_external_calls(self):
        self.require_model()
        with patch("src.ai_transport.chat_content", side_effect=AssertionError("external LLM called")):
            report, _ = evaluate.evaluate(self.db, list(self.cases.values()))
        self.assertTrue(report["gate_passed"], report["cases"])
        self.assertEqual(report["metrics"]["external_api_calls"], 0)
        self.assertFalse(report["formal_policy_acceptance"])

    def test_three_parser_materials_with_different_amounts_are_not_a_tax_conclusion(self):
        from tests.test_core_inputs import book, mapper, TAX
        from src import core_materials
        con, field_mapper = mapper()
        self.addCleanup(con.close)
        task = {"id": "task", "tenant_id": "org", "company_id": "company", "tax_no": TAX,
                "start_date": "2026-01-01", "end_date": "2026-01-31", "applicable_forms": []}
        documents = []
        for index, (kind, kwargs, cell, value) in enumerate([
                ("account_balance", {}, "D6", "980000.00"), ("vat_declaration", {}, "C6", "1000000.00"),
                ("invoice_detail", {}, "E6", "900000.00"), ("invoice_detail", {"direction": "进项"}, "E6", "400000.00")]):
            workbook = load_workbook(BytesIO(book(kind, **kwargs)))
            workbook.active[cell] = value
            stream = BytesIO()
            workbook.save(stream)
            workbook.close()
            documents.append(core_materials.parse(stream.getvalue(), "synthetic.xlsx", kind, str(index), field_mapper))
        processing = core_materials.prepare(task, documents)
        self.assertEqual(processing["status"], "ready_for_analysis")
        definitions = relation_search.rule_context(self.db, "R13")["metrics"]
        definitions["M13"]["source_filter_json"] = json.dumps({"form_type": "main", "form_name": "增值税纳税申报表主表", "item_key": "1", "column_scope": "本期数"})
        facts = metrics.collect(definitions, processing, task)
        self.assertEqual(facts["M13"]["value"], "1000000.00")
        self.assertEqual(facts["MI13"]["value"], "900000.00")
        balance = {**definitions["M13"], "source_schema": "AccountBalanceRow", "source_field": "current_debit",
                   "period_basis": "accounting_period", "source_filter_json": '{"account_code":"1001"}'}
        self.assertEqual(metrics.resolve(balance, processing, task)["value"], "980000.00")
        context, _, _ = evaluate.case_input(self.db, self.cases["K13"])
        context.update(period_start=task["start_date"], period_end=task["end_date"])
        package = bundle.build(self.db, context, "R13", facts, source_catalog=[r["id"] for r in processing["records"]], model_dir=Path(self.tmp.name) / "missing")
        self.assertEqual(package["check"]["evaluation"], "UNKNOWN")


class AnswerValidationTests(unittest.TestCase):
    setUp = RetrievalTests.setUp
    package = RetrievalTests.package
    def test_strict_amount_date_source_and_condition_validation(self):
        package = self.package()
        good = {"assessments": [{"type": "risk_clue", "text": "请结合原件进一步核对，此处仅为测试线索。", "rule_id": "R01",
                                "clause_ids": ["CL01"], "source_ids": ["SYNTHETIC-SOURCE-M01"],
                                "amounts": [{"metric_id": "M01", "value": "200000.01"}], "dates": ["2026-02-01"], "required_materials": ["原始记录"]}]}
        self.assertEqual(service.validate_answer(deepcopy(good), package), good)
        for conclusion in ("condition_not_met", "excluded"):
            bad = deepcopy(good)
            bad["assessments"][0]["type"] = conclusion
            with self.assertRaises(ValueError):
                service.validate_answer(bad, package)
        for key, value in [("clause_ids", ["invented"]), ("source_ids", ["foreign"]),
                           ("amounts", [{"metric_id": "M01", "value": "999999"}]), ("dates", ["2020-01-01"]),
                           ("text", "确定违法"), ("text", "漏报了999999元"), ("rule_id", "foreign-rule")]:
            bad = deepcopy(good)
            bad["assessments"][0][key] = value
            with self.assertRaises(ValueError):
                service.validate_answer(bad, package)
        insufficient = deepcopy(package)
        insufficient["bundle_status"] = "needs_review"
        with self.assertRaises(ValueError):
            service.validate_answer(good, insufficient)
        negative = deepcopy(good)
        negative["assessments"][0].update(type="condition_not_met", text="条件未满足，仅形成待核查线索。")
        false_package = deepcopy(package)
        false_package["check"].update(evaluation="FALSE", disposition="condition_not_met")
        with self.assertRaises(ValueError):
            service.validate_answer(negative, false_package)
        negative["assessments"][0]["text"] = "条件未满足，本条件不形成风险线索。"
        service.validate_answer(negative, false_package)


@material_key
class RetrievalHTTPTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.path = Path(self.tmp.name) / "test.db"
        db = sqlite3.connect(self.path)
        db.row_factory = sqlite3.Row
        fixtures.seed(db, Path(self.tmp.name) / "originals")
        db.execute("UPDATE Metric SET source_filter_json=? WHERE metric_id='M01'", (json.dumps({"form_type": "main", "form_name": "增值税纳税申报表主表", "item_key": "1", "column_scope": "本期数"}),))
        db.commit()
        db.close()
        self.store = Store(self.path)
        patched = patch.object(app_module, "store", self.store)
        patched.start()
        self.addCleanup(patched.stop)
        self.store.create_user("kg-admin", "Knowledge-test-2026!", "管理员", "org_admin", "org-a")
        self.store.create_user("kg-other", "Knowledge-test-2026!", "其他机构", "org_admin", "org-b")
        self.client = TestClient(app_module.app)
        self.client.__enter__()
        self.addCleanup(self.client.__exit__, None, None, None)
        self.login("kg-admin")

    def login(self, name):
        self.client.cookies.clear()
        token = self.store.authenticate(name, "Knowledge-test-2026!")[1]
        self.client.cookies.set(app_module.COOKIE_NAME, token)

    def task(self):
        from tests.test_core_inputs import TAX, book
        created = self.client.post("/api/enterprise/input-tasks", json={"company_name": "纯合成测试企业", "tax_no": TAX,
                                  "start_date": "2026-01-01", "end_date": "2026-01-31", "applicable_forms": []})
        self.assertEqual(created.status_code, 200, created.text)
        task = created.json()
        uploaded = self.client.post("/api/enterprise/input-tasks/" + task["id"] + "/files",
                                    data={"expected_revision": task["revision"], "material_type": "vat_declaration"},
                                    files={"files": ("合成.xlsx", book("vat_declaration"))})
        self.assertEqual(uploaded.status_code, 200, uploaded.text)
        return uploaded.json()

    def test_automatic_budget_uses_selected_model_capacity_and_clamps_manual_requests(self):
        from src.settings import AISettings
        from webapp import ai_keys
        settings=AISettings(enabled=False,model='capacity-test')
        with patch.object(ai_keys,'capability',return_value={'context_window':1048576,'max_output_tokens':393216}):
            profile=ai_keys.context_budget(settings,self.store)
        self.assertEqual(profile['context_window'],1048576)
        self.assertEqual(profile['input_budget'],128000)
        self.assertEqual(profile['output_reserve'],8192)
        self.assertTrue(profile['cache_tokens_included'])
        self.assertLess(profile['input_budget'],1048576)
        with patch.object(ai_keys,'capability',return_value={'context_window':32000,'max_output_tokens':4096}):
            small=ai_keys.context_budget(settings,self.store)
        self.assertEqual(small['output_reserve'],4096)
        self.assertLessEqual(small['input_budget'],32000*3//5)
        self.assertLessEqual(sum(small[k] for k in ('input_budget','output_reserve','instruction_reserve','safety_reserve')),32000)
        task=self.task()
        body={'rule_id':'R01','task_id':task['id'],'expected_revision':task['revision'],'region_code':'CN-33'}
        with patch.object(ai_keys,'context_budget',return_value=profile):
            first=self.client.post('/api/knowledge/retrieve',json=body)
            self.assertEqual(first.status_code,200,first.text)
            plan=first.json()['bundle']['batching']
            self.assertEqual(plan['soft_token_budget'],profile['input_budget'])
            self.assertFalse(plan['unsplittable'])
            self.assertEqual(plan['model_context']['capacity_source'],'provider_models')
            body.update(expected_revision=first.json()['revision'],token_budget=16000000)
            clamped=self.client.post('/api/knowledge/retrieve',json=body)
            self.assertEqual(clamped.status_code,200,clamped.text)
            self.assertEqual(clamped.json()['bundle']['batching']['soft_token_budget'],profile['input_budget'])

    def test_task_facts_history_cache_and_tenant_isolation(self):
        task = self.task()
        body = {"rule_id": "R01", "task_id": task["id"], "expected_revision": task["revision"], "region_code": "CN-33"}
        first = self.client.post("/api/knowledge/retrieve", json=body)
        self.assertEqual(first.status_code, 200, first.text)
        response = first.json()
        self.assertEqual(response["bundle"]["facts"][0]["value"], "100.0123")
        self.assertEqual(response["bundle"]["check"]["evaluation"], "FALSE")
        self.assertTrue(response["saved"])
        body["expected_revision"] = response["revision"]
        second = self.client.post("/api/knowledge/retrieve", json=body)
        self.assertEqual(second.status_code, 200, second.text)
        self.assertTrue(second.json()["cached"])
        self.assertEqual(response["revision"], second.json()["revision"])
        history = self.client.get("/api/enterprise/input-tasks/" + task["id"]).json()
        self.assertIn("R01", history["tax_checks"])
        self.assertNotIn(b"100.0123", self.path.read_bytes())
        self.login("kg-other")
        with patch("webapp.knowledge_evidence.analyze") as ai:
            body["use_ai"] = True
            denied = self.client.post("/api/knowledge/retrieve", json=body)
            self.assertEqual(denied.status_code, 404)
            ai.assert_not_called()

    def test_stale_revision_stops_paid_model_and_unrelated_update_keeps_cache(self):
        from tests.test_core_inputs import book
        task = self.task()
        body = {"rule_id": "R01", "task_id": task["id"], "expected_revision": task["revision"], "region_code": "CN-33"}
        first = self.client.post("/api/knowledge/retrieve", json=body).json()
        with patch("webapp.knowledge_evidence.analyze") as ai:
            stale = self.client.post("/api/knowledge/retrieve", json={**body, "use_ai": True})
            self.assertEqual(stale.status_code, 409)
            ai.assert_not_called()
        updated = self.client.post("/api/enterprise/input-tasks/" + task["id"] + "/files",
                                   data={"expected_revision": first["revision"], "material_type": "account_balance"},
                                   files={"files": ("合成余额.xlsx", book("account_balance"))})
        self.assertEqual(updated.status_code, 200, updated.text)
        self.assertFalse(updated.json()["tax_checks"]["R01"]["stale"])
        again = self.client.post("/api/knowledge/retrieve", json={**body, "expected_revision": updated.json()["revision"]})
        self.assertTrue(again.json()["cached"], again.text)

    def test_gate_blocks_ai_until_offline_acceptance(self):
        task = self.task()
        with patch("webapp.knowledge_evidence.acceptance_gate", return_value={"passed": False}), patch("webapp.knowledge._chat_json") as ai:
            response = self.client.post("/api/knowledge/retrieve", json={"rule_id": "R01", "task_id": task["id"],
                                        "expected_revision": task["revision"], "region_code": "CN-33", "use_ai": True})
            self.assertEqual(response.status_code, 409)
            ai.assert_not_called()

    def test_valid_ai_is_saved_and_repeated_request_does_not_call_again(self):
        task = self.task()
        model, _ = semantic_search.available()
        gate = {"passed": True, "model_fingerprints": [model.fingerprint if model else None]}
        def answer(system, payload, failure):
            return {"assessments": [{"type": "condition_not_met", "text": "合成检查条件未满足，请复核原始资料。",
                    "rule_id": "R01", "clause_ids": ["CL01"], "source_ids": payload["facts"][0]["source_ids"],
                    "amounts": [{"metric_id": "M01", "value": "100.0123"}], "dates": [], "required_materials": []}]}, "unit-test-provider"
        body = {"rule_id": "R01", "task_id": task["id"], "expected_revision": task["revision"], "region_code": "CN-33", "use_ai": True}
        with patch.object(service, "acceptance_gate", return_value=gate), patch("webapp.knowledge._chat_json", side_effect=answer) as ai:
            first = self.client.post("/api/knowledge/retrieve", json=body)
            self.assertEqual(first.status_code, 200, first.text)
            self.assertTrue(first.json()["ai"]["validated"])
            second = self.client.post("/api/knowledge/retrieve", json={**body, "expected_revision": first.json()["revision"]})
            self.assertEqual(second.status_code, 200, second.text)
            self.assertTrue(second.json()["cached"])
            self.assertEqual(ai.call_count, 1)
            from dataclasses import replace
            with patch('webapp.ai_keys.apply', side_effect=lambda settings, *args, **kwargs: replace(settings,model='another-model')):
                changed=self.client.post('/api/knowledge/retrieve',json={**body,'expected_revision':second.json()['revision']})
            self.assertEqual(changed.status_code,200,changed.text)
            self.assertFalse(changed.json().get('cached'))
            self.assertEqual(ai.call_count,2)
        saved = self.client.get("/api/enterprise/input-tasks/" + task["id"]).json()
        self.assertTrue(saved["tax_checks"]["R01"]["ai"]["validated"])

    def test_http_capacity_recovery_is_saved_and_complete_result_is_reused(self):
        from src.kg_retrieval import batching
        from webapp.knowledge import ContextLengthExceeded
        task = self.task()
        body = {'rule_id': 'R01', 'task_id': task['id'], 'expected_revision': task['revision'], 'region_code': 'CN-33'}
        retrieved = self.client.post('/api/knowledge/retrieve', json=body)
        self.assertEqual(retrieved.status_code, 200, retrieved.text)
        package = retrieved.json()['bundle']
        package['policy_evidence'][0]['text'] = '完整合成条款，不得遗漏。' * 900
        package['batching'] = bundle.split_bundle(package, 128000)
        package['content_hash'] = bundle.content_hash(package)
        model_fingerprint = package['retrieval_trace']['routes']['semantic'].get('model_fingerprint')

        def provider(system, part, failure):
            if batching.encoded_size(part) > 14000:
                raise ContextLengthExceeded()
            return RetrievalTests.context_answer(system, part, failure)

        with patch.object(metrics, 'build_task_bundle', side_effect=lambda *args, **kwargs: deepcopy(package)), \
                patch.object(service, 'acceptance_gate', return_value={'passed': True, 'model_fingerprints': [model_fingerprint]}), \
                patch('webapp.knowledge._chat_json', side_effect=provider) as ai:
            body.update(use_ai=True, expected_revision=retrieved.json()['revision'])
            first = self.client.post('/api/knowledge/retrieve', json=body)
            self.assertEqual(first.status_code, 200, first.text)
            self.assertTrue(first.json()['ai']['validated'])
            self.assertTrue(first.json()['ai']['context_splits'])
            calls = ai.call_count
            second = self.client.post('/api/knowledge/retrieve', json={**body, 'expected_revision': first.json()['revision']})
            self.assertEqual(second.status_code, 200, second.text)
            self.assertTrue(second.json()['cached'])
            self.assertEqual(ai.call_count, calls)
        saved = self.client.get('/api/enterprise/input-tasks/' + task['id']).json()
        self.assertEqual(saved['tax_checks']['R01']['ai']['context_splits'], first.json()['ai']['context_splits'])

    def test_invalid_model_output_is_downgraded_and_not_saved_as_finding(self):
        task = self.task()
        model, _ = semantic_search.available()
        gate = {"passed": True, "model_fingerprints": [model.fingerprint if model else None]}
        with patch.object(service, "acceptance_gate", return_value=gate), patch("webapp.knowledge._chat_json", return_value=({"assessments": [{"type": "invented"}]}, "unit-test-provider")):
            result = self.client.post("/api/knowledge/retrieve", json={"rule_id": "R01", "task_id": task["id"], "expected_revision": task["revision"], "region_code": "CN-33", "use_ai": True})
        self.assertEqual(result.status_code, 200, result.text)
        self.assertFalse(result.json()["ai"]["validated"])
        self.assertEqual(result.json()["ai"]["assessments"], [])
        self.assertEqual(result.json()["ai"]["status"], "needs_review")

    def test_material_update_during_ai_prevents_stale_result_write(self):
        from tests.test_core_inputs import book
        task = self.task()
        def concurrent_update(package):
            changed = self.client.post("/api/enterprise/input-tasks/" + task["id"] + "/files",
                        data={"expected_revision": task["revision"], "material_type": "account_balance"},
                        files={"files": ("合成余额.xlsx", book("account_balance"))})
            self.assertEqual(changed.status_code, 200, changed.text)
            return {"validated": True, "assessments": [], "status": "reviewed_clues"}
        with patch.object(service, "analyze", side_effect=concurrent_update):
            result = self.client.post("/api/knowledge/retrieve", json={"rule_id": "R01", "task_id": task["id"], "expected_revision": task["revision"], "region_code": "CN-33", "use_ai": True})
        self.assertEqual(result.status_code, 409, result.text)
        saved = self.client.get("/api/enterprise/input-tasks/" + task["id"]).json()
        self.assertNotIn("R01", saved.get("tax_checks", {}))

    def test_relevant_material_change_marks_previous_check_stale(self):
        from tests.test_core_inputs import book
        task = self.task()
        result = self.client.post("/api/knowledge/retrieve", json={"rule_id": "R01", "task_id": task["id"], "expected_revision": task["revision"], "region_code": "CN-33"}).json()
        workbook = load_workbook(BytesIO(book("vat_declaration")))
        workbook.active["C6"] = "200000.00"
        content = BytesIO()
        workbook.save(content)
        workbook.close()
        changed = self.client.post("/api/enterprise/input-tasks/" + task["id"] + "/files",
                    data={"expected_revision": result["revision"], "material_type": "vat_declaration"},
                    files={"files": ("更新申报.xlsx", content.getvalue())})
        self.assertEqual(changed.status_code, 200, changed.text)
        self.assertTrue(changed.json()["tax_checks"]["R01"]["stale"])


@material_key
class CoreReportTests(unittest.TestCase):
    setUp = RetrievalHTTPTests.setUp
    login = RetrievalHTTPTests.login
    task = RetrievalHTTPTests.task

    def checked(self):
        task = self.task()
        response = self.client.post("/api/knowledge/retrieve", json={"rule_id": "R01", "task_id": task["id"],
                                    "expected_revision": task["revision"], "region_code": "CN-33"})
        self.assertEqual(response.status_code, 200, response.text)
        return task["id"], response.json()

    def test_report_uses_readable_policy_citations_and_retains_frozen_ids(self):
        from jinja2 import Environment, FileSystemLoader, select_autoescape
        from webapp import core_reports

        tid, result = self.checked()
        base = '/api/enterprise/input-tasks/' + tid + '/reports'
        self.assertEqual(self.client.post(base, json={'expected_revision': result['revision']}).status_code, 200)
        report = self.client.get(base + '/1?format=json').json()
        original = deepcopy(report)
        slot = report['checks']['R01']
        cid = 'TAX-COLLECTION-LAW-2015:organized:c47f560c3d8ae679e90c:第二十五条'
        evidence = deepcopy(slot['bundle']['policy_evidence'][0])
        evidence.update(clause_id=cid, title='合成测试法规', required=True, document_no='合成文号〔2026〕1号')
        slot['bundle']['policy_evidence'] = [evidence]
        slot['bundle']['policy_leads'] = []
        slot['bundle']['check']['disposition'] = 'condition_met'
        slot['bundle']['risk_assessment']['status'] = 'RISK_CANDIDATE'
        slot['bundle']['conditions'][0]['source_clause_id'] = cid
        slot['bundle']['exceptions'] = [{'name': '合成例外', 'source_clause_ids': [cid]}]
        slot['bundle']['relationships'] = []
        slot['bundle']['retrieval_trace']['excluded'] = [{'clause_id': cid, 'reasons': ['测试过滤']}]
        internal_rule = 'R-001@2.0:converted:d30fb9986b82df0c'
        slot['bundle']['check']['rule_id'] = internal_rule
        calculation = deepcopy(slot['bundle']['facts'][0])
        calculation.update(metric_id=internal_rule + ':CALC1', name='内部子条件1计算',
                           is_calculation=True, unit='布尔数值', calculation='内部程序条件比较',
                           period_basis='rule_period')
        slot['bundle']['facts'].append(calculation)
        slot['bundle']['gaps'] = [{'code': 'rule_period'}, {'code': 'rule_unapproved', 'message': 'TaxRule 版本尚未审核。'}]
        slot['ai'] = {'validated': True, 'assessments': [{'text': '整体核查状态为 UNKNOWN。conditions 与 fact_evaluation 仅为计算事实。合成说明：请核对人员花名册。', 'clause_ids': [cid]}],
                      'usage': [{'model': 'unit-test-internal-model', 'elapsed_seconds': 9.85,
                                 'usage': {'prompt_tokens': 3256, 'completion_tokens': 1827}}]}
        report['source_files'] = [{'file_id': 'internal-proof-id', 'name': '合成证明.pdf',
                                   'purpose': '合成核对材料', 'sha256': 'ac' * 32}]
        before = deepcopy(report)
        env = Environment(loader=FileSystemLoader(ROOT / 'templates'), autoescape=select_autoescape(['html']))
        env.filters.update(report_evidence=core_reports.report_evidence, report_analysis=core_reports.report_analysis, report_materials=core_reports.report_materials, report_period=core_reports.report_period, report_gaps=core_reports.report_gaps,
                           source_location=core_reports.source_location, clause_reference=core_reports.clause_reference,
                           check_result=core_reports.check_result, report_result=core_reports.report_result,
                           material_reference=core_reports.material_reference, metric_label=core_reports.metric_label)
        html = env.get_template('core_check_report.html').render(report=report)
        self.assertIn('合成测试法规 · 第二十五条', html)
        self.assertIn('合成文号〔2026〕1号', html)
        self.assertNotIn(cid, html)
        self.assertNotIn(':organized:', html)
        self.assertNotIn(evidence['text'], html)
        self.assertNotIn(evidence['file_sha256'], html)
        self.assertIn('处理建议', html)
        self.assertIn('材料来源与核对数据', html)
        self.assertIn('合成证明.pdf', html)
        self.assertIn('合成说明：请核对人员花名册。', html)
        self.assertNotIn('适用性待核实', html)
        self.assertEqual(html.count('内容由AI生成，请核实重要信息'), 1)
        self.assertTrue(html.rstrip().endswith('内容由AI生成，请核实重要信息</footer>\n</body>\n</html>'))
        self.assertEqual(html.count('待补充/复核：本项判断依据尚待确认。'), 1)
        for hidden in (internal_rule, '内部子条件1计算', '布尔数值', '程序条件', 'rule_period',
                       'TaxRule', 'JSON', 'Token', 'unit-test-internal-model', 'SHA-256',
                       'internal-proof-id', 'ac' * 32, 'UNKNOWN', 'conditions', 'fact_evaluation'):
            self.assertNotIn(hidden, html)
        self.assertEqual(report, before)
        self.assertEqual(self.client.get(base + '/1?format=json').json(), original)
        slot['ai'].update(validated=False, reason='internal diagnostic 3256 Token')
        pending_html = env.get_template('core_check_report.html').render(report=report)
        self.assertIn('部分辅助说明尚待核实，本项仍需复核', pending_html)
        self.assertNotIn('internal diagnostic', pending_html)
        self.assertNotIn('Token', pending_html)
        self.assertIn('待核实', core_reports.clause_reference('missing', slot))
        evidence.update(clause_id='opaque', text='第一条 合成原文')
        slot['bundle']['policy_evidence'] = [evidence]
        self.assertEqual(core_reports.clause_reference('opaque', slot), '合成测试法规 · 第一条')
        evidence.update(text='<script>合成文本</script>', clause_label='<b>合成条款</b>')
        self.assertIn('&lt;b&gt;合成条款&lt;/b&gt;', env.from_string('{{ cid|clause_reference(slot) }}').render(cid='opaque', slot=slot))

    def test_scope_changes_invalidate_results_and_recheck_uses_current_period(self):
        tid, checked = self.checked()
        base = "/api/enterprise/input-tasks/" + tid
        current = self.client.get(base).json()
        scope = {key: current["task"][key] for key in ("company_name", "tax_no", "start_date", "end_date")}
        changed = self.client.post(base + "/scope", json={**scope, "company_name": "纯合成测试企业更名",
                                   "expected_revision": current["revision"], "confirm": True})
        self.assertEqual(changed.status_code, 200, changed.text)
        self.assertTrue(changed.json()["tax_checks"]["R01"]["stale"])
        self.assertEqual(self.client.post(base + "/reports", json={"expected_revision": changed.json()["revision"]}).status_code, 409)
        fresh = self.client.post("/api/knowledge/retrieve", json={"rule_id": "R01", "task_id": tid,
                                 "expected_revision": changed.json()["revision"], "region_code": "CN-33"})
        self.assertEqual(fresh.status_code, 200, fresh.text)
        self.assertFalse(fresh.json()["cached"])
        changed = self.client.post(base + "/scope", json={**scope, "end_date": "2026-02-28",
                                   "expected_revision": fresh.json()["revision"], "confirm": True})
        self.assertEqual(changed.status_code, 200, changed.text)
        self.assertTrue(changed.json()["tax_checks"]["R01"]["stale"])
        with patch("webapp.knowledge_evidence.analyze") as ai:
            rechecked = self.client.post(base + "/recheck", json={"expected_revision": changed.json()["revision"]})
            self.assertEqual(rechecked.status_code, 200, rechecked.text)
            self.assertEqual(rechecked.json()["tax_checks"]["R01"]["bundle"]["task"]["period_end"], "2026-02-28")
            self.assertFalse(rechecked.json()["tax_checks"]["R01"]["stale"])
            ai.assert_not_called()

    def test_frozen_report_html_json_pdf_cache_encryption_and_tenant_permissions(self):
        tid, result = self.checked()
        base = "/api/enterprise/input-tasks/" + tid + "/reports"
        with patch("webapp.knowledge_evidence.analyze") as ai:
            created = self.client.post(base, json={"expected_revision": result["revision"]})
            self.assertEqual(created.status_code, 200, created.text)
            repeated = self.client.post(base, json={"expected_revision": result["revision"]})
            self.assertTrue(repeated.json()["cached"])
            ai.assert_not_called()
        self.assertEqual(len(self.client.get(base).json()), 1)
        html = self.client.get(base + "/1")
        self.assertEqual(html.status_code, 200, html.text)
        self.assertIn("100.0123", html.text)
        self.assertIn("本项核对未发现异常", html.text)
        self.assertIn("仍有资料待补充", html.text)
        self.assertNotIn("纯合成工程验收样本", html.text)
        data = self.client.get(base + "/1?format=json").json()
        self.assertEqual(data['checks']['R01']['bundle']['task']['data_scope'], 'synthetic')
        self.assertEqual(data["summary"]["not_met"], 1)
        self.assertEqual(data['result']['status'], 'incomplete')
        self.assertEqual(data['result']['label'], '暂无法完成检查')
        self.assertIn(data['result']['label'], html.text)
        self.assertIn('合成.xlsx', html.text)
        self.assertFalse(data["formal_finding"])
        self.assertEqual(data["checks"]["R01"]["bundle"]["facts"][0]["value"], "100.0123")
        def pdf(html, path, **kwargs):
            Path(path).write_bytes(b"%PDF-1.7\nsynthetic-cache-test")
        with patch("webapp.core_reports.render.export_pdf", side_effect=pdf) as export:
            first = self.client.get(base + "/1?format=pdf")
            second = self.client.get(base + "/1?format=pdf")
            self.assertEqual(first.content, second.content)
            self.assertEqual(export.call_count, 1)
        self.assertIn("no-store", html.headers["cache-control"])
        self.assertNotIn(b"100.0123", self.path.read_bytes())
        self.login("kg-other")
        for url in (base, base + "/1", base + "/1?format=pdf", base + "/1?format=json"):
            self.assertEqual(self.client.get(url).status_code, 404, url)
        self.assertEqual(self.client.post(base, json={"expected_revision": result["revision"]}).status_code, 404)

    def test_complete_archive_reopen_preserves_report_and_blocks_edits(self):
        from tests.test_core_inputs import book
        from webapp import core_inputs, material_batches
        from webapp.access import AccessDenied
        tid, result = self.checked()
        base = '/api/enterprise/input-tasks/' + tid
        request = {'expected_revision': result['revision'], 'report_version': 1}
        self.assertEqual(self.client.post(base + '/complete', json=request).status_code, 409)
        report = self.client.post(base + '/reports', json={'expected_revision': result['revision']})
        self.assertEqual(report.status_code, 200, report.text)
        before = self.client.get(base + '/reports/1').content
        checked_before = self.client.get(base).json()['tax_checks']
        self.assertEqual(self.client.post(base + '/complete', json={**request, 'expected_revision': 0}).status_code, 409)
        self.assertEqual(self.client.post(base + '/complete', json={**request, 'report_version': 2}).status_code, 409)
        with patch.object(service, 'analyze') as ai:
            response = self.client.post(base + '/complete', json=request)
            self.assertEqual(response.status_code, 200, response.text)
            saved = response.json()
            self.assertEqual(saved['revision'], result['revision'] + 1)
            self.assertEqual(saved['completion']['content_hash'], report.json()['content_hash'])
            self.assertEqual(saved['completion']['input_revision'], result['revision'])
            self.assertEqual(saved['tax_checks'], checked_before)
            self.assertEqual(next(row for row in self.client.get('/api/enterprise/input-tasks').json() if row['id'] == tid)['status'], 'archived')
            self.assertEqual(self.client.get(base + '/reports/1').content, before)
            self.assertEqual(self.client.get(base + '/reports/1?format=json').json()['result']['status'], 'incomplete')
            revision = {'expected_revision': saved['revision']}
            original_url = base + '/originals/' + saved['documents'][0]['id']
            original = self.client.get(original_url).content
            self.assertEqual(self.client.request('DELETE', '/api/enterprise/materials/' + tid + '/originals/' + saved['documents'][0]['id'], json=revision).status_code, 409)
            self.assertEqual(self.client.get(original_url).content, original)
            for suffix in ('/prepare', '/scope', '/recheck', '/reports'):
                self.assertEqual(self.client.post(base + suffix, json=revision).status_code, 409, suffix)
            self.assertEqual(self.client.post('/api/knowledge/retrieve', json={**revision, 'task_id': tid, 'rule_id': 'R01', 'region_code': 'CN-33', 'use_ai': True}).status_code, 409)
            self.assertEqual(self.client.post(base + '/files', data={**revision, 'material_type': 'account_balance'}, files={'files': ('合成.xlsx', book('account_balance'))}).status_code, 409)
            actor = self.store.authenticate('kg-admin', 'Knowledge-test-2026!')[0]
            payload = deepcopy(core_inputs.latest(material_batches.read(self.store, actor, tid)))
            payload.pop('completion')
            with self.assertRaises(AccessDenied):
                material_batches.append_revision(self.store, actor, tid, saved['revision'], 'edit', payload)
            ai.assert_not_called()
        self.login('kg-other')
        self.assertEqual(self.client.post(base + '/complete', json=request).status_code, 404)
        self.assertEqual(self.client.post(base + '/reopen', json=revision).status_code, 404)
        self.login('kg-admin')
        self.assertEqual(self.client.post(base + '/reopen', json={'expected_revision': 0}).status_code, 409)
        reopened = self.client.post(base + '/reopen', json=revision)
        self.assertEqual(reopened.status_code, 200, reopened.text)
        self.assertNotIn('completion', reopened.json())
        self.assertEqual(self.client.get(base + '/reports/1').content, before)
        refreshed = self.client.post(base + '/reports', json={'expected_revision': reopened.json()['revision']})
        self.assertTrue(refreshed.json()['cached'])
        # A cached report may have an earlier input revision; identical evidence is valid.
        again = self.client.post(base + '/complete', json={'expected_revision': reopened.json()['revision'], 'report_version': 1})
        self.assertEqual(again.status_code, 200, again.text)

    def test_complete_rejects_changed_evidence_and_rolls_back_failed_event(self):
        tid, result = self.checked()
        base = '/api/enterprise/input-tasks/' + tid
        request = {'expected_revision': result['revision'], 'report_version': 1}
        self.client.post(base + '/reports', json={'expected_revision': result['revision']})
        from webapp import material_batches
        original_event = material_batches._event
        def fail_completion_event(db, batch, actor, action, detail, **kwargs):
            if action == 'core_complete':
                raise RuntimeError('event failure')
            return original_event(db, batch, actor, action, detail, **kwargs)
        with patch('webapp.material_batches._event', side_effect=fail_completion_event):
            with self.assertRaises(RuntimeError):
                self.client.post(base + '/complete', json=request)
        current = self.client.get(base).json()
        self.assertNotIn('completion', current)
        self.assertEqual(current['revision'], result['revision'])
        with self.store.connect() as db:
            db.execute("UPDATE Policy SET validity_status='unverified' WHERE policy_id='P01'")
        self.assertEqual(self.client.post(base + '/complete', json=request).status_code, 409)
        self.assertNotIn('completion', self.client.get(base).json())

    def test_supplement_recheck_updates_report_version_without_rewriting_history(self):
        from tests.test_core_inputs import book
        tid, result = self.checked()
        base = "/api/enterprise/input-tasks/" + tid
        self.assertEqual(self.client.post(base + "/reports", json={"expected_revision": result["revision"]}).status_code, 200)
        before = self.client.get(base + "/reports/1").content
        wb = load_workbook(BytesIO(book("vat_declaration")))
        wb.active["C6"] = "200000.00"
        output = BytesIO()
        wb.save(output)
        wb.close()
        updated = self.client.post(base + "/files", data={"expected_revision": result["revision"], "material_type": "vat_declaration"},
                                   files={"files": ("更新.xlsx", output.getvalue())}).json()
        self.assertTrue(updated["tax_checks"]["R01"]["stale"])
        self.assertEqual(updated["tax_check_summary"]["needs_review"], 1)
        stale = self.client.post(base + "/reports", json={"expected_revision": updated["revision"]})
        self.assertEqual(stale.status_code, 409, stale.text)
        again = self.client.post("/api/knowledge/retrieve", json={"rule_id": "R01", "task_id": tid,
                                 "expected_revision": updated["revision"], "region_code": "CN-33"})
        self.assertEqual(again.status_code, 200, again.text)
        version = self.client.post(base + "/reports", json={"expected_revision": again.json()["revision"]})
        self.assertEqual(version.status_code, 200, version.text)
        self.assertEqual(version.json()["version"], 2)
        self.assertEqual(before, self.client.get(base + "/reports/1").content)
        self.assertEqual(self.client.get(base + "/reports/2?format=json").json()["summary"]["needs_review"], 1)

    def test_report_rejects_unchecked_stale_revision_and_changed_policy(self):
        task = self.task()
        base = "/api/enterprise/input-tasks/" + task["id"] + "/reports"
        self.assertEqual(self.client.post(base, json={"expected_revision": task["revision"]}).status_code, 409)
        result = self.client.post("/api/knowledge/retrieve", json={"rule_id": "R01", "task_id": task["id"],
                                  "expected_revision": task["revision"], "region_code": "CN-33"}).json()
        self.assertEqual(self.client.post(base, json={"expected_revision": task["revision"]}).status_code, 409)
        with self.store.connect() as db:
            db.execute("UPDATE Policy SET validity_status='unverified' WHERE policy_id='P01'")
        changed = self.client.post(base, json={"expected_revision": result["revision"]})
        self.assertEqual(changed.status_code, 409, changed.text)
        self.assertEqual(self.client.get(base).json(), [])

    def test_report_accepts_legacy_gap_order_without_rechecking_or_calling_ai(self):
        with self.store.connect() as db:
            db.execute("UPDATE Policy SET validity_status='unverified' WHERE policy_id='P01'")
        original = metrics.build_task_bundle
        def legacy(*args, **kwargs):
            package = original(*args, **kwargs)
            package['gaps'].reverse()
            package['content_hash'] = 'legacy-order-dependent-hash'
            return package
        with patch.object(metrics, 'build_task_bundle', side_effect=legacy):
            tid, checked = self.checked()
        base = '/api/enterprise/input-tasks/' + tid
        with patch('webapp.knowledge_evidence.analyze') as ai:
            recheck = self.client.post(base + '/recheck', json={'expected_revision': checked['revision']})
            self.assertEqual(recheck.status_code, 200, recheck.text)
            self.assertEqual(recheck.json()['changed_rule_ids'], [])
            self.assertEqual(recheck.json()['reused_rule_ids'], ['R01'])
            created = self.client.post(base + '/reports', json={'expected_revision': checked['revision']})
            self.assertEqual(created.status_code, 200, created.text)
            ai.assert_not_called()
        report = self.client.get(base + '/reports/1?format=json').json()
        self.assertEqual(report['checks']['R01']['bundle'], checked['bundle'])

    def test_report_escapes_user_text_and_pdf_reauthorizes_after_render(self):
        tid, result = self.checked()
        base = "/api/enterprise/input-tasks/" + tid
        from webapp import material_batches, core_inputs
        actor = self.store.authenticate("kg-admin", "Knowledge-test-2026!")[0]
        view = material_batches.read(self.store, actor, tid)
        payload = deepcopy(core_inputs.latest(view))
        payload["task"]["company_name"] = '<script>alert("unsafe")</script>'
        revision = material_batches.append_revision(self.store, actor, tid, result["revision"], "edit", payload)
        self.assertEqual(self.client.post(base + "/reports", json={"expected_revision": revision}).status_code, 409)
        rechecked = self.client.post(base + "/recheck", json={"expected_revision": revision})
        self.assertEqual(rechecked.status_code, 200, rechecked.text)
        archived = self.client.post(base + "/reports", json={"expected_revision": rechecked.json()["revision"]})
        self.assertEqual(archived.status_code, 200, archived.text)
        html = self.client.get(base + "/reports/1").text
        self.assertNotIn('<script>alert("unsafe")</script>', html)
        self.assertIn("&lt;script&gt;", html)
        def revoke(html, path, **kwargs):
            Path(path).write_bytes(b"%PDF-1.7")
            with self.store.connect() as db:
                db.execute("UPDATE users SET active=0 WHERE username='kg-admin'")
        with patch("webapp.core_reports.render.export_pdf", side_effect=revoke):
            denied = self.client.get(base + "/reports/1?format=pdf")
            self.assertIn(denied.status_code, (401, 403))

    def test_bulk_recheck_preserves_unaffected_checks_and_makes_no_ai_calls(self):
        from tests.test_core_inputs import book
        with self.store.connect() as db:
            db.execute("UPDATE Metric SET source_schema='AccountBalanceRow',source_field='current_debit',period_basis='accounting_period',source_filter_json=? WHERE metric_id='M03'", ('{"account_code":"1001"}',))
        tid, first = self.checked()
        base = "/api/enterprise/input-tasks/" + tid
        task = self.client.post(base + "/files", data={"expected_revision": first["revision"], "material_type": "account_balance"},
                                files={"files": ("余额.xlsx", book("account_balance"))}).json()
        second = self.client.post("/api/knowledge/retrieve", json={"rule_id": "R03", "task_id": tid,
                                  "expected_revision": task["revision"], "region_code": "CN-33"}).json()
        wb = load_workbook(BytesIO(book("vat_declaration")))
        wb.active["C6"] = "200000.00"
        output = BytesIO()
        wb.save(output)
        wb.close()
        updated = self.client.post(base + "/files", data={"expected_revision": second["revision"], "material_type": "vat_declaration"},
                                   files={"files": ("更新.xlsx", output.getvalue())}).json()
        with patch("webapp.knowledge_evidence.analyze") as ai:
            checked = self.client.post(base + "/recheck", json={"expected_revision": updated["revision"]})
            self.assertEqual(checked.status_code, 200, checked.text)
            self.assertEqual(checked.json()["changed_rule_ids"], ["R01"])
            self.assertEqual(checked.json()["reused_rule_ids"], ["R03"])
            self.assertEqual(checked.json()["tax_checks"]["R03"]["bundle"], second["bundle"])
            self.assertEqual(checked.json()["external_api_calls"], 0)
            again = self.client.post(base + "/recheck", json={"expected_revision": checked.json()["revision"]})
            self.assertEqual(again.json()["changed_rule_ids"], [])
            self.assertEqual(again.json()["revision"], checked.json()["revision"])
            ai.assert_not_called()
        self.assertEqual(self.client.post(base + "/recheck", json={"expected_revision": updated["revision"]}).status_code, 409)
        self.login("kg-other")
        self.assertEqual(self.client.post(base + "/recheck", json={"expected_revision": checked.json()["revision"]}).status_code, 404)


class ReportPresentationTests(unittest.TestCase):
    def test_material_dates_keep_historical_boundaries_and_unknown_comparisons(self):
        from webapp.core_reports import report_period
        report = {'task': {'start_date': '2026-09-01', 'end_date': '2026-09-30'}}
        cases = {
            '存货期初余额': '期初余额：2026-09-01 期初。',
            '存货期末余额': '期末余额：2026-09-30 期末。',
            '存货本期入库成本': '要求期间：2026-09-01 至 2026-09-30。',
            '历史.上年同期收入': '要求期间：2025-09-01 至 2025-09-30。',
            '上年同期营业收入报表': '要求期间：2025-09-01 至 2025-09-30。',
            '上期收入报表': '上期：2026-08-01 至 2026-08-31。',
            '上期期末余额': '上期期末余额时点：2026-08-31。',
        }
        for name, expected in cases.items():
            with self.subTest(name=name):
                self.assertEqual(report_period(name, report), expected)
        self.assertIn('待确认', report_period('本期及可比期间收入', report))
        self.assertIn('待确认', report_period('年度.上年净利润', report))
        self.assertIn('2025-01-01 至 2025-12-31', report_period('2025-01-01至2025-12-31的利润表', report))
        self.assertIn('待确认', report_period('2025-02-30至2025-03-01的报表', report))
        self.assertIn('2025-08-01 至 2025-08-31', report_period('2025年8月利润表', report))
        self.assertIn('2025-01-01 至 2025-12-31', report_period('2025年度财务报表', report))
        self.assertIn('待确认', report_period('2025年13月利润表', report))
        self.assertIn('2025-08-31', report_period('截至2025年8月31日的余额表', report))
        annual = {'task': {'start_date': '2026-01-01', 'end_date': '2026-12-31'}}
        self.assertEqual(report_period('年度.上年净利润', annual), '要求期间：2025-01-01 至 2025-12-31。')
        self.assertEqual(report_period('年度.前年末资产', annual), '余额时点：2024-12-31 期末。')
        leap = {'task': {'start_date': '2024-02-01', 'end_date': '2024-02-29'}}
        self.assertIn('2023-02-01 至 2023-02-28', report_period('历史.上年同期收入', leap))
        custom = {'task': {'start_date': '2026-09-08', 'end_date': '2026-10-17'}}
        self.assertIn('待确认', report_period('上期收入', custom))
        fact = {'name': '历史.上年同期收入', 'metric_period': '2024年度'}
        self.assertIn('2025-09-01 至 2025-09-30', report_period(fact, report))
        self.assertEqual(fact['metric_period'], '2024年度')

    def test_material_list_deduplicates_without_losing_periods_or_project_links(self):
        from webapp.core_reports import report_materials
        report = self.report()
        report['task'] = {'start_date': '2026-09-01', 'end_date': '2026-09-30'}
        first = report['checks']['R01']
        first['bundle'].update(check={'disposition': 'needs_review', 'topic': '存货余额'},
                               facts=[], gaps=[{'code': 'policy_review'}])
        first['ai'] = {'validated': True, 'assessments': [{'type': 'insufficient_evidence',
            'required_materials': ['存货本期入库成本明细表', '存货本期入库成本明细表',
                                   '存货本期出库成本明细表', '政策条款效力与原文的复核记录']} ]}
        second = deepcopy(first)
        second['bundle']['check']['topic'] = '存货成本'
        second['ai']['assessments'][0]['required_materials'] = ['存货本期入库成本明细表', '存货期初余额凭证',
                                                             '存货期末余额凭证']
        report['checks']['R02'] = second
        third = deepcopy(first)
        third['bundle']['check']['topic'] = '收入比较'
        third['bundle']['facts'] = [{'name': '利润表.营业收入', 'state': 'known'}]
        third['ai']['assessments'][0]['required_materials'] = ['利润表.营业收入本期及可比期间数据',
                                                             '上年同期营业收入报表']
        report['checks']['R03'] = third
        original = deepcopy(report)
        result = report_materials(report)
        rows = result['uploads']
        current = next(row for row in rows if row['name'] == '存货账簿及收发存资料' and '要求期间' in row['period'])
        self.assertEqual(current['requirements'], ['存货本期入库成本明细表', '存货本期出库成本明细表'])
        self.assertEqual([p['number'] for p in current['projects']], [1, 2])
        self.assertEqual(len([row for row in rows if row['name'] == '存货账簿及收发存资料']), 3)
        self.assertTrue(any('2025-09-01 至 2025-09-30' in row['period'] for row in rows))
        self.assertTrue(any('待确认' in row['period'] for row in rows))
        self.assertFalse(any('政策条款' in text for row in rows for text in row['requirements']))
        self.assertEqual(len(result['reviews']), 1)
        for rid, item in result['by_check'].items():
            for number in item['uploads']:
                self.assertIn(rid, [p['key'] for p in rows[number - 1]['projects']])
        self.assertEqual(report, original)

    def test_business_analysis_omits_execution_diagnostics_and_preserves_material_requests(self):
        from webapp.core_reports import report_analysis
        slot = {'ai': {'validated': True, 'assessments': [
            {'text': '整体核查状态为 UNKNOWN。conditions 与 fact_evaluation 不能替代结论。'
                     '应核对个税申报人数与社保参保人数，补充对应期间的申报表。',
             'required_materials': ['人员花名册及申报记录', '补齐 rule_period 与 RuleCalculation', 'TaxRule版本审核记录']},
            {'text': '当前状态unknown、UNKOWN、TRUE、FALSE。所用规则版本尚未审核。'},
            {'text': '已知数据不足时不能仅凭差额认定问题。请核对原始凭证。'},
        ]}}
        original = deepcopy(slot)
        self.assertEqual(report_analysis(slot), [
            {'text': '应核对个税申报人数与社保参保人数，补充对应期间的申报表。',
             'required_materials': ['人员花名册及申报记录']},
            {'text': '已知数据不足时不能仅凭差额认定问题。请核对原始凭证。', 'required_materials': []},
        ])
        self.assertEqual(slot, original)
        slot['ai']['validated'] = False
        self.assertEqual(report_analysis(slot), [])

    def report(self):
        return {'checks': {'R01': {'bundle': {
            'check': {'disposition': 'condition_not_met'}, 'bundle_status': 'ready', 'gaps': [],
            'fact_assessment': {'status': 'FACT_CONFIRMED', 'amount': None},
            'risk_assessment': {'status': 'NO_RISK_SIGNAL'}}}},
            'processing': {'status': 'ready_for_analysis', 'issues': []}}

    def test_pass_requires_complete_materials_checks_and_policy_review(self):
        from webapp.core_reports import report_result
        report = self.report()
        self.assertEqual(report_result(report)['status'], 'passed')
        for change in ('materials', 'policy', 'fact', 'ai', 'stale', 'zero_with_gaps', 'no_checks', 'excluded'):
            with self.subTest(change=change):
                current = deepcopy(report)
                slot = current['checks']['R01']
                if change == 'materials':
                    current['processing']['status'] = 'needs_completion'
                elif change == 'policy':
                    current['policy_assessment'] = {'consistency': {'status': 'needs_review'}}
                elif change == 'fact':
                    slot['bundle']['fact_assessment']['status'] = 'FACT_INCOMPLETE'
                elif change == 'ai':
                    slot['ai'] = {'validated': True, 'assessments': [{'type': 'insufficient_evidence'}]}
                elif change == 'stale':
                    slot['stale'] = True
                elif change == 'zero_with_gaps':
                    slot['bundle']['fact_assessment']['amount'] = '0'
                    slot['bundle']['gaps'] = [{'code': 'policy_review'}]
                    slot['bundle']['bundle_status'] = 'needs_review'
                elif change == 'no_checks':
                    current['checks'] = {}
                elif change == 'excluded':
                    slot['bundle']['check']['disposition'] = 'excluded'
                    slot['bundle']['risk_assessment']['status'] = 'NOT_APPLICABLE'
                self.assertEqual(report_result(current)['status'], 'incomplete')

    def test_confirmed_difference_remains_visible_with_pending_risk_and_other_gaps(self):
        from webapp.core_reports import check_result, report_result
        report = self.report()
        slot = report['checks']['R01']
        slot['bundle']['fact_assessment']['amount'] = '15000.00'
        slot['bundle']['risk_assessment']['status'] = 'NEEDS_REVIEW'
        slot['bundle']['bundle_status'] = 'needs_review'
        slot['bundle']['gaps'] = [{'code': 'policy_review'}]
        report['checks']['R02'] = deepcopy(self.report()['checks']['R01'])
        report['checks']['R02']['bundle']['fact_assessment']['status'] = 'FACT_INCOMPLETE'
        before = deepcopy(report)
        self.assertEqual(check_result(slot)['label'], '已确认数据差异，问题性质待核实')
        result = report_result(report)
        self.assertEqual((result['status'], result['issues'], result['pending']), ('issues_found', 1, 2))
        self.assertEqual(report, before)

    def test_material_references_use_names_and_positions_without_inventing_missing_sources(self):
        from webapp.core_reports import material_reference
        report = {'sources': [{'record_id': 'private-id', 'file_name': '申报表.xlsx', 'source_position': '主表!C6'}],
                  'source_files': [{'file_id': 'proof-id', 'name': '零发生额证明.pdf'}]}
        self.assertEqual(material_reference(['private-id'], report), '申报表.xlsx · 主表!C6')
        self.assertEqual(material_reference(['file:proof-id'], report), '零发生额证明.pdf')
        self.assertEqual(material_reference(['unknown-id'], report), '材料来源待核对')
        self.assertEqual(material_reference([], report), '缺少材料来源')


if __name__ == "__main__":
    unittest.main()
