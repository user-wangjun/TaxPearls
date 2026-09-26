"""FR-A13: collect synthetic finance/tax pages in a localhost-only browser sandbox.

This deliberately has no configurable remote URL or production credentials.
Run: .venv/Scripts/python.exe scripts/rpa_sandbox.py --output output/rpa-sandbox.xlsx
"""
from __future__ import annotations

import argparse
import html
import sys
from contextlib import contextmanager
from decimal import Decimal, InvalidOperation
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from threading import Thread

from openpyxl import Workbook
from playwright.sync_api import sync_playwright

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from src import config, loader  # noqa: E402
from scripts.make_sample import ACCOUNTS, COMPANY, DECLARATION  # noqa: E402

ACCOUNT_HEADERS = config.COL_ACCOUNTS
DECLARATION_HEADERS = ["项目", "金额"]


def _table(headers: list[str], rows: list[tuple]) -> str:
    def cell(tag: str, value: object) -> str:
        return f"<{tag}>{html.escape(str(value))}</{tag}>"

    head = "".join(cell("th", item) for item in headers)
    body = "".join("<tr>" + "".join(cell("td", item) for item in row) + "</tr>" for row in rows)
    return f"<table id='records'><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>"


def _page(kind: str) -> bytes:
    if kind == "finance":
        title = "仿真财务软件"
        table = _table(ACCOUNT_HEADERS, ACCOUNTS)
    else:
        title = "仿真电子税务局"
        table = _table(DECLARATION_HEADERS, DECLARATION)
    fields = "".join(
        f"<div data-field='{html.escape(key)}'>{html.escape(value)}</div>"
        for key, value in COMPANY.items()
    )
    # A click is required before the records appear, to exercise UI automation.
    document = ("<!doctype html><html lang='zh'><meta charset='utf-8'>"
                f"<title>{title} · 纯合成演示</title><h1>{title} · 纯合成演示</h1>"
                "<button id='enter' onclick=\"document.querySelector('#data').hidden=false;this.remove()\">"
                "进入仿真演示</button>"
                f"<section id='data' hidden>{fields}{table}</section></html>")
    return document.encode("utf-8")


class _Handler(BaseHTTPRequestHandler):
    def do_GET(self) -> None:  # noqa: N802
        if self.path not in ("/finance", "/tax"):
            self.send_error(404)
            return
        content = _page(self.path[1:])
        self.send_response(200)
        self.send_header("Content-Type", "text/html; charset=utf-8")
        self.send_header("Content-Length", str(len(content)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(content)

    def log_message(self, format: str, *args: object) -> None:
        pass


@contextmanager
def _server():
    server = ThreadingHTTPServer(("127.0.0.1", 0), _Handler)
    worker = Thread(target=server.serve_forever, daemon=True)
    worker.start()
    try:
        yield f"http://127.0.0.1:{server.server_port}"
    finally:
        server.shutdown()
        server.server_close()
        worker.join(timeout=5)


def _read_portal(page, base_url: str, path: str, headers: list[str]) -> tuple[dict[str, str], list[list[str]]]:
    response = page.goto(base_url + path, wait_until="domcontentloaded")
    if response is None or response.status != 200 or page.url != base_url + path:
        raise ValueError("仿真门户未返回预期页面")
    page.get_by_role("button", name="进入仿真演示").click()
    if "纯合成演示" not in page.title():
        raise ValueError("来源页面缺少仿真标识")
    fields = page.locator("#data [data-field]").evaluate_all(
        "nodes => nodes.map(n => [n.getAttribute('data-field'), n.textContent.trim()])"
    )
    company = dict(fields)
    if len(fields) != len(company) or set(company) != set(config.COMPANY_FIELDS):
        raise ValueError("企业字段缺失或重复")
    observed_headers = page.locator("#records thead th").all_text_contents()
    if observed_headers != headers:
        raise ValueError("仿真门户表头与标准模板不一致")
    rows = page.locator("#records tbody tr").evaluate_all(
        "nodes => nodes.map(tr => Array.from(tr.querySelectorAll('td'), td => td.textContent.trim()))"
    )
    if not 1 <= len(rows) <= 500 or any(len(row) != len(headers) for row in rows):
        raise ValueError("采集行数或列数不符合预期")
    return company, rows


def _amount(value: str) -> float:
    try:
        number = Decimal(value)
    except InvalidOperation:
        raise ValueError("采集金额不是有效数字") from None
    if not number.is_finite() or number < 0 or number.as_tuple().exponent < -2:
        raise ValueError("采集金额须为非负且最多两位小数")
    return float(number)


def collect(output: str | Path) -> Path:
    """Drive both local mock portals and emit a loader-compatible, synthetic workbook."""
    output = Path(output).resolve()
    if output.exists():
        raise FileExistsError(f"输出文件已存在，不覆盖：{output}")
    with _server() as base_url, sync_playwright() as playwright:
        # Windows dev machines use the already-installed Chrome; Linux images
        # with Playwright's bundled Chromium use the default executable.
        browser = playwright.chromium.launch(headless=True, channel="chrome" if sys.platform == "win32" else None)
        try:
            context = browser.new_context(service_workers="block")
            context.route("**/*", lambda route: route.continue_() if route.request.url.startswith(base_url + "/") else route.abort())
            page = context.new_page()
            finance_company, accounts = _read_portal(page, base_url, "/finance", ACCOUNT_HEADERS)
            tax_company, declaration = _read_portal(page, base_url, "/tax", DECLARATION_HEADERS)
            context.close()
        finally:
            browser.close()
    if finance_company != tax_company or not all(finance_company.values()):
        raise ValueError("财务与税务门户的企业/期间身份不一致")
    if "仿真" not in finance_company["企业名称"]:
        raise ValueError("仅允许输出明确标记的仿真企业")
    seen_accounts, converted_accounts = set(), []
    for row in accounts:
        code = row[0]
        if not code or code in seen_accounts:
            raise ValueError("科目编码为空或重复")
        seen_accounts.add(code)
        converted_accounts.append([code, row[1], *(_amount(value) for value in row[2:])])
    seen_items, converted_declaration = set(), []
    for item, amount in declaration:
        if item not in config.DECLARATION_ITEMS.values() or item in seen_items:
            raise ValueError("申报项目未知或重复")
        seen_items.add(item)
        converted_declaration.append([item, _amount(amount)])
    wb = Workbook()
    company_sheet = wb.active
    company_sheet.title = config.SHEET_COMPANY
    company_sheet.append(["项目", "内容"])
    for key in config.COMPANY_FIELDS:
        company_sheet.append([key, finance_company[key]])
    account_sheet = wb.create_sheet(config.SHEET_ACCOUNTS)
    account_sheet.append(ACCOUNT_HEADERS)
    for row in converted_accounts:
        account_sheet.append(row)
    declaration_sheet = wb.create_sheet(config.SHEET_DECLARATION)
    declaration_sheet.append(DECLARATION_HEADERS)
    for row in converted_declaration:
        declaration_sheet.append(row)
    note = wb.create_sheet("采集说明")
    note.append(["来源", "范围", "性质"])
    note.append(["本机仿真财务软件页面", "科目余额表", "纯合成数据；非真实系统采集"])
    note.append(["本机仿真电子税务局页面", "增值税申报", "纯合成数据；非真实系统采集"])
    # Validate the exact bytes through the production loader before publishing the file.
    from io import BytesIO
    buffer = BytesIO()
    wb.save(buffer)
    wb.close()
    loader.load_bytes(buffer.getvalue())
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("xb") as file:
        file.write(buffer.getvalue())
    return output


def main() -> int:
    parser = argparse.ArgumentParser(description="仅采集本机纯合成财务/税务仿真门户")
    parser.add_argument("--output", required=True, help="输出标准审计工作簿路径；已存在文件不会覆盖")
    args = parser.parse_args()
    path = collect(args.output)
    print(f"已采集本机仿真数据：{path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
