"""Actual historical intervals, law editions and benchmark scope are hard gates."""
from copy import deepcopy
from dataclasses import asdict, replace
from decimal import Decimal
from io import BytesIO
from pathlib import Path
import unittest
from unittest.mock import patch

from openpyxl import Workbook

from src import config, engine, legal_periods, loader, materials, material_review, material_security, reference_scope, exercise_generator
from src.models import Company, Dataset, Metric
from src.snapshots import serialize_dataset, deserialize_dataset, serialize_findings, deserialize_findings
from webapp import enterprise_analysis
from tests.test_enterprise_materials import EnterpriseMaterialFixture

ROOT = Path(__file__).resolve().parents[1]
RULES = {r.id: r for r in engine.load_rules(ROOT / 'rules')}


def reference_dataset(kind='税负率', period='2026年度'):
    company = Company('来源测试企业', 'SOURCE-TEST', '批发和零售业', period, '江苏省', '一般纳税人', '商品销售')
    spec = reference_scope.SPECS[kind]
    ref = reference_scope.synthetic(company)[kind]
    ref.update(nature='来源参考', publisher='构造测试发布者', version='测试-v1', source='构造来源，不是真实业务验收')
    ref['metric_source'] = ref['source']
    values = ({'增值税.应纳税额': '.03', '增值税.销售额': '1'} if kind == '税负率'
              else {'利润表.营业收入': '1', '利润表.营业成本': '.7'})
    values.update({spec[0]: '.02' if kind == '税负率' else '.2', spec[1]: '.05' if kind == '税负率' else '.4'})
    return Dataset(company, [], {}, {key: Metric(key, Decimal(value), ref['source'] if key in spec[:2] else '构造原文', reference=deepcopy(ref) if key in spec[:2] else None)
                                    for key, value in values.items()})


def reference_workbook(data):
    book = Workbook()
    ws = book.active
    ws.title = config.SHEET_COMPANY
    ws.append(['项目', '内容'])
    for label, key in zip(config.COMPANY_FIELDS, ('name', 'taxpayer_id', 'industry', 'period')):
        ws.append([label, getattr(data.company, key)])
    for label, key in config.COMPANY_SCOPE_FIELDS.items():
        ws.append([label, getattr(data.company, key)])
    for name, headers in ((config.SHEET_ACCOUNTS, config.COL_ACCOUNTS), (config.SHEET_DECLARATION, ['项目', '金额'])):
        book.create_sheet(name).append(headers)
    book[config.SHEET_ACCOUNTS].append(['1001', '构造现金', 0, 0, 0, 0])
    declared = book[config.SHEET_DECLARATION]
    for key, label in config.DECLARATION_ITEMS.items():
        if key in data.metrics:
            declared.append([label, str(data.metrics[key].value)])
    supplement = book.create_sheet(config.SHEET_SUPPLEMENT)
    supplement.append(config.COL_SUPPLEMENT)
    for key, metric in data.metrics.items():
        if key not in config.DECLARATION_ITEMS:
            supplement.append([key, str(metric.value), metric.source, data.company.period, '构造口径'])
    ws = book.create_sheet(config.SHEET_REFERENCE)
    ws.append(config.COL_REFERENCE)
    for kind, spec in reference_scope.SPECS.items():
        if spec[0] in data.metrics:
            ref = data.metrics[spec[0]].reference
            ws.append([kind, *(ref[k] for k in reference_scope.FIELDS)])
    raw = BytesIO()
    book.save(raw)
    book.close()
    return raw.getvalue()


class LegalPeriods(unittest.TestCase):
    def test_invalid_calendar_years_and_history_underflow_fail_closed(self):
        for period in ('0000年度', '0000Q1', '0000H1', '0001年度'):
            data = reference_dataset()
            data.company.period = period
            self.assertEqual(engine.evaluate(RULES['R-022'], data).status, 'skipped')
        from src import periods
        for name in ('历史.上年同期收入', '年度.前年净利润'):
            with self.assertRaises(materials.InputError):
                periods.metric_period(name, '0001年度')

    def test_old_vat_period_and_new_law_do_not_mix(self):
        old = legal_periods.resolve(RULES['R-002'], '2022年度')
        new = legal_periods.resolve(RULES['R-002'], '2026年度')
        self.assertIn(legal_periods.VAT_OLD, old.references)
        self.assertNotIn(legal_periods.VAT, old.references)
        self.assertIn(legal_periods.VAT, new.references)
        self.assertNotIn(legal_periods.VAT_OLD, new.references)
        self.assertIn('暂行条例', ' '.join(old.legal_basis))

    def test_year_crossing_accounting_revision_keeps_both_segments(self):
        rule = legal_periods.resolve(RULES['R-009'], '2024年度')
        text = ' '.join(rule.legal_basis)
        self.assertIn('2024-01-01 至 2024-06-30', text)
        self.assertIn('2024-07-01 至 2024-12-31', text)
        self.assertIn(legal_periods.ACCOUNT_OLD, rule.references)
        self.assertIn(legal_periods.ACCOUNT, rule.references)
        self.assertNotIn('（2024年）', ' '.join(legal_periods.resolve(RULES['R-009'], '2024H1').legal_basis))

    def test_unreviewed_earlier_law_skips_without_new_citation(self):
        data = Dataset(Company('测试', 'TEST', '测试', '2000年度'), [], {}, {})
        finding = engine.evaluate(RULES['R-009'], data)
        self.assertEqual(finding.status, 'skipped')
        self.assertIn('尚无覆盖', finding.skip_reason)
        self.assertEqual(finding.rule.legal_basis, [])
        self.assertEqual(finding.rule.references, [])
        self.assertTrue(engine.readiness(finding.rule, data))

    def test_frozen_legal_snapshot_does_not_change_with_catalog(self):
        data = reference_dataset(period='2022年度')
        frozen = legal_periods.resolve(RULES['R-002'], data.company.period)
        before = serialize_findings([engine.evaluate(frozen, data)])
        with patch.dict(legal_periods.EDITIONS, {}, clear=True):
            after = serialize_findings([engine.evaluate(deserialize_findings(before)[0].rule, data)])
        self.assertEqual(before, after)
        other = legal_periods.resolve(frozen, '2026年度')
        self.assertIn(legal_periods.VAT, other.references)

    def test_invalid_period_is_a_limited_rule_not_exception(self):
        data = reference_dataset(period='2026年度')
        data.company.period = '未知'
        self.assertEqual(engine.evaluate(RULES['R-003'], data).status, 'skipped')


class ReferenceScope(unittest.TestCase):
    def test_manual_history_period_changes_invalidate_previous_review(self):
        doc = {'id': '0', 'kind': 'pdf', 'sha256': 'synthetic',
               'extraction': {'method': 'ai'}, 'rows': []}
        row = {'name': '历史.上年同期收入', 'value': '100', 'page': 1, 'detail': '构造原文', 'period': '2025年度'}
        before = material_security.review_requirements(doc, {'rows': [row]}, {'period': '2026年度'})
        after = material_security.review_requirements(doc, {'rows': [{**row, 'period': '2020年度'}]}, {'period': '2026年度'})
        self.assertNotEqual([r['id'] for r in before], [r['id'] for r in after])

    def test_pdf_actual_period_is_bounded_on_server(self):
        from tests.test_ai_extraction import document, CATALOG, COMPANY
        doc = document()
        selected = {'0': {'reviewed': True, 'rows': [{'name': '增值税.销售额', 'value': '1',
                    'page': 1, 'detail': '构造值', 'period': 'x' * 101}]}}
        with self.assertRaisesRegex(materials.InputError, '100'):
            materials.build_dataset([doc], selected, COMPANY, set(CATALOG))

    def test_matching_scope_executes_and_freezes_all_evidence(self):
        for kind, rid in [('税负率', 'R-003'), ('毛利率', 'R-021')]:
            data = reference_dataset(kind)
            restored = deserialize_dataset(serialize_dataset(data))
            self.assertEqual(engine.readiness(RULES[rid], restored), [])
            finding = engine.evaluate(RULES[rid], restored)
            self.assertEqual(finding.status, 'pass')
            self.assertTrue(any('发布者：构造测试发布者' in e.value for e in finding.evidence))

    def test_each_applicable_dimension_is_checked_not_free_text(self):
        for kind, rid in [('税负率', 'R-003'), ('毛利率', 'R-021')]:
            for key in reference_scope.SPECS[kind][3]:
                for wrong in ('', '另一地区/行业/身份/构成'):
                    data = reference_dataset(kind)
                    setattr(data.company, key, wrong)
                    with self.subTest(kind=kind, key=key, wrong=wrong):
                        self.assertTrue(engine.readiness(RULES[rid], data))
                        self.assertEqual(engine.evaluate(RULES[rid], data).status, 'skipped')

    def test_missing_reference_scope_or_source_never_passes(self):
        for key in ('publisher', 'source', 'version', 'definition'):
            data = reference_dataset()
            for name in reference_scope.SPECS['税负率'][:2]:
                data.metrics[name].reference[key] = ''
            self.assertEqual(engine.evaluate(RULES['R-003'], data).status, 'skipped')
        data.metrics['参考.税负率下限'].reference = None
        self.assertEqual(engine.evaluate(RULES['R-003'], data).status, 'skipped')

    def test_reference_must_cover_whole_period_and_have_right_definition(self):
        for key, value in [('start', '2026-02-01'), ('end', '2026-11-30'), ('start', '2026-02-30'),
                           ('definition', '销项税额/含税收入'), ('nature', '官方标准（未经证明）')]:
            data = reference_dataset()
            for name in reference_scope.SPECS['税负率'][:2]:
                data.metrics[name].reference[key] = value
            self.assertEqual(engine.evaluate(RULES['R-003'], data).status, 'skipped')

    def test_bounds_cannot_use_different_reference_editions(self):
        data = reference_dataset()
        data.metrics['参考.税负率上限'].reference['version'] = '另一版本'
        self.assertEqual(engine.evaluate(RULES['R-003'], data).status, 'skipped')

    def test_workbook_binds_each_numeric_origin_not_only_declared_scope(self):
        for names in [('参考.税负率下限',), ('参考.税负率上限',), reference_scope.SPECS['税负率'][:2]]:
            data = reference_dataset()
            for name in names:
                data.metrics[name].source = '另一份数值基准'
            loaded = loader.load_bytes(reference_workbook(data))
            finding = engine.evaluate(RULES['R-003'], loaded)
            self.assertEqual(finding.status, 'skipped')
            self.assertIn('来源', finding.skip_reason)
            for name in names:
                self.assertEqual(loaded.metrics[name].reference['metric_source'], '另一份数值基准')

    def test_old_scope_without_numeric_origin_is_not_inferred(self):
        data = reference_dataset()
        for name in reference_scope.SPECS['税负率'][:2]:
            data.metrics[name].reference.pop('metric_source')
        self.assertEqual(engine.evaluate(RULES['R-003'], data).status, 'skipped')

    def test_incomplete_workbook_reference_keeps_other_inputs(self):
        for key, value in [('publisher', ''), ('version', ''), ('start', ''), ('end', '2026-02-30')]:
            data = reference_dataset()
            for name in reference_scope.SPECS['税负率'][:2]:
                data.metrics[name].reference[key] = value
            loaded = loader.load_bytes(reference_workbook(data))
            self.assertEqual(loaded.metrics['增值税.销售额'].value, Decimal(1))
            self.assertEqual(engine.evaluate(RULES['R-003'], loaded).status, 'skipped')
            self.assertIn('参考范围待核对', loaded.metrics['参考.税负率下限'].detail)

    def test_changed_rule_formula_cannot_reuse_a_different_ratio_reference(self):
        rule = deepcopy(RULES['R-003'])
        rule.logic['numerator'], rule.logic['denominator'] = rule.logic['denominator'], rule.logic['numerator']
        data = reference_dataset()
        self.assertEqual(engine.evaluate(rule, data).status, 'skipped')
        self.assertTrue(engine.readiness(rule, data))

    def test_snapshot_reference_is_not_a_mutable_alias(self):
        data = reference_dataset()
        snapshot = serialize_dataset(data)
        data.metrics['参考.税负率下限'].reference['region'] = '另一省'
        self.assertEqual(snapshot['metrics']['参考.税负率下限']['reference']['region'], '江苏省')
        restored = deserialize_dataset(snapshot)
        restored.metrics['参考.税负率下限'].reference['region'] = '其他省'
        self.assertEqual(snapshot['metrics']['参考.税负率下限']['reference']['region'], '江苏省')

    def test_simulation_is_server_context_not_taxpayer_id(self):
        data = reference_dataset()
        data.company.taxpayer_id = 'TEST-GEN-SPOOF'
        for name in reference_scope.SPECS['税负率'][:2]:
            data.metrics[name].reference['nature'] = '仿真示意'
        self.assertEqual(engine.evaluate(RULES['R-003'], data).status, 'skipped')
        data.simulation = True
        self.assertEqual(engine.evaluate(RULES['R-003'], data).status, 'pass')
        self.assertTrue(deserialize_dataset(serialize_dataset(data)).simulation)

    def test_workbook_parse_and_numeric_correction_preserve_scope(self):
        raw = reference_workbook(reference_dataset())
        data = loader.load_bytes(raw)
        self.assertEqual(engine.evaluate(RULES['R-003'], data).status, 'pass')
        doc = materials.preview([('参考.xlsx', raw)], set(RULES['R-003'].inputs), capture_standard=True)[0]
        changed, _ = material_review.apply(doc, {'补充指标!B2': '.025'})
        restored = materials.build_dataset([changed], {changed['id']: {}}, {}, set(RULES['R-003'].inputs))
        self.assertEqual(engine.evaluate(RULES['R-003'], restored).status, 'pass')
        self.assertEqual(restored.company.region, '江苏省')
        self.assertIn('用户修正', restored.metrics['参考.税负率下限'].source)

    def test_enterprise_analysis_explains_mismatch_and_source_conflict(self):
        raw = reference_workbook(reference_dataset())
        docs = materials.preview([('参考.xlsx', raw)], set(RULES['R-003'].inputs))
        company = {**asdict(reference_dataset().company), 'period_start': '2026-01-01', 'period_end': '2026-12-31'}
        result = enterprise_analysis.analyze(docs, company, {}, [RULES['R-003']])
        self.assertIsNotNone(result['dataset'])
        company['region'] = '另一省'
        result = enterprise_analysis.analyze(docs, company, {}, [RULES['R-003']])
        self.assertFalse(result['can_confirm'])
        self.assertTrue(any(e['code'] == 'reference_company_scope_conflict' for e in result['feedback']['blocking']))

    def test_old_generator_digest_remains_readable_with_new_default_fields(self):
        import hashlib
        import json
        data = Dataset(Company('旧仿真', 'TEST', '仿真', '2026'), [], {}, {'x': Metric('x', Decimal(1), '原文')})
        rule = replace(RULES['R-009'], legal_context=None, legal_origin=None)
        company = asdict(data.company)
        for key in config.COMPANY_SCOPE_FIELDS.values():
            company.pop(key)
        metric = asdict(data.metrics['x'])
        metric.pop('period'); metric.pop('reference')
        definition = asdict(rule)
        definition.pop('legal_context'); definition.pop('legal_origin')
        old = dict(company=company, accounts=[], declarations={}, metrics={'x': metric}, rules=[definition])
        digest = hashlib.sha256(json.dumps(old, ensure_ascii=False, sort_keys=True, separators=(',', ':'), default=str).encode()).hexdigest()
        self.assertEqual(exercise_generator.case_digest(data, [rule]), digest)

    def test_benchmark_scope_checks_all_active_documents_and_exclusion(self):
        data = reference_dataset()
        other = deepcopy(data)
        other.company.industry = '另一行业'
        other.metrics = {name: metric for name, metric in other.metrics.items() if not name.startswith('参考.')}
        docs = materials.preview([('参考.xlsx', reference_workbook(data)),
                                  ('普通材料.xlsx', reference_workbook(other))], set(RULES['R-003'].inputs))
        company = {**asdict(data.company), 'period_start': '2026-01-01', 'period_end': '2026-12-31'}
        result = enterprise_analysis.analyze(docs, company, {}, [RULES['R-003']])
        conflicts = [e for e in result['feedback']['blocking'] if e['code'] == 'reference_company_scope_conflict']
        self.assertEqual([e['file_id'] for e in conflicts], [docs[1]['id']])
        self.assertFalse(result['can_confirm'])
        company['industry'] = '用户行业说明'
        result = enterprise_analysis.analyze(docs, company, {docs[0]['id']: {'purpose': 'excluded'}}, [RULES['R-003']])
        self.assertFalse(any(e['code'] == 'reference_company_scope_conflict' for e in result['feedback']['blocking']))
        self.assertTrue(result['can_confirm'])


class EnterpriseScopeAPI(EnterpriseMaterialFixture):
    def test_incomplete_reference_only_limits_comparison_not_confirmation(self):
        data = reference_dataset()
        for name in reference_scope.SPECS['税负率'][:2]:
            data.metrics[name].reference['publisher'] = ''
        for name, value in [('资产总额', 100), ('负债总额', 60), ('所有者权益', 40)]:
            key = '资产负债表.' + name
            data.metrics[key] = Metric(key, Decimal(value), '构造原文')
        item = self.upload(reference_workbook(data))
        self.assertTrue(item['analysis']['can_confirm'], item['analysis'])
        checks = {c['rule_id']: c for c in item['analysis']['checks']}
        self.assertFalse(checks['R-003']['ready'])
        self.assertTrue(checks['R-009']['ready'], checks['R-009'])
        response = self.confirm(item)
        self.assertEqual(response.status_code, 200, response.text)
        findings = {f.rule.id: f for f in self.store.get_audit(response.json()['audit_id'])['findings']}
        self.assertEqual(findings['R-003'].status, 'skipped')
        self.assertEqual(findings['R-009'].status, 'pass')

    def test_reference_positive_confirmation_freezes_scope_and_report(self):
        raw = reference_workbook(reference_dataset(period='2022年度'))
        item = self.upload(raw)
        check = next(c for c in item['analysis']['checks'] if c['rule_id'] == 'R-003')
        self.assertTrue(check['ready'], check)
        response = self.confirm(item)
        self.assertEqual(response.status_code, 200, response.text)
        result = response.json()
        audit = self.store.get_audit(result['audit_id'])
        finding = next(f for f in audit['findings'] if f.rule.id == 'R-003')
        self.assertEqual(finding.status, 'pass')
        self.assertEqual(audit['dataset'].company.region, '江苏省')
        self.assertFalse(audit['dataset'].simulation)
        self.assertIn('2022', finding.rule.legal_context['period'])
        html = self.client.get('/api/report/' + result['audit_id'] + '/html')
        self.assertEqual(html.status_code, 200, html.text[:200])
        self.assertIn('构造测试发布者', html.text)
        self.assertIn('未自动核定权威性', html.text)
        changed = self.client.post(self.url(item, '/analyze'), json={
            'expected_revision': self.client.get(self.url(item)).json()['revision'],
            'company': {**item['company'], 'region': '另一省'}})
        self.assertEqual(changed.status_code, 200, changed.text)
        self.assertFalse(changed.json()['analysis']['can_confirm'])
        self.assertEqual(self.client.get('/api/report/' + result['audit_id'] + '/html').content, html.content)

    def test_simulated_reference_limits_only_affected_rule_in_enterprise(self):
        data = reference_dataset()
        for key in reference_scope.SPECS['税负率'][:2]:
            data.metrics[key].reference['nature'] = '仿真示意'
        item = self.upload(reference_workbook(data))
        self.assertTrue(item['analysis']['can_confirm'])
        check = next(c for c in item['analysis']['checks'] if c['rule_id'] == 'R-003')
        self.assertFalse(check['ready'])
        self.assertIn('仿真示意', ' '.join(check['reasons']))
        response = self.confirm(item)
        self.assertEqual(response.status_code, 200, response.text)
        finding = next(f for f in self.store.get_audit(response.json()['audit_id'])['findings'] if f.rule.id == 'R-003')
        self.assertEqual(finding.status, 'skipped')


if __name__ == '__main__':
    unittest.main()
