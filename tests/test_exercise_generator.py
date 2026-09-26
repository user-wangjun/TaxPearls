from copy import deepcopy
from dataclasses import replace
from decimal import Decimal
from pathlib import Path
import os
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch
from io import BytesIO
from zipfile import ZipFile
from openpyxl import load_workbook

from fastapi.testclient import TestClient

from src import engine, training
from src import exercise_materials, loader
from src.exercise_generator import BASE, generate, materialize
from src.models import Company
from webapp import app as module
from webapp.storage import Store

ROOT = Path(__file__).resolve().parents[1]


class ExerciseGeneratorTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.rules = engine.load_rules(ROOT/"rules")

    def test_neutral_ledger_all_twenty_four_rules_pass(self):
        dataset = materialize({key:Decimal(str(value)) for key,value in BASE.items()},
                              Company("仿真", "TEST", "仿真", "2026"), self.rules)
        self.assertEqual({f.status for f in engine.run(self.rules, dataset)}, {"pass"})

    def test_all_rules_three_levels_multiple_seeds_have_real_engine_answers(self):
        for rule in self.rules:
            for level in ("near", "normal", "obvious"):
                for seed in (0, 77, 2147483647):
                    with self.subTest(rule=rule.id,level=level,seed=seed):
                        result = generate(self.rules,rule.id,seed,level)
                        answers = sorted(f.rule.id for f in engine.run(self.rules,result.dataset) if f.hit)
                        self.assertIn(rule.id,answers)
                        self.assertEqual(answers,result.metadata["standard_answer"])
                        self.assertNotIn("skipped",{f.status for f in result.findings})
                        self.assertEqual(training.score_submission(result.findings,answers)["score"],100)
                        for key,metric in result.dataset.metrics.items():
                            self.assertTrue(metric.source)
                            if key.endswith("人数"):
                                self.assertEqual(metric.value,metric.value.to_integral_value())
                                self.assertGreaterEqual(metric.value,0)

    def test_deterministic_variants_and_does_not_mutate_rules(self):
        original = deepcopy(self.rules)
        a,b = generate(self.rules,"R-020",11),generate(self.rules,"R-020",11)
        self.assertEqual(a,b)
        self.assertNotEqual(a.dataset,generate(self.rules,"R-020",12).dataset)
        self.assertEqual(self.rules,original)
        self.assertNotEqual(a.metadata["standard_answer"],["R-020"])

    def test_reads_updated_threshold_and_frozen_inputs_not_case_fixture(self):
        changed = [replace(rule,logic={**rule.logic,"threshold":.8},version="3.0") if rule.id=="R-001" else rule
                   for rule in self.rules]
        old,new = generate(self.rules,"R-001",12),generate(changed,"R-001",12)
        self.assertNotEqual(old.dataset.values(),new.dataset.values())
        finding = next(f for f in new.findings if f.rule.id=="R-001")
        self.assertGreater(finding.measured,.8)
        self.assertEqual(finding.rule.version,"3.0")

    def test_reference_inputs_are_fixed_not_moved_to_force_a_hit(self):
        result = generate(self.rules,"R-021",3)
        self.assertNotIn("参考.毛利率下限",result.metadata["changed_metrics"])
        self.assertNotIn("参考.毛利率上限",result.metadata["changed_metrics"])
        self.assertEqual(result.dataset.get("参考.毛利率下限"),Decimal(".2"))
        other=generate(self.rules,"R-024",3,year=2030)
        self.assertIn("2028-01-01 至 2028-12-31",other.dataset.metrics["年度.前年净利润"].detail)
        self.assertIn("2029-01-01 至 2029-12-31",other.dataset.metrics["历史.上年同期收入"].detail)

    def test_validation_unknown_missing_mapping_and_unsatisfiable_rule(self):
        for seed in (-1,2147483648,True,1.5):
            with self.assertRaises(ValueError):generate(self.rules,"R-001",seed)
        with self.assertRaises(ValueError):generate(self.rules,"G-001")
        with self.assertRaises(ValueError):generate(self.rules,"R-001",level="unknown")
        with self.assertRaises(ValueError):generate(self.rules,"R-001",year=1999)
        base = self.rules[0]
        impossible = replace(base,inputs={"营业收入":"仿真"},evidence=["营业收入"],
                             logic={"type":"amount_mismatch","left":"营业收入","right":"营业收入","tolerance":0})
        with self.assertRaisesRegex(ValueError,"无法反解"):generate([impossible],base.id)
        unmapped = replace(base,inputs={"新指标":"仿真"},evidence=["新指标"],
                           logic={"type":"amount_mismatch","left":"新指标","right":0,"tolerance":0})
        with self.assertRaisesRegex(ValueError,"来源映射"):generate([unmapped],base.id)

    def test_export_all_rules_levels_preserves_inputs_answers_and_literal_sources(self):
        for rule in self.rules:
            for level in ("near", "normal", "obvious"):
                with self.subTest(rule=rule.id, level=level):
                    case = generate(self.rules, rule.id, 2147483647, level, 2030)
                    content = exercise_materials.export(case.dataset)
                    restored = loader.load_bytes(content)
                    self.assertEqual(restored.values(), case.dataset.values())
                    self.assertEqual(sorted(f.rule.id for f in engine.run(self.rules, restored) if f.hit),
                                     case.metadata["standard_answer"])
                    wb = load_workbook(BytesIO(content))
                    self.assertEqual(len(wb.sheetnames), 8)
                    for ws in wb:
                        self.assertEqual(ws.sheet_state, "visible")
                        for row in ws:
                            for cell in row:
                                self.assertNotEqual(cell.data_type, "f")
                    periods = {row[0]: row[3] for row in list(wb["历史指标"].values)[1:]}
                    self.assertEqual(periods["年度.前年净利润"], "2028")
                    self.assertEqual(periods["历史.上年同期收入"], "2029")
                    wb.close()
                    with ZipFile(BytesIO(content)) as archive:
                        all_xml = b"".join(archive.read(n) for n in archive.namelist() if n.endswith(".xml"))
                    for forbidden in (b"standard_answer", b"requested_rule_id", b"changed_metrics", b"case_sha256"):
                        self.assertNotIn(forbidden, all_xml)
        case = generate(self.rules, "R-020")
        case.dataset.metrics["存货.采购加权税率"].source = "=HYPERLINK(\"https://example.invalid\")"
        content = exercise_materials.export(case.dataset)
        self.assertEqual(exercise_materials.export(case.dataset), content)

    def test_export_refuses_numeric_precision_loss_and_real_company(self):
        case = generate(self.rules, "R-020")
        case.dataset.metrics["存货.采购加权税率"].value = Decimal("0.123456789012345678901")
        with self.assertRaisesRegex(ValueError, "数值不一致"):
            exercise_materials.export(case.dataset)
        case.dataset.company.taxpayer_id = "REAL-COMPANY"
        with self.assertRaisesRegex(ValueError, "纯仿真"):
            exercise_materials.export(case.dataset)


class ExerciseWebTests(unittest.TestCase):
    def setUp(self):
        self.env=patch.dict(os.environ,{"TAXPEARLS_NOTIFICATION_EMAIL_ENABLED":"0","TAXPEARLS_AI_ENABLED":"0"})
        self.env.start();self.tmp=TemporaryDirectory();self.old_store=module.store
        self.store=Store(Path(self.tmp.name)/"exercise.db");module.store=self.store
        self.password="Exercise-test-2026!"
        self.teacher=self.store.create_user("exercise-teacher",self.password,"出题教师","teacher","school-a")
        self.student=self.store.create_user("exercise-student",self.password,"答题学生","student","school-a")
        self.client=TestClient(module.app);self.client.__enter__();self.login(self.teacher)

    def tearDown(self):
        self.client.__exit__(None,None,None);module.store=self.old_store;self.tmp.cleanup();self.env.stop()

    def login(self,user):
        _,token=self.store.authenticate(user["username"],self.password)
        self.client.cookies.clear();self.client.cookies.set(module.COOKIE_NAME,token)

    def generate(self,rule="R-020",**options):
        return self.client.post("/api/exercises",json={"rule_id":rule,"expected_version":"2.0","seed":17,**options})

    def test_catalog_and_all_twenty_four_generated_cases_atomic_report_archive(self):
        catalog=self.client.get("/api/exercises/rules")
        self.assertEqual(catalog.status_code,200);self.assertEqual(len(catalog.json()["rules"]),24)
        self.assertEqual(catalog.headers["Cache-Control"],"private, no-store")
        for rule in catalog.json()["rules"]:
            with self.subTest(rule=rule["id"]):
                response=self.generate(rule["id"])
                self.assertEqual(response.status_code,200,response.text)
                result=response.json();audit_id=result["audit_id"]
                self.assertIn(rule["id"],result["metadata"]["standard_answer"])
                self.assertEqual(len(self.store.report_versions(audit_id)),1)
                self.assertEqual(self.client.get("/api/exercises/"+audit_id).json(),result)

    def test_teacher_publish_student_can_see_all_materials_but_no_answers_before_submit(self):
        result=self.generate().json();audit_id=result["audit_id"]
        assignment=self.client.post("/api/assignments",json={"title":"生成账套实训","audit_id":audit_id,
                                                             "target_student_id":self.student["id"]})
        self.assertEqual(assignment.status_code,200,assignment.text);assignment_id=assignment.json()["id"]
        self.login(self.student)
        view=self.client.get("/api/assignments/"+assignment_id)
        self.assertEqual(view.status_code,200)
        self.assertGreater(len(view.json()["metrics"]),50)
        option_ids = [rule["id"] for rule in view.json()["rules"]]
        self.assertEqual(option_ids, sorted(option_ids))
        for forbidden in ("standard_answer","requested_rule_id","changed_metrics","seed","case_sha256","status"):
            self.assertNotIn(forbidden,view.text)
        self.assertEqual(self.client.get("/api/exercises/"+audit_id).status_code,403)
        response=self.client.post(f"/api/assignments/{assignment_id}/submit",
                                  json={"selected_rule_ids":result["metadata"]["standard_answer"]})
        self.assertEqual(response.status_code,200,response.text);self.assertEqual(response.json()["score"],100)
        wrong=self.client.post(f"/api/assignments/{assignment_id}/submit",json={"selected_rule_ids":["R-001"]})
        self.assertTrue(wrong.json()["missed"]);self.assertTrue(wrong.json()["false_positives"])

    def test_permissions_and_cross_institution_do_not_expose_answers(self):
        audit_id=self.generate().json()["audit_id"]
        foreign=self.store.create_user("foreignteacher",self.password,"外部教师","teacher","school-b")
        self.login(foreign);self.assertEqual(self.client.get("/api/exercises/"+audit_id).status_code,404)
        for role in ("student","accountant","org_admin","platform_admin"):
            user=self.store.create_user("denied-"+role,self.password,role,role,"school-a");self.login(user)
            self.assertEqual(self.client.get("/api/exercises/rules").status_code,403)
            self.assertEqual(self.generate().status_code,403)
            self.assertEqual(self.client.get("/api/exercises/"+audit_id).status_code,403)
        self.client.cookies.clear();self.assertEqual(self.generate().status_code,401)

    def test_validation_disabled_stale_and_unsatisfiable_refuse_without_saved_audit(self):
        for body in ({"seed":True},{"seed":-1},{"level":"bad"},{"year":1999},{"org_id":"evil"}):
            self.assertEqual(self.generate(**body).status_code,422)
        self.assertEqual(self.generate(expected_version="1.0").status_code,409)
        self.assertEqual(self.generate("R-999").status_code,422)
        with patch.object(self.store,"enabled_rule_ids",return_value={"R-001"}):
            self.assertEqual(self.generate().status_code,422)
        rules=engine.load_rules(ROOT/"rules");selected=next(r for r in rules if r.id=="R-020")
        impossible=replace(selected,logic={"type":"amount_mismatch","left":"利润表.营业收入",
                           "right":"利润表.营业收入","tolerance":0},
                           inputs={"利润表.营业收入":"仿真"},evidence=["利润表.营业收入"])
        with patch.object(module,"_audit_rules",return_value=[impossible]):
            self.assertEqual(self.generate().status_code,422)
        self.assertEqual(self.store.list_audits(self.teacher),[])

    def test_frozen_case_survives_rule_updates_restart_backup_and_corruption_refused(self):
        from scripts.ops_db import create_backup,restore_backup
        result=self.generate().json();audit_id=result["audit_id"]
        backup=Path(self.tmp.name)/"backup.db";restored=Path(self.tmp.name)/"restored.db"
        create_backup(self.store.path,backup);restore_backup(backup,restored)
        module.store=Store(restored)
        with patch.object(module,"_audit_rules",side_effect=AssertionError("no new rules during frozen read")), \
             patch("src.exercise_generator.generate",side_effect=AssertionError("no regenerate")):
            self.assertEqual(self.client.get("/api/exercises/"+audit_id).json(),result)
        with module.store.connect() as db:
            db.execute("UPDATE generated_exercises SET metadata_json='{}' WHERE audit_id=?",(audit_id,))
        self.assertEqual(self.client.get("/api/exercises/"+audit_id).status_code,409)

    def test_generation_record_failure_rolls_back_audit_report_and_protection(self):
        from src import render
        from webapp.report_archive import build_snapshot
        case=generate(engine.load_rules(ROOT/"rules"),"R-001")
        audit_id="rollback-exercise";when="2026-09-26 12:00:00"
        snapshot=build_snapshot({"id":audit_id,"org_id":"school-a","audited_at":when,
                                 "dataset":case.dataset,"findings":case.findings},module._org_branding("school-a"))
        bad={**case.metadata,"standard_answer":[]}
        with self.assertRaises(ValueError):
            self.store.save_audit(audit_id,self.teacher,None,case.dataset,case.findings,
                                 render.build_view_model(case.dataset,case.findings)["summary"],when,snapshot,bad)
        self.assertIsNone(self.store.get_audit(audit_id));self.assertEqual(self.store.report_versions(audit_id),[])
        with self.store.connect() as db:
            self.assertEqual(db.execute("SELECT COUNT(*) FROM generated_exercises").fetchone()[0],0)
            self.assertEqual(db.execute("SELECT COUNT(*) FROM report_protections").fetchone()[0],0)

    def test_material_download_permissions_frozen_restart_backup_and_revocation(self):
        result = self.generate().json(); audit_id = result["audit_id"]
        url = f"/api/exercises/{audit_id}/materials"
        download = self.client.get(url)
        self.assertEqual(download.status_code, 200, download.text[:200] if download.status_code != 200 else "")
        content = download.content
        self.assertEqual(download.headers["Cache-Control"], "private, no-store")
        restored = loader.load_bytes(content)
        self.assertEqual(restored.values(), self.store.get_audit(audit_id)["dataset"].values())
        assignment = self.client.post("/api/assignments", json={"title":"下载仿真材料", "audit_id":audit_id,
            "target_student_id":self.student["id"], "published":False}).json()
        assignment_url = f"/api/assignments/{assignment['id']}/materials"
        self.login(self.student)
        self.assertEqual(self.client.get(url).status_code, 403)
        self.assertEqual(self.client.get(assignment_url).status_code, 404)
        with self.store.connect() as db:
            db.execute("UPDATE assignments SET published=1 WHERE id=?", (assignment["id"],))
        self.assertEqual(self.client.get(assignment_url).content, content)
        other = self.store.create_user("material-other", self.password, "其他学生", "student", "school-a")
        self.login(other); self.assertEqual(self.client.get(assignment_url).status_code, 404)
        foreign = self.store.create_user("material-foreign", self.password, "外部教师", "teacher", "school-b")
        self.login(foreign); self.assertEqual(self.client.get(url).status_code, 404)
        self.assertEqual(self.client.get(assignment_url).status_code, 404)
        self.client.cookies.clear(); self.assertEqual(self.client.get(assignment_url).status_code, 401)
        self.login(self.teacher)
        from scripts.ops_db import create_backup, restore_backup
        backup=Path(self.tmp.name)/"materials-backup.db"; restored_db=Path(self.tmp.name)/"materials-restored.db"
        create_backup(self.store.path, backup); restore_backup(backup, restored_db)
        module.store = Store(restored_db)
        self.store = module.store
        with patch.object(module, "_audit_rules", side_effect=AssertionError("must not use new rules")), \
             patch("src.exercise_generator.generate", side_effect=AssertionError("must not regenerate")):
            self.assertEqual(self.client.get(url).content, content)
            self.assertEqual(self.client.get(url).headers["X-Material-SHA256"], download.headers["X-Material-SHA256"])
        self.login(self.student)
        with module.store.connect() as db:
            db.execute("UPDATE assignments SET published=0 WHERE id=?", (assignment["id"],))
        self.assertEqual(self.client.get(assignment_url).status_code, 404)
        self.login(self.teacher)
        with module.store.connect() as db:
            db.execute("UPDATE generated_exercises SET metadata_json='{}' WHERE audit_id=?", (audit_id,))
        self.assertEqual(self.client.get(url).status_code, 409)


if __name__ == "__main__":
    unittest.main()
