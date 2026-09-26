from __future__ import annotations

from copy import deepcopy
from decimal import Decimal
from io import BytesIO
from pathlib import Path
import tempfile
import unittest

from fastapi.testclient import TestClient
from openpyxl import Workbook

from src import config, engine, materials
from webapp import app as app_module
from webapp.storage import Store


ROOT = Path(__file__).resolve().parent.parent
COMPANY = {"name": "银行适配测试企业", "taxpayer_id": "BANK-TEST-001", "industry": "批发业", "period": "2026H1"}
KEYS = {key for rule in engine.load_rules(ROOT / "rules") for key in rule.inputs}


def selections(docs):
    return {
        doc["id"]: {
            "id": doc["id"],
            "company": deepcopy(doc["company"]),
            "rows": deepcopy(doc["rows"]),
            "reviewed": True,
        }
        for doc in docs
    }


def source_workbook(rows, headers=None, title="账户交易明细", preface=False):
    headers = headers or [
        "唯一编号", "交易日期", "借方发生额", "贷方发生额", "凭证号码",
        "摘要", "余额", "对方户名", "对方账号", "币别",
    ]
    workbook = Workbook()
    ws = workbook.active
    ws.title = title
    if preface:
        ws.append(["中国建设银行企业账户交易明细（仿真）"])
        ws.append([])
    ws.append(headers)
    for row in rows:
        ws.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def adjustment_workbook(rows, revenue=1_030_000):
    workbook = Workbook()
    ws = workbook.active
    ws.title = config.SHEET_COMPANY
    ws.append(["项目", "内容"])
    for label, key in zip(config.COMPANY_FIELDS, COMPANY):
        ws.append([label, COMPANY[key]])
    ws = workbook.create_sheet(config.SHEET_ACCOUNTS)
    ws.append(config.COL_ACCOUNTS)
    ws.append(["6001", "主营业务收入", 0, 0, revenue, 0])
    ws.append(["6051", "其他业务收入", 0, 0, 0, 0])
    ws = workbook.create_sheet(config.SHEET_BANK_ADJUSTMENTS)
    ws.append(config.COL_BANK_ADJUSTMENTS)
    for row in rows:
        ws.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


SOURCE_ROWS = [
    ["T-001", "2026-02-01", None, "1,130,000.00", "P-1", "销售回款", 1_500_000, "客户甲", "622200001111", "人民币"],
    ["T-002", "2026-02-02", None, 200_000, "P-2", "内部划转", 1_700_000, "本公司另一账户", "622200002222", "CNY"],
    ["T-003", "2026-02-03", 100_000, None, "P-3", "支付货款", 1_600_000, "供应商乙", "622200003333", "RMB"],
]

ADJUSTMENT_ROWS = [
    ["A-001", "T-001", "经营回款", "2026H1", 1_000_000, None, None, "WP-BANK-001", "已复核", "含税回款按发票和合同拆分为不含税收入"],
    ["A-002", "T-002", "内部转账", "2026H1", 0, None, None, "WP-BANK-002", "已复核", "同名账户内部划转，不属于收入"],
    ["A-003", None, "现金收入补计", "2026H1", None, 50_000, "加计", "WP-BANK-003", "已复核", "本期现金销售补计"],
    ["A-004", None, "退款折让扣减", "2026H1", None, 20_000, "扣减", "WP-BANK-004", "已复核", "本期销售退款扣减"],
]


class BankSourceParsing(unittest.TestCase):
    def test_ccb_profile_keeps_raw_receipts_separate_from_revenue(self):
        doc = materials.preview([("建行导出.xlsx", source_workbook(SOURCE_ROWS, preface=True))], KEYS)[0]
        self.assertFalse(doc["error"], doc)
        self.assertEqual(len(doc["bank_transactions"]), 3)
        self.assertEqual(doc["bank_transactions"][0]["transaction_id"], "T-001")
        self.assertTrue(doc["bank_transactions"][0]["explicit_id"])
        self.assertNotIn("622200001111", str(doc))
        data = materials.build_dataset([doc], selections([doc]), COMPANY, KEYS)
        self.assertEqual(data.get("银行.收入流水合计"), 1_330_000)
        self.assertEqual(data.get("银行.支出流水合计"), 100_000)
        self.assertIsNone(data.get("银行.调节后不含税收入"))
        self.assertIn("不得直接作为营业收入", data.detail_of("银行.收入流水合计"))

    def test_reviewed_bridge_generates_accrual_tax_exclusive_metric(self):
        docs = materials.preview([
            ("建行导出.xlsx", source_workbook(SOURCE_ROWS)),
            ("收入调节底稿.xlsx", adjustment_workbook(ADJUSTMENT_ROWS)),
        ], KEYS)
        self.assertTrue(all(not doc["error"] for doc in docs), docs)
        data = materials.build_dataset(docs, selections(docs), {}, KEYS)
        self.assertEqual(data.get("银行.经营回款已复核不含税额"), 1_000_000)
        self.assertEqual(data.get("银行.权责口径调节额"), 30_000)
        self.assertEqual(data.get("银行.调节后不含税收入"), 1_030_000)
        self.assertIn("原始收入流水 1,330,000.00 仅作对照", data.detail_of("银行.调节后不含税收入"))
        self.assertIn("WP-BANK-004", data.detail_of("银行.调节后不含税收入"))
        finding = next(item for item in engine.run(engine.load_rules(ROOT / "rules"), data) if item.rule.id == "R-015")
        self.assertEqual(finding.status, "pass")

    def test_single_amount_direction_profile_and_bounds(self):
        headers = ["交易流水号", "交易时间", "交易金额", "交易方向", "交易摘要", "币种"]
        rows = [["S-001", "2026/03/01 09:30:00", "88,000.50", "转入", "收款", "156"]]
        doc = materials.preview([("单金额导出.xlsx", source_workbook(rows, headers))], KEYS)[0]
        self.assertFalse(doc["error"], doc)
        data = materials.build_dataset([doc], selections([doc]), COMPANY, KEYS)
        self.assertEqual(data.get("银行.收入流水合计"), Decimal("88000.50"))

        foreign = deepcopy(rows)
        foreign[0][-1] = "USD"
        self.assertIn("仅支持人民币", materials.preview([("外币.xlsx", source_workbook(foreign, headers))], KEYS)[0]["error"])
        outside = deepcopy(rows)
        outside[0][1] = "2025-12-31"
        doc = materials.preview([("跨期.xlsx", source_workbook(outside, headers))], KEYS)[0]
        with self.assertRaisesRegex(materials.InputError, "不在核对期间"):
            materials.build_dataset([doc], selections([doc]), COMPANY, KEYS)

    def test_duplicate_and_incomplete_adjustments_are_rejected(self):
        duplicates = materials.preview([
            ("一.xlsx", source_workbook(SOURCE_ROWS[:1])),
            ("二.xlsx", source_workbook(SOURCE_ROWS[:1])),
        ], KEYS)
        with self.assertRaisesRegex(materials.InputError, "重复"):
            materials.build_dataset(duplicates, selections(duplicates), COMPANY, KEYS)

        source = materials.preview([("流水.xlsx", source_workbook(SOURCE_ROWS))], KEYS)[0]
        missing_rows = [row for row in ADJUSTMENT_ROWS if row[1] != "T-002"]
        adjustment = materials.preview([("缺分类.xlsx", adjustment_workbook(missing_rows))], KEYS)[0]
        with self.assertRaisesRegex(materials.InputError, "未在银行调节底稿分类"):
            materials.build_dataset([source, adjustment], selections([source, adjustment]), {}, KEYS)

        pending = deepcopy(ADJUSTMENT_ROWS)
        pending[0][8] = "待复核"
        adjustment = materials.preview([("待复核.xlsx", adjustment_workbook(pending))], KEYS)[0]
        with self.assertRaisesRegex(materials.InputError, "尚未复核"):
            materials.build_dataset([source, adjustment], selections([source, adjustment]), {}, KEYS)

        false_revenue = deepcopy(ADJUSTMENT_ROWS)
        false_revenue[1][4] = 1
        adjustment = materials.preview([("误认收入.xlsx", adjustment_workbook(false_revenue))], KEYS)[0]
        with self.assertRaisesRegex(materials.InputError, "必须为 0"):
            materials.build_dataset([source, adjustment], selections([source, adjustment]), {}, KEYS)


class BankWebFlow(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.db_path = Path(self.temp.name) / "bank.db"
        self.old_store = app_module.store
        app_module.store = Store(self.db_path)
        app_module.store.create_user("bank-admin", "bank-test-2026", "银行测试", "platform_admin", "default")
        self.client = TestClient(app_module.app)
        self.client.post("/api/login", json={"username": "bank-admin", "password": "bank-test-2026"})

    def tearDown(self):
        self.client.close()
        app_module.store = self.old_store
        self.temp.cleanup()

    def test_source_and_adjustment_commit_with_persisted_evidence(self):
        response = self.client.post(
            "/api/materials/preview",
            files=[
                ("files", ("建行导出.xlsx", source_workbook(SOURCE_ROWS))),
                ("files", ("收入调节底稿.xlsx", adjustment_workbook(ADJUSTMENT_ROWS))),
            ],
        )
        self.assertEqual(response.status_code, 200, response.text)
        draft = response.json()
        self.assertEqual(len(draft["documents"][0]["bank_transactions"]), 3)
        self.assertEqual(len(draft["documents"][1]["bank_adjustments"]), 4)
        payload = {
            "token": draft["token"], "mode": "merge", "same_scope": True,
            "company": COMPANY, "selections": list(selections(draft["documents"]).values()),
        }
        audited = self.client.post("/api/materials/audit", json=payload)
        self.assertEqual(audited.status_code, 200, audited.text)
        body = audited.json()
        self.assertFalse(body["errors"], body)
        audit = body["results"][0]["audit"]
        r015 = next(item for item in audit["findings"] if item["id"] == "R-015")
        self.assertEqual(r015["status"], "pass")
        audit_id = audit["audit_id"]
        saved = Store(self.db_path).get_audit(audit_id)["dataset"]
        self.assertEqual(saved.get("银行.调节后不含税收入"), 1_030_000)
        self.assertIn("建行导出.xlsx", saved.source_of("银行.调节后不含税收入"))
        self.assertIn("收入调节底稿.xlsx", saved.source_of("银行.调节后不含税收入"))


if __name__ == "__main__":
    unittest.main()
