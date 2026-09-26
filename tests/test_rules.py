"""规则回归、失效边界、真实 Excel 输入和证据报告的集成测试。"""
from dataclasses import replace
from decimal import Decimal
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook
import yaml

from scripts.check_rules import (ROOT, case_values, check_finding, dataset_for, read_cases, verify)
from src import config, engine, loader, render


RULES = {r.id: r for r in engine.load_rules(ROOT / "rules")}
GROUPS = {g["rule_id"]: g for g in read_cases()}


def workbook(values, rule):
    """输入适配器测试夹具；不交付或覆盖用户工作簿。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "企业信息"
    ws.append(["项目", "内容"])
    for row in zip(config.COMPANY_FIELDS, ["仿真核对企业", "SYNTHETIC", "批发和零售业", "2026年度"]):
        ws.append(row)
    ws = wb.create_sheet("科目余额表")
    ws.append(config.COL_ACCOUNTS)
    for name, spec in config.ACCOUNT_MAP.items():
        for i, code in enumerate(spec["accounts"]):
            amount = values.get(name) if i == 0 else (0 if name in values else None)
            ws.append([code, name, 0, amount if spec["side"] == "debit" else 0,
                       amount if spec["side"] == "credit" else 0, 0])
    ws = wb.create_sheet("增值税申报")
    ws.append(["项目", "金额"])
    for name, item in config.DECLARATION_ITEMS.items():
        ws.append([item, values.get(name)])
    ws = wb.create_sheet("补充指标")
    ws.append(config.COL_SUPPLEMENT)
    reserved = set(config.ACCOUNT_MAP) | set(config.DECLARATION_ITEMS)
    for name, value in values.items():
        if name not in reserved:
            ws.append([name, value, f"仿真底稿—{name}", "2026年度", rule.inputs[name]])
    return wb


class RuleCases(unittest.TestCase):
    def test_all_96_annotated_cases(self):
        report = verify()
        self.assertEqual(report["rules"], 24)
        self.assertEqual(report["cases"], 96)
        self.assertEqual(report["failed"], 0, [r for r in report["results"] if not r["ok"]])
        self.assertEqual(report["statuses"], {"pass": 48, "hit": 24, "skipped": 24})

    def test_all_cases_through_excel(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "input.xlsx"
            for rid, group in GROUPS.items():
                for case in group["cases"]:
                    with self.subTest(rule=rid, case=case["name"]):
                        wb = workbook(case_values(group, case), RULES[rid])
                        wb.save(path)
                        wb.close()
                        finding = engine.evaluate(RULES[rid], loader.load(path))
                        self.assertEqual(check_finding(finding, case), [])

    def test_original_workbook(self):
        data = loader.load(ROOT / "samples" / "样例企业-审计材料.xlsx")
        results = {f.rule.id: f for f in engine.run(list(RULES.values()), data)}
        self.assertEqual(results["R-001"].status, "hit")
        self.assertEqual(results["R-005"].measured, 41080)
        self.assertEqual(results["R-006"].status, "pass")
        self.assertEqual(results["R-002"].status, "skipped")
        self.assertEqual(results["R-003"].status, "skipped")
        self.assertNotIn("账面.应纳税额", data.metrics)

    def test_categories(self):
        self.assertEqual({r.category for r in RULES.values()}, {"内部勾稽", "外部交叉", "指标偏离"})

    def test_each_missing_dependency_skips(self):
        for rid, rule in RULES.items():
            for metric in rule.inputs:
                with self.subTest(rule=rid, metric=metric):
                    data = dataset_for(rule, GROUPS[rid]["baseline"])
                    data.metrics.pop(metric)
                    finding = engine.evaluate(rule, data)
                    self.assertEqual(finding.status, "skipped")
                    self.assertIn(metric, finding.skip_reason)
                    self.assertIn(rule.inputs[metric], finding.skip_reason)

    def test_zero_and_negative_denominators(self):
        for rid, key in [("R-001", "增值税.销售额"), ("R-003", "增值税.销售额"),
                         ("R-020", "历史.上年同期收入"), ("R-024", "年度.上年末资产")]:
            for value in (0, -1):
                with self.subTest(rule=rid, value=value):
                    data = dataset_for(RULES[rid], {**GROUPS[rid]["baseline"], key: value})
                    self.assertEqual(engine.evaluate(RULES[rid], data).status, "skipped")

    def test_both_zero_deviation_is_not_hit_or_pass(self):
        data = dataset_for(RULES["R-001"], {"营业收入": 0, "增值税.销售额": 0})
        self.assertEqual(engine.evaluate(RULES["R-001"], data).status, "skipped")

    def test_bad_metric_values_and_empty_source(self):
        rule = RULES["R-001"]
        for value in (float("nan"), float("inf"), True, "not-a-number"):
            data = dataset_for(rule, GROUPS[rule.id]["baseline"])
            data.metrics["营业收入"].value = value
            self.assertEqual(engine.evaluate(rule, data).status, "skipped")
        data = dataset_for(rule, GROUPS[rule.id]["baseline"])
        data.metrics["营业收入"].source = ""
        self.assertEqual(engine.evaluate(rule, data).status, "skipped")

    def test_decimal_boundary(self):
        rule = replace(RULES["R-005"], logic={"type": "amount_mismatch", "left": "账面.销项税额", "right": "增值税.销项税额", "tolerance": .1})
        data = dataset_for(rule, {"账面.销项税额": .3, "增值税.销项税额": .2})
        self.assertEqual(engine.evaluate(rule, data).status, "pass")
        data.metrics["账面.销项税额"].value = Decimal("0.30000001")
        self.assertEqual(engine.evaluate(rule, data).status, "hit")

    def test_ratio_both_inclusive_bounds_and_outside(self):
        rule = RULES["R-003"]
        for amount, expected in [(10000, "pass"), (45000, "pass"), (9999.99, "hit"), (45000.01, "hit")]:
            data = dataset_for(rule, {**GROUPS[rule.id]["baseline"], "增值税.应纳税额": amount})
            self.assertEqual(engine.evaluate(rule, data).status, expected)

    def test_invalid_dynamic_reference(self):
        rule = RULES["R-003"]
        data = dataset_for(rule, {**GROUPS[rule.id]["baseline"], "参考.税负率下限": .1})
        self.assertEqual(engine.evaluate(rule, data).status, "skipped")

    def test_carry_forward_and_negative_net_vat(self):
        rule = RULES["R-002"]
        values = {**GROUPS[rule.id]["baseline"], "账面.进项税额": 200000, "增值税.一般计税应纳税额": 0}
        self.assertEqual(engine.evaluate(rule, dataset_for(rule, values)).status, "pass")
        values = {**GROUPS[rule.id]["baseline"], "账面.抵扣调整": -200000, "增值税.一般计税应纳税额": 130000}
        self.assertEqual(engine.evaluate(rule, dataset_for(rule, values)).status, "pass")

    def test_combination_does_not_short_circuit_missing_data(self):
        rule = RULES["R-024"]
        values = {**GROUPS[rule.id]["baseline"], "年度.本年净利润": 1}
        values.pop("年度.前年末资产")
        self.assertEqual(engine.evaluate(rule, dataset_for(rule, values)).status, "skipped")

    def test_evidence_and_sorting(self):
        rule = RULES["R-001"]
        data = dataset_for(rule, {**GROUPS[rule.id]["baseline"], "营业收入": 1200000})
        results = engine.run(list(RULES.values()), data)
        ranks = [{"hit": 0, "pass": 1, "skipped": 2}[f.status] for f in results]
        self.assertEqual(ranks, sorted(ranks))
        self.assertEqual(results[0].rule.id, rule.id)

    def test_report_contains_formula_scope_threshold_and_sources(self):
        rule = RULES["R-011"]
        group = GROUPS[rule.id]
        data = dataset_for(rule, case_values(group, group["cases"][1]))
        findings = engine.run([rule], data)
        with tempfile.TemporaryDirectory() as directory, patch.object(render, "OUTPUT_DIR", Path(directory)):
            html, path = render.render_html(data, findings)
            self.assertTrue(path.exists())
            for text in [rule.scope, rule.threshold_basis, rule.references[0], "权益.提取盈余公积", "人工仿真底稿"]:
                self.assertIn(text, html)


class RuleValidation(unittest.TestCase):
    def invalid(self, change=None, raw=None):
        path = ROOT / "rules" / RULES["R-001"].source_file
        doc = yaml.safe_load(path.read_text(encoding="utf-8"))
        if change:
            change(doc)
        with tempfile.TemporaryDirectory() as directory:
            Path(directory, "test.yaml").write_text(raw if raw is not None else yaml.safe_dump(doc, allow_unicode=True), encoding="utf-8")
            with self.assertRaises(engine.RuleError):
                engine.load_rules(directory)

    def test_required_and_unknown_fields(self):
        for key in engine.REQUIRED_FIELDS:
            self.invalid(lambda d: d.pop(key))
        self.invalid(lambda d: d.update(threshhold=1))

    def test_bad_severity_id_category_and_types(self):
        for key, value in [("severity", "urgent"), ("severity", []), ("id", "bad"), ("id", 1),
                           ("category", "unknown"), ("logic", []), ("inputs", []),
                           ("evidence", "营业收入"), ("references", []), ("description", [])]:
            self.invalid(lambda d: d.update({key: value}))

    def test_invalid_logic(self):
        for logic in [{"type": "eval"}, {"type": "all", "conditions": []},
                      {"type": "ratio_range", "numerator": "营业收入", "denominator": "增值税.销售额", "min": 2, "max": 1}]:
            self.invalid(lambda d: d.update(logic=logic))
        for value in (-.1, True, float("inf"), float("nan"), "abc"):
            self.invalid(lambda d: d["logic"].update(threshold=value))
        self.invalid(lambda d: d["logic"].update(direction="sideways"))
        self.invalid(lambda d: d["logic"].update(unknown=1))

    def test_invalid_expressions_and_missing_evidence(self):
        for expr in ({"eval": [1, 2]}, {"sub": [1, 2, 3]}, {"add": []}, None):
            self.invalid(lambda d: d["logic"].update(left=expr))
        self.invalid(lambda d: d["evidence"].pop())
        self.invalid(lambda d: d["evidence"].append(d["evidence"][0]))

    def test_duplicate_yaml_keys_and_unsafe_yaml(self):
        self.invalid(raw="id: R-001\nid: R-002\n")
        self.invalid(raw="!!python/object/apply:os.system ['echo unsafe']")

    def test_duplicate_rule_ids(self):
        with tempfile.TemporaryDirectory() as directory:
            content = (ROOT / "rules" / RULES["R-001"].source_file).read_text(encoding="utf-8")
            for name in ("a.yaml", "b.yml"):
                Path(directory, name).write_text(content, encoding="utf-8")
            with self.assertRaisesRegex(engine.RuleError, "重复"):
                engine.load_rules(directory)


class ExcelValidation(unittest.TestCase):
    def parse(self, mutate, rule_id="R-004"):
        wb = workbook(GROUPS[rule_id]["baseline"], RULES[rule_id])
        mutate(wb)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "input.xlsx"
            wb.save(path)
            wb.close()
            return loader.load(path)

    def test_empty_value_remains_missing(self):
        data = self.parse(lambda w: setattr(w["补充指标"]["B2"], "value", None))
        self.assertEqual(engine.evaluate(RULES["R-004"], data).status, "skipped")
        data = self.parse(lambda w: setattr(w["增值税申报"]["B2"], "value", None), "R-001")
        self.assertEqual(engine.evaluate(RULES["R-001"], data).status, "skipped")

    def test_partial_account_aggregate_is_missing(self):
        data = self.parse(lambda w: w["科目余额表"].delete_rows(3), "R-001")
        self.assertNotIn("营业收入", data.metrics)

    def test_duplicates(self):
        for sheet in ("企业信息", "科目余额表", "增值税申报", "补充指标"):
            with self.subTest(sheet=sheet), self.assertRaises(loader.InputError):
                self.parse(lambda w: w[sheet].append([c.value for c in w[sheet][2]]))

    def test_invalid_cells(self):
        for value in ("abc", "NaN", "Infinity", True, "=1+2", -1, 1.5):
            with self.subTest(value=value), self.assertRaises(loader.InputError):
                self.parse(lambda w: setattr(w["补充指标"]["B2"], "value", value))

    def test_missing_source_period_or_detail(self):
        for address, value in (("C2", None), ("D2", "2025年度"), ("E2", None)):
            with self.subTest(address=address), self.assertRaises(loader.InputError):
                self.parse(lambda w: setattr(w["补充指标"][address], "value", value))

    def test_no_override_of_original_metrics(self):
        with self.assertRaises(loader.InputError):
            self.parse(lambda w: setattr(w["补充指标"]["A2"], "value", "营业收入"))

    def test_missing_sheet_and_bad_headers(self):
        for sheet in ("企业信息", "科目余额表", "增值税申报"):
            with self.subTest(sheet=sheet), self.assertRaises(loader.InputError):
                self.parse(lambda w: w.remove(w[sheet]))
        for sheet in ("企业信息", "科目余额表", "增值税申报", "补充指标"):
            with self.subTest(sheet=sheet), self.assertRaises(loader.InputError):
                self.parse(lambda w: setattr(w[sheet]["A1"], "value", "错误表头"))

    def test_invalid_file(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "invalid.xlsx"
            path.write_text("not an Excel file", encoding="utf-8")
            with self.assertRaises(loader.InputError):
                loader.load(path)


if __name__ == "__main__":
    unittest.main()
