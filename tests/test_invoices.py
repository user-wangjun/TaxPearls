from __future__ import annotations

from copy import deepcopy
from decimal import Decimal
from io import BytesIO
from pathlib import Path
import json
import tempfile
import unittest
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import Workbook

from src import engine, materials, material_provenance
from webapp import app as app_module
from webapp import material_batches
from webapp.storage import Store
from tests.enterprise_support import confirm, material_key


ROOT = Path(__file__).resolve().parent.parent
COMPANY = {"name": "发票适配测试企业", "taxpayer_id": "BUYER-TEST-001", "industry": "批发业", "period": "2026-01"}
KEYS = {key for rule in engine.load_rules(ROOT / "rules") for key in rule.inputs}


def selections(docs):
    return {
        doc["id"]: {
            "id": doc["id"],
            "company": deepcopy(doc["company"]),
            "rows": deepcopy(doc["rows"]),
            "reviewed": True,
        }
        for doc in docs
    }


def invoice_workbook(rows, headers=None, sheet="电子税务局发票查询", company=COMPANY):
    headers = headers or [
        "数电票号码", "开票时间", "发票方向", "合计金额", "合计税额", "价税合计",
        "发票状态", "用途确认状态", "用途确认", "用途确认所属期", "是否进项税额转出", "进项转出税额",
    ]
    workbook = Workbook()
    workbook.remove(workbook.active)
    if company:
        ws = workbook.create_sheet("企业信息")
        ws.append(["项目", "内容"])
        for label, key in zip(("企业名称", "纳税人识别号", "所属行业", "所属期"), company):
            ws.append([label, company[key]])
    ws = workbook.create_sheet(sheet)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def invoice_xml(
    number="24442000000000000001",
    amount="100.00",
    tax="13.00",
    total="113.00",
    issued="2026-01-10T09:30:00",
    red="false",
    confirmed="true",
    usage="抵扣税款",
    period="2026-01",
    transferred="false",
    transferred_tax="0",
):
    return f'''<?xml version="1.0" encoding="UTF-8"?>
<xbrli:xbrl xmlns:xbrli="http://www.xbrl.org/2003/instance" xmlns:einv="http://xbrl.mof.gov.cn/taxonomy/2023-12-31/einv">
  <einv:InvoiceNumber>{number}</einv:InvoiceNumber>
  <einv:SellerName>测试供应商</einv:SellerName>
  <einv:SellerIdNum>SELLER-TEST-001</einv:SellerIdNum>
  <einv:UnifiedSocialCreditCodeOfAccountingEntity>{COMPANY["taxpayer_id"]}</einv:UnifiedSocialCreditCodeOfAccountingEntity>
  <einv:NameOfAccountingEntity>{COMPANY["name"]}</einv:NameOfAccountingEntity>
  <einv:RequestTime>{issued}</einv:RequestTime>
  <einv:TotalAmWithoutTax>{amount}</einv:TotalAmWithoutTax>
  <einv:TotalTaxAm>{tax}</einv:TotalTaxAm>
  <einv:TotalTax-includedAmount>{total}</einv:TotalTax-includedAmount>
  <einv:WhetherEinvoiceIsRedEinvoice>{red}</einv:WhetherEinvoiceIsRedEinvoice>
  <einv:WhetherEinvoiceUsageHasBeenConfirmed>{confirmed}</einv:WhetherEinvoiceUsageHasBeenConfirmed>
  <einv:UsageConfirmation>{usage}</einv:UsageConfirmation>
  <einv:PeriodOfUsageConfirmation>{period}</einv:PeriodOfUsageConfirmation>
  <einv:WhetherInputVatHasBeenTransferredOut>{transferred}</einv:WhetherInputVatHasBeenTransferredOut>
  <einv:AmountOfTransferredOutInputVat>{transferred_tax}</einv:AmountOfTransferredOutInputVat>
</xbrli:xbrl>'''.encode()


class InvoiceSourceParsing(unittest.TestCase):
    def test_conflicting_signs_and_ambiguous_states_never_become_net_amounts(self):
        for amount, tax, total, status, message in (
            (-100, 13, -113, '红字', '正负号不一致'),
            (100, -13, 113, '红字', '正负号不一致'),
            (100, 13, -113, '正常', '正负号不一致'),
            (-100, -13, -113, '正常', '正常发票金额不能为负数'),
            (100, 13, 113, '已红冲', '独立红票'),
            (100, 13, 113, '部分红冲', '独立红票'),
            (100, 13, 113, '未作废', '无法识别'),
            (100, 13, 113, '非红字', '无法识别'),
            (100, 13, 113, '', '缺少明确发票状态'),
        ):
            with self.subTest(amount=amount, tax=tax, total=total, status=status):
                doc = materials.preview([('状态.xlsx', invoice_workbook([
                    ['STATE-1', '2026-01-02', '销项', amount, tax, total, status]]))], KEYS)[0]
                self.assertIn(message, doc['error'])
                self.assertEqual(doc['invoices'], [])

    def test_red_positive_and_negative_exports_subtract_once_with_separate_sources(self):
        raw = invoice_workbook([
            ['BLUE', '2026-01-02', '销项', 100, 13, 113, '正常'],
            ['RED-POS', '2026-01-03', '销项', 20, 2.6, 22.6, '红字'],
            ['RED-NEG', '2026-01-04', '销项', -10, -1.3, -11.3, '红字'],
            ['VOID', '2026-01-05', '销项', 50, 6.5, 56.5, '作废'],
            ['PURCHASE', '2026-01-06', '采购', 60, 7.8, 67.8, '正常', '已确认', '抵扣税款', '2026-01', '否', 0],
        ])
        doc = materials.preview([('正负红票.xlsx', raw)], KEYS)[0]
        self.assertFalse(doc['error'], doc)
        data = materials.build_dataset([doc], selections([doc]), {}, KEYS)
        self.assertEqual(data.get('发票.销项净额'), 70)
        self.assertEqual(data.get('发票.销项税额净额'), Decimal('9.1'))
        self.assertIn('VOID；作废已剔除', data.source_of('发票.销项净额'))
        self.assertNotIn('PURCHASE', data.source_of('发票.销项净额'))
        self.assertNotIn('BLUE', data.source_of('凭证.本期确认抵扣税额'))
        self.assertIn('原文不含税 -10', doc['invoices'][2]['detail'])

    def test_status_flags_are_checked_in_excel_and_xml(self):
        headers = ['发票号码', '开票日期', '发票方向', '合计金额', '合计税额', '价税合计', '发票状态', '红字标志', '作废标志']
        for status, red, void in (('蓝字', '是', '否'), ('红字', '否', '否'), ('作废', '否', '否'), ('正常', '否', '是')):
            with self.subTest(status=status, red=red, void=void):
                doc = materials.preview([('标志.xlsx', invoice_workbook([
                    ['FLAG', '2026-01-02', '销项', 100, 13, 113, status, red, void]], headers))], KEYS)[0]
                self.assertIn('标志冲突', doc['error'])
        for status, red, void in (('蓝字', 'true', 'false'), ('红字', 'false', 'false'), ('作废', 'false', 'false')):
            raw = invoice_xml(red=red).replace(b'</xbrli:xbrl>', (
                f'<einv:InvoiceStatus>{status}</einv:InvoiceStatus><einv:IsVoid>{void}</einv:IsVoid></xbrli:xbrl>').encode())
            doc = materials.preview([('标志.xml', raw)], KEYS)[0]
            self.assertIn('标志冲突', doc['error'])

    def test_validity_and_colour_are_separate_for_red_invoices(self):
        headers = ['发票号码', '开票日期', '发票方向', '合计金额', '合计税额', '价税合计', '发票状态', '红字标志', '作废标志']
        for amount, tax, total in ((20, 2.6, 22.6), (-20, -2.6, -22.6)):
            raw = invoice_workbook([['VALID-RED', '2026-01-02', '销项', amount, tax, total, '正常', '是', '否']], headers)
            doc = materials.preview([('有效红票.xlsx', raw)], KEYS)[0]
            self.assertFalse(doc['error'], doc)
            data = materials.build_dataset([doc], selections([doc]), {}, KEYS)
            self.assertEqual(data.get('发票.销项净额'), -20)
        raw = invoice_xml(red='true').replace(b'</xbrli:xbrl>', '<einv:InvoiceStatus>正常</einv:InvoiceStatus></xbrli:xbrl>'.encode())
        doc = materials.preview([('有效红票.xml', raw)], KEYS)[0]
        self.assertFalse(doc['error'], doc)
        self.assertEqual(doc['invoices'][0]['status'], '红字')

    def test_xml_void_and_negative_blue_are_blocked(self):
        cases = (
            (invoice_xml().replace(b'</xbrli:xbrl>', b'<einv:IsVoid>true</einv:IsVoid></xbrli:xbrl>'), '作废状态'),
            (invoice_xml(amount='-100', tax='-13', total='-113'), '正常发票金额不能为负数'),
            (invoice_xml(amount='-100', tax='13', total='-113', red='true'), '正负号不一致'),
        )
        for raw, message in cases:
            with self.subTest(message=message):
                doc = materials.preview([('错误.xml', raw)], KEYS)[0]
                self.assertIn(message, doc['error'])
                self.assertEqual(doc['invoices'], [])

    def test_deduction_needs_explicit_usage_period_and_transfer_amount(self):
        for usage, period in (('不抵扣', '2026-01'), ('抵扣税款', ''), ('', '2026-01'), ('出口退税', '2026-01'), ('抵扣税款', '2025-12')):
            with self.subTest(usage=usage, period=period):
                for name, raw in (
                    ('用途.xml', invoice_xml(usage=usage, period=period)),
                    ('用途.xlsx', invoice_workbook([['USE', '2026-01-02', '采购', 100, 13, 113, '正常', '已确认', usage, period, '否', 0]])),
                ):
                    doc = materials.preview([(name, raw)], KEYS)[0]
                    self.assertFalse(doc['error'], doc)
                    data = materials.build_dataset([doc], selections([doc]), COMPANY, KEYS)
                    self.assertEqual(data.get('发票.采购税额净额'), 13)
                    self.assertIsNone(data.get('凭证.本期确认抵扣税额'))
        for transferred, amount in (('true', ''), ('false', '2'), ('', '2'), ('true', '-2'), ('true', '14')):
            with self.subTest(transferred=transferred, amount=amount):
                doc = materials.preview([('转出.xml', invoice_xml(transferred=transferred, transferred_tax=amount))], KEYS)[0]
                self.assertIn('转出', doc['error'])
                doc = materials.preview([('转出.xlsx', invoice_workbook([
                    ['TRANSFER', '2026-01-02', '采购', 100, 13, 113, '正常', '已确认', '抵扣税款', '2026-01', transferred, amount]]))], KEYS)[0]
                self.assertIn('转出', doc['error'])

    def test_numeric_long_invoice_number_requires_new_source_text_export(self):
        for number in (24442000000000000001, 2.4442e19):
            doc = materials.preview([('长票号.xlsx', invoice_workbook([
                [number, '2026-01-02', '销项', 100, 13, 113, '正常']]))], KEYS)[0]
            self.assertIn('可能已丢失精度', doc['error'])
        doc = materials.preview([('文本票号.xlsx', invoice_workbook([
            ['24442000000000000001', '2026-01-02', '销项', 100, 13, 113, '正常']]))], KEYS)[0]
        self.assertFalse(doc['error'], doc)
        self.assertEqual(doc['invoices'][0]['number'], '24442000000000000001')

    def test_party_ids_cannot_be_overridden_by_direction_or_sheet_title(self):
        headers = ['发票号码', '开票日期', '发票方向', '合计金额', '合计税额', '价税合计', '发票状态', '销方税号', '购方税号']
        for direction, seller, buyer in (
            ('销项', 'OTHER', COMPANY['taxpayer_id']),
            ('采购', COMPANY['taxpayer_id'], 'OTHER'),
            ('销项', 'OTHER-A', 'OTHER-B'),
            ('销项', COMPANY['taxpayer_id'], COMPANY['taxpayer_id']),
        ):
            doc = materials.preview([('归属.xlsx', invoice_workbook([
                ['PARTY', '2026-01-02', direction, 100, 13, 113, '正常', seller, buyer]], headers, sheet='销项查询'))], KEYS)[0]
            self.assertIn('税号', doc['error'])

    def test_duplicate_feedback_locates_both_rows(self):
        doc = materials.preview([('重复状态.xlsx', invoice_workbook([
            ['SAME', '2026-01-02', '销项', 100, 13, 113, '正常'],
            ['SAME', '2026-01-02', '销项', 100, 13, 113, '作废']]))], KEYS)[0]
        self.assertIn('第2行', doc['error'])
        self.assertIn('第3行', doc['error'])

    def test_invoice_parser_version_and_sources_invalidate_old_rows(self):
        doc = materials.preview([('发票.xml', invoice_xml())], KEYS)[0]
        self.assertFalse(material_provenance.invoice_reanalysis_required(doc))
        for key in ('version', 'source', 'status', 'missing_program'):
            old = deepcopy(doc)
            if key == 'version':
                old.pop('invoice_parser_version')
            elif key == 'source':
                old['extraction']['local']['program']['sources']['loader.py'] = 'old'
            elif key == 'status':
                old['extraction']['local']['status'] = 'failed'
            else:
                old['extraction']['local'].pop('program')
            with self.subTest(key=key):
                self.assertTrue(material_provenance.invoice_reanalysis_required(old))
        self.assertFalse(material_provenance.invoice_reanalysis_required({'kind': 'history', 'invoices': doc['invoices']}))

    def test_source_excel_maps_red_void_and_tax_basis(self):
        rows = [
            ["S-1", "2026-01-02", "销项", 100, 13, 113, "正常", "", "", "", "", ""],
            ["S-2", "2026-01-03", "销项", 20, 2.6, 22.6, "红字", "", "", "", "", ""],
            ["S-3", "2026-01-04", "销项", 50, 6.5, 56.5, "作废", "", "", "", "", ""],
            ["P-1", "2026-01-05", "采购", 60, 7.8, 67.8, "正常", "已确认", "抵扣税款", "2026-01", "否", 0],
            ["P-2", "2026-01-06", "采购", 40, 5.2, 45.2, "正常", "未确认", "", "", "否", 0],
            ["P-3", "2026-01-07", "采购", 10, 1.3, 11.3, "正常", "已确认", "抵扣税款", "2026-01", "是", 0.3],
        ]
        doc = materials.preview([("电子税务局导出.xlsx", invoice_workbook(rows))], KEYS)[0]
        self.assertFalse(doc["error"], doc)
        self.assertEqual(len(doc["invoices"]), 6)
        data = materials.build_dataset([doc], selections([doc]), {}, KEYS)
        self.assertEqual(data.get("发票.销项净额"), 80)
        self.assertEqual(data.get("发票.销项税额净额"), Decimal("10.4"))
        self.assertEqual(data.get("发票.采购不含税净额"), 110)
        self.assertEqual(data.get("发票.采购税额净额"), Decimal("14.3"))
        self.assertEqual(data.get("凭证.本期确认抵扣税额"), Decimal("8.8"))
        self.assertIn("剔除作废 1 张", data.detail_of("发票.销项净额"))
        self.assertIn("未确认用途的采购税额不计入", data.detail_of("凭证.本期确认抵扣税额"))
        findings = engine.run(engine.load_rules(ROOT / "rules"), data)
        self.assertEqual(next(finding for finding in findings if finding.rule.id == "R-019").status, "pass")

    def test_excel_rejects_duplicate_number_and_tax_inclusive_mismatch(self):
        duplicate = [
            ["DUP-1", "2026-01-02", "销项", 100, 13, 113, "正常"],
            ["DUP-1", "2026-01-03", "销项", 20, 2.6, 22.6, "红字"],
        ]
        doc = materials.preview([("重复.xlsx", invoice_workbook(duplicate))], KEYS)[0]
        self.assertIn("发票号码重复", doc["error"])
        mismatch = [["BAD-1", "2026-01-02", "销项", 100, 13, 999, "正常"]]
        doc = materials.preview([("税额口径.xlsx", invoice_workbook(mismatch))], KEYS)[0]
        self.assertIn("价税合计", doc["error"])

    def test_mof_xbrl_xml_merge_red_and_duplicate_detection(self):
        docs = materials.preview([
            ("采购蓝票.xml", invoice_xml()),
            ("采购红票.xml", invoice_xml(
                number="24442000000000000002", amount="20", tax="2.6", total="22.6", red="true"
            )),
        ], KEYS)
        self.assertTrue(all(not doc["error"] for doc in docs), docs)
        data = materials.build_dataset(docs, selections(docs), COMPANY, KEYS)
        self.assertEqual(data.get("发票.采购不含税净额"), 80)
        self.assertEqual(data.get("发票.采购税额净额"), Decimal("10.4"))
        self.assertEqual(data.get("凭证.本期确认抵扣税额"), Decimal("10.4"))
        self.assertIn("采购蓝票.xml", data.source_of("发票.采购不含税净额"))
        self.assertIn("采购红票.xml", data.source_of("发票.采购不含税净额"))

        duplicate = materials.preview([("a.xml", invoice_xml()), ("b.xml", invoice_xml())], KEYS)
        with self.assertRaisesRegex(materials.InputError, "跨文件重复"):
            materials.build_dataset(duplicate, selections(duplicate), COMPANY, KEYS)

    def test_xml_unconfirmed_tax_is_not_deductible_and_period_is_enforced(self):
        doc = materials.preview([("未确认.xml", invoice_xml(confirmed="false", usage=""))], KEYS)[0]
        data = materials.build_dataset([doc], selections([doc]), COMPANY, KEYS)
        self.assertEqual(data.get("发票.采购税额净额"), 13)
        self.assertIsNone(data.get("凭证.本期确认抵扣税额"))

        outside = materials.preview([("跨期.xml", invoice_xml(issued="2025-12-31"))], KEYS)[0]
        with self.assertRaisesRegex(materials.InputError, "不在核对期间"):
            materials.build_dataset([outside], selections([outside]), COMPANY, KEYS)

    def test_xml_security_and_required_fields(self):
        entity = b'''<?xml version="1.0"?><!DOCTYPE x [<!ENTITY e "boom">]><x><InvoiceNumber>&e;</InvoiceNumber></x>'''
        with self.assertRaisesRegex(materials.InputError, "DTD"):
            materials.preview([("实体.xml", entity)], KEYS)
        fake = b'''<xbrl xmlns:einv="https://example.test/einv"><einv:InvoiceNumber>1</einv:InvoiceNumber></xbrl>'''
        with self.assertRaisesRegex(materials.InputError, "财政部"):
            materials.preview([("非标准.xml", fake)], KEYS)
        mismatch = invoice_xml(total="999")
        self.assertIn("价税合计", materials.preview([("错误.xml", mismatch)], KEYS)[0]["error"])


@material_key
class InvoiceWebFlow(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.old_store = app_module.store
        app_module.store = Store(Path(self.temp.name) / "invoice.db")
        self.admin = app_module.store.create_user("invoice-admin", "invoice-test-2026", "发票测试", "org_admin", "default")
        self.client = TestClient(app_module.app)
        self.client.post("/api/login", json={"username": "invoice-admin", "password": "invoice-test-2026"})

    def tearDown(self):
        self.client.close()
        app_module.store = self.old_store
        self.temp.cleanup()

    def test_duplicate_across_xml_excel_blocks_confirmation_and_exclusion_recovers(self):
        raw = invoice_workbook([['24442000000000000001', '2026-01-10', '采购', 100, 13, 113, '正常', '已确认', '抵扣税款', '2026-01', '否', 0]])
        response = self.client.post('/api/enterprise/materials', files=[
            ('files', ('清单.xlsx', raw)), ('files', ('原票.xml', invoice_xml()))])
        self.assertEqual(response.status_code, 200, response.text)
        item = response.json()
        self.assertFalse(item['analysis']['can_confirm'])
        message = ' '.join(n['message'] for n in item['analysis']['feedback']['blocking'])
        self.assertIn('清单.xlsx', message)
        self.assertIn('原票.xml', message)
        self.assertIn('跨文件重复', message)
        self.assertEqual(confirm(self.client, item).status_code, 409)
        url = '/api/enterprise/materials/' + item['id']
        xml_id = item['documents'][1]['id']
        response = self.client.post(url + '/analyze', json={
            'expected_revision': item['revision'], 'selections': {xml_id: {'purpose': 'excluded'}}})
        self.assertEqual(response.status_code, 200, response.text)
        item = response.json()
        self.assertTrue(item['analysis']['can_confirm'], item['analysis'])
        first = confirm(self.client, item)
        self.assertEqual(first.status_code, 200, first.text)
        saved = app_module.store.get_audit(first.json()['audit_id'])['dataset']
        self.assertEqual(saved.get('发票.采购税额净额'), 13)
        self.assertNotIn('原票.xml', saved.source_of('发票.采购税额净额'))
        for document, original in zip(item['documents'], (raw, invoice_xml())):
            download = self.client.get(url + '/originals/' + document['original_id'])
            self.assertEqual(download.status_code, 200, download.text)
            self.assertEqual(download.content, original)
        # Reintroducing the conflict cannot overwrite the earlier frozen result.
        current = self.client.get(url).json()
        response = self.client.post(url + '/analyze', json={
            'expected_revision': current['revision'], 'selections': {xml_id: {'purpose': 'current'}}})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertFalse(response.json()['analysis']['can_confirm'])
        repeated = confirm(self.client, item)
        self.assertEqual(repeated.status_code, 200, repeated.text)
        self.assertEqual(repeated.json()['audit_id'], first.json()['audit_id'])

    def test_failed_state_and_transfer_cannot_create_audit(self):
        for raw in (
            invoice_xml().replace(b'</xbrli:xbrl>', '<einv:InvoiceStatus>正常</einv:InvoiceStatus><einv:IsRedInvoice>true</einv:IsRedInvoice></xbrli:xbrl>'.encode()),
            invoice_xml(transferred='true', transferred_tax=''),
        ):
            response = self.client.post('/api/enterprise/materials', files={'files': ('异常.xml', raw)})
            self.assertEqual(response.status_code, 200, response.text)
            item = response.json()
            self.assertFalse(item['analysis']['can_confirm'])
            self.assertEqual(confirm(self.client, item).status_code, 409)
        self.assertEqual(app_module.store.list_audits(self.admin), [])

    def test_queued_duplicate_is_blocked_after_processing(self):
        with patch.dict('os.environ', {'TAXPEARLS_MATERIAL_QUEUE_ENABLED': '1'}):
            response = self.client.post('/api/enterprise/materials', files=[
                ('files', ('a.xml', invoice_xml())), ('files', ('b.xml', invoice_xml()))],
                data={'company': json.dumps({'name': COMPANY['name'], 'taxpayer_id': COMPANY['taxpayer_id'],
                    'period_start': '2026-01-01', 'period_end': '2026-01-31'})})
            self.assertEqual(response.status_code, 202, response.text)
            item = response.json()
            self.assertTrue(app_module.app.state.material_worker.run_once())
            response = self.client.get('/api/enterprise/materials/' + item['id'])
            self.assertEqual(response.status_code, 200, response.text)
            item = response.json()
            self.assertEqual(item['jobs'][0]['state'], 'done')
            self.assertFalse(item['analysis']['can_confirm'])
            self.assertIn('跨文件重复', ' '.join(n['message'] for n in item['analysis']['feedback']['blocking']))
            self.assertEqual(confirm(self.client, item).status_code, 409)

    def test_old_accepted_red_flush_state_is_rejected_on_source_rescan(self):
        raw = invoice_workbook([['OLD-RED', '2026-01-02', '销项', 100, 13, 113, '已红冲']])
        # Simulate the former substring classifier; retain the actual bad source.
        with patch('src.loader._invoice_status', return_value='红字'):
            response = self.client.post('/api/enterprise/materials', files={'files': ('旧蓝票.xlsx', raw)})
        self.assertEqual(response.status_code, 200, response.text)
        item = response.json()
        view = material_batches.read(app_module.store, self.admin, item['id'])
        payload = deepcopy(view['versions'][-1]['payload'])
        payload['documents'][0].pop('invoice_parser_version')
        payload['analysis']['can_confirm'] = True
        revision = material_batches.append_revision(app_module.store, self.admin, item['id'], item['revision'], 'edit', payload)
        url = '/api/enterprise/materials/' + item['id']
        response = self.client.post(url + '/analyze', json={'expected_revision': revision})
        self.assertEqual(response.status_code, 200, response.text)
        item = response.json()
        self.assertIn('独立红票', item['documents'][0]['error'])
        self.assertFalse(item['analysis']['can_confirm'])
        self.assertEqual(confirm(self.client, item).status_code, 409)

    def test_old_invoice_cache_requires_original_rescan_preserving_frozen_result(self):
        raw = invoice_workbook([['CACHE', '2026-01-02', '销项', 100, 13, 113, '正常']])
        response = self.client.post('/api/enterprise/materials', files={'files': ('留存发票.xlsx', raw)})
        self.assertEqual(response.status_code, 200, response.text)
        item = response.json()
        original_revision = item['revision']
        first = confirm(self.client, item)
        self.assertEqual(first.status_code, 200, first.text)
        url = '/api/enterprise/materials/' + item['id']
        current = self.client.get(url).json()
        view = material_batches.read(app_module.store, self.admin, item['id'])
        payload = deepcopy(next(v['payload'] for v in reversed(view['versions']) if v['kind'] in {'analysis', 'edit'}))
        doc = payload['documents'][0]
        fid = doc['id']
        doc.pop('invoice_parser_version')
        doc['invoices'][0]['amount'] = '999'
        payload['analysis']['can_confirm'] = True
        revision = material_batches.append_revision(app_module.store, self.admin, item['id'], current['revision'], 'edit', payload)
        old_versions = deepcopy(material_batches.read(app_module.store, self.admin, item['id'])['versions'])
        projected = self.client.get(url).json()
        self.assertTrue(projected['documents'][0]['invoice_reanalysis_required'])
        self.assertFalse(projected['analysis']['can_confirm'])
        self.assertIn('invoice_reanalysis_required', [n['code'] for n in projected['analysis']['feedback']['blocking']])
        self.assertEqual(confirm(self.client, projected).status_code, 409)
        refused = self.client.post(url + '/analyze', json={
            'expected_revision': revision, 'selections': {fid: {'evidence_reviews': ['old']}}})
        self.assertEqual(refused.status_code, 422, refused.text)
        with patch('src.ai_extraction.call_model', side_effect=AssertionError('invoice rescan called model')) as transport:
            response = self.client.post(url + '/analyze', json={'expected_revision': revision})
            self.assertEqual(response.status_code, 200, response.text)
            transport.assert_not_called()
        fresh = response.json()
        self.assertFalse(fresh['documents'][0]['invoice_reanalysis_required'])
        self.assertEqual(fresh['documents'][0]['invoices'][0]['amount'], '100')
        self.assertTrue(fresh['analysis']['can_confirm'], fresh['analysis'])
        self.assertEqual(fresh['selections'][fid]['evidence_reviews'], [])
        versions = material_batches.read(app_module.store, self.admin, item['id'])['versions']
        self.assertEqual(versions[:len(old_versions)], old_versions)
        repeated = self.client.post(url + '/confirm', json={'expected_revision': original_revision})
        self.assertEqual(repeated.status_code, 200, repeated.text)
        self.assertEqual(repeated.json()['audit_id'], first.json()['audit_id'])
        original = self.client.get(url + '/originals/' + fresh['documents'][0]['original_id'])
        self.assertEqual(original.content, raw)

    def test_xml_and_source_excel_commit_with_persisted_evidence(self):
        sales = invoice_workbook([
            ["S-WEB", "2026-01-02", "销项", 200, 26, 226, "正常", "", "", "", "", ""],
        ])
        response = self.client.post(
            "/api/enterprise/materials",
            files=[("files", ("销项导出.xlsx", sales)), ("files", ("采购.xml", invoice_xml()))],
        )
        self.assertEqual(response.status_code, 200, response.text)
        draft = response.json()
        self.assertEqual([doc["kind"] for doc in draft["documents"]], ["xlsx", "xml"])
        audited = confirm(self.client, draft)
        self.assertEqual(audited.status_code, 200, audited.text)
        body = audited.json()
        audit_id = body["audit_id"]
        saved = app_module.store.get_audit(audit_id)["dataset"]
        self.assertEqual(saved.get("发票.销项净额"), 200)
        self.assertEqual(saved.get("发票.采购不含税净额"), 100)
        self.assertIn("采购.xml", saved.source_of("发票.采购不含税净额"))


if __name__ == "__main__":
    unittest.main()
